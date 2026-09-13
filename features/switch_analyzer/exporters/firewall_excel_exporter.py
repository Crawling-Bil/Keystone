from __future__ import annotations

from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _join(values: Any) -> str:
    if isinstance(values, (list, tuple)):
        return ", ".join(str(v) for v in values if v not in (None, ""))
    return str(values) if values not in (None, "") else ""


def _yn(value: Any) -> str:
    return "Yes" if value else "No"


class FirewallExcelExporter:
    """Dumps one or more analyze_firewall() results (Mikrotik and/or
    Palo Alto -- the same JSON the Config Analyzer's own tabs render)
    into a single multi-sheet .xlsx workbook, one sheet PER DATA
    CATEGORY, mirroring SwitchExcelExporter's own convention in this
    same package (combine-devices-into-rows rather than a sheet-set
    per device) so a single-device export -- the common case, from the
    analyzer's own Export button -- is just a one-item list.

    Both vendors share most categories (Interfaces, Zones, Security
    Rules, NAT, Address Objects, IPsec, DHCP), so those get one shared
    sheet with a Vendor column and any vendor-only field (e.g.
    Mikrotik's rule Chain, Palo Alto's rule-level Security Profile)
    left blank on rows from the vendor that doesn't have it -- the
    same "superset columns, blank where not applicable" convention
    SwitchExcelExporter already uses across Cisco/Huawei/Aruba's own
    differing field sets. A few categories have NO shared shape at all
    (RouterOS's `/queue simple|tree` command dump vs. PAN-OS's named
    QoS profile+binding objects; Palo Alto-only service/application
    objects and security-profile objects) -- those get their own
    vendor-labeled sheet instead of forcing one artificial shape onto
    both, which is exactly the "don't reuse Palo Alto's table shape
    for Mikrotik data" complaint this exporter (and the analyzer's own
    on-screen tables) is built to avoid.
    """

    def __init__(self, results: list[dict[str, Any]] | dict[str, Any]):
        self.results: list[dict[str, Any]] = [results] if isinstance(results, dict) else list(results)

    @staticmethod
    def add_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
        sheet = workbook.create_sheet(title[:31])  # Excel sheet-name length limit
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        for column in sheet.columns:
            max_length = max((len(str(cell.value or "")) for cell in column), default=0)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max_length + 2, 60)
        if rows:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

    def export(self, path) -> None:
        self.build_workbook().save(path)

    def build_workbook(self) -> Workbook:
        workbook = Workbook()
        workbook.remove(workbook.active)

        device_info_rows: list[list[Any]] = []
        interface_rows: list[list[Any]] = []
        port_mapping_rows: list[list[Any]] = []
        vlan_rows: list[list[Any]] = []
        zone_rows: list[list[Any]] = []
        route_rows: list[list[Any]] = []
        security_rows: list[list[Any]] = []
        nat_rows: list[list[Any]] = []
        address_rows: list[list[Any]] = []
        service_object_rows: list[list[Any]] = []  # Palo Alto only
        application_group_rows: list[list[Any]] = []  # Palo Alto only
        security_profile_rows: list[list[Any]] = []  # Palo Alto only
        virtual_router_rows: list[list[Any]] = []  # Palo Alto only
        management_profile_rows: list[list[Any]] = []  # Palo Alto only
        zone_protection_profile_rows: list[list[Any]] = []  # Palo Alto only
        device_configuration_rows: list[list[Any]] = []  # Palo Alto only
        administrator_rows: list[list[Any]] = []  # Palo Alto only
        pbf_rows: list[list[Any]] = []  # Palo Alto only
        qos_mikrotik_rows: list[list[Any]] = []
        qos_paloalto_rows: list[list[Any]] = []
        ipsec_rows: list[list[Any]] = []
        dhcp_rows: list[list[Any]] = []
        finding_rows: list[list[Any]] = []

        for result in self.results:
            hostname = result.get("hostname", "")
            vendor = result.get("vendor", "")
            is_paloalto = str(vendor).strip().lower() == "palo alto"
            cards = result.get("cards") or {}
            device_info = result.get("device_info") or {}
            routing_by_vr = {
                row.get("virtual_router", ""): row for row in (result.get("routing_protocols") or [])
            }

            device_info_rows.append([
                hostname, vendor, result.get("device_type", ""),
                _yn(result.get("panorama_enabled")) if is_paloalto else "",
                device_info.get("mgmt_ip", ""), device_info.get("domain", ""),
                device_info.get("timezone", ""),
                _join(result.get("dns_servers")), _join(result.get("ntp_servers")),
                cards.get("total_interfaces", ""), cards.get("total_security_rules", ""),
                cards.get("total_nat_rules", ""), cards.get("total_ipsec_tunnels", ""),
                cards.get("total_findings", ""), cards.get("high_severity_findings", ""),
            ])

            for row in result.get("interfaces") or []:
                interface_rows.append([
                    hostname, vendor, row.get("name", ""), row.get("type", ""),
                    row.get("zone", ""), row.get("ip_address", ""),
                    row.get("comment", ""), _yn(row.get("disabled")),
                ])

            for row in result.get("port_mapping") or []:
                port_mapping_rows.append([
                    hostname, vendor, row.get("default_name", ""), row.get("name", ""),
                    row.get("zone", ""), row.get("mode", ""), row.get("link_state", ""),
                    row.get("management_profile", ""), row.get("comment", ""),
                    _yn(row.get("disabled")),
                ])

            for row in result.get("vlans") or []:
                # Reconciled column names -- Mikrotik's vlan_id/parent_interface
                # and Palo Alto's tag/vlan_interface are the same two concepts
                # (the VLAN's own tag, and what it rides on top of) under
                # different field names; row.get() covers whichever pair the
                # source vendor actually filled in.
                vlan_rows.append([
                    hostname, vendor, row.get("name", ""),
                    row.get("vlan_id", row.get("tag", "")),
                    row.get("parent_interface", row.get("vlan_interface", "")),
                    _join(row.get("interfaces")), row.get("comment", ""),
                    _yn(row.get("disabled")),
                ])

            for row in result.get("zones") or []:
                zone_rows.append([
                    hostname, vendor, row.get("name", ""), _join(row.get("interfaces")),
                    row.get("zone_protection_profile", ""),
                ])

            for row in result.get("static_routes") or []:
                vr_routing = routing_by_vr.get(row.get("virtual_router", ""), {})
                route_rows.append([
                    hostname, vendor, row.get("destination", ""), row.get("nexthop", ""),
                    row.get("metric", ""), row.get("virtual_router", ""),
                    vr_routing.get("routing_type", ""), _join(vr_routing.get("dynamic_protocols")),
                    _yn(row.get("disabled")),
                ])

            for row in result.get("security_rules") or []:
                profile_setting = row.get("profile_setting") or {}
                profile_text = ", ".join(f"{k}:{'/'.join(v)}" for k, v in profile_setting.items())
                security_rows.append([
                    hostname, vendor, row.get("chain", ""), row.get("name", ""),
                    _join(row.get("from_zone")) or "any", _join(row.get("to_zone")) or "any",
                    _join(row.get("source")) or "any", _join(row.get("destination")) or "any",
                    _join(row.get("service")) or "any", _join(row.get("application")) or "any",
                    profile_text, _join(row.get("category")), _join(row.get("source_user")),
                    row.get("action", ""), _yn(row.get("log")), _yn(row.get("disabled")),
                ])

            for row in result.get("nat_rules") or []:
                nat_rows.append([
                    hostname, vendor, row.get("chain", row.get("type", "")), row.get("name", ""),
                    row.get("type", ""), _join(row.get("source")) or "any",
                    _join(row.get("destination")) or "any", _join(row.get("service")),
                    row.get("translated_address", ""), row.get("translated_port", ""),
                    _yn(row.get("disabled")),
                ])

            for row in result.get("address_objects") or []:
                address_rows.append([hostname, vendor, row.get("name", ""), _join(row.get("members"))])

            if is_paloalto:
                for row in result.get("service_objects") or []:
                    service_object_rows.append([hostname, row.get("name", ""), _join(row.get("members"))])
                for row in result.get("application_groups") or []:
                    application_group_rows.append([hostname, row.get("name", ""), _join(row.get("members"))])
                for row in result.get("security_profiles") or []:
                    security_profile_rows.append([
                        hostname, row.get("name", ""), row.get("type", ""), row.get("rule_count", ""),
                    ])
                for row in result.get("routing_protocols") or []:
                    virtual_router_rows.append([
                        hostname, row.get("virtual_router", ""), row.get("routing_type", ""),
                        row.get("static_route_count", ""), _yn(row.get("bgp")),
                        row.get("bgp_router_id", ""), row.get("bgp_as_number", ""),
                        _yn(row.get("ospf")), row.get("ospf_router_id", ""),
                        _join(row.get("ospf_area_ids")), _yn(row.get("rip")),
                    ])
                for row in result.get("management_profiles") or []:
                    management_profile_rows.append([
                        hostname, row.get("name", ""), _join(row.get("permitted_services")),
                    ])
                for row in result.get("zone_protection_profiles") or []:
                    zone_protection_profile_rows.append([
                        hostname, row.get("name", ""), _join(row.get("protection_types")),
                    ])
                device_configuration_rows.append([
                    hostname, _yn(result.get("panorama_enabled")),
                    device_info.get("mgmt_ip", ""), device_info.get("mgmt_netmask", ""),
                    device_info.get("mgmt_gateway", ""), device_info.get("domain", ""),
                    device_info.get("timezone", ""), _join(result.get("dns_servers")),
                    _join(result.get("ntp_servers")),
                ])
                for row in result.get("administrators") or []:
                    administrator_rows.append([hostname, row.get("username", ""), row.get("role", "")])
                for row in result.get("pbf_rules") or []:
                    pbf_rows.append([
                        hostname, row.get("name", ""), _join(row.get("from_zone")) or "any",
                        _join(row.get("source")) or "any", _join(row.get("destination")) or "any",
                        _join(row.get("application")) or "any", _join(row.get("service")) or "any",
                        row.get("egress_interface", ""), row.get("nexthop", ""),
                        row.get("monitor_profile", ""), _yn(row.get("disabled")),
                    ])

            qos = result.get("qos")
            if isinstance(qos, list):  # Mikrotik: raw /queue simple|tree command dump
                for row in qos:
                    qos_mikrotik_rows.append([hostname, row.get("section", ""), row.get("line", "")])
            elif isinstance(qos, dict):  # Palo Alto: named profile + interface binding objects
                profiles = ", ".join(qos.get("profiles") or [])
                for row in qos.get("interface_bindings") or []:
                    qos_paloalto_rows.append([hostname, row.get("interface", ""), row.get("profile", ""), profiles])
                if not qos.get("interface_bindings") and profiles:
                    qos_paloalto_rows.append([hostname, "", "", profiles])

            for row in result.get("ipsec_tunnels") or []:
                ipsec_rows.append([
                    hostname, vendor, row.get("name", ""), row.get("peer_address", ""),
                    row.get("dh_group", ""), row.get("ike_encryption", ""), row.get("ike_hash", ""),
                    row.get("esp_encryption", ""), row.get("esp_authentication", ""),
                    row.get("proxy_id_local", ""), row.get("proxy_id_remote", ""),
                    _yn(row.get("disabled")),
                ])

            dhcp = result.get("dhcp") or {}
            for row in dhcp.get("servers") or []:
                dhcp_rows.append([
                    hostname, vendor, row.get("name", ""), row.get("interface", ""),
                    row.get("address_pool", ""), row.get("gateway", ""), _yn(row.get("disabled")),
                ])

            for row in result.get("findings") or []:
                finding_rows.append([
                    hostname, vendor, row.get("severity", ""), row.get("category", ""),
                    row.get("title", ""), row.get("detail", ""),
                ])

        self.add_sheet(workbook, "Device Info", [
            "Hostname", "Vendor", "Device Type", "Panorama Managed", "Management IP", "Domain",
            "Timezone", "DNS Servers", "NTP Servers", "Total Interfaces", "Total Security Rules",
            "Total NAT Rules", "Total IPsec Tunnels", "Total Findings", "High Severity Findings",
        ], device_info_rows)
        self.add_sheet(workbook, "Interfaces", [
            "Hostname", "Vendor", "Name", "Type", "Zone", "IP Address", "Comment", "Disabled",
        ], interface_rows)
        self.add_sheet(workbook, "Port Mapping", [
            "Hostname", "Vendor", "Default-Factory Name", "Current Name", "Zone", "Mode",
            "Link State", "Management Profile", "Comment", "Disabled",
        ], port_mapping_rows)
        self.add_sheet(workbook, "VLANs", [
            "Hostname", "Vendor", "Name", "VLAN ID-Tag", "Parent-VLAN Interface",
            "Member Interfaces", "Comment", "Disabled",
        ], vlan_rows)
        self.add_sheet(workbook, "Zones", [
            "Hostname", "Vendor", "Zone", "Interfaces", "Zone Protection Profile (Palo Alto)",
        ], zone_rows)
        self.add_sheet(workbook, "Static Routes", [
            "Hostname", "Vendor", "Destination", "Next Hop", "Metric", "Virtual Router",
            "Virtual Router Routing Type", "Dynamic Protocols on VR", "Disabled",
        ], route_rows)
        self.add_sheet(workbook, "Security Rules", [
            "Hostname", "Vendor", "Chain (Mikrotik)", "Name", "From Zone", "To Zone", "Source",
            "Destination", "Service", "Application", "Security Profile (Palo Alto)",
            "Category (Palo Alto)", "Source User (Palo Alto)", "Action", "Log", "Disabled",
        ], security_rows)
        self.add_sheet(workbook, "NAT Rules", [
            "Hostname", "Vendor", "Chain-Type", "Name", "Type", "Source", "Destination", "Service",
            "Translated Address", "Translated Port", "Disabled",
        ], nat_rows)
        self.add_sheet(workbook, "Address Objects", ["Hostname", "Vendor", "Name", "Members"], address_rows)
        self.add_sheet(workbook, "Service Objects (Palo Alto)", ["Hostname", "Name", "Members"], service_object_rows)
        self.add_sheet(workbook, "Application Groups (Palo Alto)", ["Hostname", "Name", "Members"], application_group_rows)
        self.add_sheet(workbook, "Security Profiles (Palo Alto)", ["Hostname", "Name", "Type", "Rule Count"], security_profile_rows)
        self.add_sheet(workbook, "Virtual Routers (Palo Alto)", [
            "Hostname", "Virtual Router", "Routing Type", "Static Route Count",
            "BGP Enabled", "BGP Router ID", "BGP AS Number",
            "OSPF Enabled", "OSPF Router ID", "OSPF Area(s)", "RIP Enabled",
        ], virtual_router_rows)
        self.add_sheet(workbook, "Management Profiles (Palo Alto)", [
            "Hostname", "Name", "Permitted Services",
        ], management_profile_rows)
        self.add_sheet(workbook, "Zone Protection Profiles (PA)", [
            "Hostname", "Name", "Protection Types",
        ], zone_protection_profile_rows)
        self.add_sheet(workbook, "Device Configuration (PA)", [
            "Hostname", "Panorama Managed", "Management IP", "Management Subnet",
            "Management Gateway", "Domain", "Timezone", "DNS Servers", "NTP Servers",
        ], device_configuration_rows)
        self.add_sheet(workbook, "Administrators (Palo Alto)", [
            "Hostname", "Username", "Role",
        ], administrator_rows)
        self.add_sheet(workbook, "PBF Rules (Palo Alto)", [
            "Hostname", "Name", "From Zone", "Source", "Destination", "Application", "Service",
            "Egress Interface", "Next Hop", "Monitor Profile", "Disabled",
        ], pbf_rows)
        self.add_sheet(workbook, "QoS (Mikrotik Queues)", ["Hostname", "Section", "Command"], qos_mikrotik_rows)
        self.add_sheet(workbook, "QoS (Palo Alto Profiles)", ["Hostname", "Interface", "Bound Profile", "Defined Profiles"], qos_paloalto_rows)
        self.add_sheet(workbook, "IPsec Tunnels", [
            "Hostname", "Vendor", "Name", "Peer Address", "DH Group", "IKE Encryption", "IKE Hash",
            "ESP Encryption", "ESP Authentication", "Proxy ID Local", "Proxy ID Remote", "Disabled",
        ], ipsec_rows)
        self.add_sheet(workbook, "DHCP", [
            "Hostname", "Vendor", "Name", "Interface", "Address Pool", "Gateway", "Disabled",
        ], dhcp_rows)
        self.add_sheet(workbook, "Findings", [
            "Hostname", "Vendor", "Severity", "Category", "Finding", "Detail",
        ], finding_rows)

        return workbook

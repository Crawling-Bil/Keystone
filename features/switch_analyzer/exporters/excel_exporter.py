from __future__ import annotations

from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..sizing import build_sizing_summary
from . import analysis_sheets


def _cell(value: Any) -> Any:
    """A list becomes comma-separated text: openpyxl cannot write a list into a cell."""
    return ", ".join(str(item) for item in value) if isinstance(value, (list, tuple)) else value


class SwitchExcelExporter:
    """Dumps one or more Switch Analyzer results (the same JSON the
    frontend renders into tabs) into a single multi-sheet .xlsx
    workbook — one sheet PER DATA CATEGORY (Interfaces, Transceivers,
    PoE, etc.), not one sheet-set per device. Every row carries a
    "Hostname" column, so combining more devices means more rows in the
    same sheets rather than more sheets — built for mapping many
    devices' data into one workbook for a sizing assessment, add
    devices to the `results` list to extend it. A single-device export
    (the common case from the analyzer's own Export button) is just a
    one-item list.
    """

    def __init__(self, results: list[dict[str, Any]] | dict[str, Any]):
        # Accept a bare single result too, so existing single-device
        # call sites don't need to remember to wrap it in a list.
        self.results: list[dict[str, Any]] = [results] if isinstance(results, dict) else list(results)

    @staticmethod
    def add_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
        sheet = workbook.create_sheet(title[:31])  # Excel sheet-name length limit
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([_cell(value) for value in row])
        for column in sheet.columns:
            max_length = max((len(str(cell.value or "")) for cell in column), default=0)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max_length + 2, 60)
        if rows:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

    def export(self, path) -> None:
        self.build_workbook().save(path)

    def build_workbook(self) -> Workbook:
        workbook = Workbook()
        workbook.remove(workbook.active)

        device_info_rows: list[list[Any]] = []
        serial_rows: list[list[Any]] = []
        stack_rows: list[list[Any]] = []
        interface_rows: list[list[Any]] = []
        transceiver_rows: list[list[Any]] = []
        port_channel_rows: list[list[Any]] = []
        vpc_mlag_summary_rows: list[list[Any]] = []
        vpc_mlag_member_rows: list[list[Any]] = []
        poe_rows: list[list[Any]] = []
        vlan_rows: list[list[Any]] = []
        neighbor_rows: list[list[Any]] = []
        route_rows: list[list[Any]] = []
        routing_table_rows: list[list[Any]] = []
        l3_interface_rows: list[list[Any]] = []
        l3_default_rows: list[list[Any]] = []
        l3_ospf_rows: list[list[Any]] = []
        l3_acl_rows: list[list[Any]] = []
        l3_acl_binding_rows: list[list[Any]] = []
        l3_finding_rows: list[list[Any]] = []
        arp_rows: list[list[Any]] = []
        mac_rows: list[list[Any]] = []
        dns_snmp_rows: list[list[Any]] = []
        dhcp_summary_rows: list[list[Any]] = []
        dhcp_pool_rows: list[list[Any]] = []
        dhcp_helper_rows: list[list[Any]] = []
        dhcp_snooping_rows: list[list[Any]] = []
        pbr_sla_rows: list[list[Any]] = []
        pbr_track_rows: list[list[Any]] = []
        pbr_policy_rows: list[list[Any]] = []
        pbr_binding_rows: list[list[Any]] = []
        pbr_eem_rows: list[list[Any]] = []
        # Findings, Health, Trunks, Spanning Tree and Logs (exporters/analysis_sheets.py)
        analysis_rows = analysis_sheets.new_rows()

        for result in self.results:
            hostname = result.get("hostname", "")
            os_version = result.get("os_version") or {}
            inventory = result.get("inventory") or {"modules": [], "stack": {}}
            stack = inventory.get("stack") or {}
            cards = result.get("cards") or {}

            serial_info = result.get("serials") or {}
            serial_by_member = {str(m.get("member")): m.get("serial", "") for m in serial_info.get("members", [])}

            device_info_rows.append([
                hostname, result.get("vendor", ""), result.get("device_type", ""),
                os_version.get("os_version", ""), os_version.get("model", ""),
                "; ".join((serial_info.get("display") or "").splitlines()),
                stack.get("mode", ""), stack.get("member_count", ""),
                cards.get("total_interfaces", ""), cards.get("interfaces_up", ""),
                cards.get("interfaces_down", ""), cards.get("interfaces_admin_down", ""),
                cards.get("interfaces_unknown", ""),
            ])

            # Hardware parts + transceivers (features/switch_analyzer/parts.py);
            # a result saved before that list existed only has the inventory
            # modules, which are written without a Kind.
            part_rows = result["parts"] if "parts" in result else inventory.get("modules", [])
            for m in part_rows:
                serial_rows.append([
                    hostname, m.get("slot", ""), m.get("kind", ""), m.get("description", ""),
                    m.get("pid", ""), m.get("serial", ""),
                ])

            for m in stack.get("members", []):
                stack_rows.append([
                    hostname, m.get("slot", ""), m.get("role", ""), m.get("mac", ""),
                    m.get("state", ""), m.get("priority", ""),
                    serial_by_member.get(str(m.get("slot")), ""),
                ])

            for i in result.get("interfaces", []):
                interface_rows.append([
                    hostname, i.get("name", ""), i.get("status", ""), i.get("protocol", ""), i.get("mode", ""),
                    i.get("vlan", ""), i.get("allowed_vlans", ""), i.get("ip_address", ""), i.get("description", ""),
                    i.get("port_security", ""), i.get("hsrp_vrrp_ip", ""),
                    i.get("status_detail", ""), i.get("status_source", ""),
                    i.get("speed", ""), i.get("duplex", ""), i.get("port_type", ""), i.get("link_since", ""), i.get("transitions", ""),
                    i.get("rx_errors", ""), i.get("tx_dropped", ""), i.get("util_in", ""), i.get("util_out", ""), i.get("mac_address", ""), i.get("mtu", ""),
                    i.get("bundle", ""), i.get("bundle_state", ""), ", ".join(m.get("port", "") for m in i.get("bundle_members") or []),
                ])

            for t in result.get("transceivers", []):
                transceiver_rows.append([
                    hostname, t.get("interface", ""), t.get("status", ""), "Yes" if t.get("present") else "No",
                    t.get("type", ""), t.get("vendor_part_number", ""), t.get("serial_number", ""),
                ])

            for p in result.get("port_channels", []):
                port_channel_rows.append([
                    hostname, p.get("name", ""), p.get("protocol", ""), p.get("status", ""),
                    p.get("description", ""), p.get("mode", ""), p.get("vlan", ""),
                    p.get("allowed_vlans", ""), p.get("ip_address", ""), _cell(p.get("members", "")),
                ])

            # source: "live" (a real show-command session), "config" (no
            # live status yet — a not-yet-deployed draft, see
            # service._apply_vpc_mlag_config_fallback), or "none" (no
            # vPC/M-LAG at all). Carried into the export so a reviewer
            # opening the workbook later doesn't mistake a draft's
            # configured-only rows for a confirmed, active vPC/M-LAG.
            vpc_mlag = result.get("vpc_mlag") or {"summary": {}, "members": [], "source": "none"}
            source = vpc_mlag.get("source", "")
            for key, value in (vpc_mlag.get("summary") or {}).items():
                vpc_mlag_summary_rows.append([hostname, key.replace("_", " ").title(), value, source])
            for m in vpc_mlag.get("members", []):
                vpc_mlag_member_rows.append([
                    hostname, m.get("id", ""), m.get("port", ""), m.get("status", ""),
                    m.get("consistency", ""), m.get("vlans", ""), source,
                ])

            for p in result.get("poe", []):
                poe_rows.append([
                    hostname, p.get("interface", ""), p.get("admin_status", ""), p.get("oper_status", ""),
                    p.get("power_watts", ""), p.get("device", ""), p.get("class", ""), p.get("max_watts", ""),
                    p.get("source", ""),
                ])

            for v in result.get("vlans", []):
                vlan_rows.append([hostname, v.get("vlan_id", ""), v.get("name", "")])

            for n in result.get("neighbors", []):
                neighbor_rows.append([
                    hostname, n.get("protocol", ""), n.get("neighbor_id", ""), n.get("local_interface", ""),
                    n.get("remote_interface", ""), n.get("platform", ""),
                ])

            for r in result.get("routes", []):
                route_rows.append([
                    hostname, r.get("destination", ""), r.get("mask", ""), r.get("next_hop", ""),
                    r.get("distance", ""), r.get("track_id", ""), r.get("vrf", ""), r.get("interface", ""),
                    r.get("description", ""),
                ])

            l3_data = result.get("l3") or {}
            for item in l3_data.get("interfaces", []):
                acl_text = ", ".join(
                    [name for name in (item.get("acl_in"), item.get("acl_out")) if name]
                    + [f"{p['policy']} ({p['direction']})" for p in item.get("traffic_policies", [])]
                    + [f"{f['acl']} ({f['direction']})" for f in item.get("traffic_filters", [])]
                )
                l3_interface_rows.append([
                    hostname, item.get("name", ""), item.get("kind", ""), item.get("address", ""),
                    ", ".join(item.get("secondary", [])), item.get("network", ""), item.get("vrf", ""),
                    item.get("status", ""), "Yes" if item.get("shutdown") else "", item.get("virtual_ip", ""),
                    ", ".join(item.get("helpers", [])), acl_text, item.get("ospf_area", ""), item.get("description", ""),
                ])
            for item in l3_data.get("default_routes", []):
                l3_default_rows.append([
                    hostname, item.get("next_hop", ""), item.get("via", ""), item.get("vrf", ""),
                    item.get("distance") if item.get("distance") is not None else "", item.get("description", ""),
                ])
            l3_ospf = l3_data.get("ospf") or {}
            for proc in l3_ospf.get("processes", []):
                l3_ospf_rows.append([
                    hostname, "Process", proc.get("id", ""), proc.get("vrf", ""), proc.get("router_id", ""),
                    ", ".join(proc.get("areas", [])), proc.get("interface_count", 0),
                    "; ".join(f"{n.get('cidr', '')} area {n.get('area', '')}" for n in proc.get("networks", [])),
                    "all" if proc.get("passive_default") else ", ".join(proc.get("passive_interfaces", [])),
                    "; ".join(proc.get("redistribute", [])),
                ])
            for row in l3_ospf.get("interfaces", []):
                l3_ospf_rows.append([
                    hostname, "Interface", row.get("process", ""), "", "", row.get("area", ""), "",
                    f"{row.get('interface', '')} {row.get('ip', '')}", "yes" if row.get("passive") else "no", row.get("how", ""),
                ])
            for item in l3_ospf.get("findings", []):
                l3_ospf_rows.append([hostname, "Finding", "", "", "", "", "", item.get("text", ""), "", ""])
            l3_acls = l3_data.get("acls") or {}
            for acl in l3_acls.get("definitions", []):
                l3_acl_rows.append([
                    hostname, acl.get("name", ""), acl.get("family", ""), acl.get("kind", ""), acl.get("rules", 0),
                    "Yes" if acl.get("applied") else "No", "; ".join(acl.get("referenced_by", [])),
                    "Yes" if acl.get("system") else "", "\n".join(acl.get("rule_lines", [])),
                ])
            for binding in l3_acls.get("bindings", []):
                l3_acl_binding_rows.append([
                    hostname, binding.get("target", ""), binding.get("acl", ""), binding.get("direction", ""),
                    binding.get("via", ""), "Yes" if binding.get("defined") else "No",
                ])
            for item in l3_data.get("findings", []):
                l3_finding_rows.append([hostname, item.get("level", ""), item.get("text", "")])

            pbr = result.get("pbr") or {}
            for s in pbr.get("sla_tests", []):
                pbr_sla_rows.append([
                    hostname, s.get("kind", ""), s.get("id", ""), s.get("test_type", ""),
                    s.get("destination", ""), s.get("frequency", ""), s.get("timeout", ""),
                    s.get("probe_count", ""),
                ])
            for t in pbr.get("track_objects", []):
                pbr_track_rows.append([hostname, t.get("track_id", ""), t.get("tracks", "")])
            for p in pbr.get("pbr_policies", []):
                pbr_policy_rows.append([
                    hostname, p.get("vendor_mechanism", ""), p.get("policy_name", ""),
                    p.get("sequence", ""), p.get("action", ""), p.get("match", ""),
                    p.get("set_next_hop", ""), p.get("tracked_by", ""),
                ])
            for b in pbr.get("pbr_bindings", []):
                pbr_binding_rows.append([
                    hostname, b.get("interface", ""), b.get("policy_name", ""),
                    b.get("direction", ""), b.get("mechanism", ""),
                ])
            for e in pbr.get("eem_dynamic_pbr", []):
                pbr_eem_rows.append([
                    hostname, e.get("applet_name", ""), e.get("route_map_name", ""), e.get("action", ""),
                    ", ".join(e.get("target_interfaces", [])), e.get("trigger_description", ""),
                ])

            for r in result.get("routing_table", []):
                routing_table_rows.append([hostname, r.get("protocol", ""), r.get("destination", ""), r.get("next_hop", ""), r.get("interface", ""),
                                           r.get("vrf", ""), r.get("distance", ""), r.get("metric", ""), r.get("age", ""), r.get("detail", "")])

            analysis_sheets.collect(result, hostname, analysis_rows)

            for a in result.get("arp_table", []):
                arp_rows.append([
                    hostname, a.get("ip_address", ""), a.get("mac_address", ""), a.get("age", ""),
                    a.get("interface", ""), a.get("type", ""), a.get("vlan", ""), a.get("state", ""),
                ])

            for m in result.get("mac_table", []):
                mac_rows.append([
                    hostname, m.get("vlan", ""), m.get("mac_address", ""), m.get("type", ""),
                    m.get("interface", ""), m.get("age", ""), "Yes" if m.get("system") else "",
                ])

            dns = result.get("dns") or {}
            snmp = result.get("snmp") or {}
            snmp_summary = snmp.get("summary") or {}
            dns_snmp_rows.append([
                hostname, ", ".join(dns.get("servers", [])), dns.get("domain", ""),
                "Yes" if snmp.get("enabled") else "No",
                ", ".join(snmp_summary.get("communities", [])),
                ", ".join(snmp_summary.get("locations", [])),
                ", ".join(snmp_summary.get("hosts", [])),
            ])

            dhcp = result.get("dhcp") or {}
            dhcp_snooping = dhcp.get("snooping") or {}
            dhcp_summary_rows.append([
                hostname,
                "Yes" if dhcp_snooping.get("enabled") else "No",
                ", ".join(dhcp_snooping.get("vlans", [])),
                "Yes" if dhcp_snooping.get("option82_enabled") else "No",
                len(dhcp.get("pools", [])),
                len(dhcp.get("helper_addresses", [])),
            ])
            for pool in dhcp.get("pools", []):
                dhcp_pool_rows.append([
                    hostname, pool.get("name", ""), pool.get("network", ""), pool.get("mask", ""),
                    pool.get("gateway", ""), pool.get("dns_servers", ""), pool.get("domain_name", ""),
                    pool.get("lease", ""), "; ".join(pool.get("options", [])),
                ])
            for helper in dhcp.get("helper_addresses", []):
                dhcp_helper_rows.append([hostname, helper.get("interface", ""), ", ".join(helper.get("servers", []))])
            trusted = set(dhcp_snooping.get("trusted_interfaces", []))
            rate_by_interface = {item.get("interface"): item.get("limit_pps") for item in dhcp_snooping.get("rate_limited_interfaces", [])}
            for interface_name in trusted | set(rate_by_interface):
                dhcp_snooping_rows.append([
                    hostname, interface_name,
                    "Trusted" if interface_name in trusted else "Untrusted",
                    rate_by_interface.get(interface_name, ""),
                ])

        self.add_sheet(workbook, "Device Info",
            ["Hostname", "Vendor", "Device Type", "OS Version", "Model", "Serial Number", "Stack Mode",
             "Stack Member Count", "Total Interfaces", "Interfaces Up", "Interfaces Down",
             "Interfaces Admin Down", "Interfaces Unknown"],
            device_info_rows)
        self._add_analysis_sheets(workbook, analysis_rows, "front")
        self.add_sheet(workbook, "Serial Numbers", ["Hostname", "Slot", "Kind", "Description", "PID", "Serial Number"], serial_rows)
        self.add_sheet(workbook, "Stack Members", ["Hostname", "Slot", "Role", "MAC Address", "State", "Priority", "Serial Number"], stack_rows)
        self.add_sheet(workbook, "Interfaces",
            ["Hostname", "Name", "Status", "Protocol", "Mode", "VLAN", "Allowed VLANs", "IP Address", "Description", "Port Security", "HSRP/VRRP IP",
             "Status Detail", "Status Source", "Speed", "Duplex", "Port Type", "Link Since", "Link Changes", "RX Errors", "TX Drops",
             "In %", "Out %", "MAC Address", "MTU", "Port-channel", "Bundle State", "Bundle Members"],
            interface_rows)
        self.add_sheet(workbook, "Transceivers",
            ["Hostname", "Interface", "Status", "Present", "Type", "Vendor Part Number", "Serial Number"],
            transceiver_rows)
        self.add_sheet(workbook, "Port-Channels",
            ["Hostname", "Name", "Protocol", "Status", "Description", "Mode", "VLAN", "Allowed VLANs", "IP Address", "Member Ports"],
            port_channel_rows)
        self.add_sheet(workbook, "VPC-MLAG Summary", ["Hostname", "Field", "Value", "Source"], vpc_mlag_summary_rows)
        self.add_sheet(workbook, "VPC-MLAG Members", ["Hostname", "ID", "Port", "Status", "Consistency", "VLANs", "Source"], vpc_mlag_member_rows)
        self.add_sheet(workbook, "PoE",
            ["Hostname", "Interface", "Admin Status", "Oper Status", "Power (W)", "Device", "Class", "Max (W)", "Source"],
            poe_rows)
        self.add_sheet(workbook, "VLANs", ["Hostname", "VLAN ID", "Name"], vlan_rows)
        self._add_analysis_sheets(workbook, analysis_rows, "layer2")
        self.add_sheet(workbook, "Neighbors",
            ["Hostname", "Protocol", "Neighbor", "Local Interface", "Remote Interface", "Platform"],
            neighbor_rows)
        self.add_sheet(workbook, "Static Routes",
            ["Hostname", "Destination", "Mask", "Next Hop", "Admin Distance", "Track ID", "VRF", "Exit Interface", "Description"], route_rows)
        self.add_sheet(workbook, "L3 Interfaces",
            ["Hostname", "Interface", "Kind", "Address", "Secondary", "Network", "VRF", "State", "Shutdown", "HSRP/VRRP IP",
             "DHCP Relay", "ACL / Policy", "OSPF Area", "Description"], l3_interface_rows)
        self.add_sheet(workbook, "L3 Default Route", ["Hostname", "Next Hop", "Configured By", "VRF", "Distance", "Description"], l3_default_rows)
        self.add_sheet(workbook, "L3 OSPF",
            ["Hostname", "Row", "Process", "VRF", "Router ID", "Area(s)", "Interfaces", "Networks / Interface", "Passive", "Redistribute / How"], l3_ospf_rows)
        self.add_sheet(workbook, "L3 ACLs",
            ["Hostname", "Name", "Family", "Kind", "Rules", "Applied", "Referenced By", "System", "Rule Text"], l3_acl_rows)
        self.add_sheet(workbook, "L3 ACL Bindings", ["Hostname", "Applied To", "ACL", "Direction", "Via", "ACL Defined"], l3_acl_binding_rows)
        self.add_sheet(workbook, "L3 Findings", ["Hostname", "Level", "Finding"], l3_finding_rows)
        self.add_sheet(workbook, "Routing Table",
            ["Hostname", "Protocol", "Destination", "Next Hop", "Interface", "VRF", "Distance", "Metric", "Age", "Detail"], routing_table_rows)
        self._add_analysis_sheets(workbook, analysis_rows, "layer3")
        self.add_sheet(workbook, "PBR SLA-NQA Tests",
            ["Hostname", "Kind", "Test", "Type", "Destination", "Frequency (s)", "Timeout (s)", "Probe Count"],
            pbr_sla_rows)
        self.add_sheet(workbook, "PBR Track Objects", ["Hostname", "Track ID", "Tracks"], pbr_track_rows)
        self.add_sheet(workbook, "PBR Policies",
            ["Hostname", "Mechanism", "Policy Name", "Sequence/Precedence", "Action", "Match", "Set Next Hop", "Tracked By"],
            pbr_policy_rows)
        self.add_sheet(workbook, "PBR Interface Bindings",
            ["Hostname", "Interface", "Policy Name", "Direction", "Mechanism"], pbr_binding_rows)
        self.add_sheet(workbook, "PBR EEM Dynamic (Cisco)",
            ["Hostname", "Applet Name", "Route-Map", "Action", "Target Interfaces", "Trigger"], pbr_eem_rows)
        self.add_sheet(workbook, "ARP Table",
            ["Hostname", "IP Address", "MAC Address", "Age", "Interface", "Type", "VLAN", "State"], arp_rows)
        self.add_sheet(workbook, "MAC Table",
            ["Hostname", "VLAN", "MAC Address", "Type", "Interface", "Age", "Switch Own Entry"], mac_rows)
        self.add_sheet(workbook, "DNS-SNMP",
            ["Hostname", "DNS Servers", "Domain Name", "SNMP Configured", "SNMP Communities", "SNMP Locations", "SNMP Trap Hosts"],
            dns_snmp_rows)
        self.add_sheet(workbook, "DHCP Summary",
            ["Hostname", "Snooping Enabled", "Snooping VLANs", "Option 82 Enabled", "Pool Count", "Relay Interface Count"],
            dhcp_summary_rows)
        self.add_sheet(workbook, "DHCP Pools",
            ["Hostname", "Pool Name", "Network", "Mask", "Gateway", "DNS Servers", "Domain", "Lease", "Options"],
            dhcp_pool_rows)
        self.add_sheet(workbook, "DHCP Relay-Helper",
            ["Hostname", "Interface", "Relay/Helper Servers"], dhcp_helper_rows)
        self.add_sheet(workbook, "DHCP Snooping Interfaces",
            ["Hostname", "Interface", "Trust State", "Rate Limit (pps)"], dhcp_snooping_rows)
        self._add_analysis_sheets(workbook, analysis_rows, "logs")

        self._add_sizing_sheet(workbook)

        return workbook

    def _add_analysis_sheets(self, workbook: Workbook, rows: dict[str, list[list[Any]]], group: str) -> None:
        for title in analysis_sheets.GROUPS[group]:
            self.add_sheet(workbook, title, analysis_sheets.HEADERS[title], rows[title])

    def _add_sizing_sheet(self, workbook: Workbook) -> None:
        """"Sizing Summary" — one row per device plus an aggregate
        totals row and a port-speed breakdown, built from sizing.py's
        build_sizing_summary(). Placed first (index 0) so it's the
        workbook's landing page, same convention as the Comparison
        sheet in export_comparison().
        """
        summary = build_sizing_summary(self.results)
        devices = summary["devices"]
        totals = summary["totals"]

        sheet = workbook.create_sheet("Sizing Summary", 0)
        headers = [
            "Hostname", "Vendor", "Model", "Serial Number", "OS Version", "Physical Ports",
            "Ports Up", "Ports Down", "Ports Admin Down", "Ports Unknown", "Trunk Ports", "Port-Channels",
            "PoE Ports", "PoE Total (W)", "PoE Source", "Transceivers",
            "VLANs", "Static Routes",
            "OSPF Configured", "vPC/M-LAG Configured", "vPC/M-LAG Source",
            "vPC/M-LAG Members", "HSRP/VRRP SVIs", "Stack Mode", "Stack Members",
        ]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)

        def yes_no(value: bool) -> str:
            return "Yes" if value else "No"

        for d in devices:
            sheet.append([
                d["hostname"], d["vendor"], d["model"], d["serial_number"], d["os_version"],
                d["physical_port_count"], d["physical_ports_up"], d["physical_ports_down"],
                d["physical_ports_admin_down"], d["physical_ports_unknown"],
                d["trunk_port_count"], d["port_channel_count"],
                d["poe_port_count"], d["poe_total_watts"], d["poe_source"],
                d["transceiver_count"],
                d["vlan_count"], d["static_route_count"],
                yes_no(d["ospf_configured"]), yes_no(d["vpc_mlag_configured"]), d["vpc_mlag_source"],
                d["vpc_mlag_member_count"], d["hsrp_vrrp_svi_count"],
                d["stack_mode"], d["stack_member_count"],
            ])

        blank_row = [""] * len(headers)
        sheet.append(blank_row)
        sheet.append([
            f"TOTALS ({totals['device_count']} device(s))", "", "", "",
            totals["physical_port_count"], totals["physical_ports_up"], totals["physical_ports_down"],
            totals["physical_ports_admin_down"], totals["physical_ports_unknown"],
            totals["trunk_port_count"], totals["port_channel_count"],
            totals["poe_port_count"], totals["poe_total_watts"], "",
            totals["transceiver_count"],
            totals["unique_vlan_count"], totals["static_route_count"],
            totals["ospf_device_count"], totals["vpc_mlag_device_count"], "",
            "", totals["hsrp_vrrp_svi_count"],
            f"{totals['stacked_device_count']} stacked", "",
        ])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)

        sheet.append(blank_row)
        sheet.append(["Note: VLANs total above is the unique VLAN-ID count across all devices (VLANs commonly repeat across switches at the same site, so a plain sum would double-count)."])
        sheet.append(blank_row)

        sheet.append(["Physical Port Count by Speed"])
        sheet[sheet.max_row][0].font = Font(bold=True)
        sheet.append(["Type", "Count"])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        for label, count in sorted(totals["port_type_totals"].items()):
            sheet.append([label, count])

        sheet.append(blank_row)
        sheet.append(["Transceiver Count by Type"])
        sheet[sheet.max_row][0].font = Font(bold=True)
        sheet.append(["Note: vendor-specific abbreviations for the same physical SFP/QSFP class are merged into one row (e.g. ArubaOS-Switch's \"1000SX\" and Cisco's \"1000BaseSX SFP\" are the same part) — see sizing.py for the confirmed aliases; anything not confirmed is kept exactly as the device reported it. Broken down by Vendor Part Number too, not just Type — the same media class can still be a different orderable SKU (a genuine Cisco GLC-SX-MMD vs. a third-party/compatible optic), so this is what to actually check against before ordering replacements."])
        sheet.append(["Type", "Vendor Part Number", "Count"])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        for row in totals["transceiver_part_totals"]:
            sheet.append([row["type"], row["vendor_part_number"], row["count"]])

        for column in sheet.columns:
            max_length = max((len(str(cell.value or "")) for cell in column), default=0)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max_length + 2, 42)
        sheet.freeze_panes = "A2"


def export_comparison(path, result_a: dict[str, Any], result_b: dict[str, Any], comparison: dict[str, Any]) -> None:
    """Two-device comparison export: the full per-device data (same
    sheets as a regular export, both devices combined via the Hostname
    column) plus one extra "Comparison" sheet up front summarizing the
    parity check — device info side-by-side, then VLAN/route/DNS-server
    presence across both devices, for the migration sign-off use case.
    """
    exporter = SwitchExcelExporter([result_a, result_b])
    workbook = exporter.build_workbook()

    device_a = comparison.get("device_a", {})
    device_b = comparison.get("device_b", {})
    field_labels = [
        ("hostname", "Hostname"), ("vendor", "Vendor"), ("model", "Model"),
        ("serial_number", "Serial Number"),
        ("os_version", "OS Version"), ("stack_mode", "Stack Mode"),
        ("stack_member_count", "Stack Member Count"),
        ("total_interfaces", "Total Interfaces"), ("interfaces_up", "Interfaces Up"),
        ("interfaces_down", "Interfaces Down"), ("interfaces_admin_down", "Interfaces Admin Down"),
        ("interfaces_unknown", "Interfaces Unknown"), ("arp_entries", "ARP Entries"),
        ("mac_entries", "MAC Entries"), ("total_vlans", "Total VLANs"),
        ("total_routes", "Total Routes"), ("neighbor_count", "Neighbor Count"),
        ("port_channel_count", "Port-Channel Count"), ("transceiver_count", "Transceiver Count"),
        ("poe_port_count", "PoE Port Count"),
        ("vpc_mlag_configured", "vPC / M-LAG?"),
        ("has_custom_routing", "Has PBR / SLA-Track / NQA?"),
    ]
    _VPC_MLAG_SOURCE_LABEL = {"live": "Yes (live)", "config": "Yes (config only)"}

    def _comparison_value(key: str, value: Any, device: dict[str, Any] | None = None) -> Any:
        if key == "has_custom_routing":
            return "Yes" if value else "No"
        if key == "vpc_mlag_configured":
            return _VPC_MLAG_SOURCE_LABEL.get((device or {}).get("vpc_mlag_source", ""), "No") if value else "No"
        return value

    comparison_rows = [
        [label, _comparison_value(key, device_a.get(key, ""), device_a), _comparison_value(key, device_b.get(key, ""), device_b)]
        for key, label in field_labels
    ]
    comparison_rows.append(["", "", ""])

    def diff_rows(title: str, diff: dict[str, list]) -> list[list]:
        rows = [[title, "", ""]]
        for value in diff.get("only_a", []):
            rows.append([value, "Only Device A", ""])
        for value in diff.get("only_b", []):
            rows.append([value, "", "Only Device B"])
        for value in diff.get("both", []):
            rows.append([value, "Both", "Both"])
        rows.append(["", "", ""])
        return rows

    comparison_rows += diff_rows("VLANs", comparison.get("vlans", {}))
    comparison_rows += diff_rows("Routes (destination)", comparison.get("routes", {}))
    comparison_rows += diff_rows("DNS Servers", comparison.get("dns_servers", {}))
    comparison_rows += diff_rows("NTP Servers", comparison.get("ntp_servers", {}))
    comparison_rows += diff_rows("Syslog Servers", comparison.get("syslog_servers", {}))
    comparison_rows += diff_rows("TACACS+ Servers", comparison.get("tacacs_servers", {}))
    comparison_rows += diff_rows("RADIUS Servers", comparison.get("radius_servers", {}))
    comparison_rows += diff_rows("SNMP Trap / Informs Hosts", comparison.get("snmp_trap_hosts", {}))
    comparison_rows += diff_rows("CDP/LLDP Neighbor Identity", comparison.get("neighbors", {}))
    comparison_rows += diff_rows("MAC addresses (learned)", comparison.get("mac_addresses", {}))
    comparison_rows += diff_rows("ARP hosts (IP = MAC)", comparison.get("arp_hosts", {}))

    sheet = workbook.create_sheet("Comparison", 0)  # first sheet
    sheet.append(["Field / Value", "Device A", "Device B"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in comparison_rows:
        sheet.append(row)
    for column in sheet.columns:
        # A stack's serial cell holds one line per member; widths are sized to
        # the longest line, not to the whole multi-line string.
        max_length = max((max((len(line) for line in str(cell.value or "").splitlines()), default=0) for cell in column), default=0)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max_length + 2, 60)
    for sheet_row in sheet.iter_rows(min_row=2):
        for cell in sheet_row:
            if isinstance(cell.value, str) and "\n" in cell.value:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"

    _add_l3_comparison_sheet(workbook, comparison.get("l3") or {})
    _add_fhrp_pair_sheet(workbook, comparison.get("fhrp") or {}, device_a.get("hostname") or "A", device_b.get("hostname") or "B")
    _add_ospf_pair_sheet(workbook, comparison.get("ospf_live") or {}, device_a.get("hostname") or "A", device_b.get("hostname") or "B")
    _add_stp_pair_sheet(workbook, comparison.get("spanning_tree") or {}, device_a.get("hostname") or "A", device_b.get("hostname") or "B")
    _add_vpc_mlag_pair_sheet(workbook, comparison.get("vpc_mlag") or {}, device_a.get("hostname") or "A", device_b.get("hostname") or "B")
    workbook.save(path)


def _add_stp_pair_sheet(workbook: Workbook, stp: dict[str, Any], name_a: str, name_b: str) -> None:
    """Sheet "STP Compare": protocol and root-bridge-visibility findings, then one row per VLAN
    where the two captures disagree about (or both claim) the root — empty when none do."""
    if not stp.get("captured") or not (stp.get("findings") or stp.get("root_mismatch")):
        return
    sheet = workbook.create_sheet("STP Compare")
    proto = stp.get("protocol") or {}
    counts = stp.get("counts") or {}
    sheet.append([f"A = {name_a}", f"B = {name_b}",
                   f"Protocol: {proto.get('a', '—')} / {proto.get('b', '—')}",
                   f"Blocking: {counts.get('a_blocking', 0)} / {counts.get('b_blocking', 0)}",
                   f"Inconsistent: {counts.get('a_broken', 0)} / {counts.get('b_broken', 0)}"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    sheet.append(["Findings"])
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, size=12)
    for item in stp.get("findings", []):
        sheet.append([item["level"], item["text"]])
    sheet.append([])
    sheet.append(["VLAN", "A root bridge", "B root bridge"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for row in stp.get("root_mismatch", []):
        sheet.append([row["vlan"], row["a_root"], row["b_root"]])
    for column in sheet.columns:
        longest = max((len(str(cell.value or "")) for cell in column), default=0)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(longest + 2, 90)
    sheet.freeze_panes = "A2"


def _add_vpc_mlag_pair_sheet(workbook: Workbook, vpc: dict[str, Any], name_a: str, name_b: str) -> None:
    """Sheet "vPC-MLAG Pair Check": domain/role/consistency checks, problems first, then one row
    per vPC / Eth-Trunk ID with its up/consistency state on each side."""
    rows, member_rows = vpc.get("rows") or [], vpc.get("member_rows") or []
    if not rows and not member_rows:
        return
    sheet = workbook.create_sheet("vPC-MLAG Pair Check")
    sheet.append([f"A = {name_a}", f"B = {name_b}"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    sheet.append(["Check", "Verdict", "Device A", "Device B", "What Differs"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for row in sorted(rows, key=lambda item: item.get("verdict") in {"Same", "Healthy split", "Different priority"}):
        sheet.append([row.get("check"), row.get("verdict"), row.get("a"), row.get("b"), row.get("reason")])
    if member_rows:
        sheet.append([])
        sheet.append(["ID", "Port", "Verdict", "A Status", "B Status", "A Consistency", "B Consistency", "What Differs"])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        for row in sorted(member_rows, key=lambda item: item.get("verdict") == "Up on both"):
            sheet.append([row.get("id"), row.get("port"), row.get("verdict"), row.get("a_status"), row.get("b_status"),
                          row.get("a_consistency"), row.get("b_consistency"), row.get("reason")])
    for column in sheet.columns:
        longest = max((len(str(cell.value or "")) for cell in column), default=0)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(longest + 2, 70)
    sheet.freeze_panes = "A3"


def _add_ospf_pair_sheet(workbook: Workbook, pair: dict[str, Any], name_a: str, name_b: str) -> None:
    """Sheet "OSPF LSDB Compare": the LSA totals of the two link-state databases per process, area and LSA type, with the verdict (problems first)."""
    rows = pair.get("rows") or []
    if not rows:
        return
    index = 1
    for name in ("L3 Comparison", "HSRP Pair Check"):
        if name in workbook.sheetnames:
            index = workbook.sheetnames.index(name) + 1
    sheet = workbook.create_sheet("OSPF LSDB Compare", index)
    cut = pair.get("incomplete") or []
    sheet.append([f"A = {name_a}", f"B = {name_b}", ("Pager cut the database of " + " and ".join(cut) + ": counts are not compared") if cut else "Both databases complete"])
    sheet.append(["Verdict", "Process", "VRF", "Area", "LSA Type", "A LSAs", "B LSAs", "A Oldest (s)", "B Oldest (s)", "What Differs"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for row in sorted(rows, key=lambda item: item.get("verdict") == "Same"):
        sheet.append([row.get("verdict"), row.get("process"), row.get("vrf"), row.get("area"), row.get("type"), row.get("a_count"), row.get("b_count"),
                      row.get("a_oldest"), row.get("b_oldest"), row.get("reason")])
    for column in sheet.columns:
        longest = max((len(str(cell.value or "")) for cell in column[1:]), default=0)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(longest + 2, 70)
    sheet.freeze_panes = "A3"


def _add_fhrp_pair_sheet(workbook: Workbook, pair: dict[str, Any], name_a: str, name_b: str) -> None:
    """Sheet "HSRP Pair Check": one row per HSRP group of the two devices with the verdict, problems first."""
    rows = pair.get("rows") or []
    if not rows:
        return
    sheet = workbook.create_sheet("HSRP Pair Check", 2 if "L3 Comparison" in workbook.sheetnames else 1)
    sheet.append([f"A = {name_a}", f"B = {name_b}", f"{pair.get('active_a', 0)} groups Active on A", f"{pair.get('active_b', 0)} groups Active on B"])
    sheet.append(["Verdict", "Interface", "Group", "Virtual IP", "A IP Address", "B IP Address", "A State", "B State", "A Priority", "B Priority", "A Preempt", "B Preempt", "What Differs"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for row in sorted(rows, key=lambda item: item.get("verdict") == "Healthy pair"):
        sheet.append([row.get("verdict"), row.get("interface"), row.get("group"), row.get("virtual_ip"), row.get("a_ip"), row.get("b_ip"),
                      row.get("a_state"), row.get("b_state"), row.get("a_priority"), row.get("b_priority"),
                      row.get("a_preempt"), row.get("b_preempt"), row.get("reason")])
    for column in sheet.columns:
        longest = max((len(str(cell.value or "")) for cell in column[1:]), default=0)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(longest + 2, 70)
    sheet.freeze_panes = "A3"


def _add_l3_comparison_sheet(workbook: Workbook, l3_cmp: dict[str, Any]) -> None:
    """Sheet "L3 Comparison": findings first, then each Layer-3 block as a table
    (gateways, default route, static routes, OSPF, ACLs, DHCP snooping)."""
    if not l3_cmp:
        return
    sheet = workbook.create_sheet("L3 Comparison", 1)
    bold = Font(bold=True)

    def block(title: str, headers: list[str], rows: list[list[Any]]) -> None:
        sheet.append([title])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, size=12)
        sheet.append(headers)
        for cell in sheet[sheet.max_row]:
            if cell.value is not None:
                cell.font = bold
        for row in rows:
            sheet.append(row)
        if not rows:
            sheet.append(["(none)"])
        sheet.append([])

    block("Findings", ["Level", "Finding"], [[f["level"], f["text"]] for f in l3_cmp.get("findings", [])])
    block("Layer-3 interfaces", ["Interface", "Kind", "Result", "A name", "A address", "B name", "B address",
                                  "A FHRP/VRRP VIP", "B FHRP/VRRP VIP", "A DHCP relay", "B DHCP relay", "A ACL/policy", "B ACL/policy"], [
        [r["label"], r["kind"], r["status"], r["a"]["name"], r["a"]["address"], r["b"]["name"], r["b"]["address"],
         r["a"]["virtual_ip"], r["b"]["virtual_ip"], ", ".join(r["a"]["helpers"]), ", ".join(r["b"]["helpers"]), r["a"]["acl"], r["b"]["acl"]]
        for r in l3_cmp.get("interfaces", [])
    ])
    default = l3_cmp.get("default_route") or {}
    block("Default route", ["Result", "A next hop", "A configured by", "B next hop", "B configured by"], [[
        default.get("status", ""), ", ".join((default.get("a") or {}).get("next_hops", [])), (default.get("a") or {}).get("via", ""),
        ", ".join((default.get("b") or {}).get("next_hops", [])), (default.get("b") or {}).get("via", ""),
    ]])
    block("Static routes", ["VRF", "Destination", "Result", "A next hop", "A distance", "B next hop", "B distance"], [
        [r["vrf"], r["cidr"], r["status"], r["a_next_hop"], r["a_distance"], r["b_next_hop"], r["b_distance"]]
        for r in l3_cmp.get("static_routes", [])
    ])
    ospf = l3_cmp.get("ospf") or {}
    oa, ob = ospf.get("a") or {}, ospf.get("b") or {}
    block("OSPF", ["Item", "Device A", "Device B"], [
        ["State", oa.get("state", ""), ob.get("state", "")],
        ["Process", oa.get("process", ""), ob.get("process", "")],
        ["Router ID", oa.get("router_id", ""), ob.get("router_id", "")],
        ["Areas", ", ".join(oa.get("areas", [])), ", ".join(ob.get("areas", []))],
        ["Networks", "\n".join(oa.get("networks", [])), "\n".join(ob.get("networks", []))],
        ["Interfaces running OSPF", "\n".join(oa.get("interfaces", [])), "\n".join(ob.get("interfaces", []))],
        ["Redistribute", "; ".join(oa.get("redistribute", [])), "; ".join(ob.get("redistribute", []))],
    ])
    block("ACLs", ["ACL", "Result", "System", "A rules", "A applied to", "B rules", "B applied to"], [
        [r["name"], r["status"], "Yes" if r["system"] else "", r["a_rules"], r["a_applied"], r["b_rules"], r["b_applied"]]
        for r in l3_cmp.get("acls", [])
    ])
    snoop = l3_cmp.get("dhcp_snooping") or {}
    block("DHCP snooping", ["Item", "Device A", "Device B"], [
        ["Enabled", "Yes" if snoop.get("a_enabled") else "No", "Yes" if snoop.get("b_enabled") else "No"],
        ["VLANs", snoop.get("a_vlans", ""), snoop.get("b_vlans", "")],
        ["Trusted interfaces", ", ".join(snoop.get("a_trusted", [])), ", ".join(snoop.get("b_trusted", []))],
    ])
    for column in sheet.columns:
        max_length = max((max((len(line) for line in str(cell.value or "").splitlines()), default=0) for cell in column), default=0)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max_length + 2, 60)
    for sheet_row in sheet.iter_rows():
        for cell in sheet_row:
            if isinstance(cell.value, str) and "\n" in cell.value:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

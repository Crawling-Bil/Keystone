"""Device inventory for an LLD project.

One row per device: Site, Hostname, Model, Role, Mgmt IP, SN, Config. The
Config column names an uploaded file (matched case-insensitively, with or
without extension); it may be blank for a design written before anything
is configured. Each device also gets a technology (switch, wireless,
firewall, router, loadbalancer) from, in order: the detected domain of its
config, keywords in Role, then the model name.
"""

from __future__ import annotations

import csv
import io
import ipaddress
import re
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any, Iterable

COLUMNS = ["Site", "Hostname", "Model", "Role", "Mgmt IP", "SN", "Config", "Rack", "Rack Size", "U", "U Height", "Status"]
STATUSES = ("New", "Existing", "Replace")
STANDARD_RACK_SIZES = (24, 27, 32, 36, 42, 45, 47, 48)
DEFAULT_RACK_SIZE = 42
TECHNOLOGIES = ["switch", "wireless", "firewall", "router", "loadbalancer"]
TECH_LABELS = {
    "switch": "Switch", "wireless": "Wireless", "firewall": "Firewall",
    "router": "Router", "loadbalancer": "Load Balancer",
}

_HEADER_ALIASES = {
    "site": "site", "site name": "site", "lokasi": "site", "location": "site",
    "hostname": "hostname", "host": "hostname", "device name": "hostname", "nama perangkat": "hostname",
    "model": "model", "type": "model", "device model": "model", "pid": "model",
    "role": "role", "function": "role", "fungsi": "role",
    "mgmt ip": "mgmt_ip", "management ip": "mgmt_ip", "ip management": "mgmt_ip", "ip mgmt": "mgmt_ip", "ip": "mgmt_ip",
    "sn": "serial", "serial": "serial", "serial number": "serial", "s/n": "serial",
    "config": "config", "config file": "config", "file": "config", "configuration": "config",
    "technology": "technology", "teknologi": "technology",
    "status": "status", "device status": "status", "kondisi": "status", "existing/new": "status",
    "rack": "rack", "rack name": "rack", "nama rack": "rack", "cabinet": "rack",
    "rack size": "rack_size", "rack u": "rack_size", "rack height": "rack_size", "ukuran rack": "rack_size",
    "u": "u", "u position": "u", "rack position": "u", "posisi u": "u", "start u": "u",
    "u height": "u_height", "height": "u_height", "u size": "u_height", "tinggi u": "u_height", "ru": "u_height",
}

_ROLE_HINTS = [
    ("loadbalancer", r"\b(lb|load ?balancer|f5|big-?ip|ltm|gtm|adc)\b"),
    ("firewall", r"\b(fw|firewall|ngfw|palo|panorama|forti(gate)?|asa)\b"),
    ("wireless", r"\b(wlc|ap|access ?point|wireless|controller|wifi|wi-fi)\b"),
    ("router", r"\b(router|rtr|wan|cedge|vedge|sd-?wan edge|ce|pe)\b"),
    ("switch", r"\b(switch|sw|core|distribution|dist|access|tor|leaf|spine|agg)\b"),
]
_MODEL_HINTS = [
    ("firewall", r"^(PA-|VM-\d|FG-|FORTIGATE|FGT|ASA)"),
    ("loadbalancer", r"(BIG-?IP|^I\d{4}|^R\d{4}|VIPRION|^F5)"),
    ("wireless", r"(C98\d\d|AIR-|^AP|AIRENGINE|^AC\d{4}|9115|9120|9130|9136|9166|^C91\d\d)"),
    ("router", r"^(ISR|ASR|C8\d{3}|C1\d{3}|NCS|AR\d{3,4}|NE\d|CISCO\d{4}|C11\d\d)"),
    ("switch", r"^(C9[2-6]\d\d|WS-C|N[3-9]K|S\d{4}|CE\d{4}|JL\d{3}|J9\d{3}|6[1-4]00|8[1-4]00|EX\d|QFX)"),
]


@dataclass
class Device:
    site: str
    hostname: str
    model: str = ""
    role: str = ""
    mgmt_ip: str = ""
    serial: str = ""
    config: str = ""
    rack: str = ""
    rack_size: int = DEFAULT_RACK_SIZE
    u: int | None = None               # lowest U the device occupies
    u_height: int = 1
    status: str = "New"                # New | Existing | Replace (rack layout colour)
    technology: str = ""
    technology_source: str = ""
    config_name: str = ""          # the uploaded file actually matched
    config_text: str = field(default="", repr=False)
    notes: list = field(default_factory=list)

    @property
    def mgmt_address(self) -> str:
        return self.mgmt_ip.split("/")[0].strip()

    @property
    def mgmt_netmask(self) -> str:
        value = self.mgmt_ip.strip()
        if "/" in value:
            try:
                return str(ipaddress.ip_interface(value).network.netmask)
            except ValueError:
                return ""
        return ""

    def public(self) -> dict[str, Any]:
        data = asdict(self)
        data.pop("config_text", None)
        data["has_config"] = bool(self.config_text)
        return data


def _norm_header(value: Any) -> str:
    return _HEADER_ALIASES.get(re.sub(r"\s+", " ", str(value or "")).strip().lower(), "")


def _rows_from_table(rows: list[list[Any]]) -> list[dict[str, str]]:
    header_index = None
    for i, row in enumerate(rows[:20]):
        keys = [_norm_header(c) for c in row]
        if "hostname" in keys and ("site" in keys or "model" in keys):
            header_index = i
            break
    if header_index is None:
        raise ValueError("Inventory needs a header row with at least Hostname and Site (or Model).")
    keys = [_norm_header(c) for c in rows[header_index]]
    out = []
    for row in rows[header_index + 1:]:
        record = {}
        for key, value in zip(keys, row):
            if key:
                record[key] = "" if value is None else str(value).strip()
        if record.get("hostname"):
            out.append(record)
    return out


def read_inventory_file(name: str, data: bytes) -> list[dict[str, str]]:
    suffix = Path(name).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook["Inventory"] if "Inventory" in workbook.sheetnames else workbook.worksheets[0]
        return _rows_from_table([list(r) for r in sheet.iter_rows(values_only=True)])
    text = data.decode("utf-8-sig", errors="ignore")
    dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else ",", delimiters=",;\t")
    return _rows_from_table([row for row in csv.reader(io.StringIO(text), dialect)])


def tech_from_text(role: str, model: str) -> tuple[str, str]:
    role_l = role.lower()
    for tech, pattern in _ROLE_HINTS:
        if re.search(pattern, role_l):
            return tech, "role"
    model_u = model.upper().strip()
    for tech, pattern in _MODEL_HINTS:
        if re.search(pattern, model_u):
            return tech, "model"
    return "", ""


_DOMAIN_TO_TECH = {"switch": "switch", "firewall": "firewall", "router": "router",
                   "load_balancer": "loadbalancer", "loadbalancer": "loadbalancer", "wireless": "wireless"}


def detect_config_domain(text: str) -> str:
    try:
        from features.configuration_studio.converter_engine.domains import classify

        guess = classify(text)
        if guess.confidence == "high":
            return _DOMAIN_TO_TECH.get(guess.domain, "")
    except Exception:
        pass
    lowered = text.lower()
    if "wlan " in lowered and ("wireless" in lowered or "ap profile" in lowered):
        return "wireless"
    if "ltm virtual" in lowered or "ltm pool" in lowered:
        return "loadbalancer"
    return ""


def _match_config(device: Device, files: dict[str, str]) -> str:
    """Return the uploaded file name for this device, or ''."""
    if not files:
        return ""
    by_lower = {name.lower(): name for name in files}
    by_stem = {Path(name).stem.lower(): name for name in files}
    wanted = device.config.strip().lower()
    if wanted:
        return by_lower.get(wanted) or by_stem.get(Path(wanted).stem.lower()) or ""
    host = device.hostname.lower()
    if host in by_stem:
        return by_stem[host]
    for stem, name in by_stem.items():
        if host and host in stem:
            return name
    # last resort: the hostname printed inside the config
    pattern = re.compile(r"^\s*(?:hostname|sysname|set deviceconfig system hostname|set hostname)\s+\"?" + re.escape(device.hostname) + r"\"?\s*$",
                         re.IGNORECASE | re.MULTILINE)
    for name, text in files.items():
        if pattern.search(text):
            return name
    return ""


def normalize_status(value, device=None) -> str:
    """Existing / New / Replace (new device in the old one's position); blank = New."""
    text = str(value or "").strip().lower()
    if not text:
        return "New"
    if "replac" in text or "pengganti" in text or "ganti" in text:
        return "Replace"
    if text.startswith(("exist", "old", "lama", "eksis")):
        return "Existing"
    if text.startswith(("new", "baru")):
        return "New"
    if device is not None:
        device.notes.append(f"Status '{value}' is not Existing, New or Replace; New is used.")
    return "New"


def _int(value) -> int | None:
    text = str(value or "").strip().upper().rstrip("U").strip()
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return -1


def _place(device: Device, record: dict) -> None:
    """Rack position: Rack, Rack Size (default 42U), U (lowest unit), U Height (default 1)."""
    size = _int(record.get("rack_size"))
    if size is not None:
        if size <= 0 or size > 60:
            device.notes.append(f"Rack Size '{record.get('rack_size')}' is not a number of units; 42U is used.")
        else:
            device.rack_size = size
    height = _int(record.get("u_height"))
    if height is not None:
        if height <= 0 or height > 20:
            device.notes.append(f"U Height '{record.get('u_height')}' is not valid; 1U is used.")
        else:
            device.u_height = height
    u = _int(record.get("u"))
    if u is not None:
        if u <= 0:
            device.notes.append(f"U '{record.get('u')}' is not a rack unit number.")
        else:
            device.u = u
    if device.rack and device.u is None:
        device.notes.append(f"In rack {device.rack} but no U given: not drawn in the rack layout.")
    if device.u is not None and not device.rack:
        device.notes.append("U given but no Rack name: not drawn in the rack layout.")
    if device.rack and device.u is not None and device.u + device.u_height - 1 > device.rack_size:
        device.notes.append(f"U{device.u}-U{device.u + device.u_height - 1} does not fit a {device.rack_size}U rack.")


def check_racks(devices: list) -> None:
    """Note devices that overlap in the same rack (same site and rack name)."""
    by_rack: dict = {}
    for device in devices:
        if device.rack and device.u is not None:
            by_rack.setdefault((device.site, device.rack), []).append(device)
    for members in by_rack.values():
        for i, a in enumerate(members):
            for b in members[i + 1:]:
                a_top, b_top = a.u + a.u_height - 1, b.u + b.u_height - 1
                if a.u <= b_top and b.u <= a_top:
                    for x, y in ((a, b), (b, a)):
                        x.notes.append(f"Overlaps {y.hostname} in {x.rack} (U{max(a.u, b.u)}-U{min(a_top, b_top)}).")


def build_devices(records: Iterable[dict[str, Any]], files: dict[str, str] | None = None) -> list[Device]:
    files = files or {}
    devices: list[Device] = []
    seen = set()
    for record in records:
        hostname = str(record.get("hostname") or "").strip()
        if not hostname:
            continue
        device = Device(
            site=str(record.get("site") or "").strip() or "Site",
            hostname=hostname,
            model=str(record.get("model") or "").strip(),
            role=str(record.get("role") or "").strip(),
            mgmt_ip=str(record.get("mgmt_ip") or "").strip(),
            serial=str(record.get("serial") or "").strip(),
            config=str(record.get("config") or "").strip(),
            rack=str(record.get("rack") or "").strip(),
        )
        _place(device, record)
        device.status = normalize_status(record.get("status"), device)
        if hostname.lower() in seen:
            device.notes.append("Duplicate hostname in the inventory.")
        seen.add(hostname.lower())
        matched = _match_config(device, files)
        if matched:
            device.config_name = matched
            device.config_text = files[matched]
        elif device.config:
            device.notes.append(f"Config file '{device.config}' was not uploaded.")
        forced = str(record.get("technology") or "").strip().lower().replace(" ", "")
        if forced in TECHNOLOGIES:
            device.technology, device.technology_source = forced, "inventory"
        if not device.technology and device.config_text:
            domain = detect_config_domain(device.config_text)
            if domain:
                device.technology, device.technology_source = domain, "config"
        if not device.technology:
            device.technology, device.technology_source = tech_from_text(device.role, device.model)
        if not device.technology:
            device.notes.append("Technology unknown: add a Role such as 'Core Switch', 'Firewall' or 'WLC'.")
        devices.append(device)
    check_racks(devices)
    return devices


def template_workbook() -> bytes:
    """An empty inventory sheet to fill in."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Inventory"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F8A7E")
    widths = [16, 26, 22, 18, 18, 18, 26, 24, 10, 6, 9, 12]
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(65 + index)].width = width
    # Examples live on their own sheet, so importing the template back never
    # turns them into devices (the importer reads the "Inventory" sheet).
    example = workbook.create_sheet("Example")
    example.append(COLUMNS)
    for cell in example[1]:
        cell.font = Font(bold=True)
    example.append(["HO", "HO-CORE-SW01", "S5755-H24T4Y2CZ", "Core Switch", "10.10.10.1/24", "2102355XXXX", "ho-core-sw01.txt", "Ruang Server Rack 1", 42, 30, 1, "New"])
    example.append(["HO", "HO-WLC01", "C9800-40", "WLC", "10.10.10.4/24", "", "", "Ruang Server Rack 1", 42, 20, 1, "Existing"])
    example.append(["DC", "DC-FW01", "PA-1410", "Firewall", "10.20.10.5/24", "0123456789", "dc-fw01.xml"])
    example.append(["DC", "DC-RTR01", "C8200-1N-4T", "WAN Router", "10.20.10.1/24", "", ""])
    example.append(["DC", "DC-LB01", "BIG-IP i5800", "Load Balancer", "10.20.10.9/24", "", ""])
    example.append([])
    example.append(["Role decides the technology when no config is uploaded: switch/core/access, WLC/AP, firewall/FW, router/WAN, LB/F5."])
    example.append(["Rack layout: Rack = rack name, Rack Size = 24/27/32/36/42/45/47/48 (blank = 42), U = lowest unit the device sits in, U Height = units it takes (blank = 1), Status = Existing / New / Replace (blank = New, sets the colour in the rack)."])
    for index, width in enumerate(widths):
        example.column_dimensions[chr(65 + index)].width = width
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()

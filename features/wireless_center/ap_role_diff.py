from __future__ import annotations

import csv
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable


EMPTY_VALUES = {"", "-", "--", "n/a", "na", "none", "null"}

FIELD_ALIASES = {
    "ap_id": {"ap id", "apid", "id"},
    "name": {"ap name", "apname", "name", "sysname", "hostname"},
    "mac": {"mac address", "macaddress", "mac", "ap mac", "apmac", "ethernet mac"},
    "serial": {"serial number", "serialnumber", "serial", "sn", "ap sn", "apsn"},
    "ip": {"ip address", "ipaddress", "management ip", "managementip", "ap ip", "ip"},
    "group": {
        "ap group", "apgroup", "group", "group name", "groupname",
        "profile region", "profileregion",
    },
    "model": {"ap type", "aptype", "ap model", "apmodel", "model", "type"},
    "version": {"system version", "systemversion", "software version", "softwareversion", "os version", "osversion", "version"},
    "status": {"status", "ap status", "apstatus", "state"},
    "location": {"location", "site", "longitude latitude", "longitudelatitude"},
    "expected_active": {
        "expected active wlc", "expectedactivewlc", "expected wlc", "expectedwlc",
        "primary wlc", "primarywlc", "active wlc", "activewlc",
    },
}

STATUS_MAP = {
    "nor": "Normal",
    "normal": "Normal",
    "joined": "Normal",
    "standby": "Standby",
    "stdby": "Standby",
    "idle": "Idle",
    "fault": "Fault",
    "config": "Config",
    "cfg": "Config",
    "config-failed": "Config-Failed",
    "config failed": "Config-Failed",
    "cfgfa": "Config-Failed",
    "cfgfai": "Config-Failed",
    "download": "Download",
    "committing": "Committing",
    "commit-failed": "Commit-Failed",
    "ver-mismatch": "Version-Mismatch",
    "version-mismatch": "Version-Mismatch",
    "type-not-match": "Type-Mismatch",
}

AP_STATUS_GUIDE = [
    {
        "status": "Normal",
        "tone": "healthy",
        "meaning": "AP sudah terhubung ke WLC dan bekerja normal. Pengguna dapat memakai Wi-Fi dari AP ini.",
        "action": "Tidak perlu tindakan.",
    },
    {
        "status": "Standby",
        "tone": "standby",
        "meaning": "AP dikenali oleh WLC cadangan dan sedang menunggu jika WLC aktif mengalami gangguan.",
        "action": "Normal jika WLC pasangannya menunjukkan status Normal.",
    },
    {
        "status": "Idle",
        "tone": "attention",
        "meaning": "Data AP sudah dikenal WLC, tetapi AP belum terhubung atau belum pernah berkomunikasi dengan WLC tersebut.",
        "action": "Periksa power AP, VLAN management, IP, dan koneksi CAPWAP.",
    },
    {
        "status": "Fault",
        "tone": "critical",
        "meaning": "AP mengalami kegagalan saat mencoba terhubung atau melakukan registrasi ke WLC.",
        "action": "Periksa konektivitas, konfigurasi, kompatibilitas versi, dan log error.",
    },
]

# Generic two-controller role pairing guide. "A" and "B" are the two WLCs
# being compared — see the "Controller A" / "Controller B" labels the user
# can rename per comparison (ap-role form's controller-label inputs).
AP_ROLE_PAIR_GUIDE = [
    {"pair": "Normal A + Standby B", "result": "Active A", "meaning": "Sehat — AP melayani pengguna melalui Controller A dan Controller B siap sebagai cadangan."},
    {"pair": "Standby A + Normal B", "result": "Active B", "meaning": "Sehat — AP melayani pengguna melalui Controller B dan Controller A siap sebagai cadangan."},
    {"pair": "Normal + Normal", "result": "Dual Active", "meaning": "Perlu diperiksa — kedua WLC melihat AP sebagai aktif."},
    {"pair": "Standby + Standby", "result": "No Active WLC", "meaning": "Critical — tidak ada WLC yang menjadi active untuk AP."},
    {"pair": "Normal + Idle/Fault", "result": "Backup Not Ready", "meaning": "AP masih aktif, tetapi jalur backup belum siap digunakan."},
]


def _clean(value: Any) -> str:
    text = str(value or "").replace("\t", "").strip()
    return "" if text.casefold() in EMPTY_VALUES else text


def normalize_mac(value: Any) -> str:
    return re.sub(r"[^0-9a-f]", "", str(value or "").casefold())


def normalize_text(value: Any) -> str:
    return re.sub(r"\s+", " ", _clean(value)).casefold()


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", _clean(value).casefold()).strip()


def normalize_status(value: Any) -> str:
    clean = normalize_text(value)
    return STATUS_MAP.get(clean, _clean(value).title() or "Unknown")


def _canonical_field(header: str) -> str | None:
    normalized = normalize_header(header)
    compact = normalized.replace(" ", "")
    for canonical, aliases in FIELD_ALIASES.items():
        if normalized in aliases or compact in aliases:
            return canonical
    return None


def _read_tabular(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if path.suffix.casefold() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        first = next(values, None)
        if first is None:
            return [], []
        headers = [_clean(value) for value in first]
        rows = []
        for row_values in values:
            if not any(_clean(value) for value in row_values):
                continue
            rows.append({
                headers[index]: _clean(value)
                for index, value in enumerate(row_values[:len(headers)])
                if headers[index]
            })
        return headers, rows

    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(handle, dialect=dialect)
        headers = [_clean(value) for value in (reader.fieldnames or [])]
        rows = []
        for raw in reader:
            rows.append({
                _clean(key): _clean(value)
                for key, value in raw.items()
                if key is not None
            })
        return headers, rows


def _record_from_raw(raw: dict[str, Any], header_map: dict[str, str | None]) -> dict[str, str]:
    record = {key: "" for key in FIELD_ALIASES}
    for header, value in raw.items():
        canonical = header_map.get(_clean(header))
        clean = _clean(value)
        if canonical and clean and not record[canonical]:
            record[canonical] = clean
    record["status"] = normalize_status(record["status"])
    return record


def _identity(record: dict[str, Any]) -> tuple[str, str]:
    mac = normalize_mac(record.get("mac"))
    if mac:
        return "MAC", mac
    serial = normalize_text(record.get("serial"))
    if serial:
        return "Serial", serial
    name = normalize_text(record.get("name"))
    if name:
        return "AP Name", name
    return "", ""


def _deduplicate(records: Iterable[dict[str, str]], label: str) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    unique: list[dict[str, str]] = []
    indexes: dict[tuple[str, str], int] = {}
    issues: list[dict[str, str]] = []
    for record in records:
        identity_type, identity_value = _identity(record)
        if not identity_value:
            issues.append({"source": label, "severity": "Critical", "message": "AP record has no MAC, serial, or AP name."})
            continue
        key = (identity_type, identity_value)
        if key not in indexes:
            indexes[key] = len(unique)
            unique.append(record)
            continue
        existing = unique[indexes[key]]
        expected_conflict = (
            existing.get("expected_active") and record.get("expected_active")
            and existing.get("expected_active") != record.get("expected_active")
        )
        if normalize_text(existing.get("status")) != normalize_text(record.get("status")) or expected_conflict:
            issues.append({
                "source": label,
                "severity": "Critical",
                "message": f"Duplicate {identity_type} '{record.get('mac') or record.get('serial') or record.get('name')}' has conflicting status.",
            })
        else:
            issues.append({
                "source": label,
                "severity": "Warning",
                "message": f"Duplicate {identity_type} '{record.get('mac') or record.get('serial') or record.get('name')}' was consolidated.",
            })
        for field, value in record.items():
            if value and not existing.get(field):
                existing[field] = value
    return unique, issues


def _parse_display_ap_all(path: Path) -> list[dict[str, str]]:
    content = path.read_text(encoding="utf-8", errors="ignore")
    row_pattern = re.compile(
        r"^\s*(\d+)\*?\s+"
        r"([0-9a-fA-F.:-]{12,17})\s+"
        r"(\S+)\s+(\S+)\s+"
        r"(\d{1,3}(?:\.\d{1,3}){3}|-)\s+"
        r"(\S+)\s+(\S+)(?:\s+.*)?$"
    )
    records = []
    for line in content.splitlines():
        match = row_pattern.match(line)
        if not match:
            continue
        ap_id, mac, name, group, ip, model, status = match.groups()
        records.append({
            "ap_id": ap_id,
            "name": name,
            "mac": mac,
            "serial": "",
            "ip": "" if ip == "-" else ip,
            "group": group,
            "model": model,
            "version": "",
            "status": normalize_status(status),
            "location": "",
            "expected_active": "",
        })
    return records


def read_ap_state(path: Path, label: str) -> dict[str, Any]:
    suffix = path.suffix.casefold()
    if suffix in {".txt", ".log"}:
        raw_records = _parse_display_ap_all(path)
        if not raw_records:
            raise ValueError(f"{label}: no AP rows found in display ap all capture.")
        records, issues = _deduplicate(raw_records, label)
        return {"label": label, "records": records, "issues": issues, "source_type": "display ap all"}
    if suffix not in {".csv", ".xlsx", ".xlsm"}:
        raise ValueError(f"{label}: format must be CSV, XLSX, TXT, or LOG.")

    headers, raw_rows = _read_tabular(path)
    header_map = {header: _canonical_field(header) for header in headers}
    detected = {value for value in header_map.values() if value}
    if not detected.intersection({"mac", "serial", "name"}):
        raise ValueError(f"{label}: AP identity column not found (MAC, serial, or AP name).")
    if "status" not in detected:
        raise ValueError(f"{label}: AP Status/State column not found.")
    normalized = [_record_from_raw(raw, header_map) for raw in raw_rows]
    records, issues = _deduplicate(normalized, label)
    if not records:
        raise ValueError(f"{label}: AP inventory contains no usable records.")
    return {"label": label, "records": records, "issues": issues, "source_type": suffix.lstrip(".").upper()}


def _normalize_expected(value: Any, a_name: str = "", b_name: str = "") -> str:
    """Map free-text 'expected active WLC' values onto the generic A/B
    controller codes. Matches against whichever display names the user
    gave this comparison (Controller A / Controller B by default, but
    renamed to anything — e.g. real site names) so the tool carries no
    project-specific vocabulary. Also accepts the bare letters
    "A"/"B" directly, so a template CSV always works even before the
    user renames the controllers."""
    clean = normalize_text(value)
    if not clean:
        return ""
    a_key = normalize_text(a_name) or "a"
    b_key = normalize_text(b_name) or "b"
    if clean == a_key or clean in {"a", "controller a"} or (a_key and a_key in clean):
        return "A"
    if clean == b_key or clean in {"b", "controller b"} or (b_key and b_key in clean):
        return "B"
    return ""


def read_expected_allocation(path: Path, a_name: str = "", b_name: str = "") -> dict[str, Any]:
    if path.suffix.casefold() not in {".csv", ".xlsx", ".xlsm"}:
        raise ValueError("Expected allocation must be CSV or XLSX.")
    headers, raw_rows = _read_tabular(path)
    header_map = {header: _canonical_field(header) for header in headers}
    detected = {value for value in header_map.values() if value}
    if "expected_active" not in detected:
        raise ValueError("Expected allocation column not found. Use 'Expected Active WLC'.")
    if not detected.intersection({"mac", "serial", "name"}):
        raise ValueError("Expected allocation needs MAC, serial, or AP name.")

    records = []
    issues = []
    for raw in raw_rows:
        record = _record_from_raw(raw, header_map)
        expected = _normalize_expected(record.get("expected_active"), a_name, b_name)
        if not expected:
            issues.append({
                "source": "Expected Allocation",
                "severity": "Warning",
                "message": f"Expected WLC for '{record.get('name') or record.get('mac') or record.get('serial') or 'unknown AP'}' must be {a_name or 'Controller A'} or {b_name or 'Controller B'}.",
            })
            continue
        record["expected_active"] = expected
        records.append(record)
    records, duplicate_issues = _deduplicate(records, "Expected Allocation")
    issues.extend(duplicate_issues)
    return {"records": records, "issues": issues}


def _indexes(records: list[dict[str, Any]]) -> dict[str, dict[str, list[int]]]:
    indexes: dict[str, dict[str, list[int]]] = {
        "mac": defaultdict(list), "serial": defaultdict(list), "name": defaultdict(list)
    }
    for index, record in enumerate(records):
        if normalize_mac(record.get("mac")):
            indexes["mac"][normalize_mac(record.get("mac"))].append(index)
        if normalize_text(record.get("serial")):
            indexes["serial"][normalize_text(record.get("serial"))].append(index)
        record_name = record.get("name") or record.get("ap_name")
        if normalize_text(record_name):
            indexes["name"][normalize_text(record_name)].append(index)
    return indexes


def _find_match(record: dict[str, Any], candidates: list[dict[str, Any]], indexes: dict[str, dict[str, list[int]]], used: set[int]) -> tuple[int | None, str]:
    lookups = [
        ("MAC", "mac", normalize_mac(record.get("mac"))),
        ("Serial", "serial", normalize_text(record.get("serial"))),
        ("AP Name", "name", normalize_text(record.get("name") or record.get("ap_name"))),
    ]
    for label, field, value in lookups:
        if not value:
            continue
        available = [index for index in indexes[field].get(value, []) if index not in used]
        if len(available) != 1:
            continue
        candidate = candidates[available[0]]
        if label == "AP Name":
            left_mac, right_mac = normalize_mac(record.get("mac")), normalize_mac(candidate.get("mac"))
            left_serial, right_serial = normalize_text(record.get("serial")), normalize_text(candidate.get("serial"))
            if (left_mac and right_mac and left_mac != right_mac) or (left_serial and right_serial and left_serial != right_serial):
                continue
        return available[0], label
    return None, ""


def _state_class(value: Any) -> str:
    status = normalize_status(value)
    if status == "Normal":
        return "active"
    if status == "Standby":
        return "standby"
    if status == "Unknown":
        return "missing"
    if status in {"Committing", "Config", "Download"}:
        return "transitional"
    return "unhealthy"


def _role(a_status: str, b_status: str) -> tuple[str, str, str]:
    a_state, b_state = _state_class(a_status), _state_class(b_status)
    if a_state == "active" and b_state == "standby":
        return "Active A", "A", "Healthy"
    if a_state == "standby" and b_state == "active":
        return "Active B", "B", "Healthy"
    if a_state == "active" and b_state == "active":
        return "Dual Active", "Both", "Critical"
    if a_state == "standby" and b_state == "standby":
        return "No Active WLC", "None", "Critical"
    if a_state == "active":
        return "Backup B Not Ready", "A", "Critical" if b_state in {"missing", "unhealthy"} else "Warning"
    if b_state == "active":
        return "Backup A Not Ready", "B", "Critical" if a_state in {"missing", "unhealthy"} else "Warning"
    if a_state == "standby":
        return "No Active — B Not Ready", "None", "Critical"
    if b_state == "standby":
        return "No Active — A Not Ready", "None", "Critical"
    return "AP Not Operational", "None", "Critical"


def _different(left: Any, right: Any, field: str) -> bool:
    if not _clean(left) or not _clean(right):
        return False
    if field == "mac":
        return normalize_mac(left) != normalize_mac(right)
    return normalize_text(left) != normalize_text(right)


def _group_mapping_status(a_group: Any, b_group: Any) -> str:
    a_value = _clean(a_group)
    b_value = _clean(b_group)
    if not a_value and not b_value:
        return "Unassigned"
    if not a_value or not b_value:
        return "Peer Group Missing"
    if normalize_text(a_value) == normalize_text(b_value):
        return "Matched"
    return "Group Mismatch"


def _mapping_group(a_group: Any, b_group: Any) -> tuple[str, str]:
    a_value = _clean(a_group)
    b_value = _clean(b_group)
    if a_value:
        return a_value, "A"
    if b_value:
        return b_value, "B"
    return "Unassigned", "None"


def _row(a: dict[str, Any] | None, b: dict[str, Any] | None, match_method: str) -> dict[str, Any]:
    a = a or {}
    b = b or {}
    a_status = normalize_status(a.get("status")) if a else "Missing"
    b_status = normalize_status(b.get("status")) if b else "Missing"
    role, active_wlc, role_severity = _role(a_status, b_status)
    mismatch_labels = {
        "name": "AP Name", "mac": "MAC", "serial": "Serial", "ip": "IP",
        "group": "AP Group", "model": "AP Model", "version": "Software Version",
    }
    mismatches = [
        label for field, label in mismatch_labels.items()
        if _different(a.get(field), b.get(field), field)
    ]
    reference = a or b
    group_status = _group_mapping_status(a.get("group"), b.get("group"))
    mapping_group, mapping_group_source = _mapping_group(a.get("group"), b.get("group"))
    group_warning = group_status in {"Group Mismatch", "Peer Group Missing"}
    severity = role_severity if role_severity != "Healthy" else ("Warning" if mismatches or group_warning else "Healthy")
    return {
        "ap_name": reference.get("name", ""),
        "mac": reference.get("mac", ""),
        "serial": reference.get("serial", ""),
        "model": reference.get("model", ""),
        "version": reference.get("version", ""),
        "match_method": match_method or "Only on one WLC",
        "a_status": a_status,
        "b_status": b_status,
        "a_ap_id": a.get("ap_id", ""),
        "b_ap_id": b.get("ap_id", ""),
        "a_ip": a.get("ip", ""),
        "b_ip": b.get("ip", ""),
        "a_group": a.get("group", ""),
        "b_group": b.get("group", ""),
        "mapping_group": mapping_group,
        "mapping_group_source": mapping_group_source,
        "group_status": group_status,
        "a_model": a.get("model", ""),
        "b_model": b.get("model", ""),
        "a_version": a.get("version", ""),
        "b_version": b.get("version", ""),
        "role": role,
        "active_wlc": active_wlc,
        "severity": severity,
        "attribute_mismatches": mismatches,
        "attribute_mismatch": bool(mismatches),
        "missing_peer": not a or not b,
        "expected_active": "",
        "allocation_status": "Not provided",
    }


def _build_group_summary(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, dict[str, Any]] = {}
    for row in rows:
        group_name = row.get("mapping_group") or "Unassigned"
        key = normalize_text(group_name) or "unassigned"
        item = groups.setdefault(key, {
            "group": group_name,
            "total_aps": 0,
            "active_a": 0,
            "active_b": 0,
            "healthy_pairs": 0,
            "role_issues": 0,
            "group_matched": 0,
            "group_mismatch": 0,
            "peer_group_missing": 0,
            "unassigned": 0,
            "normal_states": 0,
            "standby_states": 0,
            "idle_states": 0,
            "fault_states": 0,
            "critical": 0,
            "warning": 0,
        })
        item["total_aps"] += 1
        item["active_a"] += row.get("active_wlc") == "A"
        item["active_b"] += row.get("active_wlc") == "B"
        item["healthy_pairs"] += row.get("role") in {"Active A", "Active B"}
        item["role_issues"] += row.get("role") not in {"Active A", "Active B"} or bool(row.get("missing_peer"))
        status_key = {
            "Matched": "group_matched",
            "Group Mismatch": "group_mismatch",
            "Peer Group Missing": "peer_group_missing",
            "Unassigned": "unassigned",
        }.get(row.get("group_status"))
        if status_key:
            item[status_key] += 1
        for status in (row.get("a_status"), row.get("b_status")):
            state_key = {
                "Normal": "normal_states",
                "Standby": "standby_states",
                "Idle": "idle_states",
                "Fault": "fault_states",
            }.get(status)
            if state_key:
                item[state_key] += 1
        item["critical"] += row.get("severity") == "Critical"
        item["warning"] += row.get("severity") == "Warning"
    return sorted(
        groups.values(),
        key=lambda item: (item["group"] == "Unassigned", normalize_text(item["group"])),
    )


def _apply_allocation(rows: list[dict[str, Any]], allocation: dict[str, Any] | None, issues: list[dict[str, str]]) -> None:
    if not allocation:
        return
    indexes = _indexes(rows)
    used: set[int] = set()
    for expected in allocation.get("records", []):
        index, _ = _find_match(expected, rows, indexes, used)
        if index is None:
            issues.append({
                "source": "Expected Allocation",
                "severity": "Warning",
                "message": f"Expected allocation AP '{expected.get('name') or expected.get('mac') or expected.get('serial')}' was not found in either WLC inventory.",
            })
            continue
        used.add(index)
        row = rows[index]
        row["expected_active"] = expected["expected_active"]
        if row["active_wlc"] in {"A", "B"}:
            row["allocation_status"] = "Matched" if row["active_wlc"] == expected["expected_active"] else "Role Mismatch"
        else:
            row["allocation_status"] = "Unable to Validate"
        if row["allocation_status"] == "Role Mismatch" and row["severity"] == "Healthy":
            row["severity"] = "Warning"


def compare_ap_roles(
    a_inventory: dict[str, Any],
    b_inventory: dict[str, Any],
    allocation: dict[str, Any] | None = None,
) -> dict[str, Any]:
    a_records = a_inventory.get("records", [])
    b_records = b_inventory.get("records", [])
    b_indexes = _indexes(b_records)
    used_b: set[int] = set()
    rows = []
    for a in a_records:
        b_index, match_method = _find_match(a, b_records, b_indexes, used_b)
        if b_index is None:
            rows.append(_row(a, None, ""))
        else:
            used_b.add(b_index)
            rows.append(_row(a, b_records[b_index], match_method))
    for index, b in enumerate(b_records):
        if index not in used_b:
            rows.append(_row(None, b, ""))

    issues = [
        *a_inventory.get("issues", []),
        *b_inventory.get("issues", []),
        *(allocation.get("issues", []) if allocation else []),
    ]
    _apply_allocation(rows, allocation, issues)
    rows.sort(key=lambda row: (
        {"Critical": 0, "Warning": 1, "Healthy": 2}.get(row["severity"], 3),
        normalize_text(row["ap_name"]), normalize_mac(row["mac"]),
    ))

    group_summary = _build_group_summary(rows)
    summary = {
        "total_aps": len(rows),
        "active_a": sum(row["active_wlc"] == "A" for row in rows),
        "active_b": sum(row["active_wlc"] == "B" for row in rows),
        "healthy_pairs": sum(row["role"] in {"Active A", "Active B"} for row in rows),
        "role_issues": sum(row["role"] not in {"Active A", "Active B"} or row["missing_peer"] for row in rows),
        "dual_active": sum(row["role"] == "Dual Active" for row in rows),
        "no_active": sum(row["active_wlc"] == "None" for row in rows),
        "missing_peer": sum(row["missing_peer"] for row in rows),
        "attribute_mismatch": sum(row["attribute_mismatch"] for row in rows),
        "allocation_mismatch": sum(row["allocation_status"] == "Role Mismatch" for row in rows),
        "critical": sum(row["severity"] == "Critical" for row in rows),
        "warning": sum(row["severity"] == "Warning" for row in rows),
        "ap_groups": sum(item["group"] != "Unassigned" for item in group_summary),
        "group_matched": sum(row["group_status"] == "Matched" for row in rows),
        "group_mismatch": sum(row["group_status"] == "Group Mismatch" for row in rows),
        "peer_group_missing": sum(row["group_status"] == "Peer Group Missing" for row in rows),
        "unassigned_group": sum(row["group_status"] == "Unassigned" for row in rows),
    }
    return {
        "summary": summary,
        "rows": rows,
        "group_summary": group_summary,
        "input_issues": issues,
        "sources": {
            "a_records": len(a_records),
            "b_records": len(b_records),
            "allocation_records": len(allocation.get("records", [])) if allocation else 0,
            "a_source_type": a_inventory.get("source_type", ""),
            "b_source_type": b_inventory.get("source_type", ""),
        },
        "policy": {
            "identity_priority": "MAC → Serial → AP Name",
            "active_state": "Normal",
            "standby_state": "Standby",
            "capture_guidance": "Capture both controller inventories within 1–2 minutes of each other",
        },
        "status_guide": AP_STATUS_GUIDE,
        "role_pair_guide": AP_ROLE_PAIR_GUIDE,
    }

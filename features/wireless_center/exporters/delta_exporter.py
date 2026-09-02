from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="123744")
HEADER_FONT = Font(bold=True, color="F1F7F9")


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(width, 70)


def export_delta_comparison(result: dict[str, Any], path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = result.get("summary", {})
    mac_auth = result.get("mac_auth", {})
    _sheet(workbook, "Comparison Summary", ["Metric", "Value"], [
        ["Vendor", result.get("vendor", "")],
        ["Direction", "Source to Target"],
        ["Total Objects", summary.get("total_objects", 0)],
        ["Identical", summary.get("identical", 0)],
        ["Missing on Target", summary.get("missing", 0)],
        ["Value Mismatch", summary.get("changed", 0)],
        ["Extra on Target", summary.get("extra", 0)],
        ["Manual Review", summary.get("manual_review", 0)],
        ["Generated Objects", summary.get("generated_objects", 0)],
        ["Structured Parser Objects", summary.get("structured_objects", 0)],
        ["Fallback Parser Objects", summary.get("fallback_objects", 0)],
        ["MAC-auth Chain Objects", mac_auth.get("total_objects", 0)],
        ["MAC-auth Aligned", mac_auth.get("identical", 0)],
        ["MAC-auth Drift", mac_auth.get("drift", 0)],
        ["MAC-auth Manual Review", mac_auth.get("manual_review", 0)],
    ])

    headers = [
        "Category", "Object Type", "Name", "Status", "Impact",
        "Source Configuration", "Target Configuration", "Source Delta",
        "Target-only Detail", "Eligible", "Manual Review",
        "Parser Mode", "MAC-auth Dependency", "Dependency Stage",
    ]
    all_rows = result.get("rows", [])
    groups = [
        ("All Differences", {"Missing on Target", "Value Mismatch", "Extra on Target"}),
        ("Missing on Target", {"Missing on Target"}),
        ("Value Mismatch", {"Value Mismatch"}),
        ("Extra on Target", {"Extra on Target"}),
        ("Manual Review", None),
    ]
    for title, statuses in groups:
        selected = [
            item for item in all_rows
            if (item.get("manual_review") if statuses is None else item.get("status") in statuses)
        ]
        _sheet(workbook, title, headers, [[
            item.get("category", ""), item.get("object_type", ""), item.get("name", ""),
            item.get("status", ""), item.get("impact", ""),
            "\n".join(item.get("source_commands", [])),
            "\n".join(item.get("target_commands", [])),
            "\n".join(item.get("missing_commands", [])),
            "\n".join(item.get("extra_commands", [])),
            "Yes" if item.get("eligible") else "No",
            "Yes" if item.get("manual_review") else "No",
            item.get("parser_mode", "Structured"),
            "Yes" if item.get("mac_auth_dependency") else "No",
            item.get("dependency_stage", ""),
        ] for item in selected])

    mac_rows = [item for item in all_rows if item.get("mac_auth_dependency")]
    _sheet(workbook, "MAC Authentication", headers, [[
        item.get("category", ""), item.get("object_type", ""), item.get("name", ""),
        item.get("status", ""), item.get("impact", ""),
        "\n".join(item.get("source_commands", [])),
        "\n".join(item.get("target_commands", [])),
        "\n".join(item.get("missing_commands", [])),
        "\n".join(item.get("extra_commands", [])),
        "Yes" if item.get("eligible") else "No",
        "Yes" if item.get("manual_review") else "No",
        item.get("parser_mode", "Structured"), "Yes", item.get("dependency_stage", ""),
    ] for item in mac_rows])
    _sheet(workbook, "MAC Auth Dependency", [
        "Stage", "Total", "Aligned", "Missing", "Value Mismatch",
        "Target-only", "Manual Review",
    ], [[
        stage.get("stage", ""), stage.get("total", 0), stage.get("identical", 0),
        stage.get("missing", 0), stage.get("changed", 0), stage.get("extra", 0),
        stage.get("manual_review", 0),
    ] for stage in mac_auth.get("stages", [])])

    _sheet(workbook, "Implementation Config", ["Generated Source-to-Target Delta"], [
        [line] for line in result.get("implementation_config", "").splitlines()
    ])
    _sheet(workbook, "Rollback Config", ["Generated Rollback Candidate"], [
        [line] for line in result.get("rollback_config", "").splitlines()
    ])
    workbook.save(path)

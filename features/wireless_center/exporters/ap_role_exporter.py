from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = [
    "AP Name", "MAC", "Serial", "Match Method", "Controller A Status", "Controller B Status",
    "Detected Role", "Active Controller", "Expected Active Controller", "Allocation Status",
    "Severity", "Controller A AP ID", "Controller B AP ID", "Controller A IP", "Controller B IP",
    "Controller A AP Group", "Controller B AP Group", "Mapping Group", "Mapping Group Source",
    "Group Mapping Status", "Controller A Model", "Controller B Model", "Controller A Version",
    "Controller B Version", "Attribute Mismatches", "Missing Peer",
]

GROUP_SUMMARY_HEADERS = [
    "AP Group", "Total AP", "Active Controller A", "Active Controller B", "Healthy Pairs", "Role Issues",
    "Group Matched", "Group Mismatch", "Peer Group Missing", "Unassigned",
    "Normal States", "Standby States", "Idle States", "Fault States", "Critical", "Warning",
]


def _row_values(row: dict[str, Any]) -> list[Any]:
    return [
        row.get("ap_name"), row.get("mac"), row.get("serial"), row.get("match_method"),
        row.get("a_status"), row.get("b_status"), row.get("role"), row.get("active_wlc"),
        row.get("expected_active"), row.get("allocation_status"), row.get("severity"),
        row.get("a_ap_id"), row.get("b_ap_id"), row.get("a_ip"), row.get("b_ip"),
        row.get("a_group"), row.get("b_group"), row.get("mapping_group"),
        row.get("mapping_group_source"), row.get("group_status"),
        row.get("a_model"), row.get("b_model"),
        row.get("a_version"), row.get("b_version"), ", ".join(row.get("attribute_mismatches", [])),
        "Yes" if row.get("missing_peer") else "No",
    ]


def _format_sheet(worksheet, filter_enabled: bool = True) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 24
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A6075")
        cell.alignment = Alignment(vertical="center")
    worksheet.freeze_panes = "A2"
    if filter_enabled and worksheet.max_row >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions
    for column in worksheet.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 2
        worksheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(width, 12), 42)


def _add_rows(workbook: Workbook, title: str, rows: Iterable[dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append(HEADERS)
    for row in rows:
        worksheet.append(_row_values(row))
    _format_sheet(worksheet)


def _add_group_summary(workbook: Workbook, rows: Iterable[dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet("Group Summary")
    worksheet.append(GROUP_SUMMARY_HEADERS)
    for row in rows:
        worksheet.append([
            row.get("group"), row.get("total_aps"), row.get("active_a"), row.get("active_b"),
            row.get("healthy_pairs"), row.get("role_issues"), row.get("group_matched"),
            row.get("group_mismatch"), row.get("peer_group_missing"), row.get("unassigned"),
            row.get("normal_states"), row.get("standby_states"), row.get("idle_states"),
            row.get("fault_states"), row.get("critical"), row.get("warning"),
        ])
    _format_sheet(worksheet)


def export_ap_role_diff(result: dict[str, Any], path: Path) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Role Summary"
    summary_sheet.append(["Metric", "Value"])
    for key, value in result.get("summary", {}).items():
        summary_sheet.append([key.replace("_", " ").title(), value])
    summary_sheet.append([])
    summary_sheet.append(["Policy", "Value"])
    for key, value in result.get("policy", {}).items():
        summary_sheet.append([key.replace("_", " ").title(), value])
    summary_sheet.append([])
    summary_sheet.append(["Source", "Value"])
    for key, value in result.get("sources", {}).items():
        summary_sheet.append([key.replace("_", " ").title(), value])
    _format_sheet(summary_sheet, filter_enabled=False)

    rows = result.get("rows", [])
    _add_rows(workbook, "AP Role Matrix", rows)
    _add_group_summary(workbook, result.get("group_summary", []))
    _add_rows(workbook, "AP by Group", sorted(
        rows,
        key=lambda row: (str(row.get("mapping_group", "")).casefold(), str(row.get("ap_name", "")).casefold()),
    ))
    _add_rows(workbook, "Group Mismatch", (
        row for row in rows if row.get("group_status") == "Group Mismatch"
    ))
    _add_rows(workbook, "Active Controller A", (row for row in rows if row.get("active_wlc") == "A"))
    _add_rows(workbook, "Active Controller B", (row for row in rows if row.get("active_wlc") == "B"))
    _add_rows(workbook, "Role Issues", (
        row for row in rows
        if row.get("role") not in {"Active A", "Active B"} or row.get("missing_peer")
    ))
    _add_rows(workbook, "Allocation Mismatch", (
        row for row in rows if row.get("allocation_status") == "Role Mismatch"
    ))
    _add_rows(workbook, "Attribute Mismatch", (
        row for row in rows if row.get("attribute_mismatch")
    ))

    issue_sheet = workbook.create_sheet("Input Issues")
    issue_sheet.append(["Source", "Severity", "Message"])
    for issue in result.get("input_issues", []):
        issue_sheet.append([issue.get("source"), issue.get("severity"), issue.get("message")])
    _format_sheet(issue_sheet)

    guide_sheet = workbook.create_sheet("Status Guide")
    guide_sheet.append(["Type", "Status / Pair", "Detected Result", "Plain-Language Meaning", "Recommended Action"])
    for item in result.get("status_guide", []):
        guide_sheet.append([
            "AP Status", item.get("status"), "", item.get("meaning"), item.get("action"),
        ])
    for item in result.get("role_pair_guide", []):
        guide_sheet.append([
            "Controller A / Controller B Pair", item.get("pair"), item.get("result"), item.get("meaning"), "Periksa jika hasilnya bukan pasangan active/standby yang sehat.",
        ])
    _format_sheet(guide_sheet)
    guide_sheet.column_dimensions["D"].width = 62
    guide_sheet.column_dimensions["E"].width = 54
    for row in guide_sheet.iter_rows(min_row=2, max_row=guide_sheet.max_row, min_col=4, max_col=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        guide_sheet.row_dimensions[row[0].row].height = 42
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)

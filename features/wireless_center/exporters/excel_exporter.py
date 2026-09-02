from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..summaries import build_ap_device_summary


class ExcelExporter:
    def __init__(self, data, analyzer):
        self.data = data
        self.analyzer = analyzer

    @staticmethod
    def add_sheet(
        workbook,
        title,
        headers,
        rows
    ):
        sheet = workbook.create_sheet(title)

        sheet.append(headers)

        for cell in sheet[1]:
            cell.font = Font(bold=True)

        for row in rows:
            sheet.append(row)

        for column in sheet.columns:
            max_length = 0

            for cell in column:
                value = str(
                    cell.value or ""
                )

                max_length = max(
                    max_length,
                    len(value)
                )

            letter = get_column_letter(
                column[0].column
            )

            sheet.column_dimensions[
                letter
            ].width = min(
                max_length + 2,
                60
            )

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = (
            sheet.dimensions
        )

    def export(self, path):
        workbook = Workbook()

        workbook.remove(
            workbook.active
        )

        wlc = self.data["wlc"]

        self.add_sheet(
            workbook,
            "WLC Information",
            [
                "Field",
                "Value"
            ],
            [
                [
                    key.replace(
                        "_",
                        " "
                    ).title(),
                    value
                ]
                for key, value
                in wlc.items()
            ]
        )

        self.add_sheet(
            workbook,
            "AP Inventory",
            [
                "AP ID",
                "AP Name",
                "AP Group",
                "Type ID",
                "MAC",
                "Serial",
                "IP",
                "Status",
                "Version"
            ],
            [
                [
                    ap["ap_id"],
                    ap["name"],
                    ap["group"],
                    ap["type_id"],
                    ap["mac"],
                    ap["serial"],
                    ap["ip"],
                    ap["status"],
                    ap["version"]
                ]
                for ap
                in self.data["aps"]
            ]
        )

        device_summary = build_ap_device_summary(self.data["aps"])
        self.add_sheet(
            workbook,
            "AP Device Summary",
            [
                "AP Device Type", "Total AP", "Software Version", "AP Count",
                "Status", "Consistency", "Unknown Version Count",
            ],
            [
                [
                    model["model"], model["total"], version["version"],
                    version["count"], version["status_text"], model["consistency"],
                    model["unknown_version_count"],
                ]
                for model in device_summary["rows"]
                for version in model["versions"]
            ],
        )

        ssid_catalog = (
            self.analyzer
            .build_ssid_catalog()
        )

        self.add_sheet(
            workbook,
            "WLAN SSID",
            [
                "SSID",
                "Security",
                "Security Profile",
                "Authentication Profile",
                "Traffic Profile",
                "VLAN",
                "Forward Mode",
                "Deployed VAP Profiles",
                "VAP Profile Names",
            ],
            [
                [
                    item["ssid"],
                    item["security"],
                    item[
                        "security_profile"
                    ],
                    item[
                        "authentication_profile"
                    ],
                    item[
                        "traffic_profile"
                    ],
                    item[
                        "vlan"
                    ],
                    item[
                        "forward_mode"
                    ],
                    item[
                        "vap_profile_count"
                    ],
                    item[
                        "vap_profiles"
                    ],
                ]
                for item in ssid_catalog
            ]
        )

        vap_inventory = (
            self.analyzer
            .build_vap_inventory()
        )

        self.add_sheet(
            workbook,
            "VAP Profiles",
            [
                "VAP Profile", "Usage", "Mapping Status", "SSID",
                "SSID Profile", "Service VLAN", "Effective VLAN",
                "Security", "Security Profile", "Authentication Profile",
                "Traffic Profile", "Forward Mode", "AP Group / Site",
                "Site Count", "AP Count", "Radio", "WLAN ID",
                "Mapping Count", "Full VAP Configuration",
            ],
            [
                [
                    item["vap_profile"], item["usage_scope"],
                    item["deployment_status"], item["ssid"],
                    item["ssid_profile"], item["service_vlan"],
                    item["effective_vlans"], item["security"],
                    item["security_profile"], item["authentication_profile"],
                    item["traffic_profile"], item["forward_mode"],
                    item["ap_groups"], item["ap_group_count"],
                    item["ap_count"], item["radios"], item["wlan_ids"],
                    item["mapping_count"], "\n".join(item["full_config"]),
                ]
                for item in vap_inventory
            ]
        )

        mapping = (
            self.analyzer
            .build_group_mapping()
        )

        self.add_sheet(
            workbook,
            "AP Group Mapping",
            [
                "AP Group",
                "AP Count",
                "Radio",
                "WLAN ID",
                "SSID",
                "VAP Profile",
                "Security",
                "Security Profile",
                "VLAN",
                "Forward Mode"
            ],
            [
                [
                    item["ap_group"],
                    item["ap_count"],
                    item["radio"],
                    item["wlan_id"],
                    item["ssid"],
                    item["vap_profile"],
                    item["security"],
                    item[
                        "security_profile"
                    ],
                    item["vlan"],
                    item[
                        "forward_mode"
                    ]
                ]
                for item in mapping
            ]
        )

        self.add_sheet(
            workbook,
            "Local Users",
            [
                "Username",
                "Privilege",
                "Service Type",
                "State",
                "Configuration"
            ],
            [
                [
                    user["username"],
                    user["privilege"],
                    user["service_type"],
                    user["state"],
                    "\n".join(
                        user["config"]
                    )
                ]
                for user
                in self.data[
                    "local_users"
                ]
            ]
        )

        profile_details = (
            self.analyzer
            .build_ap_profile_details()
        )

        self.add_sheet(
            workbook,
            "AP Profiles",
            [
                "AP ID",
                "AP Name",
                "AP Group",
                "AP System Profile",
                "Location Profile",
                "WLAN Count",
                "WLAN Mapping"
            ],
            [
                [
                    ap["ap_id"],
                    ap["name"],
                    ap["group"],
                    ap[
                        "ap_system_profile"
                    ],
                    ap[
                        "location_profile"
                    ],
                    len(
                        ap[
                            "wlan_mappings"
                        ]
                    ),
                    "\n".join(
                        f'Radio {x["radio"]}: '
                        f'{x["ssid"]} '
                        f'({x["vap_profile"]})'
                        for x
                        in ap[
                            "wlan_mappings"
                        ]
                    )
                ]
                for ap
                in profile_details
            ]
        )

        workbook.save(path)

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COMMON_FIELDS = [
    "AP ID", "AP Name", "Serial Number", "MAC Address", "AP Model",
    "Management IP", "Location", "Switch Name", "Switch Port", "VLAN",
    "Deployment Status", "Remark",
]
CISCO_FIELDS = ["Site Tag", "Policy Tag", "RF Tag"]
HUAWEI_FIELDS = ["AP Group", "Regulatory Domain"]
VALID_STATUSES = {
    "Planned", "Received", "Staged", "Installed", "Joined", "Validated", "Failed",
    # Operational states exported by Huawei WLC/AP Info. These are valid
    # assessment inputs even though they are not deployment workflow states.
    "Normal", "Standby", "Idle", "Fault", "Config", "Config-Failed", "Download",
}

ALIASES = {
    "ap_id": {"ap id", "ap_id", "id"},
    "ap_name": {"ap name", "ap_name", "name", "hostname"},
    "serial_number": {"serial number", "serial", "sn"},
    "mac_address": {"mac address", "mac", "ap mac", "ethernet mac"},
    "ap_model": {"ap model", "model", "type", "ap type"},
    "management_ip": {"management ip", "ip address", "ip", "ap ip"},
    "location": {"location", "site", "longitude, latitude"},
    "switch_name": {"switch name", "access switch", "switch"},
    "switch_port": {"switch port", "port", "interface"},
    "vlan": {"vlan", "management vlan"},
    "deployment_status": {"deployment status", "status", "ap status"},
    "remark": {"remark", "remarks", "note", "notes"},
    "site_tag": {"site tag", "site_tag"},
    "policy_tag": {"policy tag", "policy_tag"},
    "rf_tag": {"rf tag", "rf_tag"},
    "ap_group": {"ap group", "group", "ap-group"},
    "regulatory_domain": {"regulatory domain", "country code", "domain"},
}


def vendor_fields(vendor: str) -> list[str]:
    return COMMON_FIELDS + (CISCO_FIELDS if vendor.lower() == "cisco" else HUAWEI_FIELDS)


def normalize_mac(value: Any) -> str:
    return re.sub(r"[^0-9a-f]", "", str(value or "").lower())


def normalize_text(value: Any) -> str:
    return str(value or "").strip().lower()


def _canonical_header(value: str) -> str | None:
    normalized = normalize_text(value).replace("-", " ").replace("_", " ")
    normalized = re.sub(r"\s+", " ", normalized)
    for canonical, aliases in ALIASES.items():
        if normalized in aliases:
            return canonical
    return None


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
        reader = csv.DictReader(handle)
        # Huawei eSight/WLC exports commonly prefix every header and value
        # with a tab. Normalize field names once and use the same normalized
        # keys for every row so the alias map can resolve AC6508 exports.
        headers = [str(value or "").replace("\t", "").strip() for value in (reader.fieldnames or [])]
        records = []
        for row in reader:
            records.append({
                str(key or "").replace("\t", "").strip():
                str(value or "").replace("\t", "").strip()
                for key, value in row.items()
                if key is not None
            })
        return headers, records


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(v or "").strip() for v in rows[0]]
    records = []
    for values in rows[1:]:
        if not any(v not in (None, "") for v in values):
            continue
        records.append({headers[i]: str(values[i] or "").strip() for i in range(min(len(headers), len(values)))})
    return headers, records


def read_inventory(path: Path, vendor: str) -> dict[str, Any]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        headers, raw_records = _read_xlsx(path)
    else:
        headers, raw_records = _read_csv(path)

    header_map = {header: _canonical_header(header) for header in headers}
    detected_fields = {value for value in header_map.values() if value}
    required = {"ap_name", "serial_number", "mac_address"}
    missing_identity = sorted(required - detected_fields)
    vendor_required = {"site_tag", "policy_tag", "rf_tag"} if vendor.lower() == "cisco" else {"ap_group"}
    vendor_missing = sorted(vendor_required - detected_fields)

    records: list[dict[str, str]] = []
    issues: list[dict[str, Any]] = []
    seen: dict[str, dict[str, int]] = {"ap_name": {}, "serial_number": {}, "mac_address": {}}

    for index, raw in enumerate(raw_records, start=2):
        record = {key: "" for key in ALIASES}
        for original, value in raw.items():
            canonical = header_map.get(original)
            if canonical:
                record[canonical] = str(value or "").strip()
        records.append(record)

        if not any(record.get(k) for k in ("ap_name", "serial_number", "mac_address")):
            issues.append({"row": index, "severity": "Critical", "field": "Identity", "message": "AP Name, Serial Number, dan MAC Address semuanya kosong."})
        status = record.get("deployment_status", "")
        if status and status.title() not in VALID_STATUSES:
            issues.append({"row": index, "severity": "Warning", "field": "Deployment Status", "message": f"Status '{status}' tidak termasuk status standar NES."})
        mac = normalize_mac(record.get("mac_address"))
        if record.get("mac_address") and len(mac) != 12:
            issues.append({"row": index, "severity": "Warning", "field": "MAC Address", "message": "Format MAC Address tidak valid atau tidak lengkap."})
        for field in ("ap_name", "serial_number", "mac_address"):
            value = normalize_mac(record[field]) if field == "mac_address" else normalize_text(record[field])
            if not value:
                continue
            if value in seen[field]:
                issues.append({"row": index, "severity": "Critical", "field": field.replace("_", " ").title(), "message": f"Duplikat dengan row {seen[field][value]}."})
            else:
                seen[field][value] = index

    return {
        "vendor": vendor.title(),
        "headers": headers,
        "records": records,
        "issues": issues,
        "summary": {
            "total_records": len(records),
            "critical": sum(1 for issue in issues if issue["severity"] == "Critical"),
            "warning": sum(1 for issue in issues if issue["severity"] == "Warning"),
            "missing_identity_columns": missing_identity,
            "missing_vendor_columns": vendor_missing,
            "valid": not missing_identity and not any(issue["severity"] == "Critical" for issue in issues),
        },
    }


def compare_inventory(master_records: list[dict[str, Any]], actual_aps: list[dict[str, Any]], vendor: str) -> dict[str, Any]:
    indexes = {"mac": {}, "serial": {}, "name": {}, "id": {}}
    for ap in actual_aps:
        if normalize_mac(ap.get("mac")): indexes["mac"][normalize_mac(ap.get("mac"))] = ap
        if normalize_text(ap.get("serial")): indexes["serial"][normalize_text(ap.get("serial"))] = ap
        if normalize_text(ap.get("name")): indexes["name"][normalize_text(ap.get("name"))] = ap
        if normalize_text(ap.get("ap_id")): indexes["id"][normalize_text(ap.get("ap_id"))] = ap

    rows = []
    matched_actual_ids: set[int] = set()
    counts = {"planned": len(master_records), "matched": 0, "missing": 0, "unexpected": 0, "mismatch": 0, "validated": 0}
    for master in master_records:
        actual = None
        method = ""
        lookups = [
            ("MAC", indexes["mac"], normalize_mac(master.get("mac_address"))),
            ("Serial", indexes["serial"], normalize_text(master.get("serial_number"))),
            ("AP Name", indexes["name"], normalize_text(master.get("ap_name"))),
            ("AP ID", indexes["id"], normalize_text(master.get("ap_id"))),
        ]
        for label, index, key in lookups:
            if key and key in index:
                actual, method = index[key], label
                break
        mismatch_fields: list[str] = []
        if actual:
            matched_actual_ids.add(id(actual))
            counts["matched"] += 1
            checks = [
                ("AP Name", master.get("ap_name"), actual.get("name")),
                ("Serial Number", master.get("serial_number"), actual.get("serial")),
                ("MAC Address", normalize_mac(master.get("mac_address")), normalize_mac(actual.get("mac"))),
                ("Management IP", master.get("management_ip"), actual.get("ip")),
                ("AP Model", master.get("ap_model"), actual.get("model") or actual.get("type_id")),
            ]
            if vendor.lower() == "huawei":
                checks.append(("AP Group", master.get("ap_group"), actual.get("group")))
            else:
                # Cisco carries site tag and policy tag as two separate
                # fields (cisco9800_parser.py sets both `site_tag` and
                # `policy_tag`, with `group` mirroring `policy_tag` only).
                # This used to collapse both into one check --
                # `master.get("site_tag") or master.get("policy_tag")`
                # compared against `actual.get("group")` -- which always
                # compares against the policy tag only, no matter which
                # master value won the `or`. In a normal Cisco 9800
                # deployment site tag and policy tag are two different
                # profiles by design, so whenever `site_tag` was populated
                # on the master side this flagged a false "Site/Policy Tag"
                # mismatch even when policy tag matched exactly and nothing
                # was actually wrong. Compare each tag against its own
                # actual-side field instead.
                checks.append(("Site Tag", master.get("site_tag"), actual.get("site_tag")))
                checks.append(("Policy Tag", master.get("policy_tag"), actual.get("policy_tag") or actual.get("group")))
            for label, expected, observed in checks:
                if expected and observed and normalize_text(expected) != normalize_text(observed):
                    mismatch_fields.append(label)
            joined = normalize_text(actual.get("status")) in {"joined", "normal", "nor"}
            if mismatch_fields:
                state = "Mismatch"
                counts["mismatch"] += 1
            elif joined:
                state = "Validated"
                counts["validated"] += 1
            else:
                state = "Matched - Status Review"
        else:
            state = "Missing from WLC"
            counts["missing"] += 1

        rows.append({
            **master,
            "actual_ap_name": actual.get("name", "") if actual else "",
            "actual_serial": actual.get("serial", "") if actual else "",
            "actual_mac": actual.get("mac", "") if actual else "",
            "actual_ip": actual.get("ip", "") if actual else "",
            "actual_group": actual.get("group", "") if actual else "",
            "actual_status": actual.get("status", "") if actual else "",
            "match_method": method,
            "validation_status": state,
            "mismatch_fields": ", ".join(mismatch_fields),
        })

    unexpected = []
    for ap in actual_aps:
        if id(ap) not in matched_actual_ids:
            unexpected.append(ap)
    counts["unexpected"] = len(unexpected)
    return {"summary": counts, "rows": rows, "unexpected_aps": unexpected}


def template_csv(vendor: str, quantity: int = 0, prefix: str = "AP", site: str = "", floor: str = "") -> bytes:
    output = io.StringIO()
    fields = vendor_fields(vendor)
    writer = csv.DictWriter(output, fieldnames=fields)
    writer.writeheader()
    for number in range(1, max(0, quantity) + 1):
        parts = [part for part in (site.strip(), floor.strip(), prefix.strip(), f"{number:03d}") if part]
        row = {field: "" for field in fields}
        row["AP ID"] = str(number)
        row["AP Name"] = "-".join(parts)
        row["Deployment Status"] = "Planned"
        writer.writerow(row)
    return output.getvalue().encode("utf-8-sig")


def export_as_built(payload: dict[str, Any], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "As-Built AP Inventory"
    rows = payload.get("rows", [])
    headers = [
        "AP ID", "Planned AP Name", "Actual AP Name", "Serial Number", "Actual Serial",
        "MAC Address", "Actual MAC", "Management IP", "Actual IP", "AP Model", "Location",
        "Switch Name", "Switch Port", "VLAN", "Expected Group/Tag", "Actual Group/Tag",
        "Actual Status", "Match Method", "Validation Status", "Mismatch Fields", "Remark",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        expected_group = row.get("ap_group") or " / ".join(filter(None, [row.get("site_tag"), row.get("policy_tag"), row.get("rf_tag")]))
        ws.append([
            row.get("ap_id"), row.get("ap_name"), row.get("actual_ap_name"), row.get("serial_number"), row.get("actual_serial"),
            row.get("mac_address"), row.get("actual_mac"), row.get("management_ip"), row.get("actual_ip"), row.get("ap_model"),
            row.get("location"), row.get("switch_name"), row.get("switch_port"), row.get("vlan"), expected_group,
            row.get("actual_group"), row.get("actual_status"), row.get("match_method"), row.get("validation_status"),
            row.get("mismatch_fields"), row.get("remark"),
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        ws.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)

    summary = wb.create_sheet("Validation Summary")
    summary.append(["Metric", "Value"])
    for key, value in payload.get("summary", {}).items():
        summary.append([key.replace("_", " ").title(), value])
    unexpected = wb.create_sheet("Unexpected AP")
    unexpected.append(["AP ID", "AP Name", "Serial", "MAC", "IP", "Group/Tag", "Status"])
    for ap in payload.get("unexpected_aps", []):
        unexpected.append([ap.get("ap_id"), ap.get("name"), ap.get("serial"), ap.get("mac"), ap.get("ip"), ap.get("group"), ap.get("status")])
    wb.save(path)

"""Excel workbook of the discovery: summary, devices, links, hosts, findings."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="005073")
SEVERITY_FILL = {"medium": "F8CBAD", "low": "FFE699", "info": "DDEBF7"}
STATE_FILL = {"confirmed": "C6E0B4", "single": "FFE699", "conflict": "F8CBAD", "inferred": "DDEBF7"}


def _put(sheet, row: int, column: int, value: Any):
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        cell.data_type = "s"           # a hostname or description must never run as a formula
    return cell


def _table(workbook: Workbook, title: str, rows: list[dict[str, Any]], fill_column: str = "", fills: dict[str, str] | None = None) -> None:
    sheet = workbook.create_sheet(title)
    if not rows:
        _put(sheet, 1, 1, "Nothing to show.")
        return
    columns = list(dict.fromkeys(key for row in rows for key in row))
    for c, name in enumerate(columns, 1):
        cell = _put(sheet, 1, c, name.replace("_", " ").title())
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, row in enumerate(rows, 2):
        for c, name in enumerate(columns, 1):
            value = row.get(name, "")
            if isinstance(value, (list, tuple)):
                value = ", ".join(str(v) for v in value)
            cell = _put(sheet, r, c, value)
            cell.alignment = Alignment(vertical="top", wrap_text=name in ("detail", "recommendation", "note"))
            if name == fill_column and fills and str(value) in fills:
                cell.fill = PatternFill("solid", fgColor=fills[str(value)])
    for c, name in enumerate(columns, 1):
        longest = max([len(str(name))] + [len(str(r.get(name, ""))) for r in rows[:300]])
        sheet.column_dimensions[get_column_letter(c)].width = min(max(10, longest + 2), 60)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


def to_xlsx(summary: dict[str, Any], devices: list[dict], links: list[dict], hosts: list[dict], findings: list[dict],
            behind: list[dict] | None = None, bundles: list[dict] | None = None, by_device: list[dict] | None = None,
            l3: dict[str, list[dict]] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    labels = [("captured", "Devices captured"), ("seen_only", "Seen by neighbors only"), ("links", "Links (CDP/LLDP)"), ("confirmed", "Confirmed from both ends"),
              ("single_sided", "Reported by one end only"), ("conflicts", "Conflicts"), ("routed", "Routed /30 links inferred"), ("vpc_pairs", "vPC pairs"),
              ("hosts", "Hosts from MAC tables"), ("behind_macs", "Of which behind uncaptured devices"),
              ("carried", "Devices carried from a previous run"), ("findings", "Findings")]
    _put(sheet, 1, 1, "Keystone Topology Builder").font = Font(bold=True, size=14)
    for i, (key, label) in enumerate(labels, 3):
        _put(sheet, i, 1, label)
        _put(sheet, i, 2, summary.get(key, 0)).font = Font(bold=True)
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 12
    _table(workbook, "Devices", devices)
    _table(workbook, "Links", links, "state", STATE_FILL)
    _table(workbook, "Hosts", hosts)
    if bundles:
        _table(workbook, "Bundles", bundles, "state", STATE_FILL)
    if behind:
        _table(workbook, "Behind devices", behind)
    if l3 and l3.get("links"):
        _table(workbook, "L3 links", l3["links"], "state", {"working": "C6E0B4", "configured only": "FFE699", "broken": "F8CBAD"})
        _table(workbook, "L3 routes", l3.get("routes", []))
        if l3.get("sites"):
            _table(workbook, "SD-WAN sites", l3["sites"])
        if l3.get("overlay"):
            _table(workbook, "SD-WAN overlay", l3["overlay"])
        if l3.get("unresolved"):
            _table(workbook, "L3 unresolved", l3["unresolved"])
    _table(workbook, "Findings", [{**f, "devices": ", ".join(f["devices"])} for f in findings], "severity", SEVERITY_FILL)
    if by_device:
        _table(workbook, "Findings by device", by_device)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

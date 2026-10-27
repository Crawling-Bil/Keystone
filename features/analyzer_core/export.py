"""Generic Excel exporter for analyzer results (see result.py).

One workbook per export, one sheet per section id.  When several devices
are exported together every sheet starts with a Device column so the rows
stack instead of duplicating sheet sets -- ready to filter/pivot across a
whole site, the same convention the Switch Analyzer export uses.

Cell hygiene: values starting with = + - @ are prefixed with an
apostrophe so a hostname or ACL remark can never execute as an Excel
formula, and credentials never reach this module in the first place.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .findings import SEVERITY_LABEL, severity_counts
from .result import map_as_table

HEADER_FILL = PatternFill("solid", fgColor="1F3A4D")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D0D5DA")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
SEVERITY_FILL = {
    "high": PatternFill("solid", fgColor="F8D7DA"),
    "medium": PatternFill("solid", fgColor="FFF0C9"),
    "low": PatternFill("solid", fgColor="E1EEF9"),
    "info": PatternFill("solid", fgColor="EEF1F3"),
}
_BAD_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
MAX_COL_WIDTH = 60


def _sections_of(result: dict[str, Any]) -> list[dict[str, Any]]:
    """The result's sections as tables / key-value blocks: a map (tree) is written as one row per object with its path from the root."""
    return [map_as_table(section) if section.get("kind") == "map" else section for section in result.get("sections", [])]


def safe_cell(value: Any) -> Any:
    if isinstance(value, (list, tuple)):
        value = ", ".join(str(item) for item in value)
    if isinstance(value, str):
        if value[:1] in ("=", "+", "-", "@") and not re.fullmatch(r"-?\d+(\.\d+)?", value):
            return "'" + value
        # openpyxl rejects control characters
        return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
    return value


SHEET_NAME_OVERRIDES = {
    "License: active and time-limited modules": "License modules",
    "Still disabled or forced offline": "Still disabled or offline",
}
_TRIM_WORDS = {"and", "or", "of", "the", "in", "on", "to", "for", "from", "by", "with", "a"}


def short_sheet_title(title: str) -> str:
    """A sheet name Excel accepts (31 characters) that still reads as the section title: a trailing bracket note is dropped and a
    long title is cut at a word, never in the middle of one ("Customised system settings (sys" becomes "Customised system settings")."""
    text = SHEET_NAME_OVERRIDES.get(title) or title
    if len(_BAD_SHEET_CHARS.sub("-", text)) > 31:
        text = re.sub(r"\s*\([^)]*\)", "", text).strip() or text
    text = _BAD_SHEET_CHARS.sub("-", text).strip("'")
    if len(text) <= 31:
        return text
    words = text.split()
    while words and len(" ".join(words)) > 31:
        words.pop()
    while len(words) > 1 and words[-1].lower() in _TRIM_WORDS:
        words.pop()
    return " ".join(words) or text[:31]


def _sheet_name(title: str, used: set[str]) -> str:
    base = short_sheet_title(title)[:31] or "Sheet"
    name, counter = base, 2
    while name.lower() in used:
        suffix = f" ({counter})"
        name = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(name.lower())
    return name


def _style_header(ws, width: int) -> None:
    for col in range(1, width + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = "A2"


def _autosize(ws, headers: list[str], rows: list[list[Any]]) -> None:
    for index, header in enumerate(headers, start=1):
        longest = len(str(header))
        for row in rows[:400]:
            if index - 1 < len(row):
                longest = max(longest, min(len(str(row[index - 1])), MAX_COL_WIDTH))
        ws.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), MAX_COL_WIDTH)


def _write_table(ws, headers: list[str], rows: list[list[Any]], severity_col: int | None = None) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([safe_cell(value) for value in row])
    _style_header(ws, len(headers))
    _autosize(ws, headers, rows)
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    if severity_col is not None:
        for row_index, row in enumerate(rows, start=2):
            fill = SEVERITY_FILL.get(str(row[severity_col]).lower())
            if fill:
                ws.cell(row=row_index, column=severity_col + 1).fill = fill


def export_results(results: list[dict[str, Any]], path: Path, *, analyzer_label: str = "Analysis") -> Path:
    wb = Workbook()
    used: set[str] = set()

    # ---- Summary -------------------------------------------------------
    ws = wb.active
    ws.title = _sheet_name("Summary", used)
    headers = ["Device", "Vendor", "Platform", "Source file", "High", "Medium", "Low", "Info"]
    card_labels: list[str] = []
    for result in results:
        for item in result.get("cards", []):
            if item["label"] not in card_labels:
                card_labels.append(item["label"])
    headers += card_labels
    rows: list[list[Any]] = []
    for result in results:
        counts = severity_counts(result.get("findings", []))
        card_values = {item["label"]: item["value"] for item in result.get("cards", [])}
        rows.append(
            [result.get("hostname", ""), result.get("vendor", ""), result.get("platform", ""), result.get("source_name", ""),
             counts["high"], counts["medium"], counts["low"], counts["info"]]
            + [card_values.get(label, "") for label in card_labels]
        )
    _write_table(ws, headers, rows)

    # ---- Device information -------------------------------------------
    ws = wb.create_sheet(_sheet_name("Device Info", used))
    labels: list[str] = []
    for result in results:
        for item in result.get("info", []):
            if item["label"] not in labels:
                labels.append(item["label"])
    rows = []
    for result in results:
        values = {item["label"]: item["value"] for item in result.get("info", [])}
        rows.append([result.get("hostname", "")] + [values.get(label, "") for label in labels])
    _write_table(ws, ["Device"] + labels, rows)

    # ---- Findings -------------------------------------------------------
    ws = wb.create_sheet(_sheet_name("Findings", used))
    headers = ["Device", "Severity", "Rule", "Category", "Title", "Object", "Detail", "Recommendation", "Reference"]
    rows = []
    for result in results:
        for item in result.get("findings", []):
            rows.append([
                result.get("hostname", ""), SEVERITY_LABEL.get(item["severity"], item["severity"]), item.get("id", ""), item.get("category", ""),
                item.get("title", ""), item.get("object", ""), item.get("detail", ""), item.get("recommendation", ""), item.get("reference", ""),
            ])
    # severity column holds the label; map back for fill
    _write_table(ws, headers, rows)
    for row_index, row in enumerate(rows, start=2):
        key = str(row[1]).lower()
        if key in SEVERITY_FILL:
            ws.cell(row=row_index, column=2).fill = SEVERITY_FILL[key]

    # ---- Sections ---------------------------------------------------------
    order: list[str] = []
    meta: dict[str, dict[str, Any]] = {}
    for result in results:
        for section in _sections_of(result):
            if section["id"] not in meta:
                meta[section["id"]] = section
                order.append(section["id"])
    empty: list[str] = []          # sections with no row for any device: no sheet, one line on the "Empty sections" sheet
    for sid in order:
        first = meta[sid]
        if first["kind"] == "table":
            columns = list(first["columns"])
            known = {col["key"] for col in columns}
            for result in results:
                for section in _sections_of(result):
                    if section["id"] == sid and section["kind"] == "table":
                        for col in section["columns"]:
                            if col["key"] not in known:
                                columns.append(col)
                                known.add(col["key"])
            rows = []
            for result in results:
                for section in _sections_of(result):
                    if section["id"] == sid and section["kind"] == "table":
                        for row in section["rows"]:
                            rows.append([result.get("hostname", "")] + [row.get(col["key"], "") for col in columns])
            if not rows:
                empty.append(first["title"])
                continue
            ws = wb.create_sheet(_sheet_name(first["title"], used))
            _write_table(ws, ["Device"] + [col["label"] for col in columns], rows)
        else:
            rows = []
            for result in results:
                for section in _sections_of(result):
                    if section["id"] == sid and section["kind"] == "kv":
                        for pair in section["pairs"]:
                            rows.append([result.get("hostname", ""), pair["label"], pair["value"]])
            if not rows:
                empty.append(first["title"])
                continue
            ws = wb.create_sheet(_sheet_name(first["title"], used))
            _write_table(ws, ["Device", "Item", "Value"], rows)

    if empty:
        ws = wb.create_sheet(_sheet_name("Empty sections", used))
        _write_table(ws, ["Section", "Why there is no sheet"], [[title, "Nothing was found in the capture(s) for this section."] for title in empty])

    # ---- Coverage ---------------------------------------------------------
    ws = wb.create_sheet(_sheet_name("Capture Coverage", used))
    rows = []
    for result in results:
        for item in result.get("coverage", []):
            rows.append([result.get("hostname", ""), item.get("command", ""), item.get("status", ""), item.get("lines", ""), item.get("note", "")])
    _write_table(ws, ["Device", "Command", "Status", "Output lines", "Note"], rows)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path

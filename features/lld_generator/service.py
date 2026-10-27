"""LLD Generator: files, generate, LLD Index hand-off and the legacy import.

Storage on disk (outside Git): data/lld_generator/{templates,uploads/<project id>,outputs}.
The Word engine is core.docgen; this module only wires it to the store.
"""

from __future__ import annotations

import json
import os
import re
import shutil
import sqlite3
import time
import uuid
from copy import deepcopy
from pathlib import Path

import docx
from docx.oxml.ns import qn

from core.docgen import docgen, importer, placeholders
from . import store

BASE_DIR = Path(__file__).resolve().parents[2]
DATA_DIR = BASE_DIR / "data" / "lld_generator"          # tests point this at a temporary folder
MASTER_TEMPLATE = BASE_DIR / "core" / "docgen" / "master" / "LLD Master Template.docx"
OUTPUT_TTL_DAYS = 7
CLOSING_NAMES = {"document acceptance", "penutup", "lembar pengesahan", "approval"}


class GenError(Exception):
    """A request the user can fix; the message is shown as is."""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def folder(name: str) -> Path:
    path = Path(DATA_DIR) / name
    path.mkdir(parents=True, exist_ok=True)
    return path


def uploads_dir(project_id: int) -> Path:
    return folder("uploads") / str(int(project_id))


# ----------------------------------------------------------------- templates
def add_template(name: str, filename: str, data: bytes, closing_heading: str = "", levels=(1, 2),
                 figure_seq_name: str = "Gambar", table_seq_name: str = "Tabel", legacy_id=None) -> dict:
    if not filename.lower().endswith(".docx"):
        raise GenError("The template must be a .docx file")
    stored = f"{uuid.uuid4().hex}_{os.path.basename(filename)}"
    dest = folder("templates") / stored
    dest.write_bytes(data)
    try:
        analysis = docgen.analyze_template(str(dest))
    except Exception as exc:
        dest.unlink(missing_ok=True)
        raise GenError(f"Could not read the template: {exc}") from None
    if not analysis["has_heading1"]:
        dest.unlink(missing_ok=True)
        raise GenError("The template has no 'Heading 1' paragraph. "
                       "Chapters must use Word's built-in heading styles.")
    if not closing_heading:
        for item in analysis["outline"]:
            if item["style"] == "Heading1" and item["text"].strip().lower() in CLOSING_NAMES:
                closing_heading = item["text"].strip()
                break
    max_level = max(analysis["available_heading_levels"] or [4])
    row = store.create_template(name, stored, closing_heading, list(levels), max_level, figure_seq_name, table_seq_name)
    if legacy_id is not None:
        with store.db_conn() as conn:
            conn.execute("UPDATE lld_gen_templates SET legacy_id = ? WHERE id = ?", (legacy_id, row["id"]))
    row["analysis"] = analysis
    return row


def template_path(row: dict) -> Path:
    return folder("templates") / os.path.basename(row["filename"])


def ensure_master_template() -> dict | None:
    """The bundled LLD Master Template, registered the first time it is needed."""
    if store.list_templates() or not MASTER_TEMPLATE.exists():
        return None
    return add_template("LLD Master Template", MASTER_TEMPLATE.name, MASTER_TEMPLATE.read_bytes())


def delete_template(template_id: int) -> None:
    row = store.get_template(template_id)
    if not row:
        raise GenError("Template not found", 404)
    in_use = store.count_projects_using_template(template_id)
    if in_use:
        raise GenError(f"This template is used by {in_use} project(s). Switch those projects to another template "
                       "(or delete them) first.")
    try:
        store.delete_template(template_id)
    except sqlite3.IntegrityError:
        raise GenError("This template is still used by a project.") from None
    template_path(row).unlink(missing_ok=True)


# ----------------------------------------------------------------- images
PICTURE_TYPES = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".emf", ".wmf")


def save_image(project_id: int, filename: str, data: bytes) -> dict:
    ext = os.path.splitext(os.path.basename(filename or ""))[1].lower() or ".png"
    if ext not in PICTURE_TYPES:
        raise GenError("Pictures must be PNG, JPG, GIF, BMP, EMF or WMF")
    target = uploads_dir(project_id)
    target.mkdir(parents=True, exist_ok=True)
    stored = f"{uuid.uuid4().hex}{ext}"
    (target / stored).write_bytes(data)
    return {"filename": stored, "original_name": filename}


def image_path(project_id: int, filename: str) -> Path:
    return uploads_dir(project_id) / os.path.basename(filename)


# ----------------------------------------------------------------- import
def import_workbook(project_id: int, path: str, mode: str) -> dict:
    if mode not in ("append", "replace", "update"):
        raise GenError("mode must be append, replace or update")
    try:
        items = importer.parse_index_workbook(path)
        info = importer.parse_project_sheet(path)
    except importer.ImportError_ as exc:
        raise GenError(str(exc)) from None
    try:
        stats = store.import_index(project_id, items, mode)
    except Exception as exc:  # one transaction: nothing changed
        raise GenError(f"Import failed, nothing was changed: {exc}") from None
    updated = {k: v for k, v in info.items() if v}
    if updated:
        store.update_project(project_id, **updated)
    return {"ok": True, **stats, "project_fields": sorted(updated)}


def index_key(project: dict) -> str:
    """Identity of an LLD Index project: the same customer, name and code
    rebuild into the same LLD Generator project."""
    parts = [project.get(k) or "" for k in ("customer", "project_name", "project_code")]
    return "|".join(re.sub(r"\s+", " ", p).strip().lower() for p in parts)


def from_index(workbook_path: str, lld_project: dict, template_id: int | None = None) -> dict:
    """LLD Index build -> LLD Generator project, through the same importer the
    Excel upload uses (one import path, one set of Perbarui rules).
    The project with the same index key is updated; otherwise a new one is made."""
    key = index_key(lld_project)
    existing = store.find_project_by_index_key(key)
    if existing:
        result = import_workbook(existing["id"], workbook_path, "update")
        return {"project_id": existing["id"], "created": False, "mode": "update", "import": result}
    if not template_id:
        ensure_master_template()
        templates = store.list_templates()
        master = next((t for t in templates if "master" in t["name"].lower()), templates[0] if templates else None)
        if not master:
            raise GenError("No master template yet: upload one in Templates.")
        template_id = master["id"]
    if not store.get_template(template_id):
        raise GenError("Template not found", 404)
    project = store.create_project(lld_project.get("project_name") or "LLD", lld_project.get("customer") or "", "0.1",
                                   template_id, project_code=lld_project.get("project_code") or "",
                                   company=lld_project.get("company") or "Multipolar Technology",
                                   prepared_by=lld_project.get("prepared_by") or "", lld_index_key=key)
    result = import_workbook(project["id"], workbook_path, "replace")
    return {"project_id": project["id"], "created": True, "mode": "replace", "import": result}


# ----------------------------------------------------------------- review / generate
TODO_LABEL = {"image": "picture to upload", "table": "table still TBD", "paragraph": "text still TBD"}


def _needs_work(block: dict) -> bool:
    """A picture not uploaded, or a table/text with a cell or line that only says TBD."""
    kind, content = block["block_type"], block["content"]
    if kind == "image":
        return not content.get("filename")
    if kind == "table":
        return any(str(c).strip() == "TBD" for row in content.get("rows") or [] for c in row)
    if kind == "paragraph":
        return any(str(t).strip() == "TBD" for t in content.get("paragraphs") or [])
    return False


def review(project_id: int) -> dict:
    """Per chapter: its block types and what is still to do (a picture not
    uploaded, a table or text that only says TBD). A chapter counts as done
    only when it has content and nothing left to do."""
    chapters = store.list_chapters_flat(project_id)
    parents = {c["parent_id"] for c in chapters if c["parent_id"] is not None}
    out, filled, empty, todo_total = [], 0, 0, 0
    for chapter in chapters:
        blocks = store.list_blocks(chapter["id"])
        types = [b["block_type"] for b in blocks]
        todo = [TODO_LABEL[b["block_type"]] for b in blocks if b["block_type"] in TODO_LABEL and _needs_work(b)]
        is_group = chapter["id"] in parents
        out.append({**chapter, "block_types": types, "is_group": is_group, "todo": todo})
        if not is_group:
            if not types:
                empty += 1
            elif todo:
                todo_total += 1
            else:
                filled += 1
    return {"chapters": out, "leaf_filled": filled, "leaf_empty": empty, "leaf_todo": todo_total}


# ----------------------------------------------------------------- pictures in bulk
def _words(text: str) -> set:
    return {w for w in re.split(r"[^a-z0-9]+", (text or "").lower()) if w and w not in {"the", "of", "and", "site", "png", "jpg", "jpeg"}}


def place_pictures(project_id: int, files: list[tuple[str, bytes]]) -> dict:
    """Put each picture into the empty image block whose chapter title or
    caption shares the most words with the file name ("topology-after-ho.png"
    -> "Physical Topology After Site HO"). Unmatched files are reported."""
    slots = []
    for chapter in store.list_chapters_flat(project_id):
        for block in store.list_blocks(chapter["id"]):
            if block["block_type"] == "image" and not block["content"].get("filename"):
                slots.append((block, chapter, _words(chapter["title"]) | _words(block["content"].get("caption", ""))))
    placed, unmatched = [], []
    for name, data in files:
        if os.path.splitext(name)[1].lower() not in PICTURE_TYPES:
            unmatched.append(name)                       # not a picture Word can take: report it, keep going
            continue
        want = _words(os.path.splitext(name)[0])
        best, score = None, 0.0
        for slot in slots:
            common = len(want & slot[2])
            if not common:
                continue
            s = common / len(want | slot[2]) + common * 0.01
            if s > score:
                best, score = slot, s
        if best is None or score < 0.2:
            unmatched.append(name)
            continue
        slots.remove(best)
        saved = save_image(project_id, name, data)
        block, chapter, _ = best
        store.update_block(block["id"], content={**block["content"], "filename": saved["filename"], "original_name": name})
        placed.append({"file": name, "chapter": chapter["title"], "chapter_id": chapter["id"]})
    return {"placed": placed, "unmatched": unmatched, "empty_left": len(slots)}


# ----------------------------------------------------------------- approvers (Document Approval)
MULTIPOLAR = "PT. Multipolar Technology Tbk."
DEFAULT_MULTIPOLAR = [{"name": "", "title": "Section Head"}, {"name": "{prepared_by}", "title": "Network Engineer"},
                      {"name": "", "title": "Project Manager"}]


def _clean_people(rows) -> list:
    out = []
    for r in rows or []:
        side = "multipolar" if str(r.get("side") or "").lower().startswith("m") else "customer"
        person = {k: " ".join(str(r.get(k) or "").split())[:120] for k in ("name", "title", "company")}
        if person["name"] or person["title"]:
            out.append({**person, "side": side})
    return out


def approvers(project: dict) -> list:
    """The people in Document Approval for this project: its own list, else
    the customer's list from earlier projects plus the Multipolar defaults."""
    own = json.loads(project.get("approvers") or "[]")
    if own:
        return own
    customer = [{**r, "side": "customer"} for r in store.get_approver_book(project.get("client"))]
    multipolar = store.get_setting("multipolar_approvers") or DEFAULT_MULTIPOLAR
    filled = [{"name": (r.get("name") or "").replace("{prepared_by}", project.get("prepared_by") or ""),
               "title": r.get("title") or "", "company": r.get("company") or MULTIPOLAR, "side": "multipolar"}
              for r in multipolar]
    return customer + filled


def save_approvers(project: dict, rows, multipolar_default: bool = False) -> list:
    people = _clean_people(rows)
    store.update_project(project["id"], approvers=json.dumps(people))
    customer = [{k: p[k] for k in ("name", "title", "company")} for p in people if p["side"] == "customer"]
    if customer:
        store.set_approver_book(project.get("client"), customer)       # next project for this customer starts here
    if multipolar_default:
        prepared = project.get("prepared_by") or ""
        store.set_setting("multipolar_approvers", [
            {"name": "{prepared_by}" if prepared and p["name"] == prepared else p["name"], "title": p["title"],
             "company": p["company"]} for p in people if p["side"] == "multipolar"])
    return people


def _set_cell(cell, lines) -> None:
    paragraph = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    runs = paragraph.runs
    first = runs[0] if runs else paragraph.add_run()
    for run in runs[1:]:
        run._r.getparent().remove(run._r)
    for br in first._r.findall(qn("w:br")):
        first._r.remove(br)
    first.text = lines[0] if lines else ""
    for line in lines[1:]:
        first.add_break()
        first.add_text(line)


def _fill_approvers(document, project: dict) -> None:
    """Rebuild the two Document Approval tables (customer, then Multipolar)
    from the approver list: one row per person, name and title in the first
    column, company in the second, date and signature left for signing."""
    people = approvers(project)
    if not people:
        return
    tables = [t for t in document.tables if t.rows and t.rows[0].cells[0].text.strip().lower().startswith("name and title")]
    for table, side in zip(tables, ("customer", "multipolar")):
        rows = [p for p in people if p["side"] == side]
        if not rows or len(table.rows) < 2:
            continue
        template = deepcopy(table.rows[1]._tr)
        for row in list(table.rows)[1:]:
            row._tr.getparent().remove(row._tr)
        for person in rows:
            table._tbl.append(deepcopy(template))
            cells = table.rows[-1].cells
            company = person.get("company") or (project.get("client") if side == "customer" else MULTIPOLAR) or ""
            _set_cell(cells[0], [person.get("name") or "", person.get("title") or ""])
            if len(cells) > 1:
                _set_cell(cells[1], [company])
            for cell in cells[2:]:
                _set_cell(cell, [""])


# ----------------------------------------------------------------- release history
def record_version(project: dict, new_version: str, comment: str = "") -> list:
    """Release History rows: one per version, added when the version changes."""
    history = json.loads(project.get("history") or "[]")
    if not history:
        history.append({"version": project.get("version") or "0.1", "date": placeholders.project_values(project)["date"],
                        "by": project.get("prepared_by") or "", "status": "Initial", "comment": "-"})
    if new_version and new_version != history[-1]["version"]:
        history.append({"version": new_version, "date": placeholders.project_values({})["date"],
                        "by": project.get("prepared_by") or "", "status": "Revised", "comment": comment or "-"})
    return history


def _fill_release_history(document, project: dict) -> None:
    history = json.loads(project.get("history") or "[]")
    if not history:
        return
    for table in document.tables:
        if not table.rows or table.rows[0].cells[0].text.strip().lower() not in ("rev", "rev."):
            continue
        if len(table.rows) < 2:
            return
        template = table.rows[1]._tr
        for row in list(table.rows)[1:]:
            row._tr.getparent().remove(row._tr)
        for entry in history:
            tr = deepcopy(template)
            table._tbl.append(tr)
            row = table.rows[-1]
            for cell, value in zip(row.cells, (entry["version"], entry["date"], entry["by"], entry["status"], entry["comment"])):
                paragraph = cell.paragraphs[0]
                for extra in cell.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
                runs = paragraph.runs
                if runs:
                    runs[0].text = str(value)
                    for run in runs[1:]:
                        run._r.getparent().remove(run._r)
                else:
                    paragraph.add_run(str(value))
        return


def cleanup_outputs(days: int = OUTPUT_TTL_DAYS) -> int:
    limit = time.time() - days * 86400
    removed = 0
    for path in folder("outputs").glob("*.docx"):
        try:
            if path.stat().st_mtime < limit:
                path.unlink()
                removed += 1
        except OSError:
            pass
    return removed


def generate(project_id: int) -> dict:
    project = store.get_project(project_id)
    if not project:
        raise GenError("Project not found", 404)
    template = store.get_template(project["template_id"])
    if not template:
        raise GenError("This project's template no longer exists")
    chapters = store.list_chapters_flat(project_id)
    if not chapters:
        raise GenError("The project has no chapters yet")
    cleanup_outputs()
    blocks = {c["id"]: store.list_blocks(c["id"]) for c in chapters}
    uploads = uploads_dir(project_id)
    safe = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in project["name"])
    out_name = f"[LLD] {safe} v{project['version']}.docx"
    out_path = folder("outputs") / f"{project_id}_{uuid.uuid4().hex}.docx"
    try:
        docgen.generate(template_path=str(template_path(template)), chapters=chapters, blocks_by_chapter=blocks,
                        uploads_dir=str(uploads), output_path=str(out_path),
                        closing_heading=template["closing_heading"],
                        auto_numbered_levels=json.loads(template["auto_numbered_levels"]),
                        max_heading_level=template["max_heading_level"],
                        figure_seq_name=template["figure_seq_name"], table_seq_name=template["table_seq_name"])
    except docgen.DocGenError as exc:
        raise GenError(str(exc)) from None
    document = docx.Document(str(out_path))
    _fill_release_history(document, project)
    _fill_approvers(document, project)
    missing = placeholders.fill(document, placeholders.project_values(project))
    logo = image_path(project_id, project["logo"]) if project.get("logo") else None
    has_slot = placeholders.fill_logo(document, str(logo) if logo else "")
    document.save(str(out_path))
    warnings = [f"{{{{{k}}}}} is empty: fill it in Project settings." for k in sorted(missing)]
    if has_slot and not (logo and logo.exists()):
        warnings.append("No customer logo: the cover has no logo.")
    from urllib.parse import quote
    return {"download_url": f"/lld-generator/api/download/{out_path.name}?name={quote(out_name)}", "filename": out_name,
            "warnings": warnings}


def output_path(stored_name: str) -> Path:
    return folder("outputs") / os.path.basename(stored_name)


# ----------------------------------------------------------------- legacy import
def find_legacy_folder(hint: str | None = None) -> Path | None:
    """The old Document Generator (lld_generator_v*/ with backend/storage/app.db)."""
    candidates = []
    if hint:
        candidates.append(Path(hint).expanduser())
    parent = BASE_DIR.parent / "Document_Generator"
    if parent.is_dir():
        candidates += sorted((p for p in parent.iterdir() if p.is_dir() and p.name.startswith("lld_generator")),
                             key=lambda p: [int(n) for n in re.findall(r"\d+", p.name)] or [0], reverse=True)
    for path in candidates:
        for root in (path, path.parent, path.parent.parent):
            if (root / "backend" / "storage" / "app.db").is_file():
                return root
    return None


def legacy_import(hint: str | None = None) -> dict:
    """Copy templates, projects, chapters, blocks and images from the old
    Document Generator. Read-only on the old app; projects and templates keep
    their old id (legacy_id), so a second run only adds what is new."""
    root = find_legacy_folder(hint)
    if root is None:
        raise GenError("The old Document Generator folder (with backend/storage/app.db) was not found.")
    storage = root / "backend" / "storage"
    old = sqlite3.connect(f"file:{storage / 'app.db'}?mode=ro", uri=True)
    old.row_factory = sqlite3.Row
    stats = {"folder": str(root), "templates": 0, "projects": 0, "chapters": 0, "blocks": 0, "images": 0,
             "skipped_projects": 0, "missing_files": []}
    try:
        with store.db_conn() as conn:
            done_t = {r["legacy_id"]: r["id"] for r in conn.execute(
                "SELECT id, legacy_id FROM lld_gen_templates WHERE legacy_id IS NOT NULL")}
            done_p = {r["legacy_id"] for r in conn.execute(
                "SELECT legacy_id FROM lld_gen_projects WHERE legacy_id IS NOT NULL")}
        template_map = dict(done_t)
        for t in old.execute("SELECT * FROM templates ORDER BY id"):
            if t["id"] in template_map:
                continue
            src = storage / "templates" / os.path.basename(t["filename"])
            if not src.exists():
                stats["missing_files"].append(f"template {t['filename']}")
                continue
            keys = t.keys()
            row = add_template(t["name"], src.name.split("_", 1)[-1], src.read_bytes(),
                               t["closing_heading"] or "", json.loads(t["auto_numbered_levels"] or "[1,2]"),
                               t["figure_seq_name"] if "figure_seq_name" in keys else "Gambar",
                               t["table_seq_name"] if "table_seq_name" in keys else "Tabel", legacy_id=t["id"])
            template_map[t["id"]] = row["id"]
            stats["templates"] += 1

        conn = store.get_conn()
        try:
            conn.execute("BEGIN")
            for p in old.execute("SELECT * FROM projects ORDER BY id"):
                if p["id"] in done_p:
                    stats["skipped_projects"] += 1
                    continue
                if p["template_id"] not in template_map:
                    stats["missing_files"].append(f"project {p['name']}: template not imported")
                    continue
                k = p.keys()
                get = lambda name, default="": (p[name] if name in k and p[name] is not None else default)  # noqa: E731
                ts = store.now()
                cur = conn.execute(
                    "INSERT INTO lld_gen_projects (name, client, version, template_id, project_code, company, prepared_by,"
                    " doc_date, logo, legacy_id, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (p["name"], get("client"), get("version", "0.1"), template_map[p["template_id"]], get("project_code"),
                     get("company", "Multipolar Technology") or "Multipolar Technology", get("prepared_by"),
                     get("doc_date"), get("logo"), p["id"], get("created_at", ts) or ts, ts))
                new_pid = cur.lastrowid
                stats["projects"] += 1
                src_uploads = storage / "uploads" / str(p["id"])
                if src_uploads.is_dir():
                    dest = uploads_dir(new_pid)
                    dest.mkdir(parents=True, exist_ok=True)
                    for f in src_uploads.iterdir():
                        if f.is_file():
                            shutil.copy2(f, dest / f.name)
                            stats["images"] += 1
                chapter_map = {}
                rows = old.execute("SELECT * FROM chapters WHERE project_id = ? ORDER BY level, order_index",
                                   (p["id"],)).fetchall()
                pending = list(rows)
                while pending:                                   # parents first, whatever the id order
                    rest = []
                    for c in pending:
                        if c["parent_id"] is not None and c["parent_id"] not in chapter_map:
                            rest.append(c)
                            continue
                        cur = conn.execute(
                            "INSERT INTO lld_gen_chapters (project_id, parent_id, level, order_index, title) VALUES (?, ?, ?, ?, ?)",
                            (new_pid, chapter_map.get(c["parent_id"]), c["level"], c["order_index"], c["title"]))
                        chapter_map[c["id"]] = cur.lastrowid
                        stats["chapters"] += 1
                    if len(rest) == len(pending):              # orphans: their parent is gone
                        break
                    pending = rest
                for old_cid, new_cid in chapter_map.items():
                    for b in old.execute("SELECT * FROM content_blocks WHERE chapter_id = ? ORDER BY order_index",
                                         (old_cid,)):
                        conn.execute("INSERT INTO lld_gen_blocks (chapter_id, order_index, block_type, content) VALUES (?, ?, ?, ?)",
                                     (new_cid, b["order_index"], b["block_type"], b["content"]))
                        stats["blocks"] += 1
                        if b["block_type"] == "image":
                            name = (json.loads(b["content"] or "{}").get("filename") or "")
                            if name and not (src_uploads / name).exists():
                                stats["missing_files"].append(f"image {name} ({p['name']})")
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()
    finally:
        old.close()
    return stats

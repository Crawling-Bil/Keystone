"""OSPF interface summary and link-state database (features/switch_analyzer/ospf_db.py).

The layout is that of two real Nexus 9000 captures (``show ip ospf interface brief vrf all`` and ``show ip ospf database``, one of them cut by the terminal pager:
``--More--``, ``...skipping N lines`` and the pager's help screen); every name, address, process and LSA in the fixtures is invented (RFC 5737 documentation addresses only).

Documentation: Cisco Nexus 9000 NX-API CLI Reference, OSPF Commands (``show ip ospf interface brief vrf all``); Cisco Nexus 3500 NX-API CLI Reference, OSPF Commands
(``show ip ospf database``); RFC 2328 Appendix B (``LSRefreshTime`` 1800 s, ``MaxAge`` 3600 s).
"""
from __future__ import annotations

import ipaddress
import re
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from features.switch_analyzer import capture_text, ospf_db
from features.switch_analyzer.compare import compare
from features.switch_analyzer.exporters import SwitchExcelExporter, export_comparison
from features.switch_analyzer.service import analyze

FIXTURES = Path(__file__).parent / "fixtures" / "real_captures"
CAPTURE_A = (FIXTURES / "synthetic_nxos_ospf_db_a.txt").read_text(encoding="utf-8")
CAPTURE_B = (FIXTURES / "synthetic_nxos_ospf_db_b.txt").read_text(encoding="utf-8")

_CACHE: dict[str, dict] = {}


def run(name: str, text: str) -> dict:
    if name not in _CACHE:
        _CACHE[name] = analyze(text, filename=f"{name}.txt")
    return _CACHE[name]


def ids(result: dict) -> set[str]:
    return {f["id"] for f in result["findings"]}


def one(result: dict, finding_id: str) -> dict:
    return next(f for f in result["findings"] if f["id"] == finding_id)


def drop(text: str, needle: str) -> str:
    """The capture without the one line that holds ``needle``."""
    lines = text.splitlines(keepends=True)
    hits = [i for i, line in enumerate(lines) if needle in line]
    assert len(hits) == 1, (needle, hits)
    del lines[hits[0]]
    return "".join(lines)


class FixtureHygieneTests(unittest.TestCase):
    def test_only_documentation_addresses_are_in_the_fixtures(self):
        networks = [ipaddress.ip_network(n) for n in ("192.0.2.0/24", "198.51.100.0/24", "203.0.113.0/24")]
        for text in (CAPTURE_A, CAPTURE_B):
            for match in re.finditer(r"(?<![\d.])(\d{1,3}(?:\.\d{1,3}){3})(?![\d.])", text):
                if match.group(1).startswith("0.0.0."):         # an OSPF area ID written in dotted form, not an address
                    continue
                self.assertTrue(any(ipaddress.ip_address(match.group(1)) in net for net in networks), match.group(1))

    def test_the_devices_are_invented_names(self):
        self.assertIn("NXDEMO-A#", CAPTURE_A)
        self.assertIn("NXDEMO-B#", CAPTURE_B)


class InterfaceBriefTests(unittest.TestCase):
    NX = """ OSPF Process ID 172 VRF default
 Total number of interface: 3
 Interface               ID     Area            Cost   State    Neighbors Status
 Vlan100                 1      0.0.0.0         10     DR       1         up
 Po100                   2      0.0.0.0         1      P2P      1         up
 Vlan500                 3      0.0.0.5         10     DOWN     0         down

 OSPF Process ID 173 VRF blue
 Total number of interface: 1
 Interface               ID     Area            Cost   State    Neighbors Status
 Vlan10                  4      0.0.0.1         1      BDR      2         up
""".splitlines()
    IOS = """Interface    PID   Area            IP Address/Mask    Cost  State Nbrs F/C
Vl10         1     0               192.0.2.1/24       1     DR    1/1
Gi1/0/1      1     0.0.0.5         198.51.100.1/30    10    P2P   0/0
Lo0          1     0               192.0.2.255/32     1     LOOP  0/0
""".splitlines()

    def test_nxos_rows_and_totals_per_process_and_vrf(self):
        rows, totals = ospf_db.parse_interfaces(self.NX)
        self.assertEqual(len(rows), 4)
        self.assertEqual(totals, {("172", "default"): 3, ("173", "blue"): 1})
        down = next(r for r in rows if r["interface"] == "Vlan500")
        self.assertEqual((down["state"], down["status"], down["area"], down["neighbors"]), ("DOWN", "down", "0.0.0.5", 0))
        self.assertEqual(next(r for r in rows if r["interface"] == "Vlan10")["vrf"], "blue")

    def test_ios_layout_counts_full_neighbours_and_keeps_the_address(self):
        rows, totals = ospf_db.parse_interfaces(self.IOS)
        self.assertEqual(totals, {})
        self.assertEqual([r["interface"] for r in rows], ["Vl10", "Gi1/0/1", "Lo0"])
        self.assertEqual(rows[0]["address"], "192.0.2.1/24")
        self.assertEqual((rows[0]["neighbors"], rows[0]["full"]), (1, 1))
        self.assertEqual(rows[0]["process"], "1")

    def test_other_text_is_not_a_row(self):
        rows, _ = ospf_db.parse_interfaces(["% Invalid command at '^' marker.", "    ^", "Total number of interface: 2"])
        self.assertEqual(rows, [])


class DatabaseParserTests(unittest.TestCase):
    def test_nxos_summary_tables_by_area_and_type(self):
        lines = capture_text.clean_lines(CAPTURE_A)
        view = ospf_db.read(lines)
        db = view["database"]
        self.assertEqual(db["rows"], 23)
        self.assertEqual([p["router_id"] for p in db["processes"]], ["192.0.2.1", "198.51.100.1"])
        counts = {(i["process"], i["area"], i["kind"]): i["count"] for i in db["lsdb"]}
        self.assertEqual(counts[("172", "0.0.0.0", "router")], 4)
        self.assertEqual(counts[("172", "0.0.0.0", "network")], 2)
        self.assertEqual(counts[("172", "0.0.0.0", "summary")], 3)
        self.assertEqual(counts[("172", "0.0.0.0", "asbr-summary")], 1)
        self.assertEqual(counts[("172", "", "external")], 5)
        self.assertEqual(counts[("173", "0.0.0.1", "network")], 1)
        self.assertEqual((db["skipped"], db["skip_events"], db["pager_help"]), (0, 0, 0))

    def test_router_lsas_keep_age_sequence_and_link_count(self):
        db = ospf_db.read(capture_text.clean_lines(CAPTURE_A))["database"]
        first = next(r for r in db["routers"] if r["router"] == "192.0.2.1" and r["area"] == "0.0.0.0")
        self.assertEqual((first["age"], first["seq"], first["links"]), (1203, "0x80000031", 4))
        self.assertEqual(len(db["routers"]), 7)

    def test_age_buckets_follow_rfc_2328_constants(self):
        self.assertEqual((ospf_db.LSA_REFRESH, ospf_db.MAX_AGE), (1800, 3600))
        db = ospf_db.read(capture_text.clean_lines(CAPTURE_A))["database"]
        external = next(i for i in db["lsdb"] if i["process"] == "172" and i["kind"] == "external")
        self.assertEqual(external["oldest"], 1973)
        self.assertEqual(external["over_refresh"], 2)      # 1801 s and 1973 s; 1799 s is still inside LSRefreshTime
        self.assertEqual(external["stale"], 0)
        late = next(i for i in db["lsdb"] if i["process"] == "173" and i["kind"] == "network")
        self.assertEqual((late["maxage"], late["stale"]), (1, 1))

    def test_pager_cut_is_counted_and_the_next_rows_are_not_misfiled(self):
        db = ospf_db.read(capture_text.clean_lines(CAPTURE_B))["database"]
        self.assertEqual((db["skipped"], db["skip_events"], db["pager_help"]), (9, 1, 1))
        counts = {(i["process"], i["area"], i["kind"]): i["count"] for i in db["lsdb"]}
        self.assertEqual(counts[("172", "0.0.0.0", "router")], 4)           # the 4 rows printed before the cut, nothing from after it
        self.assertNotIn(("172", "0.0.0.0", "summary"), counts)             # its header and rows fell into the cut
        self.assertEqual(len(db["routers"]), 7)

    def test_a_rejected_command_is_not_a_database(self):
        text = "SW1#show ip ospf database\n                        ^\n% Invalid command at '^' marker.\nSW1#\n"
        view = ospf_db.read(capture_text.clean_lines(text))
        self.assertEqual(view["database"]["rows"], 0)

    def test_ios_layout(self):
        lines = """            OSPF Router with ID (192.0.2.1) (Process ID 1)

                Router Link States (Area 0)

Link ID         ADV Router      Age         Seq#       Checksum Link count
192.0.2.1       192.0.2.1       600         0x80000004 0x00A1B2 2
192.0.2.2       192.0.2.2       700         0x80000003 0x00C3D4 1

                Net Link States (Area 0)

Link ID         ADV Router      Age         Seq#       Checksum
192.0.2.2       192.0.2.2       700         0x80000001 0x00E5F6

                Summary Net Link States (Area 0)

Link ID         ADV Router      Age         Seq#       Checksum
198.51.100.0    192.0.2.1       500         0x80000001 0x00AA11

                Summary ASB Link States (Area 0)

Link ID         ADV Router      Age         Seq#       Checksum
192.0.2.9       192.0.2.1       500         0x80000001 0x00BB22

                Type-5 AS External Link States

Link ID         ADV Router      Age         Seq#       Checksum Tag
203.0.113.0     192.0.2.9       100         0x80000001 0x00CC33 0
""".splitlines()
        db = ospf_db.parse_database(lines)
        self.assertEqual([(i["kind"], i["count"]) for i in db["lsdb"]],
                         [("router", 2), ("network", 1), ("summary", 1), ("asbr-summary", 1), ("external", 1)])
        self.assertEqual(db["lsdb"][-1]["area"], "")       # the AS-wide LSAs come after the areas
        self.assertEqual(db["processes"], [{"process": "1", "vrf": "", "router_id": "192.0.2.1"}])


class PagerHelpTests(unittest.TestCase):
    def test_the_help_screen_is_one_marker_and_the_command_block_stays_whole(self):
        lines = capture_text.clean_lines(CAPTURE_B)
        self.assertEqual(sum(1 for line in lines if line == capture_text.PAGER_HELP_MARKER), 1)
        self.assertFalse(any(line.startswith("<space>") for line in lines))
        block = dict((index, block) for index, block in capture_text.find_blocks(lines, ospf_db.PATTERN_DATABASE))
        self.assertEqual(len(block), 1)
        self.assertTrue(any("Router with ID (198.51.100.2)" in line for line in next(iter(block.values()))))

    def test_text_without_a_help_screen_is_untouched(self):
        lines = ["SW1#show clock", "10:00:00 UTC", "SW1#"]
        self.assertEqual(capture_text._fold_pager_help(lines), lines)


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.a = run("a", CAPTURE_A)
        self.b = run("b", CAPTURE_B)

    def test_tables_on_the_routing_tab(self):
        titles = {t["id"]: t for t in self.a["extra_tables"]}
        for key in ("ospf-interfaces", "ospf-overview", "ospf-lsdb", "ospf-routers"):
            self.assertIn(key, titles)
        self.assertEqual(len(titles["ospf-interfaces"]["rows"]), 5)
        overview = {row["process"]: row for row in titles["ospf-overview"]["rows"]}
        self.assertEqual(overview["172"]["router_id"], "192.0.2.1")
        self.assertEqual(overview["172"]["total"], 18)
        self.assertEqual(overview["172"]["interfaces"], "3 up / 4")
        self.assertEqual(overview["172"]["neighbors"], "3 Full / 3")
        self.assertEqual(overview["173"]["areas"], "0.0.0.1")

    def test_no_literal_null_in_any_ospf_table(self):
        for t in self.a["extra_tables"] + self.b["extra_tables"]:
            if t["id"].startswith("ospf-"):
                for row in t["rows"]:
                    self.assertTrue(all(value is not None for value in row.values()), (t["id"], row))

    def test_interface_down_finding(self):
        finding = one(self.a, "SW-OSPF-002")
        self.assertIn("Vlan500", finding["object"])
        self.assertEqual(finding["severity"], "low")

    def test_lsa_at_maxage_is_a_medium_finding_that_names_the_area(self):
        finding = one(self.a, "SW-OSPF-005")
        self.assertEqual(finding["severity"], "medium")
        self.assertIn("process 173", finding["object"])
        self.assertIn("0.0.0.1", finding["object"])
        self.assertIn("rfc2328", finding["reference"])

    def test_lsa_ages_between_refresh_jitter_and_maxage_are_not_flagged(self):
        self.assertNotIn("SW-OSPF-004", ids(self.a))      # the externals sit at 1801-1973 s: a late refresh, not a stale LSA

    def test_total_that_disagrees_with_the_rows_is_a_commands_disagree_finding(self):
        finding = one(self.a, "SW-CAP-007")
        self.assertEqual(finding["title"], "OSPF commands disagree")
        self.assertIn("says 5", finding["detail"])
        self.assertIn("lists 4", finding["detail"])

    def test_the_pager_cut_is_a_capture_note_and_switches_the_database_checks_off(self):
        notes = [n["text"] for n in self.b["capture_notes"]]
        self.assertTrue(any("terminal pager cut the 'show ip ospf database' output" in n and "9 lines" in n for n in notes))
        self.assertIn("SW-CAP-004", ids(self.b))
        self.assertTrue(ids(self.b).isdisjoint({"SW-OSPF-003", "SW-OSPF-004", "SW-OSPF-005"}))
        self.assertNotIn("different prompts", " ".join(notes))          # the help screen must not look like a second device

    def test_coverage_reads_both_commands(self):
        rows = {r["command"]: r for r in self.a["coverage"]["rows"]}
        self.assertEqual(rows["show ip ospf interface brief vrf all"]["status"], "analysed")
        self.assertEqual(rows["show ip ospf database"]["status"], "analysed")
        self.assertEqual((self.a["extras"]["ospf_intf"], self.a["extras"]["ospf_db"]), (5, 23))

    def test_neighbour_without_a_router_lsa_is_found_only_on_a_complete_database(self):
        text = drop(CAPTURE_A, "192.0.2.4       192.0.2.4       411")
        result = analyze(text, filename="gap.txt")
        finding = one(result, "SW-OSPF-003")
        self.assertIn("192.0.2.4 on Po100", finding["object"])
        cut = drop(CAPTURE_B, "192.0.2.4       192.0.2.4       411")
        self.assertNotIn("SW-OSPF-003", ids(analyze(cut, filename="cut.txt")))

    def test_an_lsa_that_missed_its_refresh_by_ten_minutes_is_a_low_finding(self):
        text = CAPTURE_A.replace("203.0.113.192   192.0.2.8       1801", "203.0.113.192   192.0.2.8       2600")
        result = analyze(text, filename="stale.txt")
        finding = one(result, "SW-OSPF-004")
        self.assertEqual(finding["severity"], "low")
        self.assertIn("2600", finding["detail"])

    def test_a_capture_without_ospf_database_adds_nothing(self):
        result = analyze("hostname SW9\ninterface Vlan1\n ip address 192.0.2.9 255.255.255.0\n!\n", filename="none.txt")
        self.assertFalse([f for f in result["findings"] if f["id"].startswith("SW-OSPF")])
        self.assertFalse([t for t in result["extra_tables"] if t["id"].startswith("ospf-")])


class PairTests(unittest.TestCase):
    def setUp(self):
        self.a = run("a", CAPTURE_A)
        self.b = run("b", CAPTURE_B)

    def test_a_complete_database_against_a_cut_one_is_check_not_differs(self):
        pair = compare(self.a, self.b)["ospf_live"]
        self.assertEqual(pair["incomplete"], ["B"])
        self.assertEqual(pair["counts"], {"Same": 8, "Check": 2})
        self.assertNotIn("Differs", {row["verdict"] for row in pair["rows"]})
        missing = [row for row in pair["rows"] if row["verdict"] == "Check"]
        self.assertTrue(all("pager" in row["reason"] for row in missing))

    def test_two_complete_databases_are_compared_by_count_and_by_router(self):
        other = drop(CAPTURE_A.replace("NXDEMO-A", "NXDEMO-C"), "192.0.2.4       192.0.2.4       411")
        other = drop(other, "198.51.100.128  192.0.2.4       77")
        result = analyze(other, filename="c.txt")
        pair = compare(self.a, result)["ospf_live"]
        self.assertEqual(pair["incomplete"], [])
        router = next(row for row in pair["rows"] if row["area"] == "0.0.0.0" and row["type"].startswith("Router"))
        self.assertEqual(router["verdict"], "Differs")
        self.assertIn("192.0.2.4", router["reason"])
        self.assertIn("1 LSA more on A", router["reason"])

    def test_nothing_to_compare_without_a_database(self):
        empty = analyze("hostname SW9\ninterface Vlan1\n ip address 192.0.2.9 255.255.255.0\n!\n", filename="none.txt")
        pair = compare(self.a, empty)["ospf_live"]
        self.assertEqual((pair["rows"], pair["counts"]), ([], {}))

    def test_verdict_tones_cover_every_verdict(self):
        pair = compare(self.a, self.b)["ospf_live"]
        for row in pair["rows"]:
            self.assertIn(row["verdict"].lower(), ospf_db.VERDICT_TONES)


class ExcelTests(unittest.TestCase):
    def sheet(self, workbook, name):
        return [[cell.value for cell in row] for row in workbook[name].iter_rows()]

    def test_single_device_workbook_has_the_four_ospf_sheets(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "a.xlsx"
            SwitchExcelExporter(run("a", CAPTURE_A)).export(path)
            workbook = load_workbook(path)
            interfaces = self.sheet(workbook, "OSPF Interfaces")
            processes = self.sheet(workbook, "OSPF Processes")
            lsdb = self.sheet(workbook, "OSPF LSDB")
            routers = self.sheet(workbook, "OSPF Routers")
        self.assertEqual(interfaces[0][:6], ["Hostname", "Process", "VRF", "Interface", "ID", "Area"])
        self.assertEqual(len(interfaces) - 1, 5)
        self.assertIn("Vlan500", [row[3] for row in interfaces[1:]])
        self.assertEqual({row[1]: row[3] for row in processes[1:]}, {"172": "192.0.2.1", "173": "198.51.100.1"})
        self.assertEqual(len(lsdb) - 1, 10)
        self.assertEqual(len(routers) - 1, 7)
        self.assertEqual(lsdb[0][-1], "Older Than 40 Min")

    def test_pair_workbook_has_the_lsdb_compare_sheet_with_the_pager_note(self):
        a, b = run("a", CAPTURE_A), run("b", CAPTURE_B)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "pair.xlsx"
            export_comparison(path, a, b, compare(a, b))
            workbook = load_workbook(path)
            values = self.sheet(workbook, "OSPF LSDB Compare")
        self.assertIn("Pager cut the database of B", " ".join(str(v) for v in values[0] if v))
        self.assertEqual(values[1][0], "Verdict")
        verdicts = [row[0] for row in values[2:]]
        self.assertEqual(verdicts[:2], ["Check", "Check"])           # problems first
        self.assertEqual(len(verdicts), 10)

    def test_no_compare_sheet_when_a_side_has_no_database(self):
        a = run("a", CAPTURE_A)
        empty = analyze("hostname SW9\ninterface Vlan1\n ip address 192.0.2.9 255.255.255.0\n!\n", filename="none.txt")
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "pair.xlsx"
            export_comparison(path, a, empty, compare(a, empty))
            self.assertNotIn("OSPF LSDB Compare", load_workbook(path).sheetnames)


class CollectListTests(unittest.TestCase):
    def test_the_nexus_list_ends_with_the_database_after_the_log(self):
        lines = [l.strip() for l in (Path(__file__).parent.parent / "static" / "collect-scripts" / "collect-cisco-nexus.txt").read_text(encoding="utf-8").splitlines()
                 if l.strip() and not l.startswith("!!!")]
        self.assertEqual(lines[-1], "show ip ospf database")
        self.assertLess(lines.index("show logging last 1000"), lines.index("show ip ospf database"))
        self.assertLess(lines.index("term len 0"), lines.index("show ip ospf database"))
        self.assertIn("show ip ospf interface brief vrf all", lines)

    def test_the_catalyst_list_has_both_ospf_views(self):
        lines = [l.strip() for l in (Path(__file__).parent.parent / "static" / "collect-scripts" / "collect-cisco-catalyst.txt").read_text(encoding="utf-8").splitlines()
                 if l.strip() and not l.startswith("!!!")]
        self.assertIn("show ip ospf interface brief", lines)
        self.assertIn("show ip ospf database", lines)
        self.assertLess(lines.index("show ip ospf database"), lines.index("show interface"))


if __name__ == "__main__":
    unittest.main()

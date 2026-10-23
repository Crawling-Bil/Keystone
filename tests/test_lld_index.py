"""LLD Index: inventory + configs + project form -> Document Generator Index workbook.

The workbook checks mirror the Document Generator's importer.py rules
(header row "No | Judul Heading", Level "H#", Tabel = a sheet with a header
and at least one data row, sheet names <= 31 characters). All data is
synthetic.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import openpyxl
import pytest

from features.lld_index import builder, workbook
from features.lld_index.datasets import analyze_device, config_blocks
from features.lld_index.inventory import build_devices, read_inventory_file, tech_from_text, template_workbook
from features.lld_index.service import build_index, read_config_uploads

FIXTURES = Path(__file__).parent / "fixtures"
SWITCH = (FIXTURES / "switch" / "c9300_48p_nm8x_synthetic.txt").read_text()
FORTI = (FIXTURES / "firewall" / "fortigate_synthetic_7.2.conf").read_text()
BIGIP = (FIXTURES / "loadbalancer" / "bigip_capture_synthetic.txt").read_text(errors="ignore")

INVENTORY = [
    {"site": "HO", "hostname": "SW-LAB-ACC1", "model": "C9300-48P", "role": "Access Switch", "mgmt_ip": "192.0.2.11/24", "config": "sw1.txt"},
    {"site": "HO", "hostname": "HO-CORE", "model": "S5755-H24T4Y2CZ", "role": "Core Switch", "mgmt_ip": "10.0.0.1/24"},
    {"site": "BR", "hostname": "BR-CORE", "model": "S5735-S24P4XE-V2", "role": "Core Switch"},
    {"site": "DC", "hostname": "FG", "model": "FGT60F", "role": "Firewall", "config": "fg.conf", "serial": "FGT60FTK1"},
    {"site": "DC", "hostname": "LB", "model": "BIG-IP i5800", "role": "Load Balancer", "config": "lb.txt"},
    {"site": "HO", "hostname": "HO-AP-01", "model": "C9115AXI", "role": "AP"},
]
FILES = {"sw1.txt": SWITCH, "fg.conf": FORTI, "lb.txt": BIGIP}
PROJECT = {"customer": "PT Contoh", "project_name": "Network Refresh", "goals": "Replace end-of-life switches\nNew perimeter firewall"}


def _build(techs, records=INVENTORY, files=FILES):
    return build_index(dict(PROJECT), [dict(r) for r in records], techs, files)


def _index_rows(data: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    sheet = wb["Index"]
    header = next(r for r in range(1, 30) if sheet.cell(r, 1).value == "No" and "Judul" in str(sheet.cell(r, 2).value))
    rows = [list(r)[:6] for r in sheet.iter_rows(min_row=header + 1, values_only=True) if r[1]]
    return wb, rows


# ------------------------------------------------------------------ inventory

def test_technology_comes_from_config_then_role_then_model():
    devices = build_devices(INVENTORY, FILES)
    by_host = {d.hostname: d for d in devices}
    assert (by_host["SW-LAB-ACC1"].technology, by_host["SW-LAB-ACC1"].technology_source) == ("switch", "config")
    assert (by_host["HO-CORE"].technology, by_host["HO-CORE"].technology_source) == ("switch", "role")
    assert by_host["HO-AP-01"].technology == "wireless"
    assert tech_from_text("", "PA-1410") == ("firewall", "model")


def test_config_matches_by_file_name_or_hostname_and_reports_missing():
    files = {"ho-core.cfg": "hostname HO-CORE\n!", "other.txt": "hostname X"}
    devices = build_devices([{"site": "HO", "hostname": "HO-CORE"},
                             {"site": "HO", "hostname": "SW9", "config": "missing.txt"}], files)
    assert devices[0].config_name == "ho-core.cfg"
    assert any("missing.txt" in n for n in devices[1].notes)


def test_inventory_reads_csv_and_xlsx_headers_in_any_order():
    csv_data = b"Hostname;Site;Model;SN\nSW1;HO;C9300-48P;FOC1\n"
    assert read_inventory_file("inv.csv", csv_data) == [{"hostname": "SW1", "site": "HO", "model": "C9300-48P", "serial": "FOC1"}]
    rows = read_inventory_file("inv.xlsx", template_workbook())
    assert rows == []  # the template's examples live on their own sheet


def test_config_uploads_unpack_zip_text_files_only():
    import zipfile
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("site/sw1.txt", "hostname SW1")
        archive.writestr("photo.png", b"\x89PNG")
        archive.writestr("__MACOSX/._sw1.txt", "junk")
    files = read_config_uploads([("configs.zip", buffer.getvalue()), ("fw.conf", b"config system global")])
    assert set(files) == {"sw1.txt", "fw.conf"}


# ------------------------------------------------------------------ tree

def test_switch_chapters_expand_per_site_then_hostname():
    result = _build(["switch"])
    titles = [(n["level"], n["title"]) for n in result["tree"]]
    vlan = titles.index((1, "VLAN and IP Addressing"))
    assert titles[vlan + 1:vlan + 6] == [(2, "HO"), (3, "SW-LAB-ACC1"), (3, "HO-CORE"), (2, "BR"), (3, "BR-CORE")]
    befores = [t for _, t in titles if t.startswith("Physical Topology")]
    assert befores[1:5] == ["Physical Topology Before Site HO", "Physical Topology After Site HO",
                            "Physical Topology Before Site BR", "Physical Topology After Site BR"]


def test_tables_come_from_the_config_and_stay_tbd_without_one():
    result = _build(["switch"])
    tree = result["tree"]
    vlan_rows = [n for n in tree if n["title"] == "SW-LAB-ACC1" and n["kind"] == "table"]
    assert vlan_rows[0]["source"] == "config" and vlan_rows[0]["rows"] == 2 and not vlan_rows[0]["placeholder"]
    no_config = [n for n in tree if n["title"] == "HO-CORE" and n["kind"] == "table"]
    assert all(n["placeholder"] for n in no_config)


def test_firewall_site_chapter_keeps_only_sections_with_data():
    result = _build(["firewall"])
    titles = [n["title"] for n in result["tree"]]
    site = titles.index("Site DC")
    section = titles[site + 1:]
    assert "Security Policy" in section and "IKE Gateway Configuration" in section and "DHCP Configuration" in section
    assert "QoS Profile" not in section          # optional, FortiGate fixture has none
    assert "FG" not in section                   # one firewall in the site: no hostname level


def test_access_points_get_no_chapter_of_their_own_and_switches_are_not_written_twice():
    result = _build(["switch", "wireless"])
    titles = [n["title"] for n in result["tree"]]
    assert "HO-AP-01" not in titles
    assert titles.count("SW-LAB-ACC1") == 3      # VLAN, connectivity, configuration: switch chapters only
    assert "VLAN and IP Addressing - Switch" in titles and "VLAN and IP Addressing - Wireless" in titles


def test_load_balancer_chapter_is_filled_from_the_bigip_capture():
    result = _build(["loadbalancer"])
    by_title = {n["title"]: n for n in result["tree"]}
    assert by_title["Virtual Server"]["rows"] == 13 and by_title["Pool"]["rows"] == 7


def test_narrative_uses_the_project_fields_and_lists_the_chapters():
    devices = build_devices(INVENTORY, FILES)
    datasets = {d.hostname: analyze_device(d) for d in devices}
    project = dict(PROJECT, technologies=["switch"])
    nodes = builder.build(project, devices, datasets, ["switch"])
    by_title = {n.title: n for n in nodes}
    description = " ".join(by_title["Project Description"].paragraphs)
    assert "PT Contoh" in description and "Network Refresh" in description and "• Replace end-of-life switches." in description
    scope = by_title["Document Scope"].paragraphs
    assert "• Device Overview" in scope and "• Introduction" not in scope


def test_model_heading_has_a_description_paragraph_and_a_spec_table():
    result = _build(["switch"])
    spec = next(n for n in result["tree"] if n["title"] == "Huawei S5755-H24T4Y2CZ")
    assert spec["kind"] == "paragraph"
    assert spec["more"][0]["kind"] == "table" and spec["more"][0]["source"] == "hardware catalog" and spec["more"][0]["rows"] > 4


def test_extra_blocks_are_written_as_continuation_rows():
    devices = build_devices(INVENTORY, FILES)
    datasets = {d.hostname: analyze_device(d) for d in devices}
    project = dict(PROJECT, technologies=["switch"])
    nodes = builder.build(project, devices, datasets, ["switch"])
    wb, rows = _index_rows(workbook.write(nodes, project))
    sheet = wb["Index"]
    header = next(r for r in range(1, 30) if sheet.cell(r, 1).value == "No")
    all_rows = [list(r)[:6] for r in sheet.iter_rows(min_row=header + 1, values_only=True)]
    at = next(i for i, r in enumerate(all_rows) if r[1] == "Huawei S5755-H24T4Y2CZ")
    heading, extra = all_rows[at], all_rows[at + 1]
    assert heading[3] == "Paragraf" and "Core Switch" in heading[4]
    assert extra[:3] == [None, None, None] and extra[3] == "Tabel" and extra[4] in wb.sheetnames
    assert extra[5] == "S5755-H24T4Y2CZ Specification"


# ------------------------------------------------------------------ workbook format

def test_workbook_follows_the_document_generator_importer_format():
    result = _build(["switch", "firewall", "loadbalancer"])
    devices = build_devices(INVENTORY, FILES)
    datasets = {d.hostname: analyze_device(d) for d in devices}
    project = dict(PROJECT, technologies=["switch", "firewall", "loadbalancer"])
    nodes = builder.build(project, devices, datasets, project["technologies"])
    wb, rows = _index_rows(workbook.write(nodes, project))
    assert len(rows) == len(nodes) == len(result["tree"])
    names = [s.lower() for s in wb.sheetnames]
    assert len(names) == len(set(names)) and all(len(n) <= 31 for n in wb.sheetnames)
    rows = [r for r in rows]  # heading rows only (Judul filled)
    for no, title, level, kind, isi, note in rows:
        assert str(level).startswith("H") and 1 <= int(level[1:]) <= 9
        if kind == "Tabel":
            sheet = wb[isi]
            assert sheet.max_row >= 2, f"{title}: the importer drops a table without rows"
            assert any(c.value for c in sheet[2])
        if kind == "Paragraf" and isi in wb.sheetnames:
            assert wb[isi].cell(1, 2).value == "Isi"
    levels = [int(r[2][1:]) for r in rows]
    assert all(b - a <= 1 for a, b in zip(levels, levels[1:])), "a heading level was skipped"


def test_config_blocks_group_set_lines_and_drop_prompts():
    blocks = config_blocks("sw1#show running-config\nBuilding configuration...\nhostname SW1\n!\ninterface Gi1/0/1\n description PC\n!\nend\n")
    assert [b["remarks"] for b in blocks] == ["hostname", "interface"]
    flat = config_blocks("set network interface ethernet ethernet1/1 layer3\nset network interface ethernet ethernet1/2 layer3\nset zone A network layer3 [ ethernet1/1 ]\n")
    assert [b["remarks"] for b in flat] == ["network interface", "zone"]


# ------------------------------------------------------------------ API

@pytest.fixture()
def client(tmp_path, monkeypatch):
    from app import create_app
    from features.lld_index import routes

    monkeypatch.setattr(routes, "EXPORT_DIR", tmp_path / "exports")
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def test_api_builds_and_downloads_the_workbook(client):
    payload = {"project": PROJECT, "technologies": ["switch"], "inventory": INVENTORY[:2]}
    response = client.post("/lld-index/api/build", data={"payload": json.dumps(payload),
                                                         "configs": [(io.BytesIO(SWITCH.encode()), "sw1.txt")]},
                           content_type="multipart/form-data")
    assert response.status_code == 200, response.get_json()
    result = response.get_json()["result"]
    assert result["stats"]["configs_read"] == 1
    download = client.get(result["download_url"])
    assert download.status_code == 200
    _, rows = _index_rows(download.data)
    assert rows[0][1] == "Introduction"


def test_api_refuses_an_empty_request(client):
    response = client.post("/lld-index/api/build", data={"payload": json.dumps({"technologies": [], "inventory": []})})
    assert response.status_code == 400


def test_page_and_inventory_template_open(client):
    assert 'data-tech="firewall"' in client.get("/lld-index").get_data(as_text=True)
    assert client.get("/lld-index/inventory-template.xlsx").status_code == 200


# ------------------------------------------------------------------ rack layout

RACKED = [
    {"site": "HO", "hostname": "HO-CORE", "model": "S5755-H24T4Y2CZ", "role": "Core Switch", "rack": "Rack 1", "rack_size": "42", "u": "30", "u_height": "1"},
    {"site": "HO", "hostname": "HO-FW", "model": "PA-1410", "role": "Firewall", "rack": "Rack 1", "u": "24", "u_height": "2"},
    {"site": "HO", "hostname": "HO-X", "model": "S5735-S24P4XE-V2", "role": "Access Switch", "rack": "Rack 1", "u": "25"},
    {"site": "HO", "hostname": "HO-WALL", "model": "S5735-S24P4XE-V2", "role": "Access Switch", "rack": "Wallmount", "rack_size": "12", "u": "12", "u_height": "2"},
    {"site": "BR", "hostname": "BR-CORE", "model": "S5735-S24P4XE-V2", "role": "Core Switch"},
]


def test_rack_columns_are_read_and_problems_noted():
    devices = {d.hostname: d for d in build_devices(RACKED, {})}
    assert (devices["HO-FW"].rack, devices["HO-FW"].rack_size, devices["HO-FW"].u, devices["HO-FW"].u_height) == ("Rack 1", 42, 24, 2)
    assert any("Overlaps HO-FW" in n for n in devices["HO-X"].notes)
    assert any("does not fit a 12U rack" in n for n in devices["HO-WALL"].notes)
    assert devices["BR-CORE"].u is None and not devices["BR-CORE"].notes


def test_rack_table_puts_each_device_on_its_units_with_merged_cells():
    devices = build_devices(RACKED, {})
    datasets = {d.hostname: analyze_device(d) for d in devices}
    project = dict(PROJECT, technologies=["switch", "firewall"])
    nodes = builder.build(project, devices, datasets, project["technologies"])
    titles = [(n.level, n.title, n.kind) for n in nodes]
    at = titles.index((1, "Rack Layout", "group"))
    assert titles[at + 1:at + 6] == [(2, "HO", "group"), (3, "Rack 1", "table"), (3, "Wallmount", "table"),
                                     (2, "BR", "group"), (3, "Rack Layout BR", "image")]
    rack = next(n for n in nodes if n.title == "Rack 1")
    assert rack.headers == ["U", "Rack 1 (42U)", "U"] and len(rack.rows) == 42
    by_unit = {r[0]: r[1] for r in rack.rows}
    assert by_unit["30"] == "HO-CORE (S5755-H24T4Y2CZ)" and by_unit["25"] == "HO-FW (PA-1410)" and by_unit["24"] == ""
    assert rack.merges == [[17, 1, 2, 1]]           # U25 is the 18th row from the top
    wb, rows = _index_rows(workbook.write(nodes, project))
    sheet = wb[next(r[4] for r in rows if r[1] == "Rack 1")]
    assert [str(m) for m in sheet.merged_cells.ranges] == ["B19:B20"]


def test_api_returns_the_rack_preview(client):
    payload = {"project": PROJECT, "technologies": ["switch", "firewall"], "inventory": RACKED}
    result = client.post("/lld-index/api/build", data={"payload": json.dumps(payload)}).get_json()["result"]
    rack = next(r for r in result["racks"] if r["rack"] == "Rack 1")
    assert rack["size"] == 42 and {u["hostname"] for u in rack["units"]} == {"HO-CORE", "HO-FW"}
    assert any("HO-X" in p for p in rack["problems"])


def test_rack_cells_are_coloured_like_the_reference_table():
    devices = build_devices(RACKED, {})
    datasets = {d.hostname: analyze_device(d) for d in devices}
    project = dict(PROJECT, technologies=["switch", "firewall"])
    nodes = builder.build(project, devices, datasets, project["technologies"])
    rack = next(n for n in nodes if n.title == "Rack 1")
    device_fills = [f for f in rack.fills if f[1] == 1]
    assert device_fills == [[12, 1, "DDEBF7"], [17, 1, "DDEBF7"]]          # U30 and the top of the 2U at U24-25
    assert sum(1 for f in rack.fills if f[2] == "F2F2F2") == 2 * 42
    assert not rack.more                                                   # no legend
    wb, rows = _index_rows(workbook.write(nodes, project))
    sheet = wb[next(r[4] for r in rows if r[1] == "Rack 1")]
    assert sheet.cell(row=14, column=2).fill.fgColor.rgb.endswith("DDEBF7")
    assert sheet.cell(row=2, column=1).fill.fgColor.rgb.endswith("F2F2F2")
    assert sheet.cell(row=14, column=2).alignment.horizontal == "center"


REPLACED = [
    {"site": "HO", "hostname": "HO-CORE-OLD", "model": "C3850-24T", "role": "Core Switch", "rack": "Rack 1", "u": "30", "status": "Existing"},
    {"site": "HO", "hostname": "HO-CORE", "model": "S5755-H24T4Y2CZ", "role": "Core Switch", "rack": "Rack 1", "u": "30", "status": "New"},
    {"site": "HO", "hostname": "HO-PATCH", "model": "Patch Panel 24P", "role": "Patch Panel", "rack": "Rack 1", "u": "40", "status": "keep"},
    {"site": "HO", "hostname": "HO-ACC", "model": "C9300-48P", "role": "Access Switch", "rack": "Rack 1", "u": "28"},
]


def _replacement_nodes(replacement):
    devices = build_devices(REPLACED, {})
    datasets = {d.hostname: analyze_device(d) for d in devices}
    project = dict(PROJECT, technologies=["switch"], replacement=replacement)
    return devices, project, builder.build(project, devices, datasets, ["switch"])


def test_status_column_is_read_and_phases_share_units():
    devices = {d.hostname: d for d in build_devices(REPLACED, {})}
    assert [devices[h].status for h in ("HO-CORE-OLD", "HO-CORE", "HO-PATCH", "HO-ACC")] == ["Existing", "New", "Keep", "New"]
    # the old core leaves U30 and the new one takes it: not an overlap
    assert not any("Overlaps" in n for d in devices.values() for n in d.notes)
    assert not any("Technology unknown" in n for n in devices["HO-PATCH"].notes)


def test_replacement_project_draws_each_rack_before_and_after():
    _, project, nodes = _replacement_nodes(True)
    titles = [(n.level, n.title) for n in nodes]
    at = titles.index((1, "Rack Layout"))
    assert titles[at + 1:at + 4] == [(2, "HO"), (3, "Rack Layout Before - Rack 1"), (3, "Rack Layout After - Rack 1")]
    before = {r[0]: r[1] for r in next(n for n in nodes if n.title == "Rack Layout Before - Rack 1").rows}
    after = {r[0]: r[1] for r in next(n for n in nodes if n.title == "Rack Layout After - Rack 1").rows}
    assert before["30"].startswith("HO-CORE-OLD") and before["40"].startswith("HO-PATCH") and before["28"] == ""
    assert after["30"].startswith("HO-CORE (") and after["40"].startswith("HO-PATCH") and after["28"].startswith("HO-ACC")
    names = {n.sheet_hint for n in nodes}
    assert "Rack Before Rack 1" in names and "Rack After Rack 1" in names
    workbook.write(nodes, project)                                         # sheet names stay valid


def test_existing_devices_stay_out_of_every_other_chapter():
    for replacement in (True, False):
        _, _, nodes = _replacement_nodes(replacement)
        others = [n for n in nodes if "Rack" not in n.title and "Rack" not in n.sheet_hint]
        text = json.dumps([[n.title, n.rows] for n in others])
        assert "HO-CORE-OLD" not in text and "C3850" not in text
        assert "HO-PATCH" not in text                                      # no technology: rack only
        assert "HO-CORE" in text


def test_without_replacement_the_rack_shows_the_finished_state():
    _, _, nodes = _replacement_nodes(False)
    assert not any(n.title.startswith("Rack Layout Before") for n in nodes)
    rack = {r[0]: r[1] for r in next(n for n in nodes if n.title == "Rack 1").rows}
    assert rack["30"].startswith("HO-CORE (") and rack["40"].startswith("HO-PATCH")


def test_boq_counts_only_new_devices():
    records = REPLACED + [{"site": "HO", "hostname": "HO-DIST", "model": "C9500-24Y4C", "role": "Distribution Switch", "status": "Keep"}]
    devices = build_devices(records, {})
    scope = builder.Scope(devices=devices)
    _, rows, _ = builder._builder_rows({"builder": "boq"}, scope, {}, PROJECT)
    assert {r[0] for r in rows} == {"S5755-H24T4Y2CZ", "C9300-48P"}


def test_api_rack_preview_has_phases(client):
    payload = {"project": dict(PROJECT, replacement="true"), "technologies": ["switch"], "inventory": REPLACED}
    result = client.post("/lld-index/api/build", data={"payload": json.dumps(payload)}).get_json()["result"]
    assert [(r["rack"], r["phase"]) for r in result["racks"]] == [("Rack 1", "Before"), ("Rack 1", "After")]
    old = next(d for d in result["devices"] if d["hostname"] == "HO-CORE-OLD")
    assert old["status"] == "Existing" and any("Before rack" in n for n in old["notes"])

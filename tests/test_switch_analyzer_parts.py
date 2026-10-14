"""Single Device "Serial Numbers" table: hardware parts plus every transceiver
with its PID and serial, checked against masked excerpts of the real Oct 2026
TAM captures (see tests/fixtures/real_captures/README.md).

Audit of all 86 real captures (the table against the raw `show inventory` /
`display device elabel` / `display interface transceiver` text) found these
reader bugs, each pinned below
  T1  Huawei: the serial of the LAST transceiver was overwritten by the
      `Device serial number` lines of the `display lldp neighbor` that came
      after it (12 captures). The reader scanned every line to the end of
      the file instead of stopping at the next command.
  T2  Aruba: the 2540 table reader was also applied to the Catalyst
      `show interfaces transceiver` temperature / power table and turned the
      DOM row of Gi0/1 into a transceiver of type "30.6" with serial "-6.3".
  T3  Cisco: the plural `show interfaces transceiver` was not recognised, so
      the optical readings were not joined to the SFP.
  T4  Nexus: `cisco product id is SFP-10G-SR` (the PID) was never read.
  T5  A session-log record header with no date (`16:42:21|000000000135|`)
      hid `NAME: "Fan 1"` entries, so 3 of 4 fans were missing.
  T6  Repeated runs of the same command repeated the rows.
  T7  One capture holds two devices' sessions (a copy-paste of the next
      switch's `display interface transceiver` in front of its own): rows of
      the other device must not appear, and the page must say so.
And the new table: kind per row, PID per row, transceivers listed once.
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from features.switch_analyzer import capture_text, parts, runtime_parser, serials
from features.switch_analyzer.exporters import SwitchExcelExporter
from features.switch_analyzer.service import analyze

FIXTURES = Path(__file__).parent / "fixtures" / "real_captures"


def text(*names: str) -> str:
    return "\n".join((FIXTURES / name).read_text(encoding="utf-8") for name in names)


def xcvrs(*names: str) -> list[dict]:
    return runtime_parser.parse_transceivers(capture_text.clean_lines(text(*names)))


def by_interface(rows: list[dict]) -> dict[str, dict]:
    return {row["interface"]: row for row in rows}


class HuaweiTransceiverTest(unittest.TestCase):
    FILE = "huawei_s5735_display_interface_transceiver_then_lldp.txt"

    def test_t1_last_transceiver_keeps_its_own_serial(self):
        rows = by_interface(xcvrs(self.FILE))
        self.assertEqual(list(rows), ["10GE1/0/1", "10GE1/0/2", "10GE2/0/1", "10GE2/0/2"])
        self.assertEqual(rows["10GE2/0/2"]["serial_number"], "XV324219619002")
        # the two serials printed by the LLDP neighbour that follows
        serials_found = {row["serial_number"] for row in rows.values()}
        self.assertFalse(serials_found & {"7U7192926730", "4S5472471912"})

    def test_pid_is_the_vendor_part_number(self):
        for row in xcvrs(self.FILE):
            self.assertEqual(row["pid"], "02313URF")
            self.assertEqual(row["vendor_part_number"], "02313URF")
            self.assertEqual(row["type"], "1000BASE_LX10")

    def test_blocks_without_a_command_echo_are_still_read(self):
        lines = [
            " 10GE1/0/1 transceiver information:",
            "   Transceiver Type                      :1000BASE_LX10",
            "   Vendor Part Number                    :02313URF",
            "   Manu. Serial Number                   :CW500604051220",
            "<SW1>display lldp neighbor",
            "  Device serial number             :7U7192926730",
        ]
        rows = runtime_parser.parse_transceivers(lines)
        self.assertEqual([(r["interface"], r["serial_number"], r["pid"]) for r in rows], [("10GE1/0/1", "CW500604051220", "02313URF")])

    def test_device_serial_number_field_is_never_a_transceiver_serial(self):
        lines = [
            "<SW1>display interface transceiver",
            " 10GE1/0/1 transceiver information:",
            "   Transceiver Type                      :1000BASE_LX10",
            "   Device serial number                  :7U7192926730",
        ]
        self.assertEqual(runtime_parser.parse_transceivers(lines)[0]["serial_number"], "")

    def test_t7_mixed_prompts_use_only_the_captures_own_device(self):
        lines = capture_text.clean_lines(text("huawei_s5755_mixed_device_prompts_transceiver_runs.txt"))
        rows = runtime_parser.parse_transceivers(lines)
        # 4 ports of the capture's own device, listed once although the command ran twice
        self.assertEqual([r["interface"] for r in rows], ["25GE1/0/1", "25GE1/0/2", "25GE1/0/3", "25GE1/0/4"])
        self.assertFalse(any(r["interface"].startswith("10GE") for r in rows))
        note = runtime_parser.transceiver_scope_note(lines)
        self.assertEqual(note["used_host"], "MJC-OHXNAT-AFWOJX-K9640")
        self.assertEqual(note["ignored_runs"], 1)

    def test_t7_page_says_the_file_mixes_devices(self):
        result = analyze(text("huawei_s5755_mixed_device_prompts_transceiver_runs.txt"), filename="x.txt", vendor="Huawei")
        notes = " ".join(n["text"] for n in result["capture_notes"])
        self.assertIn("2 different prompts", notes)
        self.assertIn("MJC-OHXNAT-AFWOJX-K9640", notes)
        self.assertIn("left out", notes)
        self.assertEqual(len(result["transceivers"]), 4)

    def test_single_device_capture_has_no_notes(self):
        result = analyze(text(self.FILE), filename="x.txt", vendor="Huawei")
        self.assertEqual(result["capture_notes"], [])

    def test_a_capture_whose_own_device_never_ran_the_command_keeps_all_runs(self):
        lines = [
            "<A>display version", "<A>display device", "<A>display stack",
            "<B>display interface transceiver",
            " 10GE1/0/1 transceiver information:",
            "   Transceiver Type                      :SFP",
            "   Manu. Serial Number                   :ABC123456",
        ]
        self.assertEqual([r["serial_number"] for r in runtime_parser.parse_transceivers(lines)], ["ABC123456"])
        self.assertIsNone(runtime_parser.transceiver_scope_note(lines))


class ArubaTransceiverTest(unittest.TestCase):
    def test_aruba_table_gives_product_number_and_serial(self):
        rows = xcvrs("aruba_2540_show_interfaces_transceiver.txt")
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["interface"], rows[0]["type"], rows[0]["pid"], rows[0]["serial_number"]),
                         ("25", "1000LX", "J4859D", "CN13KC6CLS"))

    def test_t2_catalyst_dom_table_is_not_an_aruba_table(self):
        # The same command name on IOS prints temperature / voltage / power.
        self.assertEqual(xcvrs("cisco_2960_show_interfaces_transceiver_dom_only.txt"), [])


class CiscoTransceiverTest(unittest.TestCase):
    def test_inventory_gives_pid_and_serial_and_dom_is_joined(self):
        rows = by_interface(xcvrs("cisco_3650_show_inventory_sfp_and_dom.txt"))
        self.assertEqual(len(rows), 6)
        # third-party optic: the device prints an empty PID, which stays empty
        self.assertEqual((rows["GigabitEthernet1/1/1"]["pid"], rows["GigabitEthernet1/1/1"]["serial_number"]), ("", "EEZ5461Y3OX"))
        self.assertNotIn("tx_power_dbm", rows["GigabitEthernet1/1/1"])
        self.assertEqual((rows["GigabitEthernet1/1/2"]["pid"], rows["GigabitEthernet1/1/2"]["serial_number"]), ("GLC-SX-MMD", "EYS860688IG"))
        self.assertEqual(rows["GigabitEthernet1/1/2"]["tx_power_dbm"], "-5.6")
        self.assertEqual(rows["GigabitEthernet2/1/3"]["pid"], "GLC-LH-SMD")

    def test_t3_plural_command_joins_the_dom_table(self):
        inventory = 'NAME: "GigabitEthernet0/1", DESCR: "1000BaseLX SFP"\nPID: GLC-LH-SMD         , VID: V01  , SN: FNS1234ABCD\n'
        raw = "SW1#show inventory\n" + inventory + "\n" + text("cisco_2960_show_interfaces_transceiver_dom_only.txt")
        rows = runtime_parser.parse_transceivers(capture_text.clean_lines(raw))
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["interface"], rows[0]["serial_number"], rows[0]["pid"]), ("GigabitEthernet0/1", "FNS1234ABCD", "GLC-LH-SMD"))
        self.assertEqual((rows[0]["tx_power_dbm"], rows[0]["rx_power_dbm"]), ("-6.3", "-9.0"))

    def test_t6_inventory_printed_twice_lists_each_optic_once(self):
        block = text("cisco_3650_show_inventory_sfp_and_dom.txt")
        rows = runtime_parser.parse_transceivers(capture_text.clean_lines(block + "\n" + block))
        self.assertEqual(len(rows), 6)
        modules = runtime_parser.parse_inventory(capture_text.clean_lines(block + "\n" + block))["modules"]
        keys = [(m["slot"], m["pid"], m["serial"]) for m in modules]
        self.assertEqual(len(keys), len(set(keys)))

    def test_t4_nexus_pid_is_the_cisco_product_id(self):
        rows = by_interface(xcvrs("nxos_show_interface_transceiver_present_ports.txt"))
        self.assertEqual(list(rows), ["Ethernet1/49", "Ethernet1/52", "Ethernet1/53"])
        self.assertEqual(rows["Ethernet1/49"]["pid"], "QSFP-40G-SR4")
        self.assertEqual(rows["Ethernet1/52"]["pid"], "SFP-10G-SR")
        # the vendor's own part number is still available
        self.assertEqual(rows["Ethernet1/52"]["vendor_part_number"], "RTXM228-551-C98")
        self.assertEqual(rows["Ethernet1/52"]["serial_number"], "ACW281112VD")

    def test_t4_nexus_third_party_optic_falls_back_to_the_part_number(self):
        lines = [
            "SW# show interface transceiver",
            "Ethernet1/1", "   transceiver is present", "   type is 10Gbase-SR", "   name is FS",
            "   part number is SFP-10GSR-85", "   serial number is S2206123456",
        ]
        row = runtime_parser.parse_transceivers(lines)[0]
        self.assertEqual((row["pid"], row["serial_number"]), ("SFP-10GSR-85", "S2206123456"))


class SessionLogPrefixTest(unittest.TestCase):
    def test_t5_dateless_record_prefix_no_longer_hides_inventory_entries(self):
        clean = capture_text.clean_lines(text("nxos_show_inventory_half_stripped_log_prefix.txt"))
        self.assertFalse(any("|0000" in line for line in clean))
        names = [m["slot"] for m in runtime_parser.parse_inventory(clean)["modules"]]
        for expected in ("Chassis", "Slot 1", "Power Supply 1", "Power Supply 2", "Fan 1", "Fan 2", "Fan 3", "Fan 4"):
            self.assertIn(expected, names)

    def test_full_record_header_is_still_removed(self):
        line = "22/07/2026 14:02:44|000000000012|TERM|000000000647|show version"
        self.assertEqual(capture_text.clean_lines(line), ["show version"])

    def test_time_only_header_is_removed(self):
        self.assertEqual(capture_text.clean_lines('16:42:21|000000000135|NAME: "Slot 1"'), ['NAME: "Slot 1"'])

    def test_a_clock_in_ordinary_text_is_left_alone(self):
        for line in ("Up Time : 12:30:45", "12:30:45 UTC Mon Oct 5 2026", "Last change 12:30:45|x"):
            self.assertEqual(capture_text.clean_lines(line), [line])


class PromptHostTest(unittest.TestCase):
    def test_host_of_each_prompt_style(self):
        host = capture_text.prompt_host
        self.assertEqual(host("<SW-1>display lldp neighbor"), "SW-1")
        self.assertEqual(host("[SW-1-GigabitEthernet0/0/1]display this"), "SW-1")
        self.assertEqual(host("SW1(config)#show running-config"), "SW1")
        self.assertEqual(host("SW1# sh ver"), "SW1")
        self.assertEqual(host("<SW-1>dis"), "SW-1")
        self.assertEqual(host("[~CE1]display y"), "CE1")

    def test_table_text_is_never_a_prompt(self):
        for line in ("Gi1/1/2  25.9  3.29  -5.6  -5.5", "<SW1>", "Port Type Product", "10GE1/0/1 transceiver information:"):
            self.assertIsNone(capture_text.prompt_host(line), line)

    def test_counts_and_dominant_host(self):
        lines = ["<A>display version", "<A>display stack", "<B>display lldp neighbor", "table text"]
        self.assertEqual(capture_text.prompt_host_counts(lines), {"A": 2, "B": 1})
        self.assertEqual(capture_text.dominant_prompt_host(lines), "A")
        self.assertIsNone(capture_text.dominant_prompt_host(["no prompt here"]))


class PartsTest(unittest.TestCase):
    def test_kinds(self):
        kind = lambda slot, desc="", stack=False: parts.classify_module(slot, desc, is_stack=stack)
        self.assertEqual(kind("Chassis"), "Chassis")
        self.assertEqual(kind("Switch 1", "WS-C3650-24TS-S"), "Chassis")
        self.assertEqual(kind("Switch 1", "WS-C3650-24TS-S", True), "Stack member")
        self.assertEqual(kind("2", "WS-C2960X-24PS-L", True), "Stack member")
        self.assertEqual(kind("c36xx Stack", "c36xx Stack"), "Stack")
        self.assertEqual(kind("StackPort1/1", "StackPort1/1"), "Stack cable")
        self.assertEqual(kind("StackAdapter1/2"), "Stack cable")
        self.assertEqual(kind("Switch 1 - FlexStackPlus Module"), "Stack module")
        self.assertEqual(kind("Switch 2 - Power Supply B"), "Power supply")
        self.assertEqual(kind("1 PWR2", "PAC80S12-CN"), "Power supply")
        self.assertEqual(kind("Fan 3", "Nexus9000 C93108TC-EX chassis Fan Module"), "Fan")
        self.assertEqual(kind("1 FAN1"), "Fan")
        self.assertEqual(kind("Switch 1 FRU Uplink Module 1"), "Module")
        self.assertEqual(kind("Slot 1", "24x10GT + 6x40G/100G Ethernet Module"), "Module")
        # a power supply's description may mention the chassis; a chassis description never makes it a PSU
        self.assertEqual(kind("Chassis", "Nexus9000 C93108TC-EX chassis"), "Chassis")
        # interface-named entries are the transceivers' own rows
        self.assertEqual(kind("GigabitEthernet1/1/2", "1000BaseSX SFP"), "")
        self.assertEqual(kind("Gi1/1/2"), "")

    def test_cisco_stack_parts_list_each_optic_once_with_pid_and_serial(self):
        result = analyze(text("cisco_3650_show_inventory_sfp_and_dom.txt"), filename="x.txt", vendor="Cisco")
        rows = result["parts"]
        self.assertTrue(all(row["kind"] for row in rows))
        self.assertEqual({row["kind"] for row in rows},
                         {"Stack", "Stack member", "Stack cable", "Power supply", "Transceiver"})
        sfps = [row for row in rows if row["kind"] == "Transceiver"]
        self.assertEqual(len(sfps), 6)
        by_slot = {row["slot"]: row for row in sfps}
        self.assertEqual((by_slot["GigabitEthernet1/1/2"]["pid"], by_slot["GigabitEthernet1/1/2"]["serial"]),
                         ("GLC-SX-MMD", "EYS860688IG"))
        self.assertEqual(by_slot["GigabitEthernet1/1/2"]["description"], "1000BaseSX SFP")
        members = [row for row in rows if row["kind"] == "Stack member"]
        self.assertEqual([(r["slot"], r["pid"], r["serial"]) for r in members],
                         [("Switch 1", "WS-C3650-24TS-S", "XOJ9477F76F"), ("Switch 2", "WS-C3650-24TS-S", "XYY8914E1D6")])
        # hardware first, transceivers after it
        kinds = [row["kind"] for row in rows]
        self.assertEqual(kinds, sorted(kinds, key=lambda k: k == "Transceiver"))
        # nothing twice
        keys = [(r["slot"], r["kind"], r["pid"], r["serial"]) for r in rows]
        self.assertEqual(len(keys), len(set(keys)))

    def test_inventory_modules_are_unchanged(self):
        result = analyze(text("cisco_3650_show_inventory_sfp_and_dom.txt"), filename="x.txt", vendor="Cisco")
        slots = [m["slot"] for m in result["inventory"]["modules"]]
        self.assertIn("GigabitEthernet1/1/2", slots)  # still there for Compare / sizing

    def test_huawei_stack_parts(self):
        raw = text("huawei_s5735_stack_display_device_elabel_brief.txt", "huawei_s5735_display_interface_transceiver_then_lldp.txt")
        result = analyze(raw, filename="x.txt", vendor="Huawei")
        rows = result["parts"]
        members = [r for r in rows if r["kind"] == "Stack member"]
        self.assertEqual(len(members), 2)
        self.assertTrue(all(r["pid"] and r["serial"] for r in members))
        self.assertEqual(len([r for r in rows if r["kind"] == "Power supply"]), 4)
        optics = [r for r in rows if r["kind"] == "Transceiver"]
        self.assertEqual([(r["slot"], r["pid"]) for r in optics], [(f"10GE{n}", "02313URF") for n in ("1/0/1", "1/0/2", "2/0/1", "2/0/2")])
        self.assertEqual(optics[-1]["serial"], "XV324219619002")

    def test_nexus_parts_include_all_four_fans(self):
        result = analyze(text("nxos_show_inventory_half_stripped_log_prefix.txt", "nxos_show_interface_transceiver_present_ports.txt"),
                         filename="x.txt", vendor="Cisco")
        rows = result["parts"]
        self.assertEqual(len([r for r in rows if r["kind"] == "Fan"]), 4)
        optics = {r["slot"]: r for r in rows if r["kind"] == "Transceiver"}
        self.assertEqual(optics["Ethernet1/52"]["pid"], "SFP-10G-SR")

    def test_aruba_chassis_gets_model_and_product_number(self):
        raw = text("aruba_2540_show_system.txt", "aruba_2540_show_modules.txt", "aruba_2540_show_interfaces_transceiver.txt")
        raw = raw.replace("YJ90UAZ1G7", "VV24TXJ3CI")  # same chassis in both commands
        result = analyze(raw, filename="x.txt", vendor="Aruba")
        chassis = [r for r in result["parts"] if r["kind"] == "Chassis"]
        self.assertEqual([(r["description"], r["pid"], r["serial"]) for r in chassis], [("2540-24G-PoE+-4SFP+", "JL356A", "VV24TXJ3CI")])
        optic = [r for r in result["parts"] if r["kind"] == "Transceiver"]
        self.assertEqual([(r["slot"], r["pid"], r["serial"]) for r in optic], [("25", "J4859D", "CN13KC6CLS")])

    def test_aruba_without_show_modules_names_the_software_revision(self):
        result = analyze(text("aruba_2540_show_system.txt"), filename="x.txt", vendor="Aruba")
        chassis = [r for r in result["parts"] if r["kind"] == "Chassis"]
        self.assertEqual(len(chassis), 1)
        self.assertEqual(chassis[0]["description"], "Software revision YC.16.06.0006")
        self.assertEqual(chassis[0]["pid"], "")

    def test_aruba_chassis_helper_splits_product_number(self):
        rows = serials.aruba_chassis(capture_text.clean_lines(text("aruba_2540_show_modules.txt")))
        self.assertEqual(rows, [{"model": "2540-24G-PoE+-4SFP+", "product_number": "JL356A", "serial": "YJ90UAZ1G7"}])

    def test_empty_capture_gives_empty_parts_and_no_notes(self):
        result = analyze("hostname SW1\n", filename="x.txt", vendor="Cisco")
        self.assertEqual(result["parts"], [])
        self.assertEqual(result["capture_notes"], [])

    def test_build_parts_without_transceivers_or_modules(self):
        self.assertEqual(parts.build_parts([], []), [])


class MaskerKeepsHardwareLabelsTest(unittest.TestCase):
    """`[PWR1]` (a section header of `display device`) looked like a bracketed
    Huawei prompt, so FAN1 / PWR1 / PWR2 were masked as hostnames and the Sub
    column of `display device elabel brief` read "GJI6 / IBM5 / LQI3"."""

    @classmethod
    def setUpClass(cls):
        import importlib.util

        spec = importlib.util.spec_from_file_location("mask_capture", FIXTURES / "mask_capture.py")
        cls.mask = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.mask)

    def test_hardware_labels_are_not_hostnames(self):
        names = self.mask.collect_names(["[PWR1]", "[FAN2]", "[PSU-1]", "<REAL-HOST-01>display version"], [])
        self.assertEqual(names, {"REAL-HOST-01"})

    def test_elabel_sub_rows_keep_their_labels(self):
        raw = "\n".join([
            "<REAL-HOST-01>display device elabel brief",
            "SlotID     Sub    Type                     SN                       P/N",
            "1          --     S5735-S24T4XE-V2         7U7192926730             98012016-003",
            "           FAN1   --                       --                       --",
            "           PWR1   PAC80S12-CN              2078374294XWR2888531     02131835",
            "[PWR1]",
        ])
        masked, _originals = self.mask.mask(raw, [])
        self.assertIn("FAN1", masked)
        self.assertIn("PWR1   PAC80S12-CN", masked)
        self.assertNotIn("2078374294XWR2888531", masked)


class ExcelSerialSheetTest(unittest.TestCase):
    def test_sheet_has_kind_and_the_transceivers(self):
        result = analyze(text("cisco_3650_show_inventory_sfp_and_dom.txt"), filename="sw1.txt", vendor="Cisco")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            SwitchExcelExporter(result).export(path)
            sheet = load_workbook(path)["Serial Numbers"]
            rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Hostname", "Slot", "Kind", "Description", "PID", "Serial Number"))
        body = {(r[1], r[2]): (r[4], r[5]) for r in rows[1:]}
        self.assertEqual(body[("GigabitEthernet1/1/2", "Transceiver")], ("GLC-SX-MMD", "EYS860688IG"))
        self.assertEqual(body[("Switch 2", "Stack member")], ("WS-C3650-24TS-S", "XYY8914E1D6"))
        self.assertEqual(len(rows) - 1, len(result["parts"]))

    def test_a_result_without_parts_still_exports_the_inventory_modules(self):
        result = analyze(text("cisco_3650_show_inventory_sfp_and_dom.txt"), filename="sw1.txt", vendor="Cisco")
        result.pop("parts")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            SwitchExcelExporter(result).export(path)
            rows = list(load_workbook(path)["Serial Numbers"].iter_rows(values_only=True))
        self.assertEqual(len(rows) - 1, len(result["inventory"]["modules"]))
        self.assertTrue(all(r[2] in (None, "") for r in rows[1:]))


if __name__ == "__main__":
    unittest.main()

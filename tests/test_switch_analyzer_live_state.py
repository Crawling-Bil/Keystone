"""Vendor detection, interface state (up / down / admin down), ARP and MAC
tables for the Switch Analyzer, checked against masked excerpts of the real
Oct 2026 TAM captures (see tests/fixtures/real_captures/README.md).

What went wrong before and what is pinned here
  V1  A Catalyst 2960S capture that had no running-config at all (only show
      commands) was reported as Aruba. The device's own `show version` /
      inventory output now decides the vendor, a model prefix (WS-C...) is
      checked against it, and a vendor picked by hand that contradicts the
      capture is reported instead of silently followed.
  S1  `shutdown` / `disable` ports were folded into "down" (Cisco, Huawei) or
      shown "up" (Aruba, Nexus). A port is now one of up / down / admin down /
      unknown, read from the live state of each platform.
  A1  ARP: 4 layouts (IOS, NX-OS, ArubaOS-Switch, Huawei VRP), including the
      messy captures (session-log records, redrawn screens, lost command
      echoes) and a table the device says is longer than what was captured.
  M1  MAC table: the same 4 layouts, switch-own CPU rows flagged.
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from features.switch_analyzer import capture_text, fingerprint, live_state, runtime_parser
from features.switch_analyzer.compare import _normalize_mac, compare
from features.switch_analyzer.exporters import SwitchExcelExporter
from features.switch_analyzer.service import analyze
from features.switch_analyzer.sizing import build_sizing_summary

FIXTURES = Path(__file__).parent / "fixtures" / "real_captures"


def text(name: str) -> str:
    return (FIXTURES / name).read_text(encoding="utf-8")


def clean(name: str) -> list[str]:
    return capture_text.clean_lines(text(name))


def states(name: str) -> dict[str, str]:
    observed = live_state.observe_interface_states(clean(name))
    return {key: live_state.resolve_state(items)["state"] for key, items in observed.items()}


def states_from(raw: str) -> dict[str, str]:
    observed = live_state.observe_interface_states(capture_text.clean_lines(raw))
    return {key: live_state.resolve_state(items)["state"] for key, items in observed.items()}


def tally(mapping: dict[str, str]) -> dict[str, int]:
    out: dict[str, int] = {}
    for value in mapping.values():
        out[value] = out.get(value, 0) + 1
    return out


CISCO_2960S = "\n".join(
    text(name)
    for name in (
        "cisco_2960s_show_version.txt",
        "cisco_2960s_show_interfaces_status.txt",
        "cisco_2960s_show_ip_interface_brief.txt",
        "cisco_2960s_show_ip_arp.txt",
        "cisco_2960s_show_mac_address_table.txt",
    )
)

# Synthetic wrapper (not a capture): a hostname and a few config lines so the
# vendor can be identified, placed around a real excerpt.
ARUBA_CONFIG = (
    "; JL356A Configuration Editor; Created on release #YC.16.06.0006\n"
    'hostname "EXAMPLE-ACCESS-1"\n'
    'vlan 1\n   name "DEFAULT_VLAN"\n   untagged 1-28\n   exit\n'
    "interface 1\n   disable\n   exit\n"
    "interface 3\n   exit\n"
)
HUAWEI_CONFIG = "sysname EXAMPLE-S5735\n#\nvlan batch 10\n#\ninterface Vlanif10\n ip address 10.255.0.254 255.255.255.0\n#\n"


class VendorDetectionTest(unittest.TestCase):
    def test_config_less_2960s_is_cisco_not_aruba(self):
        detection = fingerprint.detect(text("cisco_2960s_show_version.txt"))
        self.assertEqual(detection.vendor, "Cisco")
        self.assertEqual(detection.confidence, "high")
        self.assertEqual(detection.model_hint, "WS-C2960S-F48TS-L")

    def test_analyze_reports_cisco_for_the_whole_show_only_capture(self):
        result = analyze(CISCO_2960S, filename="show-only.txt")
        self.assertEqual(result["vendor"], "Cisco")
        self.assertEqual(result["os_version"]["model"], "WS-C2960S-F48TS-L")
        self.assertEqual(result["detection"]["method"], "device identity")
        self.assertEqual(result["detection"]["warning"], "")

    def test_a_neighbor_that_names_another_vendor_cannot_hijack_the_result(self):
        # CDP/LLDP output describes OTHER devices; only the device's own
        # identity sections may decide the vendor.
        noisy = text("cisco_2960s_show_version.txt") + "\nPlatform: Aruba JL356A\nSystem Name: ARUBA-NEIGHBOR\n"
        self.assertEqual(fingerprint.detect(noisy).vendor, "Cisco")
        self.assertEqual(analyze(noisy, filename="x.txt")["vendor"], "Cisco")

    def test_a_vendor_picked_by_hand_that_contradicts_the_capture_is_flagged(self):
        result = analyze(text("cisco_2960s_show_version.txt"), filename="x.txt", vendor="Aruba")
        self.assertIn("Cisco", result["detection"]["warning"])
        self.assertEqual(result["detection"]["method"], "selected by user")

    def test_model_prefixes(self):
        for model, vendor in (
            ("WS-C2960S-F48TS-L", "Cisco"), ("C9300-48P", "Cisco"), ("N9K-C9508", "Cisco"),
            ("JL356A", "Aruba"), ("J9779A", "Aruba"),
            ("S5735-L48T4S-A1", "Huawei"), ("S5755-28C-HI", "Huawei"),
            ("ABC123", ""),
        ):
            self.assertEqual(fingerprint.vendor_from_model(model), vendor, model)

    def test_aruba_and_huawei_are_still_identified(self):
        aruba = analyze(ARUBA_CONFIG + "\n" + text("aruba_2540_show_interface_status_redraw.txt"), filename="a.txt")
        self.assertEqual((aruba["vendor"], aruba["detection"]["method"]), ("Aruba", "device identity"))
        huawei = analyze(HUAWEI_CONFIG + text("huawei_s5735_display_interface_brief.txt"), filename="h.txt")
        self.assertEqual((huawei["vendor"], huawei["detection"]["method"]), ("Huawei", "device identity"))


class InterfaceStateTest(unittest.TestCase):
    def test_cisco_show_interfaces_status_separates_disabled_from_notconnect(self):
        result = states("cisco_2960s_show_interfaces_status.txt")
        self.assertEqual(tally(result), {"up": 7, "down": 45, "admin down": 2})
        self.assertEqual(result["fa1/0/47"], "admin down")

    def test_cisco_ip_interface_brief_administratively_down(self):
        result = states("cisco_2960s_show_ip_interface_brief.txt")
        self.assertEqual(result["fa0"], "admin down")
        self.assertEqual(result["vlan10"], "up")

    def test_cisco_combined_capture_counts_add_up_and_explain_themselves(self):
        result = analyze(CISCO_2960S, filename="c.txt")
        cards = result["cards"]
        self.assertEqual((cards["interfaces_up"], cards["interfaces_down"], cards["interfaces_admin_down"], cards["interfaces_unknown"]), (9, 45, 2, 0))
        self.assertEqual(
            cards["total_interfaces"],
            cards["interfaces_up"] + cards["interfaces_down"] + cards["interfaces_admin_down"] + cards["interfaces_unknown"],
        )
        shut = {row["name"]: row for row in result["interfaces"] if row["status"] == "admin down"}
        self.assertEqual(set(shut), {"FastEthernet1/0/47", "FastEthernet0"})
        self.assertEqual(shut["FastEthernet1/0/47"]["status_detail"], "disabled")
        self.assertIn("show interfaces status", shut["FastEthernet1/0/47"]["status_source"])

    def test_err_disabled_is_a_fault_not_an_intentional_shutdown(self):
        table = (
            "SW1#show interfaces status\n"
            "Port      Name               Status       Vlan       Duplex  Speed Type\n"
            "Gi1/0/1                      err-disabled 10           auto    auto  10/100/1000BaseTX\n"
            "Gi1/0/2                      disabled     10           auto    auto  10/100/1000BaseTX\n"
            "Gi1/0/3                      connected    10           a-full  a-1000 10/100/1000BaseTX\n"
        )
        result = live_state.observe_interface_states(capture_text.clean_lines(table))
        resolved = {key: live_state.resolve_state(items)["state"] for key, items in result.items()}
        self.assertEqual(resolved, {"ge1/0/1": "down", "ge1/0/2": "admin down", "ge1/0/3": "up"})

    def test_nxos_brief_reads_the_reason_column(self):
        result = states("nxos_show_interface_brief.txt")
        self.assertEqual(tally(result), {"up": 32, "down": 8, "admin down": 27})
        self.assertEqual(result["eth1/12"], "admin down")  # Reason: Administratively down
        self.assertEqual(result["eth1/1"], "down")  # Reason: Link not connected
        self.assertEqual(result["vlan1"], "admin down")

    def test_nxos_status_disabled_is_admin_down(self):
        result = states("nxos_show_interface_status.txt")
        self.assertEqual(result["eth1/12"], "admin down")
        self.assertEqual(result["eth1/2"], "up")

    def test_config_shutdown_is_read_per_interface_block(self):
        config = (
            "interface Ethernet1/12\n  description spare\n"
            "interface Ethernet1/13\n  shutdown\n"
            "interface Ethernet1/14\n  shutdown\n  no shutdown\n"
        )
        admin = live_state.config_admin_state(capture_text.clean_lines(config), "Cisco")
        self.assertTrue(admin[live_state.interface_key("Ethernet1/13")]["admin_down"])
        self.assertFalse(admin[live_state.interface_key("Ethernet1/12")]["admin_down"])
        self.assertFalse(admin[live_state.interface_key("Ethernet1/14")]["admin_down"])

    def test_huawei_and_aruba_config_keywords(self):
        huawei = live_state.config_admin_state(["interface GigabitEthernet0/0/1", " shutdown", "#", "interface GigabitEthernet0/0/2", " undo shutdown", "#"], "Huawei")
        self.assertTrue(huawei[live_state.interface_key("GigabitEthernet0/0/1")]["admin_down"])
        self.assertFalse(huawei[live_state.interface_key("GigabitEthernet0/0/2")]["admin_down"])
        aruba = live_state.config_admin_state(["interface 1", "   disable", "   exit", "interface 2", "   exit"], "Aruba")
        self.assertTrue(aruba[live_state.interface_key("1")]["admin_down"])
        self.assertFalse(aruba[live_state.interface_key("2")]["admin_down"])

    def test_huawei_star_down_is_admin_down(self):
        result = states("huawei_s5735_display_interface_brief.txt")
        self.assertEqual(tally(result), {"admin down": 19, "up": 11, "down": 6})

    def test_huawei_analyze_keeps_phy_state_and_explains_admin_down(self):
        result = analyze(HUAWEI_CONFIG + text("huawei_s5735_display_interface_brief.txt"), filename="h.txt")
        self.assertEqual(
            (result["cards"]["interfaces_up"], result["cards"]["interfaces_down"], result["cards"]["interfaces_admin_down"]),
            (11, 6, 19),
        )
        row = next(item for item in result["interfaces"] if item["name"] == "10GE1/0/1")
        self.assertEqual(row["status"], "admin down")
        self.assertIn("administratively down", row["status_detail"])

    def test_aruba_disable_in_config_plus_down_link_is_admin_down(self):
        result = analyze(ARUBA_CONFIG + "\n" + text("aruba_2540_show_interface_status_redraw.txt"), filename="a.txt")
        by_name = {row["name"]: row for row in result["interfaces"]}
        self.assertEqual(by_name["1"]["status"], "admin down")
        self.assertIn("disable in config", by_name["1"]["status_detail"])
        self.assertEqual(by_name["2"]["status"], "down")  # link down, nothing says it was disabled
        self.assertEqual(by_name["3"]["status"], "up")

    def test_aruba_brief_enabled_no_is_admin_down(self):
        # `show interfaces brief`: Enabled = No is the port turned off by the engineer;
        # Status Down with Enabled = Yes is a port with no link.
        table = text("aruba_2540_show_interfaces_brief.txt").replace("| No        Yes     Down   1000FDx    NA", "| No        No      Down   1000FDx    NA", 1)
        self.assertNotEqual(table, text("aruba_2540_show_interfaces_brief.txt"))
        result = states_from(table)
        self.assertEqual(result["1"], "up")
        self.assertEqual(result["2"], "admin down")
        self.assertEqual(result["3"], "down")

    def test_config_only_port_without_live_state_is_unknown_not_up(self):
        config = 'hostname "EXAMPLE-ACCESS-1"\n; JL356A Configuration Editor; Created on release #YC.16.06.0006\ninterface 4\n   exit\ninterface 5\n   disable\n   exit\n'
        result = analyze(config, filename="cfg.txt")
        by_name = {row["name"]: row["status"] for row in result["interfaces"]}
        self.assertEqual(by_name["4"], "unknown")
        self.assertEqual(by_name["5"], "admin down")
        self.assertEqual(result["cards"]["interfaces_unknown"], 1)

    def test_sizing_counts_the_four_states_separately(self):
        result = analyze(CISCO_2960S, filename="c.txt")
        summary = build_sizing_summary([result])
        totals = summary["totals"]
        self.assertEqual(totals["physical_ports_admin_down"], 2)
        self.assertEqual(
            totals["physical_port_count"],
            totals["physical_ports_up"] + totals["physical_ports_down"] + totals["physical_ports_admin_down"] + totals["physical_ports_unknown"],
        )
        self.assertEqual(summary["devices"][0]["physical_ports_admin_down"], 2)


class ArpTableTest(unittest.TestCase):
    def test_ios_dynamic_and_own_address(self):
        rows, meta = live_state.parse_arp_with_meta(clean("cisco_2960s_show_ip_arp.txt"))
        self.assertEqual(len(rows), 10)
        self.assertEqual([row["type"] for row in rows].count("static"), 1)  # age "-" = the switch's own address
        self.assertEqual(rows[0]["mac_address"], "9afa.6c2b.97c2")
        self.assertEqual(rows[0]["interface"], "Vlan10")
        self.assertTrue(meta["complete"])

    def test_session_log_with_typed_keys_and_tab_completion(self):
        # STDI/TERM/KEYS records, a half-typed `s h   i p   a r p` and the
        # command echoed twice: the table is still read once, whole.
        rows = live_state.parse_arp(clean("cisco_2960s_session_log_show_ip_arp.txt"))
        self.assertEqual(len(rows), 8)
        self.assertEqual(rows[0]["ip_address"], "10.255.0.1")
        self.assertEqual(rows[-1]["type"], "static")

    def test_nxos_header_found_table_with_lost_echo_and_truncated_incomplete(self):
        rows, meta = live_state.parse_arp_with_meta(clean("nxos_show_ip_arp_lost_echo_truncated.txt"))
        self.assertEqual(len(rows), 29)
        self.assertEqual(rows[0]["age"], "00:05:47")
        incomplete = [row for row in rows if row["type"] == "incomplete"]
        self.assertEqual(len(incomplete), 11)
        self.assertTrue(all(row["mac_address"] == "" for row in incomplete))
        # "INCOMPLET" (the capture lost the last letter) is still an incomplete entry.
        self.assertEqual(rows[-1]["type"], "incomplete")
        self.assertEqual([row["type"] for row in rows].count("static"), 4)

    def test_a_table_shorter_than_the_device_reported_is_flagged(self):
        _rows, meta = live_state.parse_arp_with_meta(clean("nxos_show_ip_arp_lost_echo_truncated.txt"))
        self.assertEqual(meta["reported"], 204)
        self.assertEqual(meta["parsed"], 29)
        self.assertFalse(meta["complete"])
        result = analyze(
            "hostname EXAMPLE-NX\nfeature vpc\n" + text("nxos_show_ip_arp_lost_echo_truncated.txt"), filename="nx.txt"
        )
        self.assertFalse(result["table_checks"]["arp"]["complete"])
        self.assertEqual(result["table_checks"]["arp"]["reported"], 204)

    def test_huawei_interface_row_and_dynamic_row(self):
        rows, meta = live_state.parse_arp_with_meta(clean("huawei_s5735_display_arp.txt"))
        self.assertEqual([row["type"] for row in rows], ["interface", "dynamic"])
        self.assertEqual((rows[1]["vlan"], rows[1]["age"], rows[1]["interface"]), ("10", "9", "Eth-Trunk1"))
        self.assertEqual((meta["reported"], meta["parsed"], meta["complete"]), (2, 2, True))

    def test_aruba_arp(self):
        rows = live_state.parse_arp(clean("aruba_2540_show_arp.txt"))
        self.assertEqual(len(rows), 1)
        self.assertRegex(rows[0]["mac_address"], r"^[0-9a-f]{6}-[0-9a-f]{6}$")  # ArubaOS-Switch notation, kept as printed
        self.assertEqual((rows[0]["interface"], rows[0]["type"]), ("25", "dynamic"))

    def test_legacy_runtime_parser_keeps_its_four_key_shape(self):
        rows = runtime_parser.parse_arp_table(text("cisco_2960s_show_ip_arp.txt").splitlines())
        self.assertEqual(len(rows), 10)
        self.assertEqual(set(rows[0]), {"ip_address", "mac_address", "age", "interface"})

    def test_arp_ends_up_in_the_analysis_result(self):
        result = analyze(CISCO_2960S, filename="c.txt")
        self.assertEqual(len(result["arp_table"]), 10)
        self.assertEqual(result["cards"]["arp_entries"], 10)


class MacTableTest(unittest.TestCase):
    def test_ios_table_flags_the_switch_own_cpu_rows(self):
        rows, meta = live_state.parse_mac_table_with_meta(clean("cisco_2960s_show_mac_address_table.txt"))
        self.assertEqual(len(rows), 62)
        self.assertEqual((meta["reported"], meta["complete"]), (62, True))
        system = [row for row in rows if row["system"]]
        self.assertEqual(len(system), 20)
        self.assertTrue(all(row["interface"] == "CPU" and row["vlan"] == "All" for row in system))
        host = next(row for row in rows if not row["system"])
        self.assertEqual((host["type"], host["interface"]), ("dynamic", "Fa1/0/48"))

    def test_nxos_table_printed_after_the_wrong_command_echo(self):
        rows = live_state.parse_mac_table(clean("nxos_show_mac_address_table_lost_echo.txt"))
        self.assertEqual(len(rows), 30)
        self.assertEqual((rows[0]["vlan"], rows[0]["interface"], rows[0]["age"]), ("201", "Eth1/7", "0"))

    def test_nxos_gateway_rows_on_the_supervisor_are_system_entries(self):
        table = (
            "SW# show mac address-table\n"
            "   VLAN     MAC Address      Type      age     Secure NTFY Ports\n"
            "---------+-----------------+--------+---------+------+----+------------------\n"
            "G  215     0011.2233.4455   static   -         F      F    sup-eth1(R)\n"
            "*  201     0011.2233.4466   dynamic  0         F      F    Eth1/3\n"
        )
        rows = live_state.parse_mac_table(capture_text.clean_lines(table))
        self.assertEqual([row["system"] for row in rows], [True, False])

    def test_huawei_table_with_reported_total(self):
        rows, meta = live_state.parse_mac_table_with_meta(clean("huawei_s5735_display_mac_address.txt"))
        self.assertEqual(len(rows), 38)
        self.assertEqual((meta["reported"], meta["source"], meta["complete"]), (38, "Total items", True))
        self.assertEqual((rows[0]["vlan"], rows[0]["interface"], rows[0]["type"], rows[0]["age"]), ("1", "Eth-Trunk1", "dynamic", "717"))

    def test_aruba_table_has_no_type_column(self):
        rows = live_state.parse_mac_table(clean("aruba_2540_show_mac_address.txt"))
        self.assertEqual(len(rows), 21)
        self.assertRegex(rows[0]["mac_address"], r"^[0-9a-f]{6}-[0-9a-f]{6}$")
        self.assertEqual((rows[0]["interface"], rows[0]["vlan"], rows[0]["type"]), ("25", "1", ""))

    def test_mac_notation_is_normalised_for_comparison(self):
        for notation in ("0011.2233.4455", "0011-2233-4455", "001122-334455", "00:11:22:33:44:55"):
            self.assertEqual(_normalize_mac(notation), "001122334455", notation)
        self.assertEqual(_normalize_mac("not a mac"), "")


class CompareHostsTest(unittest.TestCase):
    def _result(self, hostname, macs, arp):
        return {
            "hostname": hostname, "vendor": "x", "os_version": {}, "inventory": {}, "cards": {"arp_entries": len(arp), "mac_entries": len(macs)},
            "vlans": [], "routes": [], "dns": {"servers": []}, "pbr": {},
            "mac_table": macs, "arp_table": arp,
        }

    def test_hosts_match_across_vendors_and_system_rows_are_ignored(self):
        old = self._result(
            "OLD",
            [
                {"vlan": "10", "mac_address": "0011.2233.4455", "interface": "Gi1/0/1", "type": "dynamic", "system": False},
                {"vlan": "10", "mac_address": "aaaa.bbbb.cccc", "interface": "Gi1/0/2", "type": "dynamic", "system": False},
                {"vlan": "All", "mac_address": "0100.0ccc.cccc", "interface": "CPU", "type": "static", "system": True},
            ],
            [{"ip_address": "10.0.0.5", "mac_address": "0011.2233.4455", "interface": "Vlan10"}],
        )
        new = self._result(
            "NEW",
            [{"vlan": "10", "mac_address": "0011-2233-4455", "interface": "GE0/0/1", "type": "dynamic", "system": False}],
            [{"ip_address": "10.0.0.5", "mac_address": "0011-2233-4455", "interface": "Vlanif10"}],
        )
        diff = compare(old, new)
        self.assertEqual(diff["mac_addresses"]["both"], ["001122334455"])
        self.assertEqual(diff["mac_addresses"]["only_a"], ["aaaabbbbcccc"])  # the CPU row is not a host
        self.assertEqual(diff["mac_addresses"]["only_b"], [])
        self.assertEqual(diff["arp_hosts"]["both"], ["10.0.0.5 = 001122334455"])
        self.assertEqual(diff["device_a"]["arp_entries"], 1)
        self.assertEqual(diff["device_a"]["mac_entries"], 3)

    def test_summary_carries_admin_down_and_unknown_counts(self):
        old = analyze(CISCO_2960S, filename="c.txt")
        new = analyze(HUAWEI_CONFIG + text("huawei_s5735_display_interface_brief.txt"), filename="h.txt")
        diff = compare(old, new)
        self.assertEqual((diff["device_a"]["interfaces_admin_down"], diff["device_b"]["interfaces_admin_down"]), (2, 19))
        self.assertEqual(diff["device_a"]["interfaces_unknown"], 0)


class ExcelExportTest(unittest.TestCase):
    def test_workbook_has_mac_sheet_state_columns_and_arp_detail(self):
        result = analyze(CISCO_2960S, filename="c.txt")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            SwitchExcelExporter([result]).export(path)
            book = load_workbook(path)
        self.assertIn("MAC Table", book.sheetnames)
        mac = book["MAC Table"]
        self.assertEqual([c.value for c in mac[1]][:6], ["Hostname", "VLAN", "MAC Address", "Type", "Interface", "Age"])
        self.assertEqual(mac.max_row - 1, 62)
        header = [c.value for c in book["Interfaces"][1]]
        self.assertIn("Status Detail", header)
        self.assertIn("Status Source", header)
        info = [c.value for c in book["Device Info"][1]]
        self.assertIn("Interfaces Admin Down", info)
        self.assertEqual(book["ARP Table"].max_row - 1, 10)
        status_col = header.index("Status")
        statuses = [row[status_col].value for row in book["Interfaces"].iter_rows(min_row=2)]
        self.assertEqual(statuses.count("admin down"), 2)
        self.assertEqual(statuses.count("up"), 9)


if __name__ == "__main__":
    unittest.main()

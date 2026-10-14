"""Live HSRP and OSPF state of a Nexus pair (features/switch_analyzer/fhrp_live.py, ospf_live.py, fhrp_live.pair_check).

The layout is that of two real Nexus 9000 captures taken at the CLI of a vPC pair (``show hsrp``, ``show hsrp brief``, ``show hsrp summary``,
``show hsrp all``, ``show ip ospf neighbors``, ``show ip ospf neighbors summary vrf all`` and one rejected ``show ip ospf brief vrf all``);
every name, address, group number and string in the two fixtures is invented (RFC 5737 documentation addresses only).

Documentation: Cisco Nexus 9000 NX-OS Unicast Routing Configuration Guide, "Configuring HSRP", and the NX-API CLI reference for ``show hsrp`` / ``show ip ospf neighbors``.
"""
from __future__ import annotations

import ipaddress
import json
import re
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from features.switch_analyzer import capture_text, fhrp_live, ospf_live
from features.switch_analyzer.compare import compare
from features.switch_analyzer.exporters import export_comparison
from features.switch_analyzer.service import analyze

FIXTURES = Path(__file__).parent / "fixtures" / "real_captures"
CAPTURE_A = (FIXTURES / "synthetic_nxos_hsrp_ospf_vpc_a.txt").read_text(encoding="utf-8")
CAPTURE_B = (FIXTURES / "synthetic_nxos_hsrp_ospf_vpc_b.txt").read_text(encoding="utf-8")

_CACHE: dict[str, dict] = {}


def run(name: str, text: str) -> dict:
    if name not in _CACHE:
        _CACHE[name] = analyze(text, filename=f"{name}.txt")
    return _CACHE[name]


def groups_by_key(result: dict) -> dict[tuple[str, str], dict]:
    return {(g["interface"], g["group"]): g for g in result["fhrp"]["groups"]}


def finding_ids(result: dict) -> set[str]:
    return {f["id"] for f in result["findings"]}


class FixtureHygieneTests(unittest.TestCase):
    def test_only_documentation_addresses_are_in_the_fixtures(self):
        networks = [ipaddress.ip_network(n) for n in ("192.0.2.0/24", "198.51.100.0/24", "203.0.113.0/24")]
        for text in (CAPTURE_A, CAPTURE_B):
            for match in re.finditer(r"(?<![\d.])(\d{1,3}(?:\.\d{1,3}){3})(?![\d.])", text):
                self.assertTrue(any(ipaddress.ip_address(match.group(1)) in net for net in networks), match.group(1))

    def test_the_two_devices_are_invented_names(self):
        self.assertIn("NXDEMO-A#", CAPTURE_A)
        self.assertIn("NXDEMO-B#", CAPTURE_B)


class DetectionTests(unittest.TestCase):
    def test_a_capture_with_show_commands_only_is_still_recognised_as_cisco(self):
        result = run("a", CAPTURE_A)
        self.assertEqual(result["vendor"], "Cisco")
        self.assertEqual(result["detection"]["confidence"], "high")
        self.assertTrue(any("NX-OS" in item for item in result["detection"]["evidence"]))

    def test_the_hostname_comes_from_the_prompt_when_there_is_no_configuration(self):
        self.assertEqual(run("a", CAPTURE_A)["hostname"], "NXDEMO-A")
        self.assertEqual(run("b", CAPTURE_B)["hostname"], "NXDEMO-B")

    def test_crlf_line_endings_give_the_same_groups(self):
        crlf = analyze(CAPTURE_A.replace("\n", "\r\n"), filename="crlf.txt")
        self.assertEqual(len(crlf["fhrp"]["groups"]), 11)
        self.assertEqual(len(crlf["ospf_live"]["neighbors"]), 5)


class HsrpReaderTests(unittest.TestCase):
    def setUp(self):
        self.a = run("a", CAPTURE_A)
        self.b = run("b", CAPTURE_B)

    def test_every_group_is_listed_including_the_one_show_hsrp_brief_hides(self):
        a = groups_by_key(self.a)
        self.assertEqual(len(a), 11)
        disabled = a[("Vlan30", "30")]
        self.assertEqual(disabled["state"], "Disabled")
        self.assertEqual(disabled["state_detail"], "virtual IP not configured")
        self.assertEqual(disabled["sources"], ["detail"])           # `show hsrp brief` does not print a Disabled group
        self.assertEqual(sorted(a[("Vlan10", "10")]["sources"]), ["brief", "detail"])

    def test_the_fields_of_one_group(self):
        group = groups_by_key(self.a)[("Vlan10", "10")]
        self.assertEqual((group["state"], group["priority"], group["cfg_priority"], group["preempt"]), ("Active", "110", "110", "on"))
        self.assertEqual((group["virtual_ip"], group["vmac"], group["name"], group["version"]), ("192.0.2.1", "0000.0c07.ac0a", "lab-users", "1"))
        self.assertEqual((group["hello"], group["hold"]), ("3 s", "10 s"))
        self.assertEqual((group["active"], group["standby"], group["standby_priority"]), ("local", "192.0.2.254", "100"))
        self.assertEqual((group["changes"], group["last_change"]), ("5", "1w4d"))

    def test_the_interface_and_group_number_may_differ(self):
        # Vlan60 runs group 61: the table is keyed by interface + group, never by the VLAN number
        group = groups_by_key(self.a)[("Vlan60", "61")]
        self.assertEqual(group["priority"], "90")
        self.assertEqual(group["cfg_priority"], "110")
        self.assertIn(("Ethernet1/10", "70"), groups_by_key(self.a))      # a routed port is a group holder too

    def test_a_default_group_name_is_marked_as_default(self):
        group = groups_by_key(self.a)[("Vlan90", "90")]
        self.assertEqual(group["name_source"], "default")

    def test_authentication_kinds(self):
        a = groups_by_key(self.a)
        self.assertEqual(a[("Vlan10", "10")]["auth"], "none (default string)")
        self.assertTrue(a[("Vlan10", "10")]["auth_default"])
        self.assertEqual(a[("Vlan60", "61")]["auth"], "md5 (key chain)")
        self.assertEqual(a[("Vlan80", "80")]["auth"], "text")

    def test_the_plain_text_authentication_string_is_never_kept(self):
        self.assertIn("demo-not-a-secret", CAPTURE_A)
        dump = json.dumps(self.a, default=str)
        self.assertNotIn("demo-not-a-secret", dump)

    def test_the_summary_is_read_and_agrees_with_the_listed_groups(self):
        summary = self.a["fhrp"]["summary"]
        self.assertEqual((summary["total_groups"], summary["tx_fail"], summary["unknown_groups"]), (11, 0, 1000))
        check = self.a["table_checks"]["hsrp"]
        self.assertEqual((check["reported"], check["parsed"], check["complete"]), (11, 11, True))
        self.assertEqual(self.b["table_checks"]["hsrp"]["reported"], 10)

    def test_a_summary_that_reports_more_groups_than_listed_is_flagged(self):
        result = analyze(CAPTURE_A.replace("Total Groups: 11", "Total Groups: 14"), filename="x.txt")
        check = result["table_checks"]["hsrp"]
        self.assertEqual((check["reported"], check["parsed"], check["complete"]), (14, 11, False))
        self.assertIn("SW-CAP-006", finding_ids(result))

    def test_findings_of_device_a(self):
        ids = finding_ids(self.a)
        self.assertTrue({"SW-HSRP-001", "SW-HSRP-002", "SW-HSRP-003", "SW-HSRP-004", "SW-HSRP-005", "SW-HSRP-006", "SW-HSRP-007"} <= ids)
        self.assertNotIn("SW-HSRP-009", ids)                       # Tx Fail is 0

    def test_findings_of_device_b(self):
        ids = finding_ids(self.b)
        self.assertTrue({"SW-HSRP-002", "SW-HSRP-003", "SW-HSRP-005", "SW-HSRP-006"} <= ids)
        self.assertFalse({"SW-HSRP-001", "SW-HSRP-004", "SW-HSRP-007"} & ids)

    def test_a_priority_below_the_configured_one_names_the_group(self):
        item = next(f for f in self.a["findings"] if f["id"] == "SW-HSRP-007")
        self.assertIn("Vlan60", item["detail"] + item["object"])

    def test_the_table_hides_columns_nobody_uses(self):
        table = next(t for t in self.a["extra_tables"] if t["id"] == "fhrp-groups")
        keys = [c["key"] for c in table["columns"]]
        self.assertIn("state", keys)
        state = next(c for c in table["columns"] if c["key"] == "state")
        self.assertTrue(state["label"].startswith("State"))
        self.assertTrue(state.get("tones"))

    def test_hsrp_brief_alone_gives_the_state_without_the_detail(self):
        text = "SW1# show hsrp brief\n" + "\n".join(line for line in CAPTURE_A.splitlines() if line.startswith((" Interface", "  Vlan", "  Ethernet", "*:IPv6", "                     "))) + "\nSW1#\n"
        view = fhrp_live.read(capture_text.clean_lines(text))
        self.assertGreaterEqual(len(view["groups"]), 10)
        self.assertEqual(view["commands"], ["brief"])
        group = next(g for g in view["groups"] if g["interface"] == "Vlan10")
        self.assertEqual((group["state"], group["priority"], group["preempt"], group["virtual_ip"]), ("Active", "110", "on", "192.0.2.1"))

    def test_the_catalyst_brief_form_is_read_too(self):
        text = ("SW1#show standby brief\n"
                "                     P indicates configured to preempt.\n"
                "                     |\n"
                "Interface   Grp  Pri P State   Active          Standby         Virtual IP\n"
                "Vl10        10   110 P Active  local           192.0.2.3       192.0.2.1\n"
                "Vl20        20   100   Standby 192.0.2.3       local           192.0.2.65\n"
                "SW1#\n")
        view = fhrp_live.read(capture_text.clean_lines(text))
        by = {g["interface"]: g for g in view["groups"]}
        self.assertEqual(set(by), {"Vl10", "Vl20"})
        self.assertEqual((by["Vl10"]["state"], by["Vl10"]["preempt"], by["Vl10"]["standby"]), ("Active", "on", "192.0.2.3"))
        self.assertEqual((by["Vl20"]["state"], by["Vl20"]["preempt"], by["Vl20"]["active"]), ("Standby", "off", "192.0.2.3"))

    def test_a_rejected_hsrp_command_adds_nothing(self):
        text = "SW1# show hsrp\n                ^\n% Invalid command at '^' marker.\nSW1#\n"
        self.assertEqual(fhrp_live.read(capture_text.clean_lines(text))["groups"], [])


class OspfReaderTests(unittest.TestCase):
    def setUp(self):
        self.a = run("a", CAPTURE_A)
        self.b = run("b", CAPTURE_B)

    def test_neighbours_of_both_devices(self):
        a, b = self.a["ospf_live"]["neighbors"], self.b["ospf_live"]["neighbors"]
        self.assertEqual((len(a), len(b)), (5, 6))
        row = a[0]
        self.assertEqual((row["process"], row["vrf"], row["state"], row["role"], row["interface"], row["time_kind"]), ("100", "default", "FULL", "DR", "Vlan10", "up"))
        self.assertEqual(sorted({r["process"] for r in a}), ["100", "200"])

    def test_a_neighbour_that_is_not_full_is_a_high_finding_with_the_rfc_or_cisco_reference(self):
        self.assertNotIn("SW-OSPF-001", finding_ids(self.a))      # FULL, and 2-Way between two DROTHERs, are normal
        item = next(f for f in self.b["findings"] if f["id"] == "SW-OSPF-001")
        self.assertEqual(item["severity"], "high")
        self.assertIn("EXSTART", item["title"])
        self.assertIn("Po200", item["object"])
        self.assertIn("MTU", item["recommendation"])
        self.assertTrue(item["reference"].startswith("https://"))

    def test_two_way_between_drothers_is_not_a_finding_but_two_way_with_the_dr_is(self):
        row = {"process": "1", "vrf": "default", "neighbor_id": "192.0.2.9", "priority": "1", "state": "2WAY", "role": "DROTHER", "time": "1w", "time_kind": "up", "address": "192.0.2.9", "interface": "Vlan1"}
        self.assertEqual(ospf_live.findings([row]), [])
        self.assertEqual(len(ospf_live.findings([dict(row, role="DR")])), 1)

    def test_the_summary_table_has_a_total_row_per_process(self):
        table = next(t for t in self.a["extra_tables"] if t["id"] == "ospf-neighbor-summary")
        totals = [r for r in table["rows"] if r["interface"] == "All interfaces"]
        self.assertEqual({(r["process"], r["total"]) for r in totals}, {("100", 4), ("200", 1)})
        self.assertEqual(totals[0]["twoway"], 1)

    def test_the_neighbour_table_colours_the_state(self):
        table = next(t for t in self.b["extra_tables"] if t["id"] == "ospf-neighbors")
        state = next(c for c in table["columns"] if c["key"] == "state")
        self.assertEqual(state["tones"]["full"], "ok")
        self.assertEqual(state["tones"]["exstart"], "bad")
        self.assertIn("Up time", [c["label"] for c in table["columns"]])

    def test_a_neighbour_count_that_disagrees_with_the_table_is_flagged(self):
        result = analyze(CAPTURE_A.replace("Total number of neighbors: 4", "Total number of neighbors: 6"), filename="x.txt")
        self.assertTrue(any("says 6" in text for text in result["ospf_live"]["checks"]))
        self.assertIn("SW-CAP-007", finding_ids(result))
        self.assertNotIn("SW-CAP-007", finding_ids(self.a))

    def test_the_rejected_show_ip_ospf_brief_names_the_documented_spelling(self):
        row = next(r for r in self.a["coverage"]["rows"] if r["command"] == "show ip ospf brief vrf all")
        self.assertEqual(row["status"], "error")
        self.assertIn("show ip ospf interface brief vrf all", row["note"])
        self.assertNotIn("is not in the configuration", row["note"])      # the other OSPF commands answered, so the feature is on

    def test_the_ios_dead_time_layout_is_read(self):
        lines = capture_text.clean_lines("R1#show ip ospf neighbor\n\n"
                                         "Neighbor ID     Pri   State           Dead Time   Address         Interface\n"
                                         "192.0.2.2         1   FULL/DR         00:00:35    192.0.2.2       Vlan10\n"
                                         "192.0.2.3         0   FULL/  -        00:00:38    192.0.2.3       Tunnel1\n"
                                         "R1#\n")
        rows = ospf_live.read(lines, None)["neighbors"]
        self.assertEqual([(r["state"], r["role"], r["time_kind"]) for r in rows], [("FULL", "DR", "dead"), ("FULL", "", "dead")])
        self.assertEqual(rows[0]["process"], "")


class PairCheckTests(unittest.TestCase):
    def setUp(self):
        self.a = run("a", CAPTURE_A)
        self.b = run("b", CAPTURE_B)
        self.pair = compare(self.a, self.b)["fhrp"]
        self.rows = {r["group"]: r for r in self.pair["rows"]}

    def test_one_row_per_group_and_the_verdict_counts(self):
        self.assertEqual(len(self.pair["rows"]), 11)
        self.assertEqual(self.pair["counts"], {"Healthy pair": 4, "Same priority": 1, "Settings differ": 1, "Both Active": 1, "Disabled": 1, "Initial": 2, "Only on A": 1})
        self.assertEqual((self.pair["active_a"], self.pair["active_b"]), (8, 3))

    def test_a_healthy_pair_has_one_active_and_one_standby(self):
        for group in ("10", "20", "61"):
            self.assertEqual(self.rows[group]["verdict"], "Healthy pair", group)
        self.assertEqual(self.rows["80"]["verdict"], "Same priority")      # 100 / 100: the tie is broken by the higher IP address
        self.assertEqual((self.rows["10"]["a_state"], self.rows["10"]["b_state"]), ("Active", "Standby"))
        self.assertEqual((self.rows["20"]["a_state"], self.rows["20"]["b_state"]), ("Standby", "Active"))

    def test_the_same_group_on_a_different_interface_name_is_paired_by_its_number(self):
        row = self.rows["70"]
        self.assertEqual(row["interface"], "Ethernet1/10 / Ethernet1/12")
        self.assertEqual(row["verdict"], "Healthy pair")

    def test_both_active_means_the_switches_do_not_hear_each_other(self):
        row = self.rows["91"]
        self.assertEqual(row["verdict"], "Both Active")
        self.assertIn("hello", row["reason"])

    def test_disabled_and_initial_groups_are_named(self):
        self.assertEqual(self.rows["30"]["verdict"], "Disabled")
        self.assertIn("A", self.rows["30"]["reason"])
        self.assertIn("virtual IP not configured", self.rows["30"]["reason"])
        self.assertEqual(self.rows["40"]["verdict"], "Initial")
        self.assertEqual(self.rows["50"]["verdict"], "Initial")
        self.assertIn("B is Initial", self.rows["50"]["reason"])

    def test_a_group_on_one_switch_only(self):
        self.assertEqual(self.rows["60"]["verdict"], "Only on A")
        self.assertEqual(self.rows["60"]["b_state"], "—")

    def test_a_timer_that_differs_is_reported(self):
        row = self.rows["90"]
        self.assertEqual(row["verdict"], "Settings differ")
        self.assertIn("hello timer differs (A 3 s, B 1 s)", row["reason"])

    def test_equal_priorities_are_called_out_only_for_an_active_standby_pair(self):
        def group(state, priority):
            item = fhrp_live._blank("Vlan1", "1")
            item.update(state=state, priority=priority, key="vlan1")
            return item
        tie = fhrp_live.pair_check([group("Active", "100")], [group("Standby", "100")])
        self.assertEqual(tie["rows"][0]["verdict"], "Same priority")
        self.assertIn("higher IP address", tie["rows"][0]["reason"])
        initial = fhrp_live.pair_check([group("Initial", "100")], [group("Initial", "100")])
        self.assertEqual(initial["rows"][0]["verdict"], "Initial")
        self.assertNotIn("same priority", initial["rows"][0]["reason"])

    def test_a_capture_without_live_state_has_no_pair_check(self):
        empty = fhrp_live.pair_check([], groups_by_key(self.b) and self.b["fhrp"]["groups"])
        self.assertEqual(empty, {"rows": [], "counts": {}, "active_a": 0, "active_b": 0})

    def test_the_excel_comparison_has_an_hsrp_pair_sheet(self):
        comparison = compare(self.a, self.b)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "pair.xlsx"
            export_comparison(path, self.a, self.b, comparison)
            workbook = load_workbook(path)
            self.assertIn("HSRP Pair Check", workbook.sheetnames)
            sheet = workbook["HSRP Pair Check"]
            values = [[cell.value for cell in row] for row in sheet.iter_rows()]
        flat = " ".join(str(v) for row in values for v in row if v)
        self.assertIn("Both Active", flat)
        self.assertIn("NXDEMO-A", flat)


class MultiDeviceFileTests(unittest.TestCase):
    """HQ-style file: the output of a second switch was appended to the first switch's capture."""

    def test_only_the_dominant_prompt_feeds_the_hsrp_and_ospf_tables(self):
        result = analyze(CAPTURE_A + "\n" + CAPTURE_B + "\n", filename="two.txt")
        self.assertEqual(result["hostname"], "NXDEMO-A")
        self.assertEqual(len(result["fhrp"]["groups"]), 11)
        self.assertEqual(len(result["ospf_live"]["neighbors"]), 5)
        notes = " ".join(n["text"] for n in result["capture_notes"])
        self.assertIn("NXDEMO-B", notes)
        self.assertIn("left out", notes)
        self.assertTrue({"SW-CAP-003", "SW-CAP-004"} <= finding_ids(result))

    def test_the_second_device_alone_is_unchanged(self):
        self.assertEqual(len(run("b", CAPTURE_B)["fhrp"]["groups"]), 10)


class ExcelExportTests(unittest.TestCase):
    def test_single_device_workbook_has_the_hsrp_and_ospf_sheets(self):
        from features.switch_analyzer.exporters import SwitchExcelExporter

        result = run("a", CAPTURE_A)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "a.xlsx"
            SwitchExcelExporter(result).export(path)
            names = load_workbook(path).sheetnames
        for name in ("HSRP", "HSRP Summary", "OSPF Neighbours", "OSPF Neighbour Summary"):
            self.assertIn(name, names)


class CollectListTests(unittest.TestCase):
    ROOT = Path(__file__).resolve().parent.parent / "static" / "collect-scripts"

    def commands(self, name: str) -> list[str]:
        return [line.strip() for line in (self.ROOT / name).read_text(encoding="utf-8").splitlines() if line.strip() and not line.startswith("#")]

    def test_the_nexus_list_asks_for_the_documented_hsrp_and_ospf_commands(self):
        commands = self.commands("collect-cisco-nexus.txt")
        for wanted in ("show hsrp all", "show hsrp brief", "show hsrp summary", "show ip ospf neighbors vrf all", "show ip ospf neighbors summary vrf all", "show ip ospf interface brief vrf all"):
            self.assertIn(wanted, commands)
        self.assertNotIn("show ip ospf brief vrf all", commands)      # not an NX-OS command (the device answers % Invalid command)

    def test_the_catalyst_list_asks_for_the_standby_and_ospf_interface_detail(self):
        commands = self.commands("collect-cisco-catalyst.txt")
        for wanted in ("show standby brief", "show standby", "show ip ospf neighbor", "show ip ospf interface brief"):
            self.assertIn(wanted, commands)


if __name__ == "__main__":
    unittest.main()

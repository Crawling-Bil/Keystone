import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from features.wireless_center.analyzers.relationship_analyzer import RelationshipAnalyzer
from features.wireless_center.exporters.excel_exporter import ExcelExporter
from features.wireless_center.parsers.config_parser import HuaweiConfigParser
from features.wireless_center.service import analyze


HUAWEI_VAP_CONFIG = """sysname LAB-WLC
#
wlan
 security-profile name LAB-SEC
  security wpa-wpa2 psk pass-phrase cipher SyntheticSecret
 ssid-profile name LAB-SSID
  ssid LAB-WIFI
 traffic-profile name LAB-TRAFFIC
  rate-limit client up 2048
 vap-profile name LAB-SITE-A
  service-vlan vlan-id 100
  ssid-profile LAB-SSID
  security-profile LAB-SEC
  authentication-profile LAB-AUTH
  traffic-profile LAB-TRAFFIC
  forward-mode tunnel
 vap-profile name LAB-SITE-B
  service-vlan vlan-id 200
  ssid-profile LAB-SSID
  security-profile LAB-SEC
 vap-profile name LAB-UNUSED
  service-vlan vlan-id 300
  ssid-profile LAB-SSID
 ap-group name SITE-A
  radio 0
   vap-profile LAB-SITE-A wlan 1
  radio 1
   vap-profile LAB-SITE-A wlan 1
 ap-group name SITE-B
  radio 0
   vap-profile LAB-SITE-B wlan 2 service-vlan vlan-id 250
#
ap-id 1 type-id 130 ap-mac 0011-2233-4455 ap-sn SERIAL-001
 ap-name SITE-A-AP-01
 ap-group SITE-A
#
ap-id 2 type-id 130 ap-mac 0011-2233-4466 ap-sn SERIAL-002
 ap-name SITE-B-AP-01
 ap-group SITE-B
#
"""


class VAPInventoryTest(unittest.TestCase):
    def setUp(self):
        self.data = HuaweiConfigParser(HUAWEI_VAP_CONFIG).parse()
        self.analyzer = RelationshipAnalyzer(self.data)

    def test_nested_named_profiles_stop_at_the_next_sibling(self):
        profiles = {item["name"]: item["config"] for item in self.data["vap_profiles"]}
        self.assertEqual(len(profiles), 3)
        self.assertEqual(profiles["LAB-SITE-A"], [
            "service-vlan vlan-id 100",
            "ssid-profile LAB-SSID",
            "security-profile LAB-SEC",
            "authentication-profile LAB-AUTH",
            "traffic-profile LAB-TRAFFIC",
            "forward-mode tunnel",
        ])
        self.assertFalse(any("ap-group name" in line for line in profiles["LAB-SITE-B"]))

    def test_vap_inventory_resolves_site_usage_and_effective_vlan(self):
        rows = {item["vap_profile"]: item for item in self.analyzer.build_vap_inventory()}
        site_a = rows["LAB-SITE-A"]
        self.assertEqual(site_a["usage_scope"], "Site-specific")
        self.assertEqual(site_a["ap_groups"], "SITE-A")
        self.assertEqual(site_a["radios"], "0, 1")
        self.assertEqual(site_a["wlan_ids"], "1")
        self.assertEqual(site_a["ap_count"], 1)
        self.assertEqual(site_a["ssid"], "LAB-WIFI")
        self.assertEqual(site_a["authentication_profile"], "LAB-AUTH")
        self.assertEqual(site_a["traffic_profile"], "LAB-TRAFFIC")
        self.assertEqual(site_a["forward_mode"], "tunnel")

        site_b = rows["LAB-SITE-B"]
        self.assertEqual(site_b["service_vlan"], "200")
        self.assertEqual(site_b["effective_vlans"], "250")
        self.assertEqual(rows["LAB-UNUSED"]["usage_scope"], "Unassigned")
        self.assertEqual(rows["LAB-UNUSED"]["deployment_status"], "Not mapped")

    def test_security_summary_and_full_config_never_expose_secret(self):
        result = analyze(HUAWEI_VAP_CONFIG, "Huawei")
        self.assertEqual(len(result["vap_inventory"]), 3)
        self.assertNotIn("SyntheticSecret", str(result["vap_inventory"]))
        self.assertEqual(result["cards"]["vap_profiles"], 3)

    def test_excel_export_adds_vap_profiles_without_replacing_wlan_ssid(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "wireless.xlsx"
            ExcelExporter(self.data, self.analyzer).export(path)
            workbook = load_workbook(path, read_only=True)
            self.assertIn("WLAN SSID", workbook.sheetnames)
            self.assertIn("VAP Profiles", workbook.sheetnames)
            sheet = workbook["VAP Profiles"]
            self.assertEqual(sheet.max_row, 4)
            self.assertEqual(sheet["A2"].value, "LAB-SITE-A")


if __name__ == "__main__":
    unittest.main()

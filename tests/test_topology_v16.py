"""Topology Builder 1.6: one line per pair of sites on the Overview, and the other changes of this version.  Synthetic captures only (invented names, documentation addresses)."""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from topology_samples import iosxe_capture  # noqa: E402

from features.topology import labels  # noqa: E402
from features.topology.extract import Device  # noqa: E402
from features.topology.graph import Connection, Graph, Link  # noqa: E402
from features.topology.logical import logical_pages  # noqa: E402
from features.topology.scene import SITE_BUNDLE_MIN, build_pages  # noqa: E402
from features.topology.service import build_topology  # noqa: E402


def _graph(sites: dict[str, int], links: list[tuple[str, str]]) -> Graph:
    nodes = {}
    for site, count in sites.items():
        for i in range(count):
            key = f"{site}-SW{i + 1}"
            nodes[key] = Device(key=key, hostname=key, captured=True, role_hint="core")
    port: dict[str, int] = {}
    conns = []
    for a, b in links:
        port[a] = port.get(a, 0) + 1
        port[b] = port.get(b, 0) + 1
        conns.append(Connection(a, b, [Link(a, f"Eth1/{port[a]}", b, f"Eth1/{port[b]}", "confirmed", {"CDP"}, "")]))
    return Graph(nodes=nodes, links=[l for c in conns for l in c.links], connections=conns, hidden_kinds=set())


MESH = [("AAA-SW1", "BBB-SW1"), ("AAA-SW1", "BBB-SW2"), ("AAA-SW2", "BBB-SW1"), ("AAA-SW2", "BBB-SW2")]
INSIDE = [("AAA-SW1", "AAA-SW2"), ("BBB-SW1", "BBB-SW2"), ("CCC-SW1", "CCC-SW2")]


def _sites(extra: list[tuple[str, str]]) -> Graph:
    return _graph({"AAA": 2, "BBB": 2, "CCC": 2}, INSIDE + extra)


def test_four_lines_between_two_sites_become_one_line_that_keeps_every_cable():
    graph = _sites(MESH + [("BBB-SW1", "CCC-SW1"), ("BBB-SW2", "CCC-SW2")])
    overview = build_pages(graph, "prefix1")[0]
    bundles = [c for c in overview.connections if c.kind == "sitebundle"]
    assert len(bundles) == 1 and len(bundles[0].links) == 4
    assert {(c.a, c.b) for c in overview.connections if c.kind != "sitebundle" and c.a[:3] != c.b[:3]} == {("BBB-SW1", "CCC-SW1"), ("BBB-SW2", "CCC-SW2")}      # two lines stay two lines
    assert {(l.a, l.b) for l in bundles[0].links} == set(MESH)                                              # every cable, each turned to run from site AAA to site BBB
    assert labels.center_text(bundles[0], graph) == "4 links" and labels.end_texts(bundles[0]) == ("", "")
    assert len(overview.connections) == len(INSIDE) + 1 + 2 and set(overview.routes) <= set(range(len(overview.connections)))
    assert [len(v) for v in overview.bundled.values()] == [4] and len(overview.unbundled()) == len(INSIDE) + len(MESH) + 2


def test_the_limit_is_three_lines_and_a_pair_below_it_keeps_its_lines():
    assert SITE_BUNDLE_MIN == 3
    two = build_pages(_sites(MESH[:2]), "prefix1")[0]
    three = build_pages(_sites(MESH[:3]), "prefix1")[0]
    assert not any(c.kind == "sitebundle" for c in two.connections) and not two.bundled
    assert sum(c.kind == "sitebundle" for c in three.connections) == 1


def test_a_bundle_takes_the_speed_of_its_cables_and_the_site_pages_stay_whole():
    graph = _sites(MESH)
    for node in graph.nodes.values():
        node.speeds = {f"Eth1/{n}": 10000 for n in range(1, 6)}
    pages = build_pages(graph, "prefix1")
    bundle = next(c for c in pages[0].connections if c.kind == "sitebundle")
    assert labels.center_text(bundle, graph) == "4 links · 10G"
    site = next(p for p in pages if p.name == "AAA")                      # a site page lists each cable on its own
    assert sum(1 for c in site.connections if c.kind != "sitebundle" and {c.a[:3], c.b[:3]} == {"AAA", "BBB"}) == 4 and not site.bundled


def test_one_page_diagrams_are_not_bundled():
    pages = build_pages(_sites(MESH), "none")
    assert len(pages) == 1 and not pages[0].bundled and not any(c.kind == "sitebundle" for c in pages[0].connections)


def test_the_vpc_twin_still_collapses_the_separate_cables():
    graph = _graph({"AAA": 2, "BBB": 2}, [("AAA-SW1", "AAA-SW2"), ("BBB-SW1", "BBB-SW2")])
    for conn in graph.connections:
        conn.kind = "peer-link"
    for a, b in MESH:
        conn = Connection(a, b, [Link(a, "Eth1/9", b, "Eth1/9", "confirmed", {"CDP"}, "")], a_po="Po101", b_po="Po101", kind="bundle", vpc_id="101", multi_chassis=True)
        graph.connections.append(conn)
        graph.links += conn.links
    graph.vpc_pairs = [("AAA-SW1", "AAA-SW2", "10"), ("BBB-SW1", "BBB-SW2", "20")]
    pages = build_pages(graph, "prefix1")
    assert [c.kind for c in pages[0].connections].count("sitebundle") == 1
    twin = logical_pages(graph, pages)[0]
    assert [(c.kind, len(c.links)) for c in twin.connections if c.kind not in ("peer-link",)] == [("vpc", 4)]


def _routed(links: list[tuple[str, str]], hosts: list[str]) -> list[tuple[str, str]]:
    cdp = {h: [] for h in hosts}
    cfg = {h: "" for h in hosts}
    port = {h: 0 for h in hosts}
    ips = {h: f"10.0.0.{i + 1}" for i, h in enumerate(hosts)}
    for n, (a, b) in enumerate(links):
        port[a] += 1
        port[b] += 1
        pa, pb = f"GigabitEthernet1/0/{port[a]}", f"GigabitEthernet1/0/{port[b]}"
        cdp[a].append((b, ips[b], "C9300L", pa, pb))
        cdp[b].append((a, ips[a], "C9300L", pb, pa))
        cfg[a] += f"interface {pa}\n no switchport\n ip address 10.9.{n}.1 255.255.255.252\n!\n"
        cfg[b] += f"interface {pb}\n no switchport\n ip address 10.9.{n}.2 255.255.255.252\n!\n"
    return [(f"{h}.txt", iosxe_capture(h, f"FOC{i:04d}ZZZ", ips[h], cdp[h], config_extra=cfg[h])) for i, h in enumerate(hosts)]


def test_the_l3_over_l2_overview_bundle_carries_what_its_cables_carry_and_exports_open():
    hosts = ["AAA-SW1", "AAA-SW2", "BBB-SW1", "BBB-SW2", "CCC-SW1"]
    links = [("AAA-SW1", "AAA-SW2"), ("BBB-SW1", "BBB-SW2"), *MESH, ("BBB-SW1", "CCC-SW1"), ("BBB-SW2", "CCC-SW1")]
    result = build_topology(_routed(links, hosts), group="prefix1", view="all")
    physical = result["pages"][0]
    assert sum(c.kind == "sitebundle" for c in physical.connections) == 1
    over2 = next(p for p in result["l3_pages"] if p.name.startswith("Overview"))
    bundle = next(c for c in over2.connections if c.kind == "sitebundle")
    assert len(bundle.l3["carries"]) == 4 and bundle.l3["width"] == 5 and "4 cables between two sites carries" in bundle.l3["meaning"]
    assert not any(c.kind == "sitebundle" for p in result["l3_pages"] if not p.name.startswith("Overview") for c in p.connections)
    assert 'kind="sitebundle"' in result["drawio"] and "4 links" in result["drawio"]
    assert any("4 links" in s["svg"] for s in result["svgs"])
    assert len(result["l3_links"]) == 8                                   # the table still lists every routed link: nothing is lost behind the bundle


# ---- the L3 page: the same subnets on the same devices become one node ---------------------------------------------

def _vlans(count: int, devices: int = 3, vrf: str = "") -> list[tuple[str, str]]:
    hosts = [f"AGG{i + 1}" for i in range(devices)]
    files = []
    for i, host in enumerate(hosts):
        cfg = "".join(f"interface Vlan{100 + n}\n description USERS-{n}\n{' vrf forwarding ' + vrf + chr(10) if vrf else ''} ip address 10.20.{n}.{2 + i} 255.255.255.0\n!\n" for n in range(count))
        cdp = [(other, f"10.0.0.{k + 1}", "C9500", f"TenGigabitEthernet1/0/{1 + j}", f"TenGigabitEthernet1/0/{1 + (i if i < k else i - 1)}")
               for j, (k, other) in enumerate((k, h) for k, h in enumerate(hosts) if k != i)]
        files.append((f"{host}.txt", iosxe_capture(host, f"FOC000{i}XXX", f"10.0.0.{i + 1}", cdp, config_extra=cfg)))
    return files


def _l3(result):
    return next(p for p in result["l3_pages"] + result["l3_extra_pages"] if p.view == "l3route")


def test_thirty_vlans_on_the_same_three_switches_are_one_node_and_three_lines():
    result = build_topology(_vlans(30), view="l3route")
    page = _l3(result)
    segments = [k for k, n in page.graph.nodes.items() if n.role_hint == "segment"]
    assert len(segments) == 1 and page.graph.nodes[segments[0]].hostname == "30 shared subnets" and page.graph.nodes[segments[0]].mgmt_ip == "3 devices"
    assert "10.20.0.0/24 ... 10.20.29.0/24" in page.graph.nodes[segments[0]].model
    assert len(page.connections) == 3 and page.width < 1500
    assert all(c.l3["group"] == 30 and len(c.l3["carries"]) == 30 for c in page.connections)
    assert {c.l3["id"] for c in page.connections} == {"L3-G01", "L3-G02", "L3-G03"}
    subnets = {r["subnet"] for r in result["l3_links"] if r["kind"] == "Shared subnet"}
    assert len(subnets) == 30 and sum(1 for r in result["l3_links"] if r["kind"] == "Shared subnets (group)") == 3      # every subnet keeps its row, the groups are listed too


def test_the_ungrouped_page_keeps_every_line_and_a_small_set_is_not_grouped():
    off = _l3(build_topology(_vlans(30), view="l3route", group_gateways=False))
    assert len(off.connections) == 90
    few = _l3(build_topology(_vlans(4), view="l3route"))
    assert len(few.connections) == 12 and not any(c.l3["group"] for c in few.connections)


def test_the_grouped_node_is_explained_and_survives_l3_over_l2_and_the_trace():
    result = build_topology(_vlans(8), view="all")
    page = _l3(result)
    meaning = page.connections[0].l3["meaning"]
    assert "8 shared subnets" in meaning and "drawn as one node" in meaning
    over2 = result["l3_pages"][0] if result["l3_pages"][0].view == "l3over2" else result["l3_extra_pages"][0]
    assert over2.view == "l3over2" and any(c.l3.get("id", "").startswith("L3-G") for c in over2.connections)
    assert any("8 shared subnets" in s["svg"] for s in result["svgs"])
    assert "8 shared subnets" in result["drawio"]


# ---- a forced single diagram of many sites ----------------------------------------------------------------------------

class _Overview:
    def __init__(self, width: float, height: float, devices: int) -> None:
        self.width, self.height, self.nodes = width, height, {f"N{i}": None for i in range(devices)}


def _many_sites(sites: int) -> Graph:
    return _graph({f"S{n:02d}": 2 for n in range(sites)}, [(f"S{n:02d}-SW1", f"S{n:02d}-SW2") for n in range(sites)])


def test_a_forced_single_diagram_that_is_too_wide_says_so_and_names_automatic_grouping():
    from features.topology.service import grouping_notices
    notices = grouping_notices(_many_sites(20), "none", [_Overview(26000, 900, 40)])
    assert len(notices) == 1 and "26,000 x 900 px with 40 devices" in notices[0] and "Automatic" in notices[0] and "20 sites" in notices[0]
    assert grouping_notices(_many_sites(20), "none", [_Overview(3000, 900, 40)]) == []              # a readable single diagram gets no warning
    assert grouping_notices(_many_sites(20), "auto", [_Overview(26000, 900, 40)]) == []              # nobody forced it
    assert grouping_notices(_many_sites(20), "prefix1", [_Overview(26000, 900, 40)]) == []
    assert len(grouping_notices(_many_sites(20), "none", [_Overview(3000, 900, 400)])) == 1          # many devices count too


def test_without_readable_sites_the_notice_suggests_splitting_the_captures():
    from features.topology.service import grouping_notices
    flat = _graph({"X": 1, "Y": 1}, [("X-SW1", "Y-SW1")])
    text = grouping_notices(flat, "none", [_Overview(20000, 900, 2)])[0]
    assert "no readable sites" in text or "give no readable sites" in text


def test_the_notice_reaches_the_page_the_run_settings_and_the_report(monkeypatch):
    import io
    import zipfile

    from app import create_app
    from features.topology import service
    monkeypatch.setattr(service, "WIDE_DIAGRAM_PX", 100)
    hosts = [f"S{n:02d}-SW{i}" for n in range(16) for i in (1, 2)]
    links = [(f"S{n:02d}-SW1", f"S{n:02d}-SW2") for n in range(16)]
    files = [(io.BytesIO(text.encode()), name) for name, text in _routed(links, hosts)]
    app = create_app()
    app.config["TESTING"] = True
    body = app.test_client().post("/topology/api/build", data={"files": files, "group": "none", "view": "physical"}, content_type="multipart/form-data").get_json()
    assert body["ok"] and len(body["notices"]) == 1 and "Automatic" in body["notices"][0]
    pack = next(d["url"] for d in body["downloads"] if d["name"] == "keystone-topology-pack.zip")
    archive = zipfile.ZipFile(io.BytesIO(app.test_client().get(pack).data))
    names = {n.split("/", 1)[1]: n for n in archive.namelist()}
    assert "Notices" in archive.read(names["run-settings.txt"]).decode("utf-8") and "Notice." in archive.read(names["topology-report.html"]).decode("utf-8")


# ---- a switch stands over its frame of end points even when seen-only switches stand between the owners ------------------

def test_spread_shares_the_move_between_the_devices_that_wish_something():
    from features.topology.layout import _spread
    row = ["A", "f1", "f2", "f3", "B"]                                    # f1..f3 wish nothing: they follow A
    want = {"A": 0.0, "B": 2.0}                                            # four steps of room are needed, two are free
    plain = _spread(row, want, 1.0)
    shared = _spread(row, want, 1.0, None, 0.0)
    assert [round(plain[k], 2) for k in row] == [-0.4, 0.6, 1.6, 2.6, 3.6]      # the followers count as much as A: B bears most of the move
    assert [round(shared[k], 2) for k in row] == [-1.0, 0.0, 1.0, 2.0, 3.0]     # A and B each give way by one step, the order and the gaps are kept
    assert all(shared[b] - shared[a] >= 1.0 - 1e-9 for a, b in zip(row, row[1:]))


SEEN_ONLY_ROW = [(0, 0, 0), (1, 0, 0), (2, 3, 1), (3, 0, 1), (4, 0, 1), (5, 0, 0), (6, 0, 1), (7, 0, 0), (8, 0, 1), (9, 0, 0), (10, 9, 1), (11, 0, 1), (12, 9, 1),
                 (13, 0, 0), (14, 9, 1), (15, 14, 1), (16, 0, 1), (17, 0, 1), (18, 14, 1), (19, 5, 1)]          # (number, end points, captured)


def test_owners_beside_seen_only_switches_share_the_distance_to_their_frames():
    from features.topology.layout import NODE_W
    nodes = {"CORE": Device(key="CORE", hostname="CORE", captured=True, role_hint="core")}
    conns = []
    for k, ends, captured in SEEN_ONLY_ROW:
        sw = f"SW-{k:02d}"
        nodes[sw] = Device(key=sw, hostname=sw, captured=bool(captured), role_hint="access")
        conns.append(Connection("CORE", sw, [Link("CORE", f"Te1/{k}", sw, "Te1/1", "confirmed", {"CDP"}, "")]))
        for i in range(ends):
            end = f"{sw}-E{i}"
            nodes[end] = Device(key=end, hostname=end, captured=False, role_hint="ap")
            conns.append(Connection(sw, end, [Link(sw, f"Gi1/0/{i + 1}", end, "Gi0", "confirmed", {"CDP"}, "")]))
    graph = Graph(nodes=nodes, links=[l for c in conns for l in c.links], connections=conns, hidden_kinds=set(), group_endpoints=True)
    page = build_pages(graph, "none")[0]
    offsets = {t: abs(page.nodes[t].x + NODE_W / 2 - (x + w / 2)) for t, x, y, w, h in page.end_boxes if t in page.nodes}
    assert len(offsets) == 7 and max(offsets.values()) <= 600 and sum(offsets.values()) <= 1500              # before: 714 px for one switch, 1742 px together


# ---- reading big captures: faster, with the very same result ------------------------------------------------------------

def _old_acl_references(statements, definitions):
    """The reading before 1.6, a regular expression per ACL name per line: the reference the faster reading must agree with."""
    import re
    refs = {item["name"]: [] for item in definitions}
    patterns = {name: re.compile(rf"(?<![\w.\-]){re.escape(name)}(?![\w.\-])") for name in refs}
    for header, children in statements:
        if header.lower().startswith(("interface ", "line vty", "user-interface")) or header.startswith(("ip access-list", "access-list")):
            continue
        for name, pattern in patterns.items():
            hit = pattern.search(header) is not None or any(pattern.search(child) for child in children)
            if hit:
                context = header if len(header) <= 80 else header[:77] + "..."
                if context not in refs[name] and len(refs[name]) < 8:
                    refs[name].append(context)
    return refs


def test_acl_references_agree_with_the_regular_expression_reading():
    import random

    from features.switch_analyzer.l3 import _acl_references
    rng = random.Random(5)
    seen = 0
    for _ in range(300):
        names = ["".join(rng.choice("ab-._1xY") for _ in range(rng.randint(1, 4))) for _ in range(rng.randint(1, 6))]
        if rng.random() < 0.4:
            names.append(rng.choice(["A B", "a/b", "x:y", "ab.", "-a"]))               # names with a space or a slash keep the regular expression
        definitions = [{"name": n} for n in names]
        statements = []
        for _ in range(rng.randint(1, 8)):
            word = lambda: "".join(rng.choice("ab-._1 /:xY") for _ in range(rng.randint(0, 14)))      # noqa: E731
            header = rng.choice(["route-map ", "class-map match-any ", "ip nat inside source list ", "interface Gi1 ", "snmp-server community "]) + word()
            statements.append((header, [word() for _ in range(rng.randint(0, 4))]))
        got, want = _acl_references(statements, definitions), _old_acl_references(statements, definitions)
        assert got == want
        seen += sum(1 for v in want.values() if v)
    assert seen > 50                                                          # the random cases do hit names


def test_clean_lines_hands_out_a_fresh_list_each_time_and_the_same_lines():
    from features.switch_analyzer.capture_text import clean_lines
    text = "SW1#show version\r\nCisco IOS XE\r\nSW1#show clock\r\n12:00:00\r\n"
    first = clean_lines(text)
    first.append("changed by the caller")
    second = clean_lines(text)
    assert second[:2] == ["SW1#show version", "Cisco IOS XE"] and "changed by the caller" not in second and second is not first
    assert clean_lines("") == [] and clean_lines(text + "SW1#") != second


# ---------------------------------------------------------------- which captures to take next
from features.topology.priority import capture_priority, collect_lists  # noqa: E402


def _stub(graph: Graph, key: str, role: str, vendor: str = "", model: str = "", version: str = "") -> None:
    graph.nodes[key] = Device(key=key, hostname=key, captured=False, role_hint=role, vendor=vendor, model=model, os_version=version)


def _link(graph: Graph, a: str, b: str, port: int, state: str = "single") -> None:
    link = Link(a, f"Eth1/{port}", b, f"Eth9/{port}", state, {"CDP"}, "")
    graph.links.append(link)
    graph.connections.append(Connection(a, b, [link]))


def test_priority_ranks_by_links_then_neighbors_then_macs() -> None:
    graph = _graph({"AAA": 3}, [])
    for name, role in (("ONE-LINK", "switch"), ("TWO-LINKS", "switch"), ("TWO-NEIGHBORS", "switch"), ("TWO-MACS", "switch")):
        _stub(graph, name, role)
    _link(graph, "AAA-SW1", "ONE-LINK", 1)
    _link(graph, "AAA-SW1", "TWO-LINKS", 2)
    _link(graph, "AAA-SW1", "TWO-LINKS", 3)
    _link(graph, "AAA-SW1", "TWO-NEIGHBORS", 4)
    _link(graph, "AAA-SW2", "TWO-NEIGHBORS", 5)
    _link(graph, "AAA-SW1", "TWO-MACS", 6)
    _link(graph, "AAA-SW1", "TWO-MACS", 7)
    graph.hosts = [{"mac": "02aa00000001", "device": "AAA-SW1", "port": "Eth1/6", "via": "TWO-MACS", "vlan": "10", "ips": [], "note": ""}] * 5
    rows = capture_priority(graph)
    assert [r["device"] for r in rows] == ["TWO-NEIGHBORS", "TWO-MACS", "TWO-LINKS", "ONE-LINK"]
    assert [r["rank"] for r in rows] == [1, 2, 3, 4]
    assert rows[0]["captured_neighbors"] == 2 and rows[1]["behind_macs"] == 5 and rows[0]["neighbors"] == "AAA-SW1, AAA-SW2"
    assert all(r["action"] == "new capture" for r in rows)
    assert "5 MAC addresses learned behind it" in rows[1]["why"] and "1 link from 1 captured device reach" in rows[3]["why"]


def test_priority_counts_a_reported_link_before_an_inferred_one() -> None:
    graph = _graph({"AAA": 2}, [])
    _stub(graph, "GUESS", "router")
    _stub(graph, "SEEN", "router")
    for port in range(1, 4):
        _link(graph, "AAA-SW1", "GUESS", port, "inferred")
    _link(graph, "AAA-SW2", "SEEN", 4)
    rows = capture_priority(graph)
    assert [r["device"] for r in rows] == ["SEEN", "GUESS"]
    assert rows[1]["links"] == 3 and rows[1]["inferred_links"] == 3 and "inferred from ARP" in rows[1]["why"]


def test_priority_skips_endpoints_hidden_nodes_and_captured_devices() -> None:
    graph = _graph({"AAA": 2}, [("AAA-SW1", "AAA-SW2")])
    for name, role in (("PHONE-1", "phone"), ("AP-1", "ap"), ("SRV-1", "server"), ("HIDDEN-SW", "switch"), ("CLOUD-1", "cloud"), ("HOST:AAA", "host")):
        _stub(graph, name, role)
        _link(graph, "AAA-SW1", name, 1)
    graph.hidden_nodes.add("HIDDEN-SW")
    assert capture_priority(graph) == []


def test_priority_asks_for_a_new_capture_of_a_device_without_neighbor_table() -> None:
    graph = _graph({"AAA": 2}, [("AAA-SW1", "AAA-SW2")])
    graph.nodes["AAA-SW1"].coverage = {"cdp": False, "lldp": False}
    graph.nodes["AAA-SW2"].coverage = {"cdp": True, "lldp": True}
    graph.nodes["AAA-SW2"].halves = [object()]                      # has a neighbor table
    rows = capture_priority(graph)
    assert [(r["action"], r["device"]) for r in rows] == [("capture again", "AAA-SW1")]
    assert "show cdp neighbors detail" in rows[0]["why"]
    graph.nodes["AAA-SW1"].vendor = "Fortinet"                      # no CDP / LLDP exists on a FortiGate: nothing to ask for
    assert capture_priority(graph) == []


def test_priority_collect_lists_follow_the_vendor_and_exist() -> None:
    def pick(**kw):
        return collect_lists(Device(key="X", hostname="X", **kw))
    assert pick(vendor="Fortinet", model="FortiGate-401F") == ["collect-firewall-fortigate.txt"]
    assert pick(vendor="Palo Alto Networks", model="PA-3220") == ["collect-firewall-paloalto.txt"]
    assert pick(vendor="F5", model="BIG-IP i5800") == ["collect-lb-f5-bigip.txt"]
    assert pick(vendor="Cisco", model="N9K-C93180YC-FX") == ["collect-cisco-nexus.txt"]
    assert pick(vendor="Cisco", model="WS-C2960X-48FPS-L", role_hint="access") == ["collect-cisco-catalyst.txt"]
    assert pick(vendor="Cisco", model="C9200L-48T-4G") == ["collect-cisco-catalyst.txt"]
    assert pick(vendor="Cisco", model="ISR4331/K9", role_hint="router") == ["collect-router-cisco-ios-xe.txt"]
    assert pick(model="JL256A", os_version="FL.10.10.1000") == ["collect-aruba-cx.txt"]
    assert pick(model="J9728A", os_version="YA.16.10.0003") == ["collect-aruba-switch.txt"]
    assert pick(model="JL256A") == ["collect-aruba-cx.txt", "collect-aruba-switch.txt"]
    assert pick(model="S5755-28P-H") == ["collect-huawei-h-series.txt"]
    assert pick(model="S5735-L24P4X") == ["collect-huawei-s-l-e-series.txt"]
    assert pick(role_hint="firewall") == ["collect-firewall-fortigate.txt", "collect-firewall-paloalto.txt", "collect-firewall-mikrotik.txt"]
    assert pick(role_hint="access") == [] and pick(vendor="Acme", model="Z1") == []


def test_priority_reaches_the_tables_the_workbook_the_report_and_the_pack() -> None:
    import io

    from openpyxl import load_workbook

    from features.topology.excel import to_xlsx
    from features.topology.report import build_report
    from features.topology.service import to_csv, to_json
    from topology_samples import lab_dc

    result = build_topology(list(lab_dc().items()), group="none")
    rows = capture_priority(result["graph"])
    assert [r["device"] for r in rows] == ["LAB-FW1"] and rows[0]["collect_list"].startswith("collect-firewall-fortigate.txt")
    assert to_csv(rows).splitlines()[0].startswith("rank,action,device,role,vendor,model,links,inferred_links")
    book = load_workbook(io.BytesIO(to_xlsx(result["summary"], [], [], [], [], priority=rows)))
    assert "Capture next" in book.sheetnames and book["Capture next"]["C2"].value == "LAB-FW1"
    assert "Capture next" not in load_workbook(io.BytesIO(to_xlsx(result["summary"], [], [], [], []))).sheetnames
    html = build_report(result["graph"], result["summary"], result["findings"], [], [], [], [], result["svgs"], [], priority=rows)
    for heading in ("6. Captures to take next", "7. Per-device pages", "8. Method and limits"):
        assert heading in html
    assert "Captures to take next" not in build_report(result["graph"], result["summary"], result["findings"], [], [], [], [], result["svgs"], [])
    assert '"capture_priority"' in to_json(result["graph"], result["findings"])


def test_priority_is_the_same_for_the_same_captures() -> None:
    from topology_samples import lab_dc
    one = capture_priority(build_topology(list(lab_dc().items()), group="none")["graph"])
    two = capture_priority(build_topology(list(reversed(list(lab_dc().items()))), group="none")["graph"])
    assert one == two


def test_the_priority_list_reaches_the_page_the_csv_and_the_pack():
    import io
    import zipfile

    from app import create_app
    from topology_samples import lab_dc
    app = create_app()
    app.config["TESTING"] = True
    client = app.test_client()
    files = [(io.BytesIO(text.encode()), name) for name, text in lab_dc().items()]
    body = client.post("/topology/api/build", data={"files": files, "group": "none"}, content_type="multipart/form-data").get_json()
    assert body["ok"] and [r["device"] for r in body["priority"]] == ["LAB-FW1"]
    urls = {d["name"]: d["url"] for d in body["downloads"]}
    assert "capture-priority.csv" in urls
    csv_text = client.get(urls["capture-priority.csv"]).data.decode("utf-8-sig")
    assert csv_text.splitlines()[1].startswith("1,new capture,LAB-FW1,firewall")
    archive = zipfile.ZipFile(io.BytesIO(client.get(urls["keystone-topology-pack.zip"]).data))
    names = archive.namelist()
    assert any(n.endswith("tables/capture-priority.csv") for n in names)
    report = archive.read(next(n for n in names if n.endswith("topology-report.html"))).decode("utf-8")
    assert "6. Captures to take next" in report and "LAB-FW1" in report
    # every device captured and every capture holds its neighbors: nothing to ask for, so no file and an empty list
    pair = [(io.BytesIO(text.encode()), name) for name, text in _routed([("PAIR-SW1", "PAIR-SW2")], ["PAIR-SW1", "PAIR-SW2"])]
    done = client.post("/topology/api/build", data={"files": pair, "group": "none"}, content_type="multipart/form-data").get_json()
    assert done["ok"] and done["priority"] == [] and "capture-priority.csv" not in {d["name"] for d in done["downloads"]}


def test_the_version_is_1_6():
    from core.modules import TOOLS
    assert next(t for t in TOOLS if t.key == "topology").version == "1.6"

"""The L3 view through the web route: tables, downloads, report, workbook, drawio and the click panel markup. Masked router fixtures."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "router"
NAMES = ("c8200_sdwan_cedge_masked.txt", "cedge_sdwan_overlay_masked.txt", "c3825_ios_gdoi_masked.txt")


@pytest.fixture()
def client():
    from app import create_app
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def _post(client, view="l3route", hints="", names=NAMES):
    files = [(io.BytesIO((FIXTURES / n).read_bytes()), n) for n in names]
    data = {"files": files, "group": "none", "view": view, "provider_hints": hints}
    return client.post("/topology/api/build", data=data, content_type="multipart/form-data")


def test_l3_view_returns_links_routes_legend_and_downloads(client):
    body = _post(client).get_json()
    assert body["ok"] and body["l3_links"]
    first = body["l3_links"][0]
    assert first["id"].startswith("L3-") and first["meaning"] and first["state"] in ("working", "broken", "configured only")
    names = {d["name"] for d in body["downloads"]}
    assert {"l3-links.csv", "l3-routes.csv"} <= names
    assert any(p["view"] == "l3route" and "data-l3=" in p["svg"] for p in body["pages"])
    assert any(text.startswith("Dashed line") for _, text in body["legend"]["l3route"])


def test_l2_view_has_no_l3_files(client):
    body = _post(client, view="both").get_json()
    assert body["ok"] and body["l3_links"] == []
    assert "l3-links.csv" not in {d["name"] for d in body["downloads"]}


def test_pack_report_and_workbook_carry_the_l3_tables(client):
    body = _post(client).get_json()
    pack = client.get(next(d["url"] for d in body["downloads"] if d["name"] == "keystone-topology-pack.zip")).data
    archive = zipfile.ZipFile(io.BytesIO(pack))
    assert "keystone-topology-pack/tables/l3-links.csv" in archive.namelist()
    report = archive.read("keystone-topology-pack/topology-report.html").decode()
    assert "L3 links" in report and "not measured traffic" in report
    sheet_names = load_workbook(io.BytesIO(archive.read("keystone-topology-pack/topology.xlsx"))).sheetnames
    assert "L3 links" in sheet_names and "L3 routes" in sheet_names


def test_drawio_edges_of_the_l3_page_carry_meaning_and_evidence(client):
    body = _post(client).get_json()
    drawio = client.get(next(d["url"] for d in body["downloads"] if d["name"] == "topology.drawio")).data.decode()
    assert 'l3_id="L3-01"' in drawio and "meaning=" in drawio


def test_provider_hint_creates_a_named_cloud_and_is_recorded_in_run_settings(client):
    body = _post(client, hints="MBR-WSDTQN-70 Gi0/0/0 = mpls Telkom\n").get_json()
    settings = client.get(next(d["url"] for d in body["downloads"] if d["name"] == "run-settings.txt")).data.decode()
    assert "Provider hints typed for the L3 views: 1 rule(s)" in settings


def test_payload_has_howto_trace_devices_and_overlay_rows(client):
    body = _post(client).get_json()
    assert "Three pictures of the same network" in body["howto"] and "RFC 5880" in body["howto"]
    assert body["trace"]["export_id"] and body["trace"]["devices"]
    assert body["l3_overlay"] and {"device", "remote_device", "local_colour", "state"} <= set(body["l3_overlay"][0])
    assert "l3-sdwan-overlay.csv" in {d["name"] for d in body["downloads"]}


def test_trace_endpoint_walks_the_routing_tables_of_the_build(client):
    body = _post(client).get_json()
    device = body["trace"]["devices"][0]["key"]
    answer = client.post("/topology/api/trace", json={"export_id": body["trace"]["export_id"], "source": device, "destination": "203.0.113.77"}).get_json()
    assert answer["ok"] and answer["hops"] and "message" in answer
    assert client.post("/topology/api/trace", json={"export_id": "../x", "source": device, "destination": "1.1.1.1"}).status_code == 404
    assert client.post("/topology/api/trace", json={"export_id": body["trace"]["export_id"], "source": device, "destination": "nope"}).get_json()["ok"] is False


def test_drawio_of_an_l3_build_has_the_how_to_read_page_and_the_l2_build_does_not(client):
    body = _post(client).get_json()
    drawio = client.get(next(d["url"] for d in body["downloads"] if d["name"] == "topology.drawio")).data.decode()
    assert "How to read this diagram" in drawio
    plain = _post(client, view="both").get_json()
    assert "How to read this diagram" not in client.get(next(d["url"] for d in plain["downloads"] if d["name"] == "topology.drawio")).data.decode()

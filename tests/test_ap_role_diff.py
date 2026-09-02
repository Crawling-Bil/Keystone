import tempfile
import unittest
from pathlib import Path

from features.wireless_center.ap_role_diff import (
    compare_ap_roles,
    read_ap_state,
    read_expected_allocation,
)
from features.wireless_center.exporters.ap_role_exporter import export_ap_role_diff


class APRoleDiffTest(unittest.TestCase):
    def inventory(self, records):
        return {"records": records, "issues": [], "source_type": "TEST"}

    def ap(self, name, mac, status, **values):
        return {
            "ap_id": values.get("ap_id", "1"),
            "name": name,
            "mac": mac,
            "serial": values.get("serial", ""),
            "ip": values.get("ip", "192.0.2.10"),
            "group": values.get("group", "GROUP-A"),
            "model": values.get("model", "AirEngine5776-26"),
            "version": values.get("version", "V600R023C10SPC200"),
            "status": status,
            "location": "",
            "expected_active": "",
        }

    def temp_file(self, content, suffix):
        handle = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
        handle.write(content.encode("utf-8"))
        handle.close()
        self.addCleanup(Path(handle.name).unlink, missing_ok=True)
        return Path(handle.name)

    def test_healthy_roles_are_classified_in_both_directions(self):
        a = self.inventory([
            self.ap("AP-A", "0011-2233-4401", "nor"),
            self.ap("AP-B", "0011-2233-4402", "standby"),
        ])
        b = self.inventory([
            self.ap("AP-A", "0011-2233-4401", "standby"),
            self.ap("AP-B", "0011-2233-4402", "normal"),
        ])
        result = compare_ap_roles(a, b)
        roles = {row["ap_name"]: row["role"] for row in result["rows"]}
        self.assertEqual(roles, {"AP-A": "Active A", "AP-B": "Active B"})
        self.assertEqual(result["summary"]["healthy_pairs"], 2)
        self.assertEqual(result["summary"]["critical"], 0)

    def test_dual_active_no_active_and_missing_peer_are_critical(self):
        a = self.inventory([
            self.ap("AP-DUAL", "0011-2233-4401", "normal"),
            self.ap("AP-NONE", "0011-2233-4402", "standby"),
            self.ap("AP-MISSING", "0011-2233-4403", "normal"),
        ])
        b = self.inventory([
            self.ap("AP-DUAL", "0011-2233-4401", "normal"),
            self.ap("AP-NONE", "0011-2233-4402", "standby"),
        ])
        result = compare_ap_roles(a, b)
        rows = {row["ap_name"]: row for row in result["rows"]}
        self.assertEqual(rows["AP-DUAL"]["role"], "Dual Active")
        self.assertEqual(rows["AP-NONE"]["role"], "No Active WLC")
        self.assertEqual(rows["AP-MISSING"]["role"], "Backup B Not Ready")
        self.assertTrue(rows["AP-MISSING"]["missing_peer"])
        self.assertEqual(result["summary"]["critical"], 3)

    def test_serial_fallback_exposes_attribute_mismatch(self):
        a = self.inventory([self.ap("AP-01", "0011-2233-4401", "normal", serial="SERIAL-1", group="GROUP-A")])
        b = self.inventory([self.ap("AP-01", "0011-2233-4499", "standby", serial="SERIAL-1", group="GROUP-B")])
        result = compare_ap_roles(a, b)
        row = result["rows"][0]
        self.assertEqual(row["match_method"], "Serial")
        self.assertEqual(row["role"], "Active A")
        self.assertEqual(set(row["attribute_mismatches"]), {"MAC", "AP Group"})
        self.assertEqual(row["severity"], "Warning")

    def test_expected_allocation_finds_role_mismatch(self):
        a = self.inventory([self.ap("AP-01", "0011-2233-4401", "standby")])
        b = self.inventory([self.ap("AP-01", "0011-2233-4401", "normal")])
        allocation = {"records": [{"name": "AP-01", "mac": "0011-2233-4401", "serial": "", "expected_active": "A"}], "issues": []}
        result = compare_ap_roles(a, b, allocation)
        self.assertEqual(result["rows"][0]["allocation_status"], "Role Mismatch")
        self.assertEqual(result["summary"]["allocation_mismatch"], 1)

    def test_name_only_allocation_matches_when_inventory_identity_is_limited(self):
        a = self.inventory([self.ap("AP-01", "", "normal")])
        b = self.inventory([self.ap("AP-01", "", "standby")])
        allocation = {"records": [{"name": "AP-01", "mac": "", "serial": "", "expected_active": "A"}], "issues": []}
        result = compare_ap_roles(a, b, allocation)
        self.assertEqual(result["rows"][0]["allocation_status"], "Matched")

    def test_tab_prefixed_huawei_csv_and_allocation_are_read(self):
        inventory = self.temp_file(
            "\tAP ID,\tAP name,\tStatus,\tMAC address,\tAP group,\tIP address,\tAP type,\tSystem version,\tSerial Number\n"
            "\t1,\tAP-01,\tstandby,\t0011-2233-4401,\tGROUP-A,\t192.0.2.10,\tAirEngine5776-26,\tV600R023C10SPC200,\tSERIAL-1\n",
            ".csv",
        )
        allocation_file = self.temp_file(
            "AP Name,MAC Address,Expected Active WLC\nAP-01,0011-2233-4401,Controller A\n",
            ".csv",
        )
        state = read_ap_state(inventory, "Controller A")
        allocation = read_expected_allocation(allocation_file, "Controller A", "Controller B")
        self.assertEqual(state["records"][0]["status"], "Standby")
        self.assertEqual(state["records"][0]["version"], "V600R023C10SPC200")
        self.assertEqual(allocation["records"][0]["expected_active"], "A")

    def test_huawei_group_name_header_is_read(self):
        inventory = self.temp_file(
            '"\tAP ID","\tAP MAC","\tAP name","\tGroup name","\tStatus"\n'
            '"\t1","\t0011-2233-4401","\tAP-01","\tGROUP-WAREHOUSE","\tnor"\n',
            ".csv",
        )
        state = read_ap_state(inventory, "Controller A")
        self.assertEqual(state["records"][0]["group"], "GROUP-WAREHOUSE")
        self.assertEqual(state["records"][0]["status"], "Normal")

    def test_group_summary_and_mapping_statuses(self):
        a = self.inventory([
            self.ap("AP-MATCH", "0011-2233-4401", "normal", group="GROUP-A"),
            self.ap("AP-DRIFT", "0011-2233-4402", "normal", group="GROUP-B"),
            self.ap("AP-ONLY", "0011-2233-4403", "normal", group="GROUP-C"),
        ])
        b = self.inventory([
            self.ap("AP-MATCH", "0011-2233-4401", "standby", group="GROUP-A"),
            self.ap("AP-DRIFT", "0011-2233-4402", "standby", group="GROUP-X"),
        ])
        result = compare_ap_roles(a, b)
        rows = {row["ap_name"]: row for row in result["rows"]}
        self.assertEqual(rows["AP-MATCH"]["group_status"], "Matched")
        self.assertEqual(rows["AP-DRIFT"]["group_status"], "Group Mismatch")
        self.assertEqual(rows["AP-ONLY"]["group_status"], "Peer Group Missing")
        self.assertEqual(result["summary"]["ap_groups"], 3)
        self.assertEqual(result["summary"]["group_matched"], 1)
        self.assertEqual(result["summary"]["group_mismatch"], 1)
        self.assertEqual(result["summary"]["peer_group_missing"], 1)
        groups = {item["group"]: item for item in result["group_summary"]}
        self.assertEqual(groups["GROUP-B"]["group_mismatch"], 1)
        self.assertEqual(groups["GROUP-C"]["peer_group_missing"], 1)

    def test_display_ap_all_text_is_read(self):
        capture = self.temp_file(
            "<WLC-A> display ap all\n"
            "ID MAC Name Group IP Type State STA Uptime\n"
            "0* 0011-2233-4401 AP-01 GROUP-A 192.0.2.10 AirEngine5776-26 nor 5 1D:2H\n",
            ".txt",
        )
        state = read_ap_state(capture, "Controller A")
        self.assertEqual(len(state["records"]), 1)
        self.assertEqual(state["records"][0]["status"], "Normal")
        self.assertEqual(state["records"][0]["model"], "AirEngine5776-26")

    def test_excel_export_contains_role_workbook(self):
        a = self.inventory([self.ap("AP-01", "0011-2233-4401", "normal")])
        b = self.inventory([self.ap("AP-01", "0011-2233-4401", "standby")])
        result = compare_ap_roles(a, b)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "role-diff.xlsx"
            export_ap_role_diff(result, output)
            from openpyxl import load_workbook
            workbook = load_workbook(output, read_only=True, data_only=True)
            self.assertIn("AP Role Matrix", workbook.sheetnames)
            self.assertIn("Group Summary", workbook.sheetnames)
            self.assertIn("AP by Group", workbook.sheetnames)
            self.assertIn("Group Mismatch", workbook.sheetnames)
            self.assertIn("Status Guide", workbook.sheetnames)
            self.assertEqual(len(workbook.sheetnames), 12)
            self.assertEqual(workbook["AP Role Matrix"].max_row, 2)
            self.assertEqual(workbook["Group Summary"].max_row, 2)
            self.assertEqual(workbook["AP by Group"].max_row, 2)
            self.assertEqual(workbook["Status Guide"].max_row, 10)

    def test_plain_language_status_guide_is_in_result(self):
        a = self.inventory([self.ap("AP-01", "0011-2233-4401", "normal")])
        b = self.inventory([self.ap("AP-01", "0011-2233-4401", "standby")])
        result = compare_ap_roles(a, b)
        guide = {item["status"]: item for item in result["status_guide"]}
        self.assertEqual(set(guide), {"Normal", "Standby", "Idle", "Fault"})
        self.assertIn("bekerja normal", guide["Normal"]["meaning"])
        self.assertIn("WLC cadangan", guide["Standby"]["meaning"])
        self.assertIn("CAPWAP", guide["Idle"]["action"])
        self.assertIn("registrasi", guide["Fault"]["meaning"])


if __name__ == "__main__":
    unittest.main()

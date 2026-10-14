"""Tests for features/analyzer_core -- the vendor-neutral building blocks shared
by the Router, Firewall and Load Balancer analyzers."""

from __future__ import annotations

import io
import json
import re
from pathlib import Path

import pytest
from flask import Flask
from openpyxl import load_workbook

from features.analyzer_core import collect
from features.analyzer_core.blueprint import AnalyzerSpec, create_analyzer_blueprint, decode_bytes
from features.analyzer_core.capture import CISCO_PROMPT, command_matches, normalize_command, single_block_capture, split_capture
from features.analyzer_core.export import export_results, safe_cell
from features.analyzer_core.findings import dedupe_findings, finding, severity_counts, sort_findings
from features.analyzer_core.redaction import HIDDEN, contains_secret, redact_line, secret_strength
from features.analyzer_core.result import card, kv_section, new_result, table_section, validate_result

REPO = Path(__file__).resolve().parents[1]


# --------------------------------------------------------------------- capture
CAPTURE = """\
edge1#show version
Cisco IOS XE Software, Version 17.9.3a
edge1#sh ip int br
Interface              IP-Address      OK? Method Status                Protocol
Loopback0              192.0.2.1       YES NVRAM  up                    up
edge1#show running-config
hostname edge1
interface Loopback0
 ip address 192.0.2.1 255.255.255.255
edge1#show ip rou
% Invalid input detected at '^' marker.
edge1#
"""


def test_split_capture_finds_commands_and_host():
    capture = split_capture(CAPTURE, (CISCO_PROMPT,))
    assert capture.has_prompts
    assert [block.command for block in capture.blocks] == ["show version", "sh ip int br", "show running-config", "show ip rou"]
    assert capture.hostname() == "edge1"


def test_abbreviated_commands_match_their_template():
    capture = split_capture(CAPTURE, (CISCO_PROMPT,))
    block = capture.first("show ip interface brief")
    assert block is not None and block.command == "sh ip int br"
    assert "Loopback0" in block.text


def test_errored_command_is_not_usable():
    capture = split_capture(CAPTURE, (CISCO_PROMPT,))
    assert capture.find("show ip route") == []          # '% Invalid input' output must never be parsed as a table
    assert capture.find("show ip route", usable=False)  # but it is still visible for the coverage list


@pytest.mark.parametrize(
    "typed, template, expected",
    [
        ("show ip int br", "show ip interface brief", True),
        ("sh run", "show running-config", True),
        ("show ip bgp neigh", "show ip bgp neighbors", True),
        ("show ip bgp summary", "show ip bgp neighbors", False),
        ("show version", "show ip route", False),
    ],
)
def test_command_matches_token_prefixes(typed, template, expected):
    assert (command_matches(normalize_command(typed), template) is not None) is expected


def test_single_block_capture_wraps_bare_config():
    capture = single_block_capture("hostname r1\ninterface Gi0/0\n")
    assert len(capture.blocks) == 1 and capture.blocks[0].command == "show running-config"


# ------------------------------------------------------------------- redaction
@pytest.mark.parametrize(
    "line, secret",
    [
        ("enable secret 5 $1$abcd$Zxy12345", "$1$abcd$Zxy12345"),
        ("username admin privilege 15 password 7 070C285F4D06", "070C285F4D06"),
        ("snmp-server community Sup3rSecret RW", "Sup3rSecret"),
        ("tacacs-server key 7 0822455D0A16", "0822455D0A16"),
        ("crypto isakmp key MyPsk address 203.0.113.9", "MyPsk"),
        (" neighbor 10.0.0.1 password 0 bgpsecret", "bgpsecret"),
        ("    set psksecret ENC AbCdEf123456==", "AbCdEf123456=="),
    ],
)
def test_redact_line_hides_secret_values(line, secret):
    assert contains_secret(line)
    assert secret not in redact_line(line)
    assert HIDDEN in redact_line(line)


def test_redact_line_keeps_ordinary_lines():
    assert redact_line("interface GigabitEthernet0/0") == "interface GigabitEthernet0/0"
    assert not contains_secret("ip address 10.0.0.1 255.255.255.0")


@pytest.mark.parametrize("code, strength", [("0", "weak"), ("7", "weak"), ("5", "weak"), ("9", "strong"), ("8", "strong"), (None, "weak")])
def test_secret_strength_classes(code, strength):
    assert secret_strength(code)[1] == strength


# ---------------------------------------------------------------------- result
def make_result(**extra):
    result = new_result("router", hostname="r1", vendor="Cisco", platform="ASR", source_name="r1.txt", title="r1")
    result["cards"] = [card("Routes", 5, "info", "x", "t1")]
    result["findings"] = [finding("X-1", "high", "Cat", "Title", obj="obj", detail="d", recommendation="r", reference="https://example.invalid/doc")]
    result["sections"] = [
        table_section("t1", "Table", [("name", "Name"), ("n", "Count", "num"), ("st", "State", "status")],
                      [{"name": "a", "n": 1, "st": "up", "_tone": "ok", "_tones": {"st": "ok", "bogus": "ok", "n": "not-a-tone"}}], group="G"),
        kv_section("k1", "Pairs", [("A", "1"), ("B", "2", "warn")], group="G"),
    ]
    result.update(extra)
    return result


def test_table_section_keeps_only_known_columns_and_tones():
    section = make_result()["sections"][0]
    row = section["rows"][0]
    assert set(row) == {"name", "n", "st", "_tone", "_tones"}
    assert row["_tones"] == {"st": "ok"}


def test_validate_result_accepts_good_and_rejects_bad():
    assert validate_result(make_result()) == []
    broken = make_result()
    broken["schema"] = "nope"
    broken["cards"][0]["tone"] = "purple"
    problems = validate_result(broken)
    assert problems and any("schema" in p.lower() for p in problems)


def test_card_rejects_unknown_tone():
    with pytest.raises(ValueError):
        card("x", 1, "purple")


def test_findings_are_sorted_and_deduplicated():
    items = [
        finding("A", "low", "c", "t1", obj="o"),
        finding("B", "high", "c", "t2", obj="o"),
        finding("A", "low", "c", "t1", obj="o"),
        finding("C", "medium", "c", "t3", obj="o"),
    ]
    out = sort_findings(dedupe_findings(items))
    assert [f["id"] for f in out] == ["B", "C", "A"]
    assert severity_counts(out) == {"high": 1, "medium": 1, "low": 1, "info": 0}


# ---------------------------------------------------------------------- export
def test_safe_cell_blocks_formula_injection():
    for value in ("=cmd|' /C calc'!A0", "+1+1", "-2+3", "@SUM(A1)"):
        assert str(safe_cell(value)).startswith("'")
    assert safe_cell("plain") == "plain"
    assert safe_cell(7) == 7


def test_export_writes_expected_sheets(tmp_path):
    target = export_results([make_result(), make_result(hostname="r2", source_name="r2.txt")], tmp_path / "out.xlsx", analyzer_label="Router Analyzer")
    wb = load_workbook(target)
    assert wb.sheetnames[0] == "Summary"
    assert "Findings" in wb.sheetnames and "Device Info" in wb.sheetnames
    summary = [[c.value for c in row] for row in wb["Summary"].iter_rows()]
    assert any("r1" in [str(v) for v in row] for row in summary) and any("r2" in [str(v) for v in row] for row in summary)
    findings = [[c.value for c in row] for row in wb["Findings"].iter_rows(values_only=False)]
    assert any("X-1" in [str(v) for v in row] for row in findings)


# --------------------------------------------------------------------- collect
def test_collect_discover_reads_description_and_count(tmp_path):
    (tmp_path / "collect-router-cisco-ios-xe.txt").write_text("!!! Cisco IOS-XE router capture\nterminal length 0\nshow version\n\nshow running-config\n", encoding="utf-8")
    (tmp_path / "collect-firewall-forti.txt").write_text("get system status\n", encoding="utf-8")
    items = collect.discover("collect-router", tmp_path)
    assert len(items) == 1
    item = items[0]
    assert item["file"] == "collect-router-cisco-ios-xe.txt"
    assert item["commands"] == 3
    assert item["description"] == "Cisco IOS-XE router capture"
    assert item["title"] == "Cisco IOS-XE"
    assert item["url"].endswith("/collect-router-cisco-ios-xe.txt")


def test_collect_discover_empty_when_nothing_matches(tmp_path):
    assert collect.discover("collect-router", tmp_path) == []
    assert collect.discover("collect-router", tmp_path / "does-not-exist") == []


def test_decode_bytes_handles_bom_and_utf16():
    assert decode_bytes("﻿hello".encode("utf-8")) == "hello"
    assert decode_bytes("hello".encode("utf-16")) == "hello"


# ------------------------------------------------------------------- blueprint
@pytest.fixture
def client(tmp_path):
    def fake_analyze(content, filename="", platform="auto"):
        if "boom" in content:
            raise ValueError("cannot read this")
        result = make_result(hostname=(re.search(r"hostname (\S+)", content) or [None, "dev"])[1])
        return result

    template_dir = tmp_path / "templates"
    template_dir.mkdir()
    (template_dir / "page.html").write_text("page {{ marker }} {{ collect_scripts|length }}", encoding="utf-8")
    spec = AnalyzerSpec(key="demo", name="Demo Analyzer", url_prefix="/demo", template="page.html", analyze=fake_analyze, collect_prefix="collect-demo",
                        base_dir=tmp_path, page_context={"marker": "ok"})
    app = Flask(__name__, template_folder=str(template_dir))
    app.register_blueprint(create_analyzer_blueprint(spec))
    return app.test_client()


def test_blueprint_page_renders(client):
    assert client.get("/demo/").get_data(as_text=True).startswith("page ok")


def test_blueprint_analyze_multi_file_and_partial_failure(client):
    data = {"config_files": [(io.BytesIO(b"hostname alpha"), "a.txt"), (io.BytesIO(b"boom"), "b.txt"), (io.BytesIO(b"hostname gamma"), "c.txt")]}
    body = client.post("/demo/api/analyze", data=data, content_type="multipart/form-data").get_json()
    assert body["ok"] and [r["hostname"] for r in body["results"]] == ["alpha", "gamma"]
    assert len(body["errors"]) == 1 and body["errors"][0]["source"] == "b.txt"


def test_blueprint_analyze_validation_errors(client):
    assert client.post("/demo/api/analyze", data={}).status_code == 400
    response = client.post("/demo/api/analyze", data={"config_files": (io.BytesIO(b"boom"), "b.txt")}, content_type="multipart/form-data")
    assert response.status_code == 400 and not response.get_json()["ok"]
    assert response.get_json()["error"] == "cannot read this"                   # a ValueError is a message for the engineer, shown without the class name


def test_blueprint_unexpected_errors_keep_their_type(tmp_path):
    spec = AnalyzerSpec(key="bad", name="Bad Analyzer", url_prefix="/bad", template="page.html", analyze=lambda content, filename="", platform="auto": {}["missing"], collect_prefix="collect-bad", base_dir=tmp_path)
    app = Flask(__name__)
    app.register_blueprint(create_analyzer_blueprint(spec))
    body = app.test_client().post("/bad/api/analyze", data={"config_text": "x"}).get_json()
    assert body["error"].startswith("KeyError")


def test_blueprint_pasted_text_and_file_limit(client):
    assert client.post("/demo/api/analyze", data={"config_text": "hostname pasted"}).get_json()["results"][0]["hostname"] == "pasted"
    many = {"config_files": [(io.BytesIO(b"hostname x"), f"{i}.txt") for i in range(17)]}
    assert client.post("/demo/api/analyze", data=many, content_type="multipart/form-data").status_code == 400


def test_blueprint_export_and_download_roundtrip(client):
    results = [make_result(), make_result(hostname="r2")]
    exported = client.post("/demo/api/export", json={"results": results}).get_json()
    assert exported["ok"] and exported["download_name"].endswith(".xlsx") and "combined" in exported["download_name"]
    download = client.get(exported["download_url"])
    assert download.status_code == 200 and download.data[:2] == b"PK"
    wb = load_workbook(io.BytesIO(download.data))
    assert "Findings" in wb.sheetnames


def test_blueprint_export_rejects_bad_payloads(client):
    assert client.post("/demo/api/export", json={}).status_code == 400
    assert client.post("/demo/api/export", json={"results": [{"schema": "nope"}]}).status_code == 400
    assert client.get("/demo/download/../../etc/passwd").status_code in (400, 404)
    assert client.get("/demo/download/not-a-valid-id").status_code == 400


def test_blueprint_collect_scripts_endpoint(client):
    body = client.get("/demo/api/collect-scripts").get_json()
    assert body["ok"] and body["scripts"] == []


def test_table_section_carries_filter_button_definitions():
    facets = [{"key": "vrf", "label": "VRF", "prefix": "vrf "}]
    section = table_section("t", "Table", [("vrf", "VRF"), ("net", "Network")], [{"vrf": "global", "net": "10.0.0.0/8"}], facets=facets)
    assert section["facets"] == facets
    assert table_section("t", "Table", [("vrf", "VRF")], [{"vrf": "global"}])["facets"] == []


def test_facet_buttons_are_built_by_the_shared_kit():
    kit = (Path(__file__).resolve().parents[1] / "static" / "js" / "analyzer_kit.js").read_text(encoding="utf-8")
    assert "facetDefs" in kit and "paintFacets" in kit
    coverage = kit[kit.index("coverageView() {"):]
    assert "facets:" in coverage[:1500] and "not parsed" in coverage[:1500]


def test_kit_declares_the_filter_button_row_before_it_assembles_the_table_block():
    """buildTable used ``facetBox`` (a const) before its declaration: a temporal-dead-zone error that broke every router
    table, and that no Python test could see. Found in the browser; this keeps the order."""
    kit = (Path(__file__).resolve().parents[1] / "static" / "js" / "analyzer_kit.js").read_text(encoding="utf-8")
    declared = kit.index("const facetBox =")
    assert declared < kit.index("wrap.append(...[facetBox"), "facetBox must be declared before it is added to the table block"
    assert kit.count("const facetBox =") == 1 and kit.count("const facetDefs =") == 1


def test_kit_never_hands_a_missing_element_to_append():
    """``Node.append(null)`` writes the text "null": every table without filter buttons showed it above the toolbar, and so did a facet with a
    single value. Found in the browser. Both places must drop the missing element."""
    kit = (Path(__file__).resolve().parents[1] / "static" / "js" / "analyzer_kit.js").read_text(encoding="utf-8")
    assert "wrap.append(...[facetBox, tools, scroller, pager].filter(Boolean))" in kit
    paint = kit[kit.index("function paintFacets()"):kit.index("function filtered()")]
    assert "return null;" in paint and ".filter(Boolean));" in paint

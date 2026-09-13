"""Smoke tests for FirewallExcelExporter's newer sheets (Virtual
Routers, Management Profiles, Zone Protection Profiles, Device
Configuration, Administrators, PBF Rules) -- added alongside the
routing-detail and new-object-type parsing in paloalto.py/
paloalto_xml.py. Runs the exporter against a real analyze_firewall()
result rather than a hand-built dict, so it exercises the actual
dashboard shape firewall_service.py produces.
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from features.switch_analyzer.firewall_service import analyze_firewall
from features.switch_analyzer.exporters.firewall_excel_exporter import FirewallExcelExporter

PALOALTO_SET_CONFIG = """
set deviceconfig system hostname EXPORT-FW01
set network interface ethernet ethernet1/1 layer3
set network interface ethernet ethernet1/1 layer3 ip 203.0.113.5/29
set zone Outside network layer3 [ ethernet1/1 ]
set zone Outside network zone-protection-profile ZPP-Edge
set network virtual-router default interface [ ethernet1/1 ]
set network virtual-router default protocol bgp enable yes
set network virtual-router default protocol bgp router-id 203.0.113.5
set network virtual-router default protocol bgp local-as 65030
set network profiles interface-management-profile Outside-Mgmt ping yes
set network profiles zone-protection-profile ZPP-Edge flood tcp-syn enable yes
set mgt-config users fw-admin permissions role-based superuser yes
set rulebase pbf rules PBF-Backup from zone [ Outside ]
set rulebase pbf rules PBF-Backup source [ any ]
set rulebase pbf rules PBF-Backup destination [ any ]
set rulebase pbf rules PBF-Backup application [ any ]
set rulebase pbf rules PBF-Backup service [ any ]
set rulebase pbf rules PBF-Backup action forward egress-interface ethernet1/1
set rulebase pbf rules PBF-Backup action forward nexthop ip-address 203.0.113.1
set rulebase pbf rules PBF-Backup disabled no
"""


class FirewallExcelExporterNewSheetsTests(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = Path(tempfile.mkdtemp())
        self.addCleanup(__import__("shutil").rmtree, self.tmp_dir, ignore_errors=True)

        result = analyze_firewall(PALOALTO_SET_CONFIG, vendor="Palo Alto")
        out_path = self.tmp_dir / "export.xlsx"
        FirewallExcelExporter([result]).export(out_path)
        self.workbook = load_workbook(out_path)

    def test_all_new_sheets_are_present(self):
        expected = {
            "Virtual Routers (Palo Alto)",
            "Management Profiles (Palo Alto)",
            "Zone Protection Profiles (PA)",
            "Device Configuration (PA)",
            "Administrators (Palo Alto)",
            "PBF Rules (Palo Alto)",
        }
        self.assertTrue(expected.issubset(set(self.workbook.sheetnames)))

    def test_virtual_router_sheet_shows_routing_type(self):
        sheet = self.workbook["Virtual Routers (Palo Alto)"]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        row = dict(zip([c.value for c in sheet[1]], rows[0]))
        self.assertEqual(row["Routing Type"], "Dynamic")
        self.assertEqual(row["BGP Enabled"], "Yes")
        self.assertEqual(row["BGP Router ID"], "203.0.113.5")
        self.assertEqual(row["BGP AS Number"], "65030")

    def test_administrator_and_pbf_rows_populated(self):
        admin_sheet = self.workbook["Administrators (Palo Alto)"]
        admin_rows = list(admin_sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(admin_rows[0][1:], ("fw-admin", "superuser"))

        pbf_sheet = self.workbook["PBF Rules (Palo Alto)"]
        pbf_rows = list(pbf_sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(pbf_rows[0][1], "PBF-Backup")

    def test_zone_sheet_shows_protection_profile(self):
        zone_sheet = self.workbook["Zones"]
        header = [c.value for c in zone_sheet[1]]
        self.assertIn("Zone Protection Profile (Palo Alto)", header)
        rows = list(zone_sheet.iter_rows(min_row=2, values_only=True))
        row = dict(zip(header, rows[0]))
        self.assertEqual(row["Zone Protection Profile (Palo Alto)"], "ZPP-Edge")


if __name__ == "__main__":
    unittest.main()

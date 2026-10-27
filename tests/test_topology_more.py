"""Topology Builder: LLDP identity fields, multi-device files, ZIP, hosts, routed links, layout, routing, Excel. Synthetic only."""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest

sys.path.insert(0, str(Path(__file__).parent))
from topology_samples import iosxe_capture, lab_dc, nxos_arp, nxos_capture, nxos_mac_table  # noqa: E402

from features.switch_analyzer.service import analyze  # noqa: E402
from features.topology import hosts as hosts_module  # noqa: E402
from features.topology.excel import to_xlsx  # noqa: E402
from features.topology.extract import mac12  # noqa: E402
from features.topology.ingest import IngestError, to_captures  # noqa: E402
from features.topology.layout_io import read_layout  # noqa: E402
from features.topology.service import build_topology, devices_table, hosts_table, links_table  # noqa: E402
from features.topology.split import split_devices  # noqa: E402


# ---- LLDP identity fields ----------------------------------------------------------------------------------------

def _neighbors(capture: str, header: str = "") -> list[dict]:
    return analyze(header + capture, filename="x.txt")["neighbors"]


NXOS_HEAD = "LAB-AGG1# show version\nCisco Nexus Operating System (NX-OS) Software\n  NXOS: version 10.3(5)\n  cisco Nexus9000 N9K-C93180YC-FX3 Chassis\n  Device name: LAB-AGG1\n\n"
IOSXE_HEAD = "LAB-ACC1#show version\nCisco IOS XE Software, Version 17.09.04\nCisco IOS Software [Cupertino], Catalyst L3 Switch Software\nLAB-ACC1 uptime is 1 day\n\n"
HUAWEI_HEAD = "<LAB-HW1>display version\nHuawei Versatile Routing Platform Software\nVRP (R) software, Version 8.150 (CE6850 V200R005C10SPC800)\nHUAWEI CE6850-48S4Q-EI uptime is 1 day\n\n"
ARUBA_HEAD = "LAB-A# show version\nImage stamp:  /ws/swbuildm/rel_xxx\n                 WC.16.10.0012\nBoot Image: Primary\n\nLAB-A# show system\n Status and Counters - General System Information\n  System Name        : LAB-A\n\n"


NXOS_LLDP = """LAB-AGG1# show lldp neighbors detail
Capability codes:
  (R) Router, (B) Bridge, (T) Telephone, (C) DOCSIS Cable Device
  (W) WLAN Access Point, (P) Repeater, (S) Station, (O) Other
Device ID            Local Intf      Hold-time  Capability  Port ID

Chassis id: 0011.2233.4455
Port id: Ethernet1/49
Local Port id: Eth1/1
Port Description: LINK LAB-AGG1
System Name: LAB-TOR1
System Description: Cisco Nexus Operating System (NX-OS) Software 10.5(6)
TAC support: http://www.cisco.com/tac
Time remaining: 115 seconds
System Capabilities: B, R
Enabled Capabilities: B, R
Management Address: 10.99.5.200
Management Address IPV6: not advertised
Vlan ID: 1
System UUID: FDO00000099

Chassis id: 0011.2233.4466
Port id: Ethernet1/50
Local Port id: Eth1/2
Port Description: LINK LAB-AGG1
System Name: LAB-TOR1
System Description: Cisco Nexus Operating System (NX-OS) Software 10.5(6)
Time remaining: 116 seconds
System Capabilities: B, R
Management Address: 10.99.5.200
Vlan ID: 1
System UUID: FDO00000099

LAB-AGG1# 
"""


def test_nxos_lldp_detail_keeps_chassis_serial_and_management_address():
    rows = [r for r in _neighbors(NXOS_LLDP, NXOS_HEAD) if r["protocol"] == "LLDP"]
    assert [r["local_interface"] for r in rows] == ["Eth1/1", "Eth1/2"]          # two blocks, no divider lines between them
    first = rows[0]
    assert first["neighbor_id"] == "LAB-TOR1" and first["remote_interface"] == "Ethernet1/49"
    assert first["chassis_id"] == "0011.2233.4455" and first["management_address"] == "10.99.5.200"
    assert first["serial"] == "FDO00000099" and first["capabilities"] == "B, R" and first["port_description"] == "LINK LAB-AGG1"
    assert first["platform"].startswith("Cisco Nexus Operating System") and "TAC support" not in first["platform"]


def test_iosxe_lldp_detail_reads_the_management_address_block():
    capture = """LAB-ACC1#show lldp neighbors detail
------------------------------------------------
Local Intf: Gi1/0/1
Chassis id: aabb.ccdd.0001
Port id: Gi1/0/9
Port Description: GigabitEthernet1/0/9
System Name: LAB-ACC2

System Description:
Cisco IOS Software [Dublin], Catalyst L3 Switch Software

Time remaining: 99 seconds
System Capabilities: B,R
Enabled Capabilities: B,R
Management Addresses:
    IP: 10.99.0.22
Auto Negotiation - supported, enabled
Physical media capabilities:
------------------------------------------------
LAB-ACC1#
"""
    row = next(r for r in _neighbors(capture, IOSXE_HEAD) if r["protocol"] == "LLDP")
    assert row["management_address"] == "10.99.0.22" and row["chassis_id"] == "aabb.ccdd.0001" and row["capabilities"] == "B,R"


def test_huawei_lldp_detail_keeps_chassis_id_and_management_address():
    capture = """<LAB-HW1>display lldp neighbor
XGE0/0/1 has 1 neighbor(s):

Neighbor index                     :1
Chassis type                       :MAC Address
Chassis ID                         :aabb-ccdd-0001
Port ID subtype                    :Interface Name
Port ID                            :XGE0/0/5
Port description                   :uplink
System name                        :LAB-HW2
System description                 :Huawei Switch
Management address type            :ipv4
Management address value           :10.88.0.12
Port VLAN ID(PVID)                 :10

XGE0/0/2 has 0 neighbor(s)
<LAB-HW1>"""
    row = next(r for r in _neighbors(capture, HUAWEI_HEAD) if r["protocol"] == "LLDP")
    assert row["neighbor_id"] == "LAB-HW2" and row["remote_interface"] == "XGE0/0/5"
    assert row["chassis_id"] == "aabb-ccdd-0001" and row["management_address"] == "10.88.0.12"
    assert row["native_vlan"] == "10" and row["port_description"] == "uplink"


def test_arubaos_switch_lldp_detail_keeps_chassis_and_address():
    capture = """LAB-A# show lldp info remote-device detail

 LLDP Remote Device Information Detail

  Local Port   : 1
  ChassisType  : mac-address
  ChassisId    : 00 11 22 33 44 55
  PortType     : local
  PortId       : 25
  SysName      : LAB-ARUBA2
  System Descr : test
  PortDescr    : 25

  System Capabilities Supported  : bridge, router
  System Capabilities Enabled    : bridge

  Remote Management Address
     Type    : ipv4
     Address : 10.88.1.2

LAB-A#
"""
    row = next(r for r in _neighbors(capture, ARUBA_HEAD) if r["protocol"] == "LLDP")
    assert row["neighbor_id"] == "LAB-ARUBA2" and row["chassis_id"] == "00 11 22 33 44 55"
    assert row["management_address"] == "10.88.1.2" and row["capabilities"] == "bridge"


def test_rows_without_the_new_fields_keep_their_old_shape():
    row = _neighbors(iosxe_capture("LAB-X", "FOC00000042", "10.99.0.42", cdp=[("LAB-Y", "10.99.0.43", "C9300L", "Gi1/0/1", "Gi1/0/1")]))[0]
    assert "chassis_id" not in row and row["protocol"] == "CDP"


def test_chassis_id_identifies_an_unnamed_neighbor_that_matches_a_captured_stack_member():
    assert mac12("7cad.4f17.1200") == mac12("7C:AD:4F:17:12:00") == "7cad4f171200" and mac12("nonsense") == ""


# ---- multi-device files and ZIP ----------------------------------------------------------------------------------

def test_a_file_with_several_device_sessions_is_split():
    files = lab_dc()
    parts = split_devices("everything.txt", "\n".join(files.values()))
    assert len(parts) == 3
    result = build_topology(parts)
    assert result["summary"]["captured"] == 3 and result["summary"]["confirmed"] >= 3


def test_one_device_file_and_a_stray_prompt_do_not_split():
    capture = lab_dc()["LAB-AGG1.txt"] + "\nSomeone#\nOther# show clock\n"
    assert len(split_devices("a.txt", capture)) == 1


def _zip(entries: dict[str, str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, text in entries.items():
            archive.writestr(name, text)
    return buffer.getvalue()


def test_zip_members_are_read_with_safe_names():
    files = lab_dc()
    data = _zip({"site/LAB-AGG1.txt": files["LAB-AGG1.txt"], "../../evil/LAB-AGG2.txt": files["LAB-AGG2.txt"], "readme.pdf": "x", "__MACOSX/._a.txt": "x", ".hidden.txt": "x"})
    captures = to_captures([("bundle.zip", data)], 100)
    assert sorted(name for name, _ in captures) == ["LAB-AGG1.txt", "LAB-AGG2.txt"]


def test_bad_or_oversized_archives_are_refused_with_a_plain_message():
    with pytest.raises(IngestError, match="not a valid ZIP"):
        to_captures([("x.zip", b"not a zip")], 100)
    many = _zip({f"d{i}.txt": "x" for i in range(600)})
    with pytest.raises(IngestError, match="limit"):
        to_captures([("x.zip", many)], 100)
    with pytest.raises(IngestError, match="More than 2"):
        to_captures([(f"{n}", t.encode()) for n, t in lab_dc().items()], 2)


# ---- coverage ----------------------------------------------------------------------------------------------------

def test_coverage_columns_and_the_missing_command_finding():
    result = build_topology(list(lab_dc().items()))
    rows = {r["device"]: r for r in devices_table(result["graph"]) if r["captured"] == "yes"}
    assert rows["LAB-AGG1"]["has_cdp"] == "yes" and rows["LAB-AGG1"]["has_lldp"] == "yes"
    assert rows["LAB-ACC1"]["has_lldp"] == "no"
    finding = next(f for f in result["findings"] if f["id"] == "TOPO-013" and "lldp" in f["title"])
    assert "LAB-ACC1" in finding["devices"]


# ---- hosts -------------------------------------------------------------------------------------------------------

def _dc_with_tables(mac_rows, arp_rows=(), cdp_extra=()):
    files = lab_dc()
    extra = nxos_mac_table("LAB-AGG1", mac_rows) + nxos_arp("LAB-AGG1", list(arp_rows))
    files["LAB-AGG1.txt"] = nxos_capture(
        "LAB-AGG1", "FDO00000001", "10.99.0.11",
        cdp=[("LAB-ACC1.lab.example(FOC00000001)", "10.99.0.21", "C9300L-48P-4X", "Ethernet1/1", "GigabitEthernet1/0/23"),
             ("LAB-AGG2.lab.example(FDO00000002)", "10.99.0.12", "N9K-C93180YC-FX3", "Ethernet1/47", "Ethernet1/47"),
             ("LAB-AGG2.lab.example(FDO00000002)", "10.99.0.12", "N9K-C93180YC-FX3", "Ethernet1/48", "Ethernet1/48")],
        lldp=[("LAB-FW1", "Eth1/40", "port5")], port_channels=[(101, ["Eth1/1"]), (240, ["Eth1/47", "Eth1/48"])],
        vpc=("240", "primary", "240", "10.99.0.12"), extra=extra)
    return files


def test_macs_on_edge_ports_become_hosts_and_macs_behind_uplinks_do_not():
    files = _dc_with_tables(
        [("10", "0050.56aa.0001", "Eth1/10"), ("10", "0050.56aa.0002", "Po101"),            # Po101 is the link to ACC1: not an edge port
         ("10", "0050.56aa.0003", "Eth1/40"),                                               # port toward the firewall: not an edge port
         ("10", "0000.0c07.ac0a", "Eth1/11"), ("10", "0100.5e00.0001", "Eth1/12")],         # HSRP virtual MAC and a multicast MAC: ignored
        [("10.99.10.5", "0050.56aa.0001", "Vlan10")])
    result = build_topology(list(files.items()))
    hosts = hosts_table(result["graph"])
    assert [(h["mac"], h["ip_address"], h["port"], h["behind"]) for h in hosts] == [
        ("0050.56aa.0001", "10.99.10.5", "Eth1/10", ""),
        ("0050.56aa.0003", "", "Eth1/40", "LAB-FW1")]            # listed, but marked as living behind the uncaptured firewall
    assert not any(h["mac"] == "0050.56aa.0002" for h in hosts)  # behind the captured switch: that switch lists it


def test_hosts_are_drawn_only_when_asked_and_many_macs_collapse():
    macs = [("10", f"0050.56bb.{n:04x}", "Eth1/20") for n in range(1, 13)]
    files = _dc_with_tables(macs + [("10", "0050.56aa.0001", "Eth1/10")], [("10.99.10.5", "0050.56aa.0001", "Vlan10")])
    plain = build_topology(list(files.items()))
    drawn = build_topology(list(files.items()), include_hosts=True)
    assert not any(k.startswith(("HOST:", "HOSTS:")) for k in plain["graph"].nodes) and plain["summary"]["hosts"] == 13
    keys = [k for k in drawn["graph"].nodes if k.startswith(("HOST:", "HOSTS:"))]
    assert len(keys) == 2 and any(drawn["graph"].nodes[k].hostname == "12 hosts" for k in keys)
    assert any(f["id"] == "TOPO-007" and "12" in f["detail"] for f in drawn["findings"])
    assert "10.99.10.5" in drawn["drawio"]
    assert drawn["summary"]["links"] == plain["summary"]["links"]          # hosts are not counted as CDP/LLDP links


def test_a_mac_on_two_switches_edge_ports_is_reported_once_and_drawn_once():
    files = _dc_with_tables([("10", "0050.56aa.0001", "Eth1/10")])
    files["LAB-ACC1.txt"] = iosxe_capture("LAB-ACC1", "FOC00000001", "10.99.0.21", cdp=[
        ("LAB-AGG1.lab.example(FDO00000001)", "10.99.0.11", "N9K-C93180YC-FX3", "GigabitEthernet1/0/23", "Ethernet1/1"),
        ("LAB-AGG2.lab.example(FDO00000002)", "10.99.0.12", "N9K-C93180YC-FX3", "GigabitEthernet1/0/24", "Ethernet1/1")],
        port_channels=[(40, ["Gi1/0/23", "Gi1/0/24"])],
        extra="LAB-ACC1#show mac address-table\n          Mac Address Table\n-------------------------------------------\n\nVlan    Mac Address       Type        Ports\n"
              "----    -----------       --------    -----\n  10    0050.56aa.0001    DYNAMIC     Gi1/0/10\nTotal Mac Addresses for this criterion: 1\n")
    result = build_topology(list(files.items()))
    assert result["summary"]["hosts"] == 1 and any(f["id"] == "TOPO-014" for f in result["findings"])


# ---- routed links ------------------------------------------------------------------------------------------------

def _routed_pair(ip_b="10.99.9.2", mask_b="255.255.255.252"):
    a = nxos_capture("LAB-RT1", "FDO00000011", "10.99.0.31", cdp=[], config_extra="interface Ethernet1/48\n  no switchport\n  ip address 10.99.9.1/30\n")
    b = iosxe_capture("LAB-RT2", "FOC00000012", "10.99.0.32", cdp=[],
                      config_extra=f"interface GigabitEthernet1/0/48\n no switchport\n ip address {ip_b} {mask_b}\n!\n")
    return [("a.txt", a), ("b.txt", b)]


def test_a_shared_slash_30_becomes_a_routed_link_between_two_captured_devices():
    result = build_topology(_routed_pair())
    routed = [c for c in result["graph"].connections if c.kind == "routed"]
    assert len(routed) == 1 and routed[0].subnet == "10.99.9.0/30" and routed[0].state == "inferred"
    assert {routed[0].ip_a, routed[0].ip_b} == {"10.99.9.1", "10.99.9.2"}
    assert result["summary"]["routed"] == 1 and result["summary"]["links"] == 0
    assert 'strokeColor=#0070c0;strokeWidth=2;dashed=1' in result["drawio"]
    assert build_topology(_routed_pair(), include_routed=False)["summary"]["routed"] == 0


def test_a_routed_subnet_on_an_existing_cdp_link_annotates_it_instead_of_drawing_twice():
    a = nxos_capture("LAB-RT1", "FDO00000011", "10.99.0.31", cdp=[("LAB-RT2", "10.99.0.32", "C9300L-48P-4X", "Ethernet1/48", "GigabitEthernet1/0/48")],
                     config_extra="interface Ethernet1/48\n  no switchport\n  ip address 10.99.9.1/30\n")
    b = iosxe_capture("LAB-RT2", "FOC00000012", "10.99.0.32", cdp=[("LAB-RT1", "10.99.0.31", "N9K-C93180YC-FX3", "GigabitEthernet1/0/48", "Ethernet1/48")],
                      config_extra="interface GigabitEthernet1/0/48\n no switchport\n ip address 10.99.9.2 255.255.255.252\n!\n")
    result = build_topology([("a.txt", a), ("b.txt", b)])
    assert [c.kind for c in result["graph"].connections] == ["link"] and result["graph"].connections[0].subnet == "10.99.9.0/30"
    assert any(r["subnet"] == "10.99.9.0/30" for r in links_table(result["graph"]))


def test_different_masks_on_a_point_to_point_subnet_are_reported():
    result = build_topology(_routed_pair(ip_b="10.99.9.2", mask_b="255.255.255.248"))
    assert any(f["id"] == "TOPO-012" for f in result["findings"])


def test_a_peer_address_known_from_a_neighbor_table_completes_a_routed_link():
    a = nxos_capture("LAB-RT1", "FDO00000011", "10.99.0.31", cdp=[("LAB-WAN", "10.99.9.2", "ASR1002-X", "Ethernet1/10", "Gi0/0/0")],
                     config_extra="interface Ethernet1/48\n  no switchport\n  ip address 10.99.9.1/30\n")
    result = build_topology([("a.txt", a)])
    assert any(c.kind == "routed" and "LAB-WAN" in (c.a, c.b) for c in result["graph"].connections)


# ---- layout ------------------------------------------------------------------------------------------------------

def _move(drawio: str, node_key: str, x: float, y: float) -> str:
    root = ET.fromstring(drawio)
    for obj in root.iter("object"):
        if obj.get("node_key") == node_key:
            geo = obj.find("mxCell").find("mxGeometry")
            geo.set("x", str(x))
            geo.set("y", str(y))
    return ET.tostring(root, encoding="unicode")


def test_layout_from_a_previous_drawio_is_kept_and_new_devices_do_not_overlap():
    files = lab_dc()
    first = build_topology(list(files.items()))
    moved = _move(first["drawio"], "LAB-ACC1", 1234, 567)
    layout = read_layout(moved.encode())
    assert layout["Overview"]["LAB-ACC1"] == pytest.approx((1234 - (76 - 76) / 2, 567))
    files["LAB-NEW.txt"] = iosxe_capture("LAB-NEW", "FOC00000777", "10.99.0.77", cdp=[
        ("LAB-AGG1.lab.example(FDO00000001)", "10.99.0.11", "N9K-C93180YC-FX3", "GigabitEthernet1/0/1", "Ethernet1/20")])
    second = build_topology(list(files.items()), layout=layout)
    placed = second["pages"][0].nodes
    assert (placed["LAB-ACC1"].x, placed["LAB-ACC1"].y) == pytest.approx((1234, 567))
    assert second["summary"]["layout_kept"] >= 3
    spots = [(p.x, p.y) for p in placed.values()]
    assert len(spots) == len(set(spots))
    assert all(abs(placed["LAB-NEW"].x - p.x) > 20 or abs(placed["LAB-NEW"].y - p.y) > 20 for k, p in placed.items() if k != "LAB-NEW")


def test_layout_reader_refuses_hostile_or_foreign_files():
    with pytest.raises(ValueError, match="DOCTYPE"):
        read_layout(b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "b">]><mxfile/>')
    with pytest.raises(ValueError, match="not a draw.io"):
        read_layout(b"<html/>")
    with pytest.raises(ValueError, match="readable"):
        read_layout(b"not xml at all")
    assert read_layout(b'<mxfile><diagram name="p">AAAA==</diagram></mxfile>') == {}      # compressed pages are skipped, not fatal


# ---- routing -----------------------------------------------------------------------------------------------------

def test_a_line_that_would_cross_a_device_is_bent_around_it():
    from features.topology.layout import Placed
    from features.topology.scene import Page, _route
    from features.topology.graph import Connection, Link
    page = Page("t", nodes={"A": Placed(0, 0, 2), "B": Placed(230, 0, 2), "C": Placed(460, 0, 2)})
    page.connections = [Connection("A", "C", [Link("A", "e1", "C", "e1", "confirmed")]), Connection("A", "B", [Link("A", "e2", "B", "e2", "confirmed")])]
    _route(page)
    assert 0 in page.routes and 1 not in page.routes           # A-C passes through B, A-B is clear
    (x1, y1), (x2, y2) = page.routes[0]
    assert abs(y1 - 38) >= 60 and abs(y1 - y2) < 0.5            # a flat run clear of the row, over the device in between
    assert x1 < 230 < x2


def test_parallel_connections_between_two_devices_are_spread_apart():
    files = lab_dc()
    result = build_topology(list(files.items()))
    graph, page = result["graph"], result["pages"][0]
    pairs = {}
    for i, c in enumerate(page.connections):
        pairs.setdefault(frozenset((c.a, c.b)), []).append(i)
    multi = [ix for ix in pairs.values() if len(ix) > 1]
    for ix in multi:
        assert len({tuple(page.routes.get(i, [])) for i in ix}) == len(ix)


def test_children_follow_their_parents_so_a_simple_two_row_graph_has_no_crossing():
    from types import SimpleNamespace
    from features.topology.layout import _arrange
    rows = {1: ["a", "b"], 3: ["x", "y"]}
    neighbours = {"a": {"y"}, "b": {"x"}, "x": {"b"}, "y": {"a"}}
    spot = _arrange(SimpleNamespace(nodes={}), rows, [1, 3], neighbours, set(), {1: 1, 3: 1})
    assert (spot["a"] < spot["b"]) == (spot["y"] < spot["x"])        # same order above and below: no crossing
    assert abs(spot["y"] - spot["a"]) < 1e-9 and abs(spot["x"] - spot["b"]) < 1e-9


def test_uplink_only_devices_sit_over_the_device_they_serve():
    from types import SimpleNamespace
    from features.topology.layout import _arrange
    rows = {0: ["R1", "R2", "R3"], 1: ["C1", "C2"]}               # R1 serves C2, R2 and R3 serve C1: natural order would cross
    neighbours = {"R1": {"C2"}, "R2": {"C1"}, "R3": {"C1"}, "C1": {"R2", "R3"}, "C2": {"R1"}}
    _arrange(SimpleNamespace(nodes={}), rows, [0, 1], neighbours, set(), {0: 1, 1: 1})
    assert rows[0] == ["R2", "R3", "R1"] and rows[1] == ["C1", "C2"]


# ---- Excel and routes --------------------------------------------------------------------------------------------

def test_excel_workbook_has_the_sheets_and_never_stores_a_formula():
    from openpyxl import load_workbook
    result = build_topology(list(lab_dc().items()))
    graph = result["graph"]
    devices = devices_table(graph)
    devices[0]["device"] = "=HYPERLINK(\"http://x\")"
    data = to_xlsx(result["summary"], devices, links_table(graph), hosts_table(graph), result["findings"])
    workbook = load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["Summary", "Devices", "Links", "Hosts", "Findings"]
    cell = workbook["Devices"]["A2"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")


@pytest.fixture()
def client():
    from app import create_app
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def test_route_accepts_a_zip_a_previous_layout_and_the_new_options(client):
    files = lab_dc()
    first = client.post("/topology/api/build", data={"files": [(io.BytesIO(_zip({k: v for k, v in files.items()})), "all.zip")], "group": "none"},
                        content_type="multipart/form-data").get_json()
    assert first["ok"] and first["summary"]["captured"] == 3
    drawio = client.get(f"/topology/api/export/{first['export_id']}/topology.drawio").data.decode()
    moved = _move(drawio, "LAB-ACC1", 900, 900)
    second = client.post("/topology/api/build", data={
        "files": [(io.BytesIO(t.encode()), n) for n, t in files.items()], "previous": (io.BytesIO(moved.encode()), "topology.drawio"),
        "include_hosts": "1", "include_routed": "1"}, content_type="multipart/form-data").get_json()
    assert second["ok"] and second["summary"]["layout_kept"] >= 3
    xlsx = client.get(f"/topology/api/export/{second['export_id']}/topology.xlsx")
    assert xlsx.status_code == 200 and xlsx.data[:2] == b"PK"
    assert client.get(f"/topology/api/export/{second['export_id']}/hosts.csv").status_code == 200


def test_route_reports_a_bad_previous_diagram_and_a_bad_zip(client):
    files = lab_dc()
    good = [(io.BytesIO(t.encode()), n) for n, t in files.items()]
    bad_layout = client.post("/topology/api/build", data={"files": good, "previous": (io.BytesIO(b"<html/>"), "x.drawio")}, content_type="multipart/form-data")
    assert bad_layout.status_code == 400 and "draw.io" in bad_layout.get_json()["error"]
    bad_zip = client.post("/topology/api/build", data={"files": [(io.BytesIO(b"nope"), "x.zip")]}, content_type="multipart/form-data")
    assert bad_zip.status_code == 400 and "ZIP" in bad_zip.get_json()["error"]


# ---- regression: first row of a table lacks columns the later rows have ------------------------------------------

def _first_row_uncaptured():
    """A neighbor named AAA-... sorts before every captured device, so the first devices-table row has no coverage columns."""
    files = lab_dc()
    files["LAB-AGG1.txt"] = nxos_capture(
        "LAB-AGG1", "FDO00000001", "10.99.0.11",
        cdp=[("AAA-UNCAPTURED.lab.example(FOC00000999)", "10.99.0.99", "C9300L-48P-4X", "Ethernet1/20", "GigabitEthernet1/0/1"),
             ("LAB-ACC1.lab.example(FOC00000001)", "10.99.0.21", "C9300L-48P-4X", "Ethernet1/1", "GigabitEthernet1/0/23")])
    return list(files.items())


def test_devices_table_has_the_same_columns_in_every_row_and_csv_never_fails():
    from features.topology.service import to_csv
    result = build_topology(_first_row_uncaptured())
    rows = devices_table(result["graph"])
    assert rows[0]["device"] == "AAA-UNCAPTURED" and rows[0]["captured"].startswith("no")
    assert len({tuple(r) for r in rows}) == 1 and rows[0]["has_cdp"] == ""
    text = to_csv(rows)
    assert text.splitlines()[0].count(",") == text.splitlines()[1].count(",")
    # rows that really differ are still written, with every column in the header
    assert to_csv([{"a": 1}, {"a": 2, "b": 3}]).splitlines()[0] == "a,b"


def test_route_builds_when_the_first_device_is_not_captured_and_all_downloads_open(client):
    files = _first_row_uncaptured()
    response = client.post("/topology/api/build", data={"files": [(io.BytesIO(t.encode()), n) for n, t in files]}, content_type="multipart/form-data")
    payload = response.get_json()
    assert response.status_code == 200 and payload["ok"]
    assert all(set(row) == set(payload["devices"][0]) for row in payload["devices"])
    for download in payload["downloads"]:
        got = client.get(download["url"])
        assert got.status_code == 200 and len(got.data) > 0, download["name"]


def test_an_unexpected_failure_reaches_the_page_as_a_readable_json_error(client, monkeypatch):
    from features.topology import routes

    def boom(*args, **kwargs):
        raise RuntimeError("synthetic failure")
    monkeypatch.setattr(routes, "build_topology", boom)
    response = client.post("/topology/api/build", data={"files": [(io.BytesIO(b"x"), "a.txt")]}, content_type="multipart/form-data")
    assert response.status_code == 500 and response.is_json
    assert "RuntimeError" in response.get_json()["error"] and "terminal" in response.get_json()["error"]


def test_cdp_port_ids_that_contain_a_space_are_kept_whole():
    capture = """LAB-ACC1#show cdp neighbors detail
-------------------------
Device ID: SEP00000000AAAA
Entry address(es):
  IP address: 10.99.20.5
Platform: Cisco IP Phone 8851,  Capabilities: Host Phone
Interface: GigabitEthernet1/0/3,  Port ID (outgoing port): Port 1
Holdtime : 150 sec

-------------------------
Device ID: LAB-WAAS1
Entry address(es):
  IP address: 10.99.20.6
Platform: OE7541,  Capabilities: Host
Interface: TwentyFiveGigE1/0/40,  Port ID (outgoing port): GigabitEthernet 0/0
Holdtime : 157 sec

LAB-ACC1#
"""
    rows = {r["neighbor_id"]: r for r in _neighbors(capture, IOSXE_HEAD) if r["protocol"] == "CDP"}
    assert rows["SEP00000000AAAA"]["remote_interface"] == "Port 1"
    assert rows["LAB-WAAS1"]["remote_interface"] == "GigabitEthernet 0/0"
    from features.topology.names import canon_port
    assert canon_port("GigabitEthernet 0/0") == "Gi0/0" and canon_port("Port 1") == "Port1"


# ---- second round: DRC findings (null names, fabric interconnects, vPC duplicates, version fallback) -----------------

def test_a_neighbor_with_a_null_system_name_is_named_by_its_chassis_mac():
    from features.topology.extract import _halves
    halves = _halves([
        {"local_interface": "Eth1/22", "neighbor_id": "null", "chassis_id": "0c48.c67e.2f74", "remote_interface": "0c48.c67e.2f74", "protocol": "LLDP"},
        {"local_interface": "Eth1/23", "neighbor_id": "N/A", "chassis_id": "", "remote_interface": "Gi0/1", "protocol": "LLDP"}])
    assert [(h.remote_name, h.remote_port) for h in halves] == [("0c48.c67e.2f74", "")]     # no name and no chassis: nothing to draw


def _fi_dc(mac_rows):
    files = lab_dc()
    extra = nxos_mac_table("LAB-AGG1", mac_rows)
    files["LAB-AGG1.txt"] = nxos_capture(
        "LAB-AGG1", "FDO00000001", "10.99.0.11",
        cdp=[("LAB-ACC1.lab.example(FOC00000001)", "10.99.0.21", "C9300L-48P-4X", "Ethernet1/1", "GigabitEthernet1/0/23"),
             ("LAB-AGG2.lab.example(FDO00000002)", "10.99.0.12", "N9K-C93180YC-FX3", "Ethernet1/47", "Ethernet1/47"),
             ("LAB-AGG2.lab.example(FDO00000002)", "10.99.0.12", "N9K-C93180YC-FX3", "Ethernet1/48", "Ethernet1/48"),
             ("FI-UCS-LAB-A(SSI00000001)", "10.99.0.31", "UCS-FI-6248UP", "Ethernet1/15", "Eth1/17")],
        port_channels=[(101, ["Eth1/1"]), (240, ["Eth1/47", "Eth1/48"]), (10, ["Eth1/15"])],
        vpc=("240", "primary", "240", "10.99.0.12"), extra=extra)
    return files


def test_macs_behind_an_uncaptured_fabric_interconnect_are_listed_but_not_drawn_or_flagged():
    macs = [("32", f"0050.56cc.{n:04x}", "Po10") for n in range(1, 15)]
    result = build_topology(list(_fi_dc(macs).items()), include_hosts=True)
    graph = result["graph"]
    assert graph.nodes["FI-UCS-LAB-A"].role_hint == "fabric"
    rows = [h for h in hosts_table(graph) if h["port"] == "Po10"]
    assert len(rows) == 14 and all(h["behind"] == "FI-UCS-LAB-A" and h["note"] == "behind FI-UCS-LAB-A" for h in rows)
    assert not any(f["id"] == "TOPO-007" for f in result["findings"])                      # not "a port with no network neighbor"
    assert not any(n.role_hint == "host" for n in graph.nodes.values())                    # nothing drawn for the virtual machines


def test_a_mac_behind_a_vpc_port_channel_is_not_reported_twice_between_the_peers():
    files = _fi_dc([("7", "0009.0fa0.000e", "Po101")])
    files["LAB-AGG2.txt"] = nxos_capture(
        "LAB-AGG2", "FDO00000002", "10.99.0.12",
        cdp=[("LAB-ACC1.lab.example(FOC00000001)", "10.99.0.21", "C9300L-48P-4X", "Ethernet1/1", "GigabitEthernet1/0/24"),
             ("LAB-AGG1.lab.example(FDO00000001)", "10.99.0.11", "N9K-C93180YC-FX3", "Ethernet1/47", "Ethernet1/47"),
             ("LAB-AGG1.lab.example(FDO00000001)", "10.99.0.11", "N9K-C93180YC-FX3", "Ethernet1/48", "Ethernet1/48")],
        port_channels=[(101, ["Eth1/1"]), (240, ["Eth1/47", "Eth1/48"])], vpc=("240", "secondary", "240", "10.99.0.11"),
        extra=nxos_mac_table("LAB-AGG2", [("7", "0009.0fa0.000e", "Po101")]))
    result = build_topology(list(files.items()))
    assert not any(f["id"] == "TOPO-014" for f in result["findings"])


def test_software_version_is_read_from_the_text_when_the_analyzer_misses_it():
    from features.topology.extract import os_from_text
    assert os_from_text("show version\nshow inventory\nSW1# show version\n  NXOS: version 10.5(6) [Maintenance Release]\n") == "10.5(6)"
    assert os_from_text("Cisco IOS XE Software, Version 17.12.5\n") == "17.12.5"
    assert os_from_text("nothing useful here") == ""


def test_an_advertised_platform_is_split_into_model_version_and_a_safe_vendor():
    from features.topology.names import split_platform
    assert split_platform("FortiGate-1800F v7.2.9") == ("FortiGate-1800F", "7.2.9", "Fortinet")
    assert split_platform("N9K-C93180YC-FX3") == ("N9K-C93180YC-FX3", "", "Cisco")
    assert split_platform("UCS-FI-6248UP") == ("UCS-FI-6248UP", "", "Cisco")
    assert split_platform("Mystery-Box 2000") == ("Mystery-Box 2000", "", "")

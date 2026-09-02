import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from features.wireless_center.config_delta import compare_configs
from features.wireless_center.exporters.delta_exporter import export_delta_comparison


HUAWEI_SOURCE = """sysname WLC-SOURCE
#
interface Vlanif100
 ip address 10.10.100.2 255.255.255.0
#
security-profile name CORP-SEC
 security wpa2 psk pass-phrase cipher SuperSecret
#
ssid-profile name CORP-SSID
 ssid CORP-WIFI
#
vap-profile name CORP-VAP
 service-vlan vlan-id 100
 ssid-profile CORP-SSID
 security-profile CORP-SEC
#
ap-group name HO-F1
 radio 0
  vap-profile CORP-VAP wlan 1
#
"""

HUAWEI_TARGET = """sysname WLC-TARGET
#
interface Vlanif100
 ip address 10.10.100.3 255.255.255.0
#
vap-profile name CORP-VAP
 service-vlan vlan-id 200
 ssid-profile CORP-SSID
 security-profile CORP-SEC
#
ap-group name TARGET-ONLY
 radio 0
#
"""

CISCO_CONFIG = """Cisco IOS XE Software, Version 17.9.4
hostname WLC9800
wlan CORP 10 CORP-WIFI
 security dot1x authentication-list ISE
 no shutdown
!
"""

HUAWEI_MAC_SOURCE = """sysname LAB-WLC-A
#
aaa
 authentication-scheme MAC-RADIUS
  authentication-mode radius
 accounting-scheme MAC-RADIUS
  accounting-mode radius
 domain LAB-DOMAIN
  authentication-scheme MAC-RADIUS
  accounting-scheme MAC-RADIUS
  radius-server MAC-RADIUS
 local-user 00-11-22-33-44-55 password cipher LabOnlySecret
#
radius-server template MAC-RADIUS
 radius-server authentication 192.0.2.10 1812
 radius-server shared-key cipher LabRadiusSecret
#
mac-access-profile name MAC-PROFILE
 mac-authen username mac-address format with-hyphen
#
authentication-profile name MAC-AUTH
 mac-access-profile MAC-PROFILE
 authentication-scheme MAC-RADIUS
 accounting-scheme MAC-RADIUS
 radius-server MAC-RADIUS
#
wlan
 security-profile name OPEN-SEC
 ssid-profile name LAB-SSID
  ssid LAB-WIFI
 vap-profile name LAB-VAP
  service-vlan vlan-id 100
  ssid-profile LAB-SSID
  security-profile OPEN-SEC
  authentication-profile MAC-AUTH
 ap-group name LAB-GROUP
  radio 0
   vap-profile LAB-VAP wlan 1
#
"""

HUAWEI_MAC_TARGET = """sysname LAB-WLC-B
#
aaa
 authentication-scheme MAC-RADIUS
  authentication-mode radius
 domain LAB-DOMAIN
  authentication-scheme MAC-RADIUS
  radius-server MAC-RADIUS
#
radius-server template MAC-RADIUS
 radius-server authentication 192.0.2.20 1812
 radius-server shared-key cipher DifferentSecret
#
mac-access-profile name MAC-PROFILE
 mac-authen username mac-address format without-hyphen
#
authentication-profile name MAC-AUTH
 mac-access-profile MAC-PROFILE
 authentication-scheme MAC-RADIUS
 radius-server MAC-RADIUS
#
wlan
 security-profile name OPEN-SEC
 ssid-profile name LAB-SSID
  ssid LAB-WIFI
 vap-profile name LAB-VAP
  service-vlan vlan-id 200
  ssid-profile LAB-SSID
  security-profile OPEN-SEC
  authentication-profile MAC-AUTH
#
"""

HUAWEI_WHITELIST_SOURCE = """sysname LAB-WLC-A
#
wlan
 sta-whitelist-profile name LAB-CLIENTS
  sta-mac 0011-2233-4455
  sta-mac 0011-2233-4466
 wids-whitelist-profile name LAB-WIDS
  sta-mac 0011-2233-4477
 vap-profile name LAB-VAP
  ssid-profile LAB-SSID
  sta-access-mode whitelist LAB-CLIENTS
#
"""

HUAWEI_WHITELIST_TARGET = """sysname LAB-WLC-B
#
wlan
 sta-whitelist-profile name LAB-CLIENTS
  sta-mac 0011-2233-4455
 wids-whitelist-profile name LAB-WIDS
 vap-profile name LAB-VAP
  ssid-profile LAB-SSID
  sta-access-mode whitelist LAB-CLIENTS
#
"""

HUAWEI_COMPREHENSIVE_SOURCE = """sysname LAB-WLC-A
#
acl number 3001
 description USER-ACCESS
 rule 5 permit ip source 192.0.2.0 0.0.0.255
 rule 10 deny ip
acl number 3002
 rule 5 permit udp destination-port eq dns
#
rrm-profile name BRANCH-RRM
 smart-roam enable
#
snmp-agent usm-user version v3 monitor authentication-mode sha2-256 SourceAuthMaterial
snmp-agent usm-user version v3 monitor privacy-mode aes128 SourcePrivacyMaterial
#
"""

HUAWEI_COMPREHENSIVE_TARGET = """sysname LAB-WLC-B
#
acl number 3001
 description USER-ACCESS
 rule 5 permit ip source 198.51.100.0 0.0.0.255
 rule 20 deny ip
#
rrm-profile name BRANCH-RRM
#
snmp-agent usm-user version v3 monitor authentication-mode sha2-256 TargetAuthMaterial
snmp-agent usm-user version v3 monitor privacy-mode aes128 TargetPrivacyMaterial
#
"""


class WLCConfigDeltaTest(unittest.TestCase):
    def test_directional_huawei_delta(self):
        result = compare_configs(HUAWEI_SOURCE, HUAWEI_TARGET)
        self.assertEqual(result["vendor"], "Huawei")
        self.assertGreaterEqual(result["summary"]["missing"], 2)
        self.assertGreaterEqual(result["summary"]["changed"], 2)
        self.assertGreaterEqual(result["summary"]["extra"], 1)
        self.assertIn("ssid-profile name CORP-SSID", result["implementation_config"])
        self.assertIn("service-vlan vlan-id 100", result["implementation_config"])
        self.assertNotIn("TARGET-ONLY", result["implementation_config"])

    def test_secrets_are_masked_and_not_generated(self):
        result = compare_configs(HUAWEI_SOURCE, HUAWEI_TARGET, "Huawei")
        security = next(
            row for row in result["rows"] if row["object_type"] == "Security Profile"
        )
        self.assertTrue(security["manual_review"])
        self.assertIn("<masked>", " ".join(security["source_commands"]))
        self.assertNotIn("SuperSecret", str(result))
        self.assertNotIn("SuperSecret", result["implementation_config"])

    def test_huawei_passphrase_without_cipher_keyword_is_never_exposed(self):
        numeric_secret = "123456789012345678901234567890123456789012345678"
        source = HUAWEI_SOURCE.replace(
            "pass-phrase cipher SuperSecret",
            f"pass-phrase {numeric_secret} aes",
        )
        result = compare_configs(source, HUAWEI_TARGET, "Huawei")
        self.assertNotIn(numeric_secret, str(result))
        self.assertNotIn(numeric_secret, result["implementation_config"])
        security = next(
            row for row in result["rows"] if row["object_type"] == "Security Profile"
        )
        self.assertIn("<masked>", " ".join(security["source_commands"]))
        self.assertTrue(security["manual_review"])

    def test_extra_target_is_report_only(self):
        result = compare_configs(HUAWEI_SOURCE, HUAWEI_TARGET)
        extra = next(row for row in result["rows"] if row["name"] == "TARGET-ONLY")
        self.assertEqual(extra["status"], "Extra on Target")
        self.assertFalse(extra["eligible"])
        self.assertNotIn("undo ap-group name TARGET-ONLY", result["implementation_config"])

    def test_rejects_cross_vendor(self):
        with self.assertRaisesRegex(ValueError, "Cross-vendor"):
            compare_configs(HUAWEI_SOURCE, CISCO_CONFIG)

    def test_cisco_delta_uses_cisco_wrapper(self):
        target = "Cisco IOS XE Software, Version 17.9.4\nhostname WLC9800-TARGET\n"
        result = compare_configs(CISCO_CONFIG, target, "Cisco")
        self.assertIn("configure terminal", result["implementation_config"])
        self.assertIn("wlan CORP 10 CORP-WIFI", result["implementation_config"])
        self.assertTrue(result["implementation_config"].rstrip().endswith("end"))

    def test_mac_auth_dependency_chain_covers_aaa_to_ap_delivery(self):
        result = compare_configs(HUAWEI_MAC_SOURCE, HUAWEI_MAC_TARGET, "Huawei")
        dependencies = {
            row["object_type"] for row in result["rows"]
            if row["mac_auth_dependency"]
        }
        self.assertTrue({
            "RADIUS Template", "Authentication Scheme", "Accounting Scheme",
            "AAA Domain", "Local AAA User", "MAC Access Profile",
            "Authentication Profile", "Security Profile", "SSID Profile",
            "VAP Profile", "AP Group",
        }.issubset(dependencies))
        self.assertEqual(
            [stage["stage"] for stage in result["mac_auth"]["stages"]],
            ["AAA & RADIUS", "Access Profiles", "Authentication Policy", "WLAN Binding", "AP Delivery"],
        )
        self.assertGreater(result["mac_auth"]["drift"], 0)

    def test_huawei_nested_context_and_radio_indentation_are_preserved(self):
        result = compare_configs(HUAWEI_MAC_SOURCE, HUAWEI_MAC_TARGET, "Huawei")
        generated = result["implementation_config"]
        self.assertIn("aaa\n accounting-scheme MAC-RADIUS", generated)
        self.assertIn("wlan\n ap-group name LAB-GROUP", generated)
        self.assertIn("  radio 0\n   vap-profile LAB-VAP wlan 1", generated)
        self.assertTrue(generated.rstrip().endswith("return"))

    def test_mac_auth_secrets_and_local_users_never_generate(self):
        result = compare_configs(HUAWEI_MAC_SOURCE, HUAWEI_MAC_TARGET, "Huawei")
        rendered = str(result)
        self.assertNotIn("LabOnlySecret", rendered)
        self.assertNotIn("LabRadiusSecret", rendered)
        self.assertNotIn("DifferentSecret", rendered)
        self.assertNotIn("local-user 00-11-22-33-44-55", result["implementation_config"])
        local_user = next(
            row for row in result["rows"] if row["object_type"] == "Local AAA User"
        )
        self.assertTrue(local_user["manual_review"])

    def test_excel_export_includes_mac_auth_detail_and_stage_summary(self):
        result = compare_configs(HUAWEI_MAC_SOURCE, HUAWEI_MAC_TARGET, "Huawei")
        with TemporaryDirectory() as directory:
            path = Path(directory) / "delta.xlsx"
            export_delta_comparison(result, path)
            workbook = load_workbook(path, read_only=True)
            self.assertIn("MAC Authentication", workbook.sheetnames)
            self.assertIn("MAC Auth Dependency", workbook.sheetnames)
            self.assertEqual(len(workbook.sheetnames), 10)
            self.assertEqual(
                workbook["MAC Auth Dependency"].max_row,
                len(result["mac_auth"]["stages"]) + 1,
            )

    def test_sta_and_wids_whitelist_entries_are_compared_and_generated(self):
        result = compare_configs(
            HUAWEI_WHITELIST_SOURCE,
            HUAWEI_WHITELIST_TARGET,
            "Huawei",
        )
        sta = next(
            row for row in result["rows"]
            if row["object_type"] == "STA Whitelist Profile"
        )
        self.assertEqual(sta["status"], "Value Mismatch")
        self.assertEqual(sta["missing_commands"], ["sta-mac 0011-2233-4466"])
        self.assertTrue(sta["eligible"])
        self.assertTrue(sta["mac_auth_dependency"])
        self.assertEqual(sta["dependency_stage"], "Access Profiles")

        wids = next(
            row for row in result["rows"]
            if row["object_type"] == "WIDS Whitelist Profile"
        )
        self.assertEqual(wids["missing_commands"], ["sta-mac 0011-2233-4477"])
        generated = result["implementation_config"]
        self.assertIn("wlan\n sta-whitelist-profile name LAB-CLIENTS", generated)
        self.assertIn("  sta-mac 0011-2233-4466", generated)
        self.assertIn("wlan\n wids-whitelist-profile name LAB-WIDS", generated)
        self.assertIn("  sta-mac 0011-2233-4477", generated)

    def test_comprehensive_fallback_captures_acl_and_unknown_sections(self):
        result = compare_configs(
            HUAWEI_COMPREHENSIVE_SOURCE,
            HUAWEI_COMPREHENSIVE_TARGET,
            "Huawei",
        )
        acl_3001 = next(
            row for row in result["rows"]
            if row["object_type"] == "ACL" and row["name"] == "acl number 3001"
        )
        self.assertEqual(acl_3001["status"], "Value Mismatch")
        self.assertIn(
            "rule 5 permit ip source 192.0.2.0 0.0.0.255",
            acl_3001["missing_commands"],
        )
        self.assertIn("undo rule 5", result["implementation_config"])
        self.assertIn(
            "rule 5 permit ip source 198.51.100.0 0.0.0.255",
            result["rollback_config"],
        )
        acl_3002 = next(
            row for row in result["rows"]
            if row["object_type"] == "ACL" and row["name"] == "acl number 3002"
        )
        self.assertEqual(acl_3002["status"], "Missing on Target")
        self.assertIn("acl number 3002", result["implementation_config"])
        self.assertIn("rule 5 permit udp destination-port eq dns", result["implementation_config"])

        rrm = next(
            row for row in result["rows"]
            if row["name"] == "rrm-profile name BRANCH-RRM"
        )
        self.assertEqual(rrm["parser_mode"], "Fallback")
        self.assertEqual(rrm["missing_commands"], ["smart-roam enable"])
        self.assertIn("rrm-profile name BRANCH-RRM", result["implementation_config"])
        self.assertGreater(result["summary"]["fallback_objects"], 0)

    def test_fallback_masks_snmp_usm_secrets_and_excludes_them_from_cli(self):
        result = compare_configs(
            HUAWEI_COMPREHENSIVE_SOURCE,
            HUAWEI_COMPREHENSIVE_TARGET,
            "Huawei",
        )
        rendered = str(result)
        for secret in (
            "SourceAuthMaterial", "SourcePrivacyMaterial",
            "TargetAuthMaterial", "TargetPrivacyMaterial",
        ):
            self.assertNotIn(secret, rendered)
        usm_rows = [
            row for row in result["rows"]
            if row["name"].startswith("snmp-agent usm-user")
        ]
        self.assertTrue(usm_rows)
        self.assertTrue(all(row["manual_review"] for row in usm_rows))
        self.assertNotIn("snmp-agent usm-user", result["implementation_config"])

    def test_cisco_fallback_captures_named_acl(self):
        source = CISCO_CONFIG + """ip access-list extended BRANCH-USERS
 permit ip 192.0.2.0 0.0.0.255 any
 deny ip any any
!
"""
        target = CISCO_CONFIG + """ip access-list extended BRANCH-USERS
 deny ip any any
!
"""
        result = compare_configs(source, target, "Cisco")
        acl = next(
            row for row in result["rows"]
            if row["object_type"] == "ACL" and "BRANCH-USERS" in row["name"]
        )
        self.assertEqual(acl["status"], "Value Mismatch")
        self.assertIn("permit ip 192.0.2.0 0.0.0.255 any", acl["missing_commands"])
        self.assertIn("ip access-list extended BRANCH-USERS", result["implementation_config"])

    # ------------------------------------------------------------------
    # Generated-command ORDER must respect real cross-object references,
    # not just the static per-type priority table. Regression coverage
    # for a bug report: pushing a full generated delta failed because an
    # object referencing another (e.g. "configure profile B" needs
    # "profile A" to already exist) was generated in the wrong order --
    # Portal Access Profile (priority 35) referencing a Free Rule
    # Template (priority 36) is a real case of this, since the static
    # table alone puts the referencer first.
    # ------------------------------------------------------------------

    def test_referenced_object_is_pushed_before_the_object_that_references_it(self):
        source = """
free-rule-template name FreeRule1
 rule 1 permit ip
#
portal-access-profile name Portal1
 free-rule-template FreeRule1
#
"""
        result = compare_configs(source, "", vendor="Huawei")
        cli = result["implementation_config"]
        free_rule_idx = cli.find("free-rule-template name FreeRule1")
        portal_idx = cli.find("portal-access-profile name Portal1")
        self.assertNotEqual(free_rule_idx, -1)
        self.assertNotEqual(portal_idx, -1)
        self.assertLess(
            free_rule_idx, portal_idx,
            "free-rule-template must be pushed before the portal-access-profile "
            "that references it, regardless of their relative priority/name order",
        )

    def test_rollback_undoes_the_dependent_before_the_dependency(self):
        source = """
free-rule-template name FreeRule1
 rule 1 permit ip
#
portal-access-profile name Portal1
 free-rule-template FreeRule1
#
"""
        result = compare_configs(source, "", vendor="Huawei")
        rollback = result["rollback_config"]
        free_rule_idx = rollback.find("free-rule-template name FreeRule1")
        portal_idx = rollback.find("undo portal-access-profile name Portal1")
        self.assertNotEqual(free_rule_idx, -1)
        self.assertNotEqual(portal_idx, -1)
        self.assertLess(
            portal_idx, free_rule_idx,
            "rollback must undo the portal-access-profile (the dependent) before "
            "undoing the free-rule-template it depends on -- the exact reverse of "
            "the apply order",
        )

    def test_unrelated_objects_still_follow_the_priority_table_order(self):
        # Sanity check that the new topological pass doesn't disturb the
        # existing, already-correct ordering when there's no cross
        # reference at all between two objects -- it should fall back to
        # exactly the same priority/category/name order as before.
        source = """
vap-profile name VAP-A
 service-vlan vlan-id 100
#
ssid-profile name SSID-A
 ssid CorpWifi
#
"""
        result = compare_configs(source, "", vendor="Huawei")
        cli = result["implementation_config"]
        ssid_idx = cli.find("ssid-profile name SSID-A")
        vap_idx = cli.find("vap-profile name VAP-A")
        self.assertNotEqual(ssid_idx, -1)
        self.assertNotEqual(vap_idx, -1)
        self.assertLess(
            ssid_idx, vap_idx,
            "with no reference between them, SSID Profile (priority 50) must still "
            "come before VAP Profile (priority 60) exactly as the static table says",
        )


if __name__ == "__main__":
    unittest.main()

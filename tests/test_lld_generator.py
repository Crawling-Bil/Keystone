"""LLD Generator (the Document Generator inside Keystone).

Covers the moved engine (importer, placeholders), the Flask API that replaced
the FastAPI app (same paths and shapes), the hand-off from LLD Index and the
one-time import from the old Document Generator. Storage goes to tmp_path.
"""

from __future__ import annotations

import io
import json
import os
import sqlite3
import time
from datetime import date
from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn

from core.docgen import importer, placeholders
from features.lld_generator import service, store

FIXTURES = Path(__file__).parent / "fixtures" / "lld_generator"
INDEX = FIXTURES / "keystone_index.xlsx"
LOGO = FIXTURES / "logo.png"
MASTER = service.MASTER_TEMPLATE


@pytest.fixture()
def client(tmp_path, monkeypatch):
    from app import create_app
    from features.lld_index import routes as lld_index_routes

    monkeypatch.setattr(store, "DB_PATH", str(tmp_path / "nes.db"))
    monkeypatch.setattr(service, "DATA_DIR", tmp_path / "lld_generator")
    monkeypatch.setattr(lld_index_routes, "EXPORT_DIR", tmp_path / "exports")
    from features.lld_index import saved
    monkeypatch.setattr(saved, "DB_PATH", str(tmp_path / "nes.db"))
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


@pytest.fixture()
def project(client):
    templates = client.get("/lld-generator/api/templates").get_json()      # registers the bundled master
    assert [t["name"] for t in templates] == ["LLD Master Template"]
    assert templates[0]["closing_heading"] == "Document Acceptance"
    p = client.post("/lld-generator/api/projects", json={"name": "Draft", "template_id": templates[0]["id"]}).get_json()
    return p


def _import(client, pid, mode, path=INDEX):
    with open(path, "rb") as f:
        return client.post(f"/lld-generator/api/projects/{pid}/import-excel",
                           data={"mode": mode, "file": (f, "index.xlsx")}, content_type="multipart/form-data")


def _chapters(client, pid):
    return client.get(f"/lld-generator/api/projects/{pid}/chapters").get_json()


# ----------------------------------------------------------------- engine
def test_index_workbook_is_read_with_blocks_merges_fills_aligns():
    items = importer.parse_index_workbook(str(INDEX))
    model = next(i for i in items if i["title"] == "Palo Alto PA-1410")
    assert [b["block_type"] for b in model["blocks"]] == ["paragraph", "table"]
    rack = next(i for i in items if i["title"] == "Rack Layout After - Rack Server HO")["blocks"][0]["content"]
    assert len(rack["rows"]) == 42 and rack["merges"] and rack["fills"] and rack["aligns"]
    assert importer.parse_project_sheet(str(INDEX))["client"] == "PT Sinar Logistik"


def test_placeholders_across_split_runs_and_missing_values():
    values = placeholders.project_values({"name": "Core Refresh", "client": "PT Contoh", "version": "1.0"},
                                         today=date(2026, 10, 24))
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("LLD – {{cust")
    p.add_run("omer}} – {{project_").bold = True
    p.add_run("name}} {{project_code}}")
    assert placeholders.fill(d, values) == {"project_code"}
    assert p.text == "LLD – PT Contoh – Core Refresh {{project_code}}" and values["date"] == "24 Oct 2026"


def test_bundled_master_template_is_fully_placeholdered():
    d = docx.Document(str(MASTER))
    assert {"project_name", "customer", "version", "prepared_by", "date", "year"} <= placeholders.remaining(d)
    assert placeholders.fill_logo(d, str(LOGO))


# ----------------------------------------------------------------- API
def test_page_is_in_the_sidebar_and_spa_routes(client):
    html = client.get("/lld-generator").get_data(as_text=True)
    assert 'class="lg-page"' in html and "lld_generator.js" in html and 'href="/lld-generator"' in html
    assert ' id="' not in html.split('class="lg-page"', 1)[1].split("</section>", 1)[0]       # no global ids
    assert '"/lld-generator"' in Path("static/js/spa-router.js").read_text()


def test_import_fills_project_and_errors_use_detail(client, project):
    res = _import(client, project["id"], "replace").get_json()
    assert res["chapters"] == 91 and "client" in res["project_fields"]
    p = client.get(f"/lld-generator/api/projects/{project['id']}").get_json()
    assert (p["name"], p["client"], p["project_code"]) == (
        "Perimeter Firewall and Core Switch Replacement", "PT Sinar Logistik", "7SINL26010001")
    missing = client.get("/lld-generator/api/projects/999")
    assert missing.status_code == 404 and missing.get_json() == {"detail": "Project not found"}
    bad = client.post("/lld-generator/api/projects", json={"name": "X", "template_id": 999})
    assert bad.status_code == 400 and "template_id" in bad.get_json()["detail"]


def test_chapter_and_block_crud(client, project):
    pid = project["id"]
    root = client.post(f"/lld-generator/api/projects/{pid}/chapters", json={"title": "Intro"}).get_json()
    child = client.post(f"/lld-generator/api/projects/{pid}/chapters", json={"title": "Scope", "parent_id": root["id"]}).get_json()
    assert child["level"] == 2
    block = client.post(f"/lld-generator/api/chapters/{child['id']}/blocks",
                        json={"block_type": "paragraph", "content": {"paragraphs": ["a"]}}).get_json()
    client.put(f"/lld-generator/api/blocks/{block['id']}", json={"content": {"paragraphs": ["b"]}})
    assert client.get(f"/lld-generator/api/chapters/{child['id']}/blocks").get_json()[0]["content"] == {"paragraphs": ["b"]}
    client.put(f"/lld-generator/api/chapters/{child['id']}", json={"title": "Scope of Work"})
    client.delete(f"/lld-generator/api/chapters/{root['id']}")
    assert _chapters(client, pid) == []                                         # cascade


def test_update_keeps_images_and_manual_text_and_failure_changes_nothing(client, project, monkeypatch):
    pid = project["id"]
    _import(client, pid, "replace")
    by = {c["title"]: c for c in _chapters(client, pid)}
    with open(LOGO, "rb") as f:
        up = client.post(f"/lld-generator/api/projects/{pid}/upload-image", data={"file": (f, "topo.png")},
                         content_type="multipart/form-data").get_json()
    img = client.get(f"/lld-generator/api/chapters/{by['Physical Topology After Site HO']['id']}/blocks").get_json()[0]
    client.put(f"/lld-generator/api/blocks/{img['id']}", json={"content": {"filename": up["filename"], "caption": "Topo"}})
    assert client.get(f"/lld-generator/api/projects/{pid}/images/{up['filename']}").status_code == 200
    res = _import(client, pid, "update").get_json()
    assert res["updated"] == 91 and res["kept_blocks"] >= 1
    img = client.get(f"/lld-generator/api/chapters/{by['Physical Topology After Site HO']['id']}/blocks").get_json()[0]
    assert img["content"]["filename"] == up["filename"]

    before = _chapters(client, pid)

    def boom(*args, **kwargs):
        raise RuntimeError("disk full")

    monkeypatch.setattr(store, "_merge_blocks", boom)
    res = _import(client, pid, "update")
    assert res.status_code == 400 and "nothing was changed" in res.get_json()["detail"]
    assert _chapters(client, pid) == before


def test_generate_fills_master_logo_acceptance_and_cleans_outputs(client, project):
    pid = project["id"]
    _import(client, pid, "replace")
    res = client.post(f"/lld-generator/api/projects/{pid}/generate").get_json()
    assert res["warnings"] == ["No customer logo: the cover has no logo."]
    with open(LOGO, "rb") as f:
        logo = client.post(f"/lld-generator/api/projects/{pid}/upload-image", data={"file": (f, "logo.png")},
                           content_type="multipart/form-data").get_json()
    client.put(f"/lld-generator/api/projects/{pid}", json={"logo": logo["filename"], "version": "1.0"})
    res = client.post(f"/lld-generator/api/projects/{pid}/generate").get_json()
    assert res["warnings"] == [] and res["filename"].startswith("[LLD] Perimeter Firewall")
    d = docx.Document(io.BytesIO(client.get(res["download_url"]).data))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "{{" not in text and "Perimeter Firewall and Core Switch Replacement" in text
    heads = [p.text for p in d.paragraphs if p.style.name == "Heading 1"]
    assert heads[-1] == "Document Acceptance"
    assert any(e.get("descr") == "Customer logo" for e in d.element.body.iter(qn("wp:docPr")))
    old = service.folder("outputs") / ("1_" + "0" * 32 + ".docx")
    old.write_bytes(b"x")
    past = time.time() - 8 * 86400
    os.utime(old, (past, past))
    assert service.cleanup_outputs() == 1
    assert client.get("/lld-generator/api/download/..%2Fnes.db").status_code == 404


# ----------------------------------------------------------------- from LLD Index
def test_lld_index_opens_and_reopens_the_same_project(client):
    from features.lld_index.routes import EXPORT_DIR  # noqa: F401  (patched by the fixture)
    from tests.test_lld_index import FILES, INVENTORY, PROJECT

    def build():
        payload = {"project": dict(PROJECT, project_code="X1"), "technologies": ["switch"], "inventory": INVENTORY}
        return client.post("/lld-index/api/build", data={"payload": json.dumps(payload)}).get_json()["result"]

    first = build()
    opened = client.post("/lld-generator/api/from-index", json={"export_id": first["export_id"], "project": first["project"]}).get_json()
    assert opened["ok"] and opened["created"] and opened["url"] == f"/lld-generator?project={opened['project_id']}"
    project = client.get(f"/lld-generator/api/projects/{opened['project_id']}").get_json()
    assert (project["name"], project["client"], project["project_code"]) == ("Network Refresh", "PT Contoh", "X1")
    chapters = _chapters(client, opened["project_id"])
    assert chapters and chapters[0]["title"] == "Introduction"

    second = build()
    again = client.post("/lld-generator/api/from-index", json={"export_id": second["export_id"], "project": second["project"]}).get_json()
    assert not again["created"] and again["project_id"] == opened["project_id"] and again["mode"] == "update"
    assert len(client.get("/lld-generator/api/projects").get_json()) == 1
    gone = client.post("/lld-generator/api/from-index", json={"export_id": "0" * 32, "project": {}})
    assert gone.status_code == 404


# ----------------------------------------------------------------- legacy import
def _old_app(root: Path) -> Path:
    """A Document Generator v1.0.5 storage folder with one project."""
    storage = root / "lld_generator_v1.0.5" / "backend" / "storage"
    for sub in ("templates", "uploads/7", "outputs"):
        (storage / sub).mkdir(parents=True)
    (storage / "templates" / "abc_LLD Master Template.docx").write_bytes(MASTER.read_bytes())
    (storage / "uploads" / "7" / "topo.png").write_bytes(LOGO.read_bytes())
    db = sqlite3.connect(storage / "app.db")
    db.executescript("""
      CREATE TABLE templates (id INTEGER PRIMARY KEY, name TEXT, filename TEXT, closing_heading TEXT,
        auto_numbered_levels TEXT, max_heading_level INTEGER, figure_seq_name TEXT, table_seq_name TEXT, created_at TEXT);
      CREATE TABLE projects (id INTEGER PRIMARY KEY, name TEXT, client TEXT, version TEXT, project_code TEXT,
        company TEXT, prepared_by TEXT, doc_date TEXT, logo TEXT, template_id INTEGER, created_at TEXT, updated_at TEXT);
      CREATE TABLE chapters (id INTEGER PRIMARY KEY, project_id INTEGER, parent_id INTEGER, level INTEGER,
        order_index INTEGER, title TEXT);
      CREATE TABLE content_blocks (id INTEGER PRIMARY KEY, chapter_id INTEGER, order_index INTEGER, block_type TEXT, content TEXT);
      INSERT INTO templates VALUES (3, 'Master', 'abc_LLD Master Template.docx', 'Document Acceptance', '[1,2]', 4, 'Gambar', 'Tabel', 'x');
      INSERT INTO projects VALUES (7, 'Old WLC', 'TAM', '0.2', 'C7', 'Multipolar Technology', 'Aldi', '', '', 3, 'x', 'x');
      INSERT INTO chapters VALUES (21, 7, 20, 2, 0, 'Child first by id');
      INSERT INTO chapters VALUES (20, 7, NULL, 1, 0, 'Parent');
      INSERT INTO content_blocks VALUES (1, 21, 0, 'image', '{"filename": "topo.png", "caption": "Topo"}');
      INSERT INTO content_blocks VALUES (2, 21, 1, 'paragraph', '{"paragraphs": ["hello"]}');
    """)
    db.commit()
    db.close()
    return root / "lld_generator_v1.0.5"


def test_legacy_import_copies_everything_once(client, tmp_path):
    folder = _old_app(tmp_path / "Document_Generator")
    res = client.post("/lld-generator/api/legacy-import", json={"folder": str(folder)}).get_json()
    assert (res["templates"], res["projects"], res["chapters"], res["blocks"], res["images"]) == (1, 1, 2, 2, 1)
    assert res["missing_files"] == []
    projects = client.get("/lld-generator/api/projects").get_json()
    assert [(p["name"], p["client"], p["version"]) for p in projects] == [("Old WLC", "TAM", "0.2")]
    chapters = _chapters(client, projects[0]["id"])
    assert [(c["title"], c["level"]) for c in chapters] == [("Parent", 1), ("Child first by id", 2)]
    blocks = client.get(f"/lld-generator/api/chapters/{chapters[1]['id']}/blocks").get_json()
    assert client.get(f"/lld-generator/api/projects/{projects[0]['id']}/images/{blocks[0]['content']['filename']}").status_code == 200
    again = client.post("/lld-generator/api/legacy-import", json={"folder": str(folder)}).get_json()
    assert (again["projects"], again["templates"], again["skipped_projects"]) == (0, 0, 1)
    assert sqlite3.connect(folder / "backend" / "storage" / "app.db").execute("SELECT COUNT(*) FROM projects").fetchone()[0] == 1


def test_review_lists_what_is_left_to_do(client, project):
    pid = project["id"]
    _import(client, pid, "replace")
    review = client.get(f"/lld-generator/api/projects/{pid}/review").get_json()
    todo = {c["title"]: c["todo"] for c in review["chapters"] if c["todo"]}
    assert todo["Physical Topology After Site HO"] == ["picture to upload"]
    assert "table still TBD" in todo["Hostname Scheme"]
    assert review["leaf_todo"] == sum(1 for c in review["chapters"] if c["todo"] and not c["is_group"])
    assert review["leaf_filled"] + review["leaf_todo"] + review["leaf_empty"] == sum(1 for c in review["chapters"] if not c["is_group"])


def test_pictures_go_to_the_chapter_their_name_matches(client, project):
    pid = project["id"]
    _import(client, pid, "replace")
    data = LOGO.read_bytes()
    files = [(io.BytesIO(data), "physical-topology-after-HO.png"), (io.BytesIO(data), "rack layout DC.png"),
             (io.BytesIO(data), "holiday.png")]
    res = client.post(f"/lld-generator/api/projects/{pid}/pictures", data={"files": files},
                      content_type="multipart/form-data").get_json()
    assert sorted(p["chapter"] for p in res["placed"]) == ["Physical Topology After Site HO", "Rack Layout DC"]
    assert res["unmatched"] == ["holiday.png"]
    review = client.get(f"/lld-generator/api/projects/{pid}/review").get_json()
    assert not next(c for c in review["chapters"] if c["title"] == "Physical Topology After Site HO")["todo"]


def test_version_changes_fill_the_release_history(client, project):
    pid = project["id"]
    _import(client, pid, "replace")
    client.put(f"/lld-generator/api/projects/{pid}", json={"version": "1.0", "change_note": "After site survey"})
    client.put(f"/lld-generator/api/projects/{pid}", json={"version": "1.1"})
    history = json.loads(client.get(f"/lld-generator/api/projects/{pid}").get_json()["history"])
    assert [(h["version"], h["status"], h["comment"]) for h in history] == [
        ("0.1", "Initial", "-"), ("1.0", "Revised", "After site survey"), ("1.1", "Revised", "-")]
    res = client.post(f"/lld-generator/api/projects/{pid}/generate").get_json()
    d = docx.Document(io.BytesIO(client.get(res["download_url"]).data))
    rev = next(t for t in d.tables if t.rows[0].cells[0].text.strip().lower().startswith("rev"))
    assert [r.cells[0].text for r in rev.rows[1:]] == ["0.1", "1.0", "1.1"]
    assert rev.rows[2].cells[4].text == "After site survey" and rev.rows[2].cells[2].text == "Mohammad Nur Alkadri"


def test_approvers_fill_document_approval_and_are_remembered(client, project):
    pid = project["id"]
    _import(client, pid, "replace")
    first = client.get(f"/lld-generator/api/projects/{pid}/approvers").get_json()
    assert first["from"] == "remembered" and [a["title"] for a in first["approvers"]] == ["Section Head", "Network Engineer", "Project Manager"]
    assert first["approvers"][1]["name"] == "Mohammad Nur Alkadri"                    # Prepared by
    people = [{"side": "customer", "name": "Budi Santoso", "title": "IT Manager"},
              {"side": "customer", "name": "Sari Dewi", "title": "Network Lead", "company": "PT Sinar Logistik (HO)"},
              {"side": "multipolar", "name": "Fiqih", "title": "Section Head"},
              {"side": "multipolar", "name": "Mohammad Nur Alkadri", "title": "Network Engineer"},
              {"side": "multipolar", "name": "", "title": ""}]                       # empty rows are dropped
    client.put(f"/lld-generator/api/projects/{pid}/approvers", json={"approvers": people, "multipolar_default": True})
    res = client.post(f"/lld-generator/api/projects/{pid}/generate").get_json()
    d = docx.Document(io.BytesIO(client.get(res["download_url"]).data))
    cust, mp = [t for t in d.tables if t.rows[0].cells[0].text.strip().lower().startswith("name and title")]
    assert [(r.cells[0].text, r.cells[1].text) for r in cust.rows[1:]] == [
        ("Budi Santoso\nIT Manager", "PT Sinar Logistik"), ("Sari Dewi\nNetwork Lead", "PT Sinar Logistik (HO)")]
    assert [r.cells[0].text for r in mp.rows[1:]] == ["Fiqih\nSection Head", "Mohammad Nur Alkadri\nNetwork Engineer"]
    # a new project for the same customer starts with the same people; Multipolar defaults apply everywhere
    templates = client.get("/lld-generator/api/templates").get_json()
    other = client.post("/lld-generator/api/projects", json={"name": "Phase 2", "client": "pt  sinar logistik",
                                                          "prepared_by": "Rina", "template_id": templates[0]["id"]}).get_json()
    again = client.get(f"/lld-generator/api/projects/{other['id']}/approvers").get_json()["approvers"]
    assert [a["name"] for a in again if a["side"] == "customer"] == ["Budi Santoso", "Sari Dewi"]
    assert [a["name"] for a in again if a["side"] == "multipolar"] == ["Fiqih", "Rina"]

"""Tests for the Switch Analyzer's two-device comparison module:
compare.compare() (the VLAN/route/DNS set-diff + device summary logic)
and exporters.export_comparison() (the .xlsx export built on top of it).
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from features.switch_analyzer.compare import compare
from features.switch_analyzer.exporters import export_comparison

RESULT_A = {
    "hostname": "OLD-CISCO-01",
    "vendor": "Cisco",
    "os_version": {"os_version": "15.2(7)E3", "model": "WS-C3650-48P"},
    "inventory": {"stack": {"mode": "Standalone", "member_count": 1}},
    "cards": {
        "total_interfaces": 48, "interfaces_up": 30, "interfaces_down": 18,
        "neighbor_count": 2, "port_channel_count": 1, "transceiver_count": 4, "poe_port_count": 48,
    },
    "vlans": [{"vlan_id": 10, "name": "DATA"}, {"vlan_id": 20, "name": "VOICE"}],
    "routes": [{"destination": "10.0.0.0/8"}, {"destination": "0.0.0.0/0"}],
    "dns": {"servers": ["10.1.1.1", "10.1.1.2"]},
}

RESULT_B = {
    "hostname": "NEW-HUAWEI-01",
    "vendor": "Huawei",
    "os_version": {"os_version": "V200R023", "model": "S5720"},
    "inventory": {"stack": {"mode": "Stack", "member_count": 2}},
    "cards": {
        "total_interfaces": 48, "interfaces_up": 28, "interfaces_down": 20,
        "neighbor_count": 2, "port_channel_count": 1, "transceiver_count": 4, "poe_port_count": 48,
    },
    "vlans": [{"vlan_id": 10, "name": "DATA"}, {"vlan_id": 30, "name": "MGMT"}],
    "routes": [{"destination": "10.0.0.0/8"}, {"destination": "172.16.0.0/12"}],
    "dns": {"servers": ["10.1.1.1"]},
}


class CompareLogicTest(unittest.TestCase):
    def test_vlan_route_dns_set_diff(self):
        result = compare(RESULT_A, RESULT_B)
        self.assertEqual(result["vlans"], {"only_a": ["20"], "only_b": ["30"], "both": ["10"]})
        self.assertEqual(
            result["routes"],
            {"only_a": ["0.0.0.0/0"], "only_b": ["172.16.0.0/12"], "both": ["10.0.0.0/8"]},
        )
        self.assertEqual(result["dns_servers"], {"only_a": ["10.1.1.2"], "only_b": [], "both": ["10.1.1.1"]})

    def test_device_summary_pulls_hostname_vendor_and_counts(self):
        result = compare(RESULT_A, RESULT_B)
        self.assertEqual(result["device_a"]["hostname"], "OLD-CISCO-01")
        self.assertEqual(result["device_a"]["vendor"], "Cisco")
        self.assertEqual(result["device_a"]["total_interfaces"], 48)
        self.assertEqual(result["device_b"]["hostname"], "NEW-HUAWEI-01")
        self.assertEqual(result["device_b"]["model"], "S5720")

    def test_device_summary_carries_stack_member_count(self):
        # Single Device already showed this (renderResult's KV list
        # auto-labels whatever key it's handed); Compare Devices used a
        # separate hardcoded field list that had simply never been
        # given this key at all, so a stacked device's member count
        # silently vanished the moment it was compared against another
        # device — confirmed against your own live screenshot of
        # GTOPAS-MKS-SWCODI-S5755 (a real 2-member stack).
        result = compare(RESULT_A, RESULT_B)
        self.assertEqual(result["device_a"]["stack_member_count"], 1)
        self.assertEqual(result["device_b"]["stack_member_count"], 2)

    def test_missing_optional_sections_do_not_raise(self):
        minimal_a = {"hostname": "A", "vendor": "Cisco"}
        minimal_b = {"hostname": "B", "vendor": "Huawei"}
        result = compare(minimal_a, minimal_b)
        self.assertEqual(result["vlans"], {"only_a": [], "only_b": [], "both": []})
        self.assertEqual(result["device_a"]["total_interfaces"], 0)

    def test_pbr_summary_surfaced_side_by_side_not_diffed(self):
        # PBR/SLA-Track/NQA is Cisco route-map/IP-SLA vs. Huawei
        # traffic-policy/NQA — different vendor mechanisms with no
        # shared naming to set-diff (same reason this module already
        # skips a per-interface diff), so both devices' full detail is
        # exposed side by side under comparison["pbr"] instead, and a
        # quick has_custom_routing flag lands in each device_summary.
        result_a = {**RESULT_A, "pbr": {
            "has_custom_routing": True,
            "sla_tests": [{"kind": "IP SLA", "id": "IP SLA 2"}],
            "track_objects": [], "pbr_policies": [], "pbr_bindings": [], "eem_dynamic_pbr": [],
        }}
        result_b = {**RESULT_B, "pbr": {
            "has_custom_routing": False,
            "sla_tests": [], "track_objects": [], "pbr_policies": [], "pbr_bindings": [], "eem_dynamic_pbr": [],
        }}
        result = compare(result_a, result_b)
        self.assertTrue(result["device_a"]["has_custom_routing"])
        self.assertFalse(result["device_b"]["has_custom_routing"])
        self.assertEqual(result["pbr"]["a"]["sla_tests"], [{"kind": "IP SLA", "id": "IP SLA 2"}])
        self.assertEqual(result["pbr"]["b"]["sla_tests"], [])

    def test_missing_pbr_key_defaults_to_false_and_empty_dict(self):
        # Neither RESULT_A nor RESULT_B (module fixtures above) carries
        # a "pbr" key -- must not raise, and must read as "no custom
        # routing" rather than crash or silently mis-report True.
        result = compare(RESULT_A, RESULT_B)
        self.assertFalse(result["device_a"]["has_custom_routing"])
        self.assertFalse(result["device_b"]["has_custom_routing"])
        self.assertEqual(result["pbr"], {"a": {}, "b": {}})


class CompareExportTest(unittest.TestCase):
    def test_export_writes_comparison_sheet_first_plus_combined_data_sheets(self):
        comparison = compare(RESULT_A, RESULT_B)
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "compare.xlsx"
            export_comparison(path, RESULT_A, RESULT_B, comparison)
            self.assertTrue(path.exists())

            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames[0], "Comparison")
            comparison_sheet = workbook["Comparison"]
            column_a_values = [row[0].value for row in comparison_sheet.iter_rows()]
            self.assertIn("VLANs", column_a_values)
            self.assertIn("Routes (destination)", column_a_values)
            self.assertIn("DNS Servers", column_a_values)
            # Stack Member Count row — same gap as the web UI's compare
            # table (see test_device_summary_carries_stack_member_count).
            stack_count_row_index = column_a_values.index("Stack Member Count")
            stack_count_row = list(comparison_sheet.iter_rows())[stack_count_row_index]
            self.assertEqual(stack_count_row[1].value, 1)
            self.assertEqual(stack_count_row[2].value, 2)

            # Combined per-device sheets still present, with both hostnames.
            self.assertIn("Device Info", workbook.sheetnames)
            device_info_hostnames = [row[0].value for row in workbook["Device Info"].iter_rows(min_row=2)]
            self.assertEqual(device_info_hostnames, ["OLD-CISCO-01", "NEW-HUAWEI-01"])

            vlans_sheet = workbook["Comparison"]
            header = [cell.value for cell in vlans_sheet[1]]
            self.assertEqual(header, ["Field / Value", "Device A", "Device B"])

            # PBR / SLA-Track / NQA -- both the quick side-by-side flag
            # on the summary sheet and the combined per-device detail
            # sheets (built from build_workbook(), same mechanism as
            # every other per-device sheet).
            pbr_flag_row_index = column_a_values.index("Has PBR / SLA-Track / NQA?")
            pbr_flag_row = list(comparison_sheet.iter_rows())[pbr_flag_row_index]
            self.assertEqual(pbr_flag_row[1].value, "No")
            self.assertEqual(pbr_flag_row[2].value, "No")
            for sheet_name in (
                "PBR SLA-NQA Tests", "PBR Track Objects", "PBR Policies",
                "PBR Interface Bindings", "PBR EEM Dynamic (Cisco)",
            ):
                self.assertIn(sheet_name, workbook.sheetnames)


if __name__ == "__main__":
    unittest.main()

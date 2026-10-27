"""Excel export: readable sheet names (31 characters at most, cut at a word) and no empty sheets. Synthetic data."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from features.analyzer_core.export import export_results, short_sheet_title
from features.analyzer_core.result import table_section


def test_long_titles_lose_their_bracket_note_or_are_cut_at_a_word():
    assert short_sheet_title("Customised system settings (sys db)") == "Customised system settings"
    assert short_sheet_title("Address and port lists (network / AFM)") == "Address and port lists"
    assert short_sheet_title("Custom files outside the TMSH configuration") == "Custom files outside the TMSH"
    assert short_sheet_title("License: active and time-limited modules") == "License modules"
    assert short_sheet_title("Interfaces") == "Interfaces"
    for title in ("Operating-system access rules (hosts.allow / hosts.deny)", "Route counters from the router (show ip route summary)", "Sections present in one capture only"):
        assert 0 < len(short_sheet_title(title)) <= 31


def test_a_section_without_rows_gets_no_sheet_but_is_listed(tmp_path: Path):
    cols = [("name", "Name")]
    result = {"hostname": "dev1", "findings": [], "coverage": [], "info": [], "cards": [],
              "sections": [table_section("full", "Customised system settings (sys db)", cols, [{"name": "x"}]),
                           table_section("none", "Persistence profiles", cols, [])]}
    out = export_results([result], tmp_path / "x.xlsx")
    names = load_workbook(out).sheetnames
    assert "Customised system settings" in names and "Persistence profiles" not in names
    ws = load_workbook(out)["Empty sections"]
    assert [row[0] for row in ws.iter_rows(min_row=2, values_only=True)] == ["Persistence profiles"]


def test_no_empty_sections_sheet_when_everything_has_rows(tmp_path: Path):
    result = {"hostname": "dev1", "findings": [], "coverage": [], "info": [], "cards": [],
              "sections": [table_section("full", "Things", [("name", "Name")], [{"name": "x"}])]}
    assert "Empty sections" not in load_workbook(export_results([result], tmp_path / "y.xlsx")).sheetnames

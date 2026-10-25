"""Work Plan and Work Report forms in the LLD Generator.

The forms fill the Multipolar Word templates (FM-Impl-039, FM-HW-005) from the
project's shared data plus each form's own fields.
"""
from __future__ import annotations

import io

import docx
import pytest
from docx.oxml.ns import qn

from core.docgen import forms
from features.lld_generator import forms_service, service, store

API = "/lld-generator/api"


@pytest.fixture()
def client(tmp_path, monkeypatch):
    from app import create_app
    monkeypatch.setattr(store, "DB_PATH", str(tmp_path / "nes.db"))
    monkeypatch.setattr(service, "DATA_DIR", tmp_path / "lld_generator")
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


@pytest.fixture()
def pid(client):
    tid = client.get(f"{API}/templates").get_json()[0]["id"]
    p = client.post(f"{API}/projects", json={"name": "Pengadaan Obsolete Server", "client": "Bank Muamalat",
                                             "project_code": "7BMUT25003001", "prepared_by": "Rudi Bachtiar",
                                             "template_id": tid}).get_json()
    return p["id"]


SHARED = {"city": "Jakarta", "site_address": "Jl. Prof. Dr. Satrio Kav. 18",
          "description": "Pengadaan infrastructure untuk refreshment teknologi.",
          "sow": ["Instalasi 2 unit Server S1022 di DC", "Instalasi 2 unit Switch Aruba CX8100 di DC"],
          "support_type": "silver", "support_period": "3 Years After BAST Date",
          "team": {"pm_mlpt": "Rona S. Pranata", "pm_customer": "Shandy Firmansyah",
                   "mlpt": ["Rudi Bachtiar", "Bayu Aji P", "M Kamaludin"], "customer": ["Masgara"], "ops_pic": ""}}


def _text(data: bytes) -> str:
    d = docx.Document(io.BytesIO(data))
    return "\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in d.element.body.iter(qn("w:p")))


def _download(client, result) -> bytes:
    res = client.get(result["download_url"])
    assert res.status_code == 200
    return res.data


def test_templates_hold_no_sample_data():
    for path in forms.TEMPLATES.values():
        text = _text(path.read_bytes())
        for word in ("Muamalat", "Rudi", "Shandy", "Satrio", "Yanuar"):
            assert word not in text


def test_work_plan_round_trip(client, pid):
    form = client.get(f"{API}/projects/{pid}/forms/wpd").get_json()
    assert form["saved"] is False
    assert [p["task"] for p in form["data"]["plan"]] == ["Planning / Design", "Implementation", "Closing"]
    data = {**form["data"],
            "assumption": ["Akses data center disediakan customer"],
            "risk": [{"issue": "Hardware not detected", "level": "Low", "mitigation": "Check hardware", "owner": "MLPT"},
                     {"issue": "", "level": "", "mitigation": "", "owner": ""}],
            "csf": ["Working permit siap"],
            "plan": [{"task": "Planning / Design", "start": "27 Oktober 2025", "finish": "26 November 2025"},
                     {"task": "Implementation", "start": "27 November 2025", "finish": "10 Desember 2025"},
                     {"task": "Closing", "start": "11 Desember 2025", "finish": "11 Desember 2025"}],
            "unknown_key": "dropped"}
    saved = client.put(f"{API}/projects/{pid}/forms/wpd", json={"shared": SHARED, "data": data}).get_json()
    assert saved["saved"] and "unknown_key" not in saved["data"]
    assert saved["shared"]["team"]["mlpt"] == ["Rudi Bachtiar", "Bayu Aji P", "M Kamaludin"]

    result = client.post(f"{API}/projects/{pid}/forms/wpd/generate").get_json()
    assert result["filename"] == "7BMUT25003001 - Work Plan Document.docx"
    body = _download(client, result)
    text = _text(body)
    assert forms.leftover_tokens(body) == []
    for expected in ("Pengadaan Obsolete Server", "7BMUT25003001", "Bank Muamalat", "Jakarta",
                     "Instalasi 2 unit Switch Aruba CX8100 di DC", "Hardware not detected", "Working permit siap",
                     "26 November 2025", "1. Rudi Bachtiar", "3. M Kamaludin", "1. Masgara", "Rona S. Pranata",
                     "Rudi Bachtiar", "3 Years After BAST Date"):
        assert expected in text, expected
    risk_rows = [r for t in docx.Document(io.BytesIO(body)).tables for r in t.rows if r.cells[0].text == "Hardware not detected"]
    assert len(risk_rows) == 1                     # the blank risk row was dropped
    assert client.get(f"{API}/projects/{pid}/documents").get_json()[0]["generated_at"]


def test_service_box_follows_the_choice(client, pid):
    client.put(f"{API}/projects/{pid}/forms/wpd", json={"shared": SHARED, "data": {}})
    body = _download(client, client.post(f"{API}/projects/{pid}/forms/wpd/generate").get_json())
    d = docx.Document(io.BytesIO(body))
    nums = {}
    for p in d.element.body.iter(qn("w:p")):
        text = "".join(t.text or "" for t in p.iter(qn("w:t")))
        num = p.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
        for label in ("Hardware Warranty", "Silver Maintenance", "Gold Maintenance"):
            if text.strip() == label and num is not None:
                nums[label] = num.get(qn("w:val"))
    assert nums["Hardware Warranty"] == nums["Gold Maintenance"] != nums["Silver Maintenance"]


def test_work_report_starts_from_the_work_plan(client, pid):
    plan = [{"task": "Planning / Design", "start": "27 Oktober 2025", "finish": "26 November 2025"},
            {"task": "Implementation", "start": "27 November 2025", "finish": "10 Desember 2025"},
            {"task": "Closing", "start": "11 Desember 2025", "finish": "11 Desember 2025"}]
    client.put(f"{API}/projects/{pid}/forms/wpd", json={"shared": SHARED, "data": {"plan": plan}})
    report = client.get(f"{API}/projects/{pid}/forms/wrd").get_json()["data"]
    assert report["start"] == "27 Oktober 2025" and report["completed"] == "11 Desember 2025"
    assert report["deliverables"] == SHARED["sow"]
    assert [s["task"] for s in report["schedule"]] == ["Planning", "Implementation", "Closing"]
    assert report["schedule"][1]["actual_finish"] == "10 Desember 2025"
    assert {n["text"] for n in report["notes"]} == {"As scheduled"}
    assert report["documents"][0] == {"type": "Low Level Design Document", "date": "", "by": "Shandy Firmansyah"}

    report.update({"lessons": ["Cek bug firmware bawaan pabrik"], "attendees": ["All Team Storage", "All Team Server"],
                   "support_signer": "Yanuar S"})
    client.put(f"{API}/projects/{pid}/forms/wrd", json={"data": report})
    body = _download(client, client.post(f"{API}/projects/{pid}/forms/wrd/generate").get_json())
    text = _text(body)
    assert forms.leftover_tokens(body) == []
    for expected in ("Jl. Prof. Dr. Satrio Kav. 18", "Cek bug firmware bawaan pabrik", "1. All Team Storage",
                     "2. All Team Server", "As scheduled", "Yanuar S", "Low Level Design Document",
                     "Support Mechanism", "Instalasi 2 unit Server S1022 di DC"):
        assert expected in text, expected


def test_empty_form_keeps_its_shape(client, pid):
    body = _download(client, client.post(f"{API}/projects/{pid}/forms/wrd/generate").get_json())
    assert forms.leftover_tokens(body) == []
    assert "PROJECT SCHEDULE REPORT" in _text(body)


def test_unknown_kind_is_404(client, pid):
    assert client.get(f"{API}/projects/{pid}/forms/xyz").status_code == 404
    assert client.get(f"{API}/projects/999/forms/wpd").status_code == 404

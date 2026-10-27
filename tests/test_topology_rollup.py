"""Topology Builder: findings rolled up per device (result, workbook, report, pack). Synthetic captures only."""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from topology_samples import lab_dc  # noqa: E402

from features.topology.excel import to_xlsx  # noqa: E402
from features.topology.findings import findings_by_device  # noqa: E402
from features.topology.service import build_topology, devices_table  # noqa: E402
from tests.test_topology_v05 import _graph_of  # noqa: E402


def _finding(code: str, severity: str, *devices: str) -> dict:
    return {"id": code, "severity": severity, "title": code, "detail": "", "devices": list(devices), "recommendation": "", "reference": ""}


def test_roll_up_counts_every_device_a_finding_names_and_sorts_worst_first():
    graph = _graph_of(["SW-2", "SW-10", "SW-1"], [])
    rows = findings_by_device(graph, [_finding("TOPO-002", "medium", "SW-10", "SW-2"), _finding("TOPO-002", "medium", "SW-10", "SW-1"),
                                      _finding("TOPO-001", "info", "SW-2"), _finding("TOPO-004", "info", "GHOST")])
    assert [r["device"] for r in rows] == ["SW-10", "SW-2", "SW-1", "GHOST"]
    first = rows[0]
    assert (first["medium"], first["low"], first["info"], first["total"], first["findings"]) == (2, 0, 0, 2, "TOPO-002 x2")
    assert {r["device"]: r["captured"] for r in rows} == {"SW-10": "yes", "SW-1": "yes", "SW-2": "yes", "GHOST": "no"}
    assert rows[1]["findings"] == "TOPO-001, TOPO-002"


def test_a_device_without_findings_has_no_row():
    assert findings_by_device(_graph_of(["A", "B"], []), []) == []


@pytest.fixture(scope="module")
def result():
    return build_topology(list(lab_dc().items()), name="Lab")


def test_build_result_holds_the_roll_up_and_it_adds_up(result):
    rows = result["findings_by_device"]
    pairs = sum(len(f["devices"]) for f in result["findings"])
    assert rows and sum(r["total"] for r in rows) == pairs


def test_workbook_has_a_findings_by_device_sheet(result):
    graph = result["graph"]
    data = to_xlsx(result["summary"], devices_table(graph), [], [], result["findings"], by_device=result["findings_by_device"])
    sheet = load_workbook(io.BytesIO(data))["Findings by device"]
    assert [c.value for c in sheet[1]][:2] == ["Device", "Captured"]
    assert sheet.max_row == len(result["findings_by_device"]) + 1


def test_report_and_pack_carry_the_roll_up(tmp_path, monkeypatch):
    from flask import Flask
    from features.topology import routes
    monkeypatch.setattr(routes, "EXPORT_DIR", tmp_path)
    app = Flask("t", template_folder="templates")
    app.register_blueprint(routes.bp)
    app.context_processor(lambda: {"app_version": "t"})
    client = app.test_client()
    body = client.post("/topology/api/build", data={"files": [(io.BytesIO(text.encode()), name) for name, text in lab_dc().items()]}, content_type="multipart/form-data").get_json()
    assert body["ok"] and body["findings_by_device"]
    pack = zipfile.ZipFile(io.BytesIO(client.get(next(d["url"] for d in body["downloads"] if d["name"] == "keystone-topology-pack.zip")).data))
    assert "keystone-topology-pack/tables/findings-by-device.csv" in pack.namelist()
    html = pack.read("keystone-topology-pack/topology-report.html").decode()
    assert "Findings by device" in html and "Findings that name this device" in html

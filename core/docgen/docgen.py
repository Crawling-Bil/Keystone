"""
Engine generate LLD (.docx) dari:
  - master template (.docx) — sumber cover/branding/heading style/section
  - struktur bab (tree, arbitrary depth) + content block per bab (paragraf/tabel/gambar)

Prinsip kerja (hasil trial-and-error di sesi chat, digeneralisasi supaya tidak
terikat ke satu template/project tertentu):

1. Front matter (cover, TOC, dst) = semua paragraf SEBELUM paragraf pertama
   yang stylenya "Heading 1". Ini titik potong body lama dibuang & body baru
   disisipkan.
2. Kalau template punya bagian penutup (mis. "Document Acceptance") yang mau
   dipertahankan apa adanya, itu ditandai lewat `closing_heading` (cocok exact
   dengan teks heading-nya) — semua konten SEBELUM heading itu (dan SESUDAH
   Heading 1 pertama) dianggap "bab lama" dan dibuang, lalu bab baru
   disisipkan tepat sebelum heading penutup itu.
3. Section properties (sectPr) yang berlaku di titik Heading 1 pertama
   (headerReference/footerReference/ukuran halaman) di-capture dan
   dipasang ulang di paragraf terakhir yang disisipkan, supaya header/footer
   & ukuran halaman body baru tetap konsisten dengan template.
4. Heading level yang ada di `auto_numbered_levels` template dianggap
   auto-numbering oleh Word sendiri (lewat multilevel list di style-nya) —
   teks heading ditulis polos. Level lainnya diberi prefix nomor manual
   (dihitung dari struktur tree, format "1.2.3").
"""
import json
import os
from copy import deepcopy

import docx
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches
from docx.oxml import OxmlElement


PLACEHOLDER_TEXT = "Konten belum diisi."


class DocGenError(Exception):
    pass


def qtext(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def find_first_by_style(body, style_id):
    for el in list(body):
        if el.tag.endswith("}p"):
            pPr = el.find(qn("w:pPr"))
            if pPr is not None:
                pStyle = pPr.find(qn("w:pStyle"))
                if pStyle is not None and pStyle.get(qn("w:val")) == style_id:
                    return el
    return None


def find_by_text_after(body, start_el, text):
    found_start = False
    for el in list(body):
        if el is start_el:
            found_start = True
            continue
        if not found_start:
            continue
        if el.tag.endswith("}p") and qtext(el).strip() == text:
            return el
    return None


def capture_sect_pr(body, after_el):
    """Cari sectPr efektif pertama mulai dari after_el ke depan: bisa berupa
    sectPr yang menempel di pPr sebuah paragraf (section break), atau (kalau
    tidak ketemu sampai akhir) sectPr final milik <w:body>."""
    found_start = False
    for el in list(body):
        if el is after_el:
            found_start = True
        if not found_start:
            continue
        if el.tag.endswith("}p"):
            pPr = el.find(qn("w:pPr"))
            if pPr is not None:
                sectPr = pPr.find(qn("w:sectPr"))
                if sectPr is not None:
                    return deepcopy(sectPr)
    # fallback: body-level trailing sectPr
    for el in list(body):
        if el.tag.endswith("}sectPr"):
            return deepcopy(el)
    raise DocGenError("Template tidak punya sectPr sama sekali (file rusak?).")


def collect_between(body, start_el, end_el):
    """Elemen dari start_el (inklusif) sampai sebelum end_el (eksklusif).
    Kalau end_el None, ambil sampai elemen terakhir SEBELUM sectPr final body."""
    children = list(body)
    out = []
    started = False
    for el in children:
        if el is start_el:
            started = True
        if not started:
            continue
        if end_el is not None and el is end_el:
            break
        if el.tag.endswith("}sectPr"):
            break
        out.append(el)
    return out


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def add_seq_field(paragraph, seq_name, cached_number="1"):
    """Sisipkan field Word asli `SEQ <seq_name> \\* ARABIC` ke sebuah paragraf.
    Ini yang bikin List of Pictures/List of Table di depan dokumen bisa
    auto-update (mereka nyari field SEQ dengan nama yang sama)."""
    run = paragraph.add_run()
    r = run._r
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" SEQ {seq_name} \\* ARABIC "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = cached_number
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r.append(fld_begin)
    run2 = paragraph.add_run()
    run2._r.append(instr)
    run3 = paragraph.add_run()
    run3._r.append(fld_sep)
    run3._r.append(txt)
    run4 = paragraph.add_run()
    run4._r.append(fld_end)


def add_caption_paragraph(doc, anchor_el, seq_name, label, text, number, keep_with_next=False):
    """Buat paragraf caption ('Gambar 1 xxx' / 'Tabel 1 xxx') pakai style
    Caption kalau template punya, else Normal+italic. Diselipkan sebelum anchor_el.

    `number` = urutan caption ini (1-based) di antara caption dengan `seq_name`
    yang sama, dipakai sebagai cached text SEQ field-nya. BUG lama: parameter
    ini nggak ada, jadi add_seq_field() selalu jalan dengan default
    cached_number="1" — tiap caption keluar sebagai "Gambar 1"/"Tabel 1" di
    teks yang tampil, dan List of Pictures/List of Table di depan dokumen pun
    ngikut cache yang sama sampai field-nya di-update manual di Word."""
    used_caption_style = True
    try:
        p = doc.add_paragraph(style="Caption")
    except KeyError:
        used_caption_style = False
        p = doc.add_paragraph(style="Normal")

    run = p.add_run(f"{label} ")
    if not used_caption_style:
        run.italic = True
    add_seq_field(p, seq_name, cached_number=str(number))
    if text:
        tail = p.add_run(f" {text}")
        if not used_caption_style:
            tail.italic = True
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    anchor_el.addprevious(p._p)
    return p._p


def set_cell_text(cell, text, bold=False, color=None, size=9, font="Arial"):
    cell.text = ""
    p = cell.paragraphs[0]
    lines = str(text).split("\n") if text is not None else [""]
    for i, line in enumerate(lines):
        if i > 0:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = font
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def _col_widths(headers, rows, total_twips):
    """Lebar kolom proporsional isi: kolom pendek sempit, kolom panjang lebar.
    Return (widths_scaled_to_total, natural_sum) — natural_sum dipakai buat
    deteksi apakah tabel butuh landscape."""
    n = len(headers)
    weights = []
    for c in range(n):
        longest = max(len(w) for w in (str(headers[c]).split() or [""]))
        for r in rows[:200]:
            v = str(r[c]) if c < len(r) and r[c] is not None else ""
            for line in v.split("\n"):
                for w in line.split():
                    longest = max(longest, min(len(w), 30))
        weights.append(max(4, min(longest + 2, 32)))
    natural_sum_twips = sum(weights) * 110  # kira-kira 110 twips per karakter (Arial 9pt)
    total_w = sum(weights)
    scaled = [int(total_twips * w / total_w) for w in weights]
    return scaled, natural_sum_twips


def table_needs_landscape(headers, rows, total_twips):
    if len(headers) >= 8:
        return True
    _, natural = _col_widths(headers, rows, total_twips)
    return natural > total_twips * 1.25


def make_landscape_sectpr(base_sectpr):
    sp = deepcopy(base_sectpr)
    pgSz = sp.find(qn("w:pgSz"))
    if pgSz is not None:
        w = pgSz.get(qn("w:w"))
        h = pgSz.get(qn("w:h"))
        pgSz.set(qn("w:w"), h)
        pgSz.set(qn("w:h"), w)
        pgSz.set(qn("w:orient"), "landscape")
    pgMar = sp.find(qn("w:pgMar"))
    if pgMar is not None:
        # tukar top/bottom jadi left/right ala Word standar
        top, bottom = pgMar.get(qn("w:top")), pgMar.get(qn("w:bottom"))
        left, right = pgMar.get(qn("w:left")), pgMar.get(qn("w:right"))
        pgMar.set(qn("w:top"), left)
        pgMar.set(qn("w:bottom"), right)
        pgMar.set(qn("w:left"), top)
        pgMar.set(qn("w:right"), bottom)
    return sp


def build_table(doc, headers, rows, merges=None, total_twips=9616, header_color="C00000", fills=None, aligns=None):
    n_cols = len(headers)
    table = doc.add_table(rows=1, cols=max(n_cols, 1))
    table.style = "Table Grid"

    widths = _col_widths(headers, rows, total_twips)[0]
    # fixed layout + tblW + tblGrid supaya lebar kolom dihormati Word
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total_twips))
    tblW.set(qn("w:type"), "dxa")
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))

    def set_tc_width(cell, w):
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = tcPr.find(qn("w:tcW"))
        if tcW is None:
            tcW = OxmlElement("w:tcW")
            tcPr.append(tcW)
        tcW.set(qn("w:w"), str(w))
        tcW.set(qn("w:type"), "dxa")

    for c, h in enumerate(headers):
        cell = table.rows[0].cells[c]
        set_cell_text(cell, h, bold=True, color="FFFFFF", size=9)
        set_cell_shading(cell, header_color)
        set_tc_width(cell, widths[c])
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)
    for r in rows:
        row_cells = table.add_row().cells
        for c in range(n_cols):
            val = r[c] if c < len(r) else ""
            set_cell_text(row_cells[c], "" if val is None else val, size=9)
            set_tc_width(row_cells[c], widths[c])

    # merged cells: [r, c, nrows, ncols], r=0 baris header
    for m in (merges or []):
        try:
            r0, c0, nr, nc = m
            a = table.cell(r0, c0)
            b = table.cell(min(r0 + nr - 1, len(table.rows) - 1), min(c0 + nc - 1, n_cols - 1))
            if a._tc is b._tc:
                continue
            merged = a.merge(b)
            # buang paragraf kosong sisa merge (tiap sel bawaan 1 paragraf kosong)
            paras = merged.paragraphs
            keep_one = [p for p in paras if p.text.strip()]
            for p in paras:
                if not p.text.strip() and (keep_one or p is not paras[0]):
                    p._p.getparent().remove(p._p)
        except Exception:
            continue

    # warna sel dari Excel: [r, c, "RRGGBB"], r=0 baris header (mis. status rack)
    for f in (fills or []):
        try:
            r0, c0, color = f
            if 0 < r0 < len(table.rows) and 0 <= c0 < n_cols and color:
                set_cell_shading(table.cell(r0, c0), str(color))
        except (ValueError, TypeError, IndexError):
            continue

    # perataan teks per sel dari Excel: [r, c, "center"|"right"]
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
    from docx.shared import Pt
    for f in (aligns or []):
        try:
            r0, c0, how = f
            if 0 < r0 < len(table.rows) and 0 <= c0 < n_cols:
                value = WD_ALIGN_PARAGRAPH.CENTER if how == "center" else WD_ALIGN_PARAGRAPH.RIGHT
                cell = table.cell(r0, c0)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for para in cell.paragraphs:
                    para.alignment = value
        except (ValueError, TypeError, IndexError):
            continue
    if aligns:
        # tabel grid (rack layout): header rata tengah, tinggi baris seragam
        for cell in table.rows[0].cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        last = len(table.rows) - 1
        for i, row in enumerate(table.rows):
            row.height = Pt(14)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            trPr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            trPr.append(cant)
            for cell in row.cells:
                for para in cell.paragraphs:
                    pf = para.paragraph_format
                    pf.space_before = pf.space_after = Pt(0)
                    pf.line_spacing = 1.0
                    # satu tabel rack satu halaman: semua baris nempel ke baris berikutnya
                    pf.keep_with_next = i < last
    return table


def compute_number(counters, level):
    return ".".join(str(counters[i]) for i in range(1, level + 1) if counters.get(i, 0) > 0)


def generate(
    template_path,
    chapters,          # list[dict] flat, depth-first order, keys: id, level, title
    blocks_by_chapter,  # dict[chapter_id] -> list[dict(block_type, content)]
    uploads_dir,
    output_path,
    closing_heading="",
    auto_numbered_levels=(1, 2),
    max_heading_level=4,
    figure_seq_name="Gambar",
    table_seq_name="Tabel",
):
    if not os.path.exists(template_path):
        raise DocGenError(f"Template tidak ditemukan: {template_path}")

    doc = docx.Document(template_path)
    body = doc.element.body

    first_h1 = find_first_by_style(body, "Heading1")
    if first_h1 is None:
        raise DocGenError(
            "Template tidak punya paragraf dengan style 'Heading 1'. "
            "Pastikan template pakai built-in Word heading styles."
        )

    closing_el = None
    if closing_heading:
        closing_el = find_by_text_after(body, first_h1, closing_heading)
        if closing_el is None:
            raise DocGenError(
                f"Heading penutup '{closing_heading}' tidak ditemukan setelah Heading 1 pertama di template."
            )

    sect_pr_template = capture_sect_pr(body, first_h1)

    # hitung lebar area konten dari sectPr template (buat lebar tabel portrait & landscape)
    total_twips = 9616
    landscape_total_twips = 14400
    try:
        pgSz = sect_pr_template.find(qn("w:pgSz"))
        pgMar = sect_pr_template.find(qn("w:pgMar"))
        if pgSz is not None and pgMar is not None:
            pw, ph = int(pgSz.get(qn("w:w"))), int(pgSz.get(qn("w:h")))
            mtop, mbot = int(pgMar.get(qn("w:top"))), int(pgMar.get(qn("w:bottom")))
            mleft, mright = int(pgMar.get(qn("w:left"))), int(pgMar.get(qn("w:right")))
            total_twips = pw - mleft - mright
            landscape_total_twips = ph - mtop - mbot
    except Exception:
        pass
    landscape_sectpr_template = make_landscape_sectpr(sect_pr_template)

    to_delete = collect_between(body, first_h1, closing_el)
    if not to_delete:
        raise DocGenError("Tidak ada konten body lama yang bisa dibuang — cek template.")

    # anchor = elemen tempat bab baru disisipkan SEBELUM dia
    if closing_el is not None:
        anchor = closing_el
    else:
        # tidak ada closing section -> kita butuh anchor stabil sebelum sectPr final body
        anchor = doc.add_paragraph()._p  # placeholder kosong, akan jadi elemen terakhir
        # taruh placeholder ini di posisi to_delete[0] SEBELUM dihapus, supaya urutan tetap benar
        to_delete[0].addprevious(anchor)

    for el in to_delete:
        el.getparent().remove(el)

    auto_levels = set(auto_numbered_levels)
    counters = {i: 0 for i in range(1, max_heading_level + 1)}
    last_el = None
    orientation = "portrait"
    fig_counter = 0
    tbl_counter = 0

    def switch_orientation(target):
        nonlocal orientation, last_el
        if orientation == target:
            return
        # sectPr yang mengakhiri section SEBELUMNYA dipasang di paragraf kosong ini
        close_sectpr = sect_pr_template if orientation == "portrait" else landscape_sectpr_template
        brk = doc.add_paragraph()
        anchor.addprevious(brk._p)
        brk._p.get_or_add_pPr().append(deepcopy(close_sectpr))
        last_el = brk._p
        orientation = target
    for index, chapter in enumerate(chapters):
        level = min(chapter["level"], max_heading_level)
        counters[level] += 1
        for deeper in range(level + 1, max_heading_level + 1):
            counters[deeper] = 0

        title = chapter["title"]
        if level not in auto_levels:
            number = compute_number(counters, level)
            text = f"{number} {title}"
        else:
            text = title

        style_name = f"Heading {level}"
        try:
            p = doc.add_paragraph(text, style=style_name)
        except KeyError:
            raise DocGenError(
                f"Template tidak punya style '{style_name}' (dipakai untuk bab level {level})."
            )
        if level == 1:
            p.paragraph_format.page_break_before = True
            if orientation != "portrait":
                switch_orientation("portrait")
        anchor.addprevious(p._p)
        last_el = p._p

        blocks = blocks_by_chapter.get(chapter["id"], [])
        has_children = index + 1 < len(chapters) and chapters[index + 1]["level"] > chapter["level"]
        if not blocks and has_children:
            continue  # judul bab induk: isinya ada di sub-bab, tanpa placeholder
        if not blocks:
            ph = doc.add_paragraph(style="Normal")
            run = ph.add_run(PLACEHOLDER_TEXT)
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            ph.paragraph_format.space_after = Pt(8)
            anchor.addprevious(ph._p)
            last_el = ph._p
            continue

        for block in blocks:
            btype = block["block_type"]
            content = block["content"]
            if btype == "paragraph":
                texts = content.get("paragraphs") or [PLACEHOLDER_TEXT]
                for t in texts:
                    bp = doc.add_paragraph(t, style="Normal")
                    bp.paragraph_format.space_after = Pt(8)
                    anchor.addprevious(bp._p)
                    last_el = bp._p

            elif btype == "table":
                headers = content.get("headers", [])
                rows = content.get("rows", [])
                needs_ls = table_needs_landscape(headers, rows, total_twips)
                switch_orientation("landscape" if needs_ls else "portrait")
                cur_width = landscape_total_twips if orientation == "landscape" else total_twips

                caption = content.get("caption", "")
                if caption:
                    tbl_counter += 1
                    add_caption_paragraph(doc, anchor, table_seq_name, table_seq_name, caption, tbl_counter, keep_with_next=True)

                table = build_table(doc, headers, rows, merges=content.get("merges"), total_twips=cur_width,
                                    fills=content.get("fills"), aligns=content.get("aligns"))
                anchor.addprevious(table._tbl)
                # spasi setelah tabel
                spacer = doc.add_paragraph(style="Normal")
                spacer.paragraph_format.space_after = Pt(8)
                anchor.addprevious(spacer._p)
                last_el = spacer._p

            elif btype == "image":
                filename = content.get("filename")
                caption = content.get("caption", "")
                img_path = os.path.join(uploads_dir, filename) if filename else None
                ip = doc.add_paragraph()
                ip.alignment = 1  # center
                if caption:
                    ip.paragraph_format.keep_with_next = True
                anchor.addprevious(ip._p)
                if img_path and os.path.exists(img_path):
                    run = ip.add_run()
                    run.add_picture(img_path, width=Inches(6))
                else:
                    r = ip.add_run(f"[Gambar tidak ditemukan: {filename}]")
                    r.italic = True
                last_el = ip._p
                if caption:
                    fig_counter += 1
                    cap_el = add_caption_paragraph(doc, anchor, figure_seq_name, figure_seq_name, caption, fig_counter)
                    last_el = cap_el
                spacer = doc.add_paragraph(style="Normal")
                spacer.paragraph_format.space_after = Pt(8)
                anchor.addprevious(spacer._p)
                last_el = spacer._p
            else:
                raise DocGenError(f"Tipe block tidak dikenal: {btype}")

    switch_orientation("portrait")

    if last_el is None:
        # tidak ada bab sama sekali -> pakai anchor placeholder yang sudah dibuat
        last_el = anchor if closing_el is None else None
        if last_el is None:
            raise DocGenError("Tidak ada bab untuk digenerate. Tambahkan minimal satu bab dulu.")

    pPr = last_el.get_or_add_pPr() if last_el.tag.endswith("}p") else None
    if pPr is None:
        # last_el adalah tabel -> butuh paragraf kosong ekstra buat menampung sectPr
        extra = doc.add_paragraph()
        anchor.addprevious(extra._p)
        pPr = extra._p.get_or_add_pPr()

    pPr.append(sect_pr_template)

    _set_update_fields_on_open(doc)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def _set_update_fields_on_open(doc):
    """Tandai dokumen supaya Word recalculate SEMUA field otomatis begitu
    dibuka (setara Ctrl+A -> F9), tanpa user harus tahu/lakuin itu manual.

    Ini backstop buat SEQ caption & List of Pictures/List of Table: fix di
    add_caption_paragraph() bikin cached text-nya langsung benar dari awal,
    tapi kalau user sebelumnya pernah ngedit dokumen lain di Word & field
    itu somehow nggak ke-refresh (atau ada viewer yang baca cached text
    tanpa recompute), ini jaring pengamannya — Word akan selalu hitung ulang
    nomor yang benar begitu file dibuka, termasuk TOC "List of Pictures"/
    "List of Table" di depan yang nggak disentuh langsung oleh docgen."""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is not None:
        return
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    # Word strict soal urutan elemen di settings.xml (skema CT_Settings):
    # updateFields harus sebelum elemen-elemen ini, kalau salah posisi Word
    # bisa nolak buka file ("unreadable content").
    after = ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
             "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
             "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
             "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")
    for child in settings:
        local = child.tag.split("}")[-1]
        if local in after:
            child.addprevious(el)
            return
    settings.append(el)


def analyze_template(template_path):
    """Dipakai waktu upload template: cari heading style apa saja yang ada,
    dan tunjukkan daftar heading (levelnya) yang muncul setelah Heading 1
    pertama, supaya user gampang milih `closing_heading` di UI."""
    doc = docx.Document(template_path)
    body = doc.element.body

    available_levels = []
    for lvl in range(1, 10):
        try:
            doc.styles[f"Heading {lvl}"]
            available_levels.append(lvl)
        except KeyError:
            break

    first_h1 = find_first_by_style(body, "Heading1")
    outline = []
    if first_h1 is not None:
        started = False
        for el in list(body):
            if el is first_h1:
                started = True
            if not started or not el.tag.endswith("}p"):
                continue
            pPr = el.find(qn("w:pPr"))
            style_id = None
            if pPr is not None:
                pStyle = pPr.find(qn("w:pStyle"))
                if pStyle is not None:
                    style_id = pStyle.get(qn("w:val"))
            if style_id and style_id.startswith("Heading"):
                text = qtext(el).strip()
                if text:
                    outline.append({"style": style_id, "text": text})

    return {
        "has_heading1": first_h1 is not None,
        "available_heading_levels": available_levels,
        "outline": outline,
    }

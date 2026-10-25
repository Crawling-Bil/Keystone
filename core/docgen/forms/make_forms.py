"""Build the blank Work Plan and Work Report templates from filled samples.

Run once when Multipolar revises a form:

    python core/docgen/forms/make_forms.py WorkPlan.docx WorkReport.docx

The samples are the .doc files converted to .docx (LibreOffice:
``soffice --headless --convert-to docx file.doc``). Sample values are replaced
by placeholders and extra list rows are dropped, so no customer data ends up
in the templates. Placeholder rules are in core/docgen/forms.py.
"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent


def tcs(row):
    return row._tr.findall(qn("w:tc"))


def set_tc(tc, text):
    """Replace a cell's text, keeping the first run's formatting."""
    paragraphs = tc.findall(qn("w:p"))
    first = paragraphs[0]
    for extra in paragraphs[1:]:
        tc.remove(extra)
    runs = first.findall(qn("w:r"))
    for r in runs[1:]:
        first.remove(r)
    for child in list(first):                       # hyperlinks, fields, bookmarks with text
        if child.tag not in (qn("w:pPr"), qn("w:r")):
            first.remove(child)
    if runs:
        run = runs[0]
        for child in list(run):
            if child.tag != qn("w:rPr"):
                run.remove(child)
    else:
        from docx.oxml import OxmlElement
        run = OxmlElement("w:r")
        first.append(run)
    from docx.oxml import OxmlElement
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run.append(t)


def cell(table, r, c, text):
    set_tc(tcs(table.rows[r])[c], text)


def drop(table, *rows):
    for r in sorted(rows, reverse=True):
        tr = table.rows[r]._tr
        tr.getparent().remove(tr)


def signatures(table, r0):
    for i, key in enumerate(("prepared", "acknowledged", "customer")):
        cell(table, r0, 2 * i + 1, f": {{{{{key}_name}}}}")
        cell(table, r0 + 1, 2 * i + 1, f": {{{{{key}_title}}}}")
        cell(table, r0 + 2, 2 * i + 1, f": {{{{{key}_date}}}}")


def team(table, pm_row, first_row, last_row):
    cell(table, pm_row, 1, "{{pm_mlpt}}")
    cell(table, pm_row, 2, "{{pm_customer}}")
    cell(table, first_row, 0, "{{team.label}}")
    cell(table, first_row, 1, "{{team.mlpt}}")
    cell(table, first_row, 2, "{{team.customer}}")
    nxt = first_row + 1                          # continuation row: label cell is merged into the first
    cell(table, nxt, 0, "")
    cell(table, nxt, 1, "{{team.mlpt}}")
    cell(table, nxt, 2, "{{team.customer}}")
    drop(table, *range(nxt + 1, last_row + 1))


def work_plan(src, dst):
    d = Document(src)
    t0, t1, t2, t3, t4 = d.tables
    cell(t1, 0, 2, "{{project_name}}"); cell(t1, 0, 6, "{{doc_date}}")
    cell(t1, 1, 2, "{{project_code}}"); cell(t1, 1, 6, "{{customer}}")
    cell(t1, 2, 6, "{{city}}")
    # t2: description, SoW, service support, assumption, risk
    cell(t2, 1, 0, "{{description}}")
    cell(t2, 5, 0, "{{sow}}")
    cell(t2, 11, 1, "{{svc_hw}} Hardware Warranty")
    cell(t2, 11, 2, "{{svc_silver}} Silver Maintenance")
    cell(t2, 11, 3, "{{svc_gold}} Gold Maintenance")
    cell(t2, 12, 1, "{{support_period}}")
    cell(t2, 15, 0, "{{assumption}}")
    cell(t2, 22, 0, "{{risk.issue}}"); cell(t2, 22, 1, "{{risk.level}}")
    cell(t2, 22, 2, "{{risk.mitigation}}"); cell(t2, 22, 3, "{{risk.owner}}")
    drop(t2, 23, 24, 16, 17, 18, 6, 7, 8, 9, 2, 3)
    cell(t3, 1, 0, "{{csf}}")
    drop(t3, 2, 3)
    # t4: team, plan, resource, signatures (bottom up so row numbers hold)
    signatures(t4, 28)
    cell(t4, 20, 0, "{{res.name}}"); cell(t4, 20, 1, "{{res.skill}}")
    drop(t4, 21, 22, 23, 24)
    cell(t4, 14, 0, "{{plan.task}}"); cell(t4, 14, 1, "{{plan.start}}"); cell(t4, 14, 2, "{{plan.finish}}")
    drop(t4, 15, 16)
    cell(t4, 10, 1, "{{ops_pic}}")
    team(t4, 3, 4, 9)
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    d.save(dst)


def work_report(src, dst):
    d = Document(src)
    t0, t1, t2, t3, t4 = d.tables
    cell(t0, 0, 2, "{{project_name}}"); cell(t0, 0, 6, "{{doc_date}}")
    cell(t0, 1, 2, "{{project_code}}"); cell(t0, 1, 6, "{{customer}}")
    cell(t0, 2, 2, "{{start}}"); cell(t0, 2, 6, "{{site_address}}")
    cell(t0, 3, 2, "{{completed}}"); cell(t0, 3, 6, "")
    # t1 bottom up: team, schedule note, schedule report, lesson, deliverables, description
    team(t1, 24, 25, 30)
    cell(t1, 18, 0, "{{note.task}}"); cell(t1, 18, 1, "{{note.text}}")
    drop(t1, 19, 20)
    for c, key in enumerate(("task", "plan_start", "plan_finish", "actual_start", "actual_finish")):
        cell(t1, 13, c, f"{{{{sched.{key}}}}}")
    drop(t1, 14, 15)
    cell(t1, 8, 0, "{{lesson}}")
    cell(t1, 3, 0, "{{deliverable}}")
    drop(t1, 4, 5, 6)
    cell(t1, 1, 0, "{{description}}")
    # t2: training, documentation (survey stays as printed)
    cell(t2, 9, 0, "{{docs.type}}"); cell(t2, 9, 1, "{{docs.date}}"); cell(t2, 9, 2, "{{docs.by}}")
    drop(t2, 10, 11)
    cell(t2, 2, 1, "{{training_date}}")
    cell(t2, 3, 0, "{{att.label}}"); cell(t2, 3, 1, "{{att.name}}")
    cell(t2, 4, 0, ""); cell(t2, 4, 1, "{{att.name}}")
    drop(t2, 5)
    signatures(t3, 3)
    cell(t4, 1, 0, "{{svc_hw}} Hardware Warranty / {{svc_silver}} Silver Maintenance / {{svc_gold}} Gold Maintenance / Periode : {{support_period}}")
    cell(t4, 2, 2, "{{customer_name}}"); cell(t4, 2, 4, "{{support_signer}}")
    cell(t4, 3, 2, "{{doc_date}}"); cell(t4, 3, 4, "{{doc_date}}")
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    d.save(dst)


def dedupe_zip(path):
    """LibreOffice output can carry docProps/core.xml twice once python-docx
    saves it; Word then asks to repair. Keep the last copy of every part."""
    import zipfile
    with zipfile.ZipFile(path) as z:
        items = {}
        for info in z.infolist():
            items[info.filename] = (info, z.read(info))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, (info, data) in items.items():
            z.writestr(name, data)


if __name__ == "__main__":
    wpd, wrd = sys.argv[1:3]
    work_plan(wpd, HERE / "Work Plan Document.docx")
    work_report(wrd, HERE / "Work Report Document.docx")
    for name in ("Work Plan Document.docx", "Work Report Document.docx"):
        dedupe_zip(HERE / name)
    print("written to", HERE)

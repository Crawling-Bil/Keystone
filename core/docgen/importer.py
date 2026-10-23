"""
Import struktur bab + konten dari file "LLD Content Index" (.xlsx) —
format yang sama dengan Excel Index yang dipakai sebelum tool ini ada:

Sheet "Index" berisi kolom:
  No | Judul Heading | Level (H1..H9) | Tipe Konten | Isi / Nama Sheet Rujukan | Catatan

- Baris tanpa Tipe Konten dianggap heading pengelompokan (tetap diimpor
  sebagai bab, tanpa isi).
- Tipe "Paragraf": kalau kolom Isi cocok dengan nama sheet di workbook,
  isi paragraf dibaca dari sheet itu (kolom ke-2, baris data); kalau tidak,
  teks di kolom Isi dipakai langsung sebagai satu paragraf.
- Tipe "Tabel" / "Tabel Config": kolom Isi harus nama sheet; header + baris
  data dibaca dari sheet itu.
- Tipe "Gambar": dibuat block gambar kosong (file diupload manual lewat app)
  dengan caption dari sheet rujukan kalau ada.
- Baris dengan Judul Heading KOSONG tapi Tipe Konten terisi = isi tambahan
  untuk heading di atasnya. Jadi satu heading bisa punya beberapa isi
  (mis. paragraf deskripsi + tabel spesifikasi + gambar), urut sesuai baris.
- Baris contoh (sel pertama mengandung "(contoh)") otomatis dilewati.
- Nomor manual di depan judul (mis. "4.2.1 AP Inventory ...") otomatis
  dibuang — penomoran diurus ulang oleh app/Word.
"""
import re

import openpyxl


class ImportError_(Exception):
    pass


NUM_PREFIX = re.compile(r"^\d+(\.\d+)*\.?\s+")


def _clean_title(title):
    return NUM_PREFIX.sub("", str(title).strip())


def _is_example_row(cells):
    first = str(cells[0] or "")
    return "(contoh)" in first.lower()


def _sheet_data_start(ws):
    """Cari baris header secara adaptif: lewati baris catatan 'Referensi untuk'
    dan baris kosong, baris non-kosong pertama berikutnya = header."""
    start = 1
    a1 = str(ws.cell(row=1, column=1).value or "")
    if a1.startswith("Referensi untuk"):
        start = 2
    for r in range(start, min(ws.max_row, 20) + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, min(ws.max_column, 30) + 1)]
        if any(v is not None and str(v).strip() != "" for v in vals):
            return r
    return start


def _clean_val(v):
    if v is None:
        return ""
    return str(v).replace("\xa0", " ").strip()


def _read_paragraphs(ws):
    header_row = _sheet_data_start(ws)
    texts = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        if _is_example_row(row):
            continue
        # ambil kolom ke-2 kalau ada (format: No | Isi | Catatan), else kolom 1
        val = row[1] if len(row) > 1 and row[1] not in (None, "") else row[0]
        if val not in (None, ""):
            texts.append(str(val))
    return texts


def _read_table(ws):
    header_row = _sheet_data_start(ws)
    header_cells = list(next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))
    while header_cells and header_cells[-1] in (None, ""):
        header_cells.pop()
    headers = [_clean_val(h) for h in header_cells]
    ncols = len(headers)
    rows = []
    row_map = {}  # excel rownum -> index dalam `rows`
    rnum = header_row
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        rnum += 1
        cells = list(row[:ncols])
        if all(v is None or str(v).strip() == "" for v in cells):
            continue
        if _is_example_row(cells):
            continue
        row_map[rnum] = len(rows)
        rows.append([_clean_val(v) for v in cells])

    # merged ranges dalam area data -> [r, c, nrows, ncols] (0 = baris header)
    merges = []
    for rng in ws.merged_cells.ranges:
        if rng.max_col > ncols or rng.min_row < header_row:
            continue
        if rng.min_row == rng.max_row and rng.min_col == rng.max_col:
            continue
        if rng.min_row == header_row:
            r0 = 0
            span_rows = 1  # merge di header cuma horizontal yang didukung
        else:
            if rng.min_row not in row_map:
                continue
            r0 = row_map[rng.min_row] + 1
            r_end = rng.max_row
            while r_end > rng.min_row and r_end not in row_map:
                r_end -= 1
            span_rows = row_map[r_end] - row_map[rng.min_row] + 1 if r_end in row_map else 1
        merges.append([r0, rng.min_col - 1, span_rows, rng.max_col - rng.min_col + 1])

    # warna sel baris data (fill solid) -> [r, c, "RRGGBB"], r=0 baris header
    # (header tetap pakai warna header template). Dipakai mis. untuk status
    # device di rack layout (Existing / New / Replace).
    fills, aligns = [], []
    for excel_row, r in [(k, v + 1) for k, v in row_map.items()]:
        for c in range(1, ncols + 1):
            cell = ws.cell(row=excel_row, column=c)
            color = _solid_fill(cell)
            if color:
                fills.append([r, c - 1, color])
            horizontal = getattr(cell.alignment, "horizontal", None)
            if horizontal in ("center", "right"):
                aligns.append([r, c - 1, horizontal])
    return headers, rows, merges, fills, aligns


def _solid_fill(cell):
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return ""
    rgb = getattr(fill.fgColor, "rgb", None)
    if not isinstance(rgb, str) or len(rgb) < 6:
        return ""  # warna theme/indexed tidak dibawa
    color = rgb[-6:].upper()
    return "" if color in ("000000", "FFFFFF") else color


def _read_image_captions(ws):
    header_row = _sheet_data_start(ws)
    captions = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        if _is_example_row(row):
            continue
        # format Gambar sheet: No | Nama Gambar | Deskripsi/Caption | Path
        cap = None
        if len(row) > 2 and row[2] not in (None, ""):
            cap = str(row[2])
        elif len(row) > 1 and row[1] not in (None, ""):
            cap = str(row[1])
        captions.append(cap or "")
    return captions


def _build_blocks(wb, tipe, isi, catatan):
    """Block(s) for one Index row: Paragraf, Tabel / Tabel Config, Gambar."""
    blocks = []
    tipe_l = tipe.lower()
    if tipe_l.startswith("paragraf"):
        if isi and isi in wb.sheetnames:
            texts = _read_paragraphs(wb[isi])
            if texts:
                blocks.append({"block_type": "paragraph", "content": {"paragraphs": texts}})
        elif isi:
            blocks.append({"block_type": "paragraph", "content": {"paragraphs": [isi]}})
    elif tipe_l.startswith("tabel"):
        if isi and isi in wb.sheetnames:
            headers, rows, merges, fills, aligns = _read_table(wb[isi])
            if headers and rows:
                content = {"headers": headers, "rows": rows}
                if merges:
                    content["merges"] = merges
                if fills:
                    content["fills"] = fills
                if aligns:
                    content["aligns"] = aligns
                if catatan:
                    content["caption"] = catatan
                blocks.append({"block_type": "table", "content": content})
    elif tipe_l.startswith("gambar"):
        captions = _read_image_captions(wb[isi]) if (isi and isi in wb.sheetnames) else []
        if captions:
            for cap in captions:
                blocks.append({"block_type": "image", "content": {"filename": "", "caption": cap}})
        else:
            blocks.append({"block_type": "image", "content": {"filename": "", "caption": catatan}})
    return blocks


def parse_index_workbook(path):
    """Return list of dicts:
    {level, title, blocks: [ {block_type, content}, ... ]}
    urut sesuai Index (depth-first order tersirat dari urutan baris)."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise ImportError_(f"File tidak bisa dibaca sebagai .xlsx: {e}")

    if "Index" not in wb.sheetnames:
        raise ImportError_("Workbook tidak punya sheet bernama 'Index'.")
    ws = wb["Index"]

    # cari baris header: sel pertama == 'No' dan sel kedua mengandung 'Judul'
    header_rownum = None
    for r in range(1, min(ws.max_row, 30) + 1):
        a = str(ws.cell(row=r, column=1).value or "").strip().lower()
        b = str(ws.cell(row=r, column=2).value or "").strip().lower()
        if a == "no" and "judul" in b:
            header_rownum = r
            break
    if header_rownum is None:
        raise ImportError_(
            "Tidak ketemu baris header Index (kolom pertama 'No', kolom kedua 'Judul Heading')."
        )

    items = []
    for row in ws.iter_rows(min_row=header_rownum + 1, values_only=True):
        no, judul, level_s, tipe, isi, catatan = (list(row) + [None] * 6)[:6]
        if _is_example_row([no]):
            continue
        tipe = str(tipe).strip() if tipe not in (None, "") else ""
        isi = str(isi).strip() if isi not in (None, "") else ""
        catatan = str(catatan).strip() if catatan not in (None, "") else ""

        if judul in (None, "") or str(judul).strip() == "":
            # baris lanjutan: isi tambahan untuk heading sebelumnya
            if tipe and items:
                items[-1]["blocks"].extend(_build_blocks(wb, tipe, isi, catatan))
            continue
        judul = str(judul).strip()

        # level: dari kolom Level ("H3"), fallback dari pola kolom No
        level = None
        if level_s not in (None, ""):
            m = re.match(r"[hH](\d+)", str(level_s).strip())
            if m:
                level = int(m.group(1))
        if level is None:
            no_s = str(no or "").strip()
            if no_s.lower().startswith("bab"):
                level = 1
            elif re.match(r"^\d+(\.\d+)*$", no_s):
                level = no_s.count(".") + 1
            else:
                level = 1

        blocks = _build_blocks(wb, tipe, isi, catatan) if tipe else []
        items.append({"level": level, "title": _clean_title(judul), "blocks": blocks})

    if not items:
        raise ImportError_("Sheet Index kosong — tidak ada baris bab yang bisa diimpor.")
    return items


# sheet "Project" (dari Keystone LLD Index): Field | Value
_PROJECT_FIELDS = {
    "project name": "name", "nama project": "name", "project": "name",
    "customer": "client", "client": "client",
    "project code": "project_code", "kode project": "project_code",
    "company": "company", "implemented by": "company", "perusahaan": "company",
    "prepared by": "prepared_by", "writer": "prepared_by", "disusun oleh": "prepared_by",
    "version": "version", "versi": "version",
}


def parse_project_sheet(path):
    """Data project dari sheet "Project" kalau ada; {} kalau nggak ada."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        return {}
    if "Project" not in wb.sheetnames:
        return {}
    out = {}
    for row in wb["Project"].iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        key = _PROJECT_FIELDS.get(str(row[0]).strip().lower())
        value = "" if len(row) < 2 or row[1] is None else str(row[1]).strip()
        if key and value:
            out[key] = value[:300]
    return out
